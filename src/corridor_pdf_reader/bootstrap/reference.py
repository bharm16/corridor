"""Build the answer key: what each printed sheet displays, cell by cell.

Only the workbook is read here. Display strings come from the workbook's own
number formats (SSF, no recalculation); hidden rows and columns, unprinted
sheets and cells outside a print area are dropped and counted. The PDF reader
never receives this output.
"""

from __future__ import annotations

import argparse
import datetime
import json
import re
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
from multiprocessing import Pool
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "paired_trial" / "vendor"))

from corridor_pdf_reader.bootstrap.corpus import Pair, load_pairs  # noqa: E402

FMT = ROOT / "bootstrap" / "fmt.mjs"
# Excel's built-in formats 14 and 22 are locale short dates: the file stores the
# generic codes below, but a US Excel prints m/d/yyyy. Reference correction,
# logged in bootstrap/README.md.
LOCALE_FORMATS = {"mm-dd-yy": "m/d/yyyy", "m/d/yy h:mm": "m/d/yyyy h:mm"}
LOCALE_FORMAT_KEYS = {14: "m/d/yyyy", 22: "m/d/yyyy h:mm"}
RANGE = re.compile(r"\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)")
ROWS = re.compile(r"\$?(\d+):\$?(\d+)")


def column_number(letters: str) -> int:
    number = 0
    for letter in letters:
        number = number * 26 + ord(letter) - 64
    return number


def parse_ranges(text: str | None) -> list[list[int]] | None:
    if not text:
        return None
    found = [
        [int(r1), column_number(c1), int(r2), column_number(c2)]
        for c1, r1, c2, r2 in RANGE.findall(text)
    ]
    return found or None


def uncached_formulas(path: Path) -> dict[str, set[str]]:
    """Formula cells the workbook stores without a cached result, per sheet xml.

    A PDF printed from Excel shows their results, but no reader of the file can
    know them; the scorer treats those coordinates as unverifiable.
    """
    import zipfile

    found: dict[str, set[str]] = {}
    with zipfile.ZipFile(path) as archive:
        try:
            workbook = archive.read("xl/workbook.xml").decode("utf-8", "replace")
            relationships = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8", "replace")
        except KeyError:
            return found
        targets = dict(re.findall(r'Id="([^"]+)"[^>]*Target="/?(?:xl/)?([^"]+)"', relationships))
        targets.update(dict(re.findall(r'Target="/?(?:xl/)?([^"]+)"[^>]*Id="([^"]+)"', relationships)[::-1]))
        for name, rid in re.findall(r'<sheet\b[^>]*name="([^"]*)"[^>]*r:id="([^"]+)"', workbook):
            member = targets.get(rid)
            if not member or "xl/" + member.lstrip("/") not in archive.namelist():
                continue
            xml = archive.read("xl/" + member.lstrip("/")).decode("utf-8", "replace")
            cells = set()
            for match in re.finditer(r'<c\b([^>]*?)(?:/>|>(.*?)</c>)', xml, re.S):
                body = match.group(2) or ""
                if "<f" in body and "<v" not in body:
                    ref = re.search(r'r="([A-Z]+\d+)"', match.group(1))
                    if ref:
                        cells.add(ref.group(1))
            found[name.replace("&amp;", "&").replace("&quot;", '"').replace("&apos;", "'").replace("&lt;", "<").replace("&gt;", ">")] = cells
    return found


def read_xlsx(path: Path) -> tuple[list[dict[str, Any]], bool]:
    import openpyxl

    book = openpyxl.load_workbook(path, data_only=True, keep_links=False)
    date1904 = book.epoch.year == 1904
    uncached = uncached_formulas(path)
    sheets = []
    for sheet in book.worksheets:
        # A row of (almost) zero height prints nothing, hidden flag or not.
        hidden_rows = {
            r for r, d in sheet.row_dimensions.items() if d.hidden or (d.height is not None and d.height < 1.0)
        }
        hidden_cols: set[int] = set()
        for letter, dimension in sheet.column_dimensions.items():
            if dimension.hidden:
                low = dimension.min or column_number(letter)
                high = dimension.max or low
                hidden_cols.update(range(low, high + 1))
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None or value == "":
                    continue
                raw: Any
                serial: float | None
                if isinstance(value, bool):
                    kind, raw, serial = "bool", value, None
                elif isinstance(value, (int, float)):
                    kind, raw, serial = "number", value, None
                elif isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
                    kind, raw = "date", str(value)
                    try:
                        serial = openpyxl.utils.datetime.to_excel(value, book.epoch)
                    except Exception:
                        serial = None
                else:
                    kind, raw, serial = "text", str(value), None
                cells.append(
                    {
                        "r": cell.row,
                        "c": cell.column,
                        "v": raw,
                        "k": kind,
                        "f": LOCALE_FORMATS.get(cell.number_format, cell.number_format),
                        "s": serial,
                        "h": cell.row in hidden_rows or cell.column in hidden_cols,
                    }
                )
        header_footer = []
        for part in (
            sheet.oddHeader,
            sheet.oddFooter,
            sheet.evenHeader,
            sheet.evenFooter,
            sheet.firstHeader,
            sheet.firstFooter,
        ):
            for side in (part.left, part.center, part.right):
                if side.text:
                    header_footer.append(side.text)
        titles = ROWS.search(sheet.print_title_rows or "")
        sheets.append(
            {
                "name": sheet.title,
                "hidden": sheet.sheet_state != "visible",
                "uncached": sorted(
                    [openpyxl.utils.cell.coordinate_to_tuple(a1) for a1 in uncached.get(sheet.title, ())]
                ),
                "cells": cells,
                "merged": [
                    [r.min_row, r.min_col, r.max_row, r.max_col]
                    for r in sheet.merged_cells.ranges
                ],
                "print_area": parse_ranges(str(sheet.print_area) if sheet.print_area else None),
                "print_titles": [int(titles[1]), int(titles[2])] if titles else None,
                "header_footer": header_footer,
            }
        )
    book.close()
    return sheets, date1904


def read_xls(path: Path) -> tuple[list[dict[str, Any]], bool]:
    import xlrd

    book = xlrd.open_workbook(str(path), formatting_info=True, on_demand=False)
    date1904 = bool(book.datemode)
    named: dict[tuple[str, int], list[list[int]]] = {}
    for name in book.name_obj_list:
        if name.name.lower() not in ("print_area", "print_titles"):
            continue
        try:
            area = name.area2d()
        except Exception:
            continue
        # area2d gives half-open 0-based bounds: sheet, rlo, rhi, clo, chi
        named.setdefault((name.name.lower(), name.scope), []).append(
            [area[1] + 1, area[3] + 1, area[2], area[4]]
        )
    sheets = []
    for index, sheet in enumerate(book.sheets()):
        cells = []
        for r in range(sheet.nrows):
            for c in range(sheet.row_len(r)):
                cell = sheet.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) or cell.value == "":
                    continue
                format_key = book.xf_list[cell.xf_index].format_key
                fmt = LOCALE_FORMAT_KEYS.get(format_key) or book.format_map[format_key].format_str
                raw: Any
                serial: float | None
                if cell.ctype == xlrd.XL_CELL_DATE:
                    kind, raw, serial = "date", str(cell.value), cell.value
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    kind, raw, serial = "number", cell.value, None
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    kind, raw, serial = "bool", bool(cell.value), None
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    kind, raw, serial = "error", xlrd.error_text_from_code.get(cell.value, "#ERR"), None
                else:
                    kind, raw, serial = "text", str(cell.value), None
                info = sheet.rowinfo_map.get(r)
                hidden = bool(
                    getattr(info, "hidden", False)
                    or (info is not None and getattr(info, "height", 1) < 20)
                    or getattr(sheet.colinfo_map.get(c), "hidden", False)
                )
                cells.append({"r": r + 1, "c": c + 1, "v": raw, "k": kind, "f": fmt, "s": serial, "h": hidden})
        header_footer = []
        for text in (getattr(sheet, "header_str", b""), getattr(sheet, "footer_str", b"")):
            if text:
                header_footer.append(text.decode("latin-1") if isinstance(text, bytes) else str(text))
        titles = named.get(("print_titles", index))
        sheets.append(
            {
                "name": sheet.name,
                "hidden": bool(sheet.visibility),
                "cells": cells,
                "merged": [[rlo + 1, clo + 1, rhi, chi] for rlo, rhi, clo, chi in sheet.merged_cells],
                "print_area": named.get(("print_area", index)),
                "print_titles": [titles[0][0], titles[0][2]] if titles else None,
                "header_footer": header_footer,
            }
        )
    book.release_resources()
    return sheets, date1904


def add_display(sheets: list[dict[str, Any]], date1904: bool) -> None:
    """Attach the display string every cell prints as; None when SSF cannot format it."""
    items: list[dict[str, Any]] = []
    slots: dict[tuple[str, Any], int] = {}
    for sheet in sheets:
        for cell in sheet["cells"]:
            if cell["k"] == "text" or cell["k"] == "error":
                cell["d"] = cell["v"]
            elif cell["k"] == "bool":
                cell["d"] = "TRUE" if cell["v"] else "FALSE"
            else:
                value = cell["s"] if cell["k"] == "date" else cell["v"]
                if value is None:
                    cell["d"] = None
                    continue
                slot = (cell["f"] or "General", value)
                if slot not in slots:
                    slots[slot] = len(items)
                    items.append({"f": slot[0], "v": value, "d": date1904})
                cell["_slot"] = slots[slot]
    if items:
        result = subprocess.run(
            ["node", str(FMT)],
            input=json.dumps(items),
            capture_output=True,
            text=True,
            timeout=900,
            check=True,
        )
        displays = json.loads(result.stdout)
        for sheet in sheets:
            for cell in sheet["cells"]:
                if "_slot" in cell:
                    cell["d"] = displays[cell.pop("_slot")]


def inside(r: int, c: int, ranges: list[list[int]] | None) -> bool:
    return ranges is None or any(r1 <= r <= r2 and c1 <= c <= c2 for r1, c1, r2, c2 in ranges)


def build_reference(pair: Pair) -> dict[str, Any]:
    path = pair.book_path
    # The suffix lies on some files; the bytes decide (OLE2 container = BIFF).
    with open(path, "rb") as handle:
        magic = handle.read(8)
    if magic.startswith(b"\xd0\xcf\x11\xe0"):
        sheets, date1904 = read_xls(path)
    elif path.suffix.lower() == ".xls":
        # A zip container behind an .xls suffix: openpyxl refuses the suffix, not the bytes.
        with tempfile.TemporaryDirectory(dir=ROOT / "tmp") as scratch:
            copy = Path(scratch) / (path.stem + ".xlsx")
            shutil.copyfile(path, copy)
            sheets, date1904 = read_xlsx(copy)
    else:
        sheets, date1904 = read_xlsx(path)
    add_display(sheets, date1904)
    # Sheet names can carry trailing spaces that the manifest trimmed.
    by_name = {sheet["name"].strip(): sheet for sheet in sheets}
    stats = {"cells": 0, "hidden_dropped": 0, "outside_print_area": 0, "unformatted": 0, "merged_nonanchor_dropped": 0, "uncached_formulas": 0}
    errors = []
    printed = []
    for name in pair.printed_sheets:
        sheet = by_name.get(name)
        if sheet is None:
            errors.append(f"printed sheet {name!r} not found in workbook")
            continue
        anchors = {(m[0], m[1]) for m in sheet["merged"]}
        covered = {
            (r, c)
            for r1, c1, r2, c2 in sheet["merged"]
            for r in range(r1, r2 + 1)
            for c in range(c1, c2 + 1)
        } - anchors
        kept = []
        # Text of cells the print area leaves out, kept only to explain what
        # overflows into the printed area; never a cell the PDF must show.
        outside_text: list[str] = []
        titles = sheet["print_titles"]
        columns = [[1, c1, 1_048_576, c2] for _, c1, _, c2 in sheet["print_area"]] if sheet["print_area"] else None
        for cell in sheet["cells"]:
            r, c = cell["r"], cell["c"]
            if cell["h"]:
                stats["hidden_dropped"] += 1
                continue
            # Rows to repeat at top print on every page even when the print
            # area starts below them; only the print area's columns limit them.
            repeated = bool(titles) and titles[0] <= r <= titles[1] and inside(r, c, columns)
            if not repeated and not inside(r, c, sheet["print_area"]):
                stats["outside_print_area"] += 1
                if cell["d"] is not None and str(cell["d"]).strip():
                    outside_text.append(str(cell["d"]))
                continue
            if (r, c) in covered:
                stats["merged_nonanchor_dropped"] += 1
                continue
            if cell["d"] is None:
                stats["unformatted"] += 1
                continue
            if not str(cell["d"]).strip():
                continue
            kept.append([r, c, str(cell["d"]), cell["k"]])
            stats["cells"] += 1
        stats["uncached_formulas"] += len(sheet.get("uncached", []))
        printed.append(
            {
                "name": name,
                "cells": kept,
                "outside_text": outside_text,
                "uncached": [
                    [r, c]
                    for r, c in sheet.get("uncached", [])
                    if inside(r, c, sheet["print_area"]) or (titles and titles[0] <= r <= titles[1])
                ],
                "merged": sheet["merged"],
                "print_area": sheet["print_area"],
                "print_titles": sheet["print_titles"],
                "header_footer": sheet["header_footer"],
            }
        )
    return {"key": pair.key, "spreadsheet": pair.spreadsheet, "sheets": printed, "stats": stats, "errors": errors}


def build_one(task: tuple[Pair, Path]) -> dict[str, Any]:
    pair, output = task
    started = time.perf_counter()
    try:
        reference = build_reference(pair)
        (output / f"{pair.key}.json").write_text(json.dumps(reference) + "\n")
        return {"key": pair.key, "seconds": round(time.perf_counter() - started, 2), **reference["stats"], "errors": reference["errors"]}
    except Exception as exc:
        return {"key": pair.key, "seconds": round(time.perf_counter() - started, 2), "error": f"{type(exc).__name__}: {exc}", "trace": traceback.format_exc()[-1500:]}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True, help="new directory")
    parser.add_argument("--jobs", type=int, default=4)
    parser.add_argument("--keys", nargs="*", help="restrict to these pair keys")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    pairs = load_pairs()
    if args.keys:
        pairs = [pair for pair in pairs if pair.key in set(args.keys)]
    receipts = []
    with Pool(args.jobs) as pool:
        for receipt in pool.imap_unordered(build_one, [(pair, args.output) for pair in pairs]):
            receipts.append(receipt)
            if "error" in receipt:
                print("ERROR", receipt["key"], receipt["error"], flush=True)
    receipts.sort(key=lambda r: r["key"])
    (args.output / "receipts.json").write_text(json.dumps(receipts, indent=1) + "\n")
    failed = [r for r in receipts if "error" in r]
    cells = sum(r.get("cells", 0) for r in receipts)
    print(f"{len(receipts)} workbooks, {cells} reference cells, {len(failed)} failed -> {args.output}")


if __name__ == "__main__":
    main()
