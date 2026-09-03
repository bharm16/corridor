"""The Corridor operations reading of one workbook offered for Adopt Baseline (#509).

Onboarding has two bounded readings and keeping them apart is the whole point
of this module.  **Corridor operations** proves that the customer's exact file
can be parsed and returned safely: which parser read it, which worksheets it
holds, what is hidden or computed, how each printed column maps to a canonical
field, which columns no canonical field names, how a source-row key is built,
which populated cells no released transformation can type, and whether the
adopted values survive a round trip back through the mapping.  Only after that
does anything reach the **coordinator**, and what reaches them is material
project questions alone (``baseline_adoption``).

Why a separate module rather than another branch of ``extract_sheet``.  The
native extractor answers "what conflicts does this document propose"; it
produces Extracted Proposals, drops a row that misses a required field, and
reports unmapped columns as a list of headings beside a Candidate.  Adopt
Baseline answers a different question — "can this exact file become the
accepted record, and what would a person be agreeing to" — so nothing may be
dropped: a retired row, a row missing a required field, an unknown column, and
a populated cell no transformation can type are each carried through to the
adoption act with a reason.  Sharing the extractor's loop was tried first and
each of those four cases had to be smuggled back out of it through an accounting
receipt that was designed for a different question.

Two readers are deliberately run over the same bytes.  ``sheets.read_workbook``
supplies the header row and the column mapping, because it already owns the
published forms' exact column names (ADR-0005, ADR-0009).
``source_segments.spreadsheet_segments`` supplies the exact per-cell text and
locator that will become the Source Segments a Fact materializes from, because
a baseline value must reproduce from the cell the record cites and not from a
second rendering of it.  The round-trip check is the two agreeing; where they
disagree the file does not adopt, and that is an operations failure the
coordinator never sees.

Nothing here touches the database, calls a model, or writes anything.  It reads
bytes and reports.
"""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from corridor.fact_types import FACT_TYPE_CONTRACTS, SINGLE_VALUED_FACT_TYPES
from corridor.materializer import FactValidationError, validated_scalar_value
from corridor.sheets import (
    NoConflictSheet,
    ConflictSheet,
    conflict_sheet,
    read_workbook,
)
from corridor.source_segments import spreadsheet_segments
from corridor.vocabulary import (
    MIN_ROW_FIELDS,
    REQUIRED,
    is_retired_row,
    is_sequencing_header,
)


# What produced this reading.  Recorded on the adoption receipt, so a workbook
# read by a later importer is a different reading of the same file and the two
# are never pooled — the same argument `extract_sheet.PROMPT_VERSION` makes.
IMPORTER_IDENTITY = "corridor.adopt-baseline"
IMPORTER_VERSION = "baseline_workbook_v1"
PARSER = "openpyxl:data_only"

# How a source row is named.  Stated rather than implied because the key is
# stored on the adoption and has to mean the same thing when a later revision
# of the same workbook is compared against it: the adopted worksheet's own
# name, and the worksheet row number the customer sees in Excel.  It is
# deliberately the same key `facts.subject_key` carries for a structured cell,
# so an adopted baseline value and its source row are one lookup apart.
SOURCE_ROW_KEY_RULE = "sheet_name!worksheet_row_number"

# The fields ADR-0076 calls material, expressed in this vocabulary's canonical
# names.  A populated material value that no released transformation can type
# is a coordinator question; the same failure on any other field stays with
# Corridor operations.  `cost responsibility` and `agreement or permit status`
# have no canonical column and so appear nowhere here.
MATERIAL_FIELDS = frozenset(
    {
        "external_org",          # Utility Owner
        "utility_id",            # conflict identity
        "size",                  # facility size
        "utility_type",          # facility type
        "material",              # facility material
        "station_from",          # location, when it affects matching
        "station_to",
        "offset_from",
        "offset_to",
        "resolution_strategy",   # resolution method
        "committed_date",        # Promised For
        "need_date",             # Required By linkage
        "marked_resolution",     # completion or closure
        "applies_to",            # Applies To
    }
)

# The canonical fields a baseline value can be captured for: the scalar Fact
# types whose released contract accepts a spreadsheet cell.  `applies_to` and
# `closure_result` are absent on purpose — their values are references into a
# Project Record that this very act is establishing, so a populated `Applies
# To` in an adopted workbook is reported as an unmapped material value rather
# than resolved against a record that does not exist yet.
BASELINE_FACT_FIELDS = tuple(
    name
    for name in SINGLE_VALUED_FACT_TYPES
    if "spreadsheet_cell" in FACT_TYPE_CONTRACTS[name].accepted_segment_kinds
)

# The roles a printed heading may carry *out* of this workbook: the
# utility-management system's own record identifier for the conflict, and a
# link to a controlled document about it.  They are preserved on source-row
# identity so #527 and #528 can offer a read-only deep link without re-reading
# the file, and they are kept apart because they are different records.
#
# **No heading is assigned either role from its spelling alone (#597).**  This
# module used to default four guessed names — `UCM Record ID`, `Document
# Control No.`, `Record URL`, `Document Link` — and no customer form was ever
# read for them.  A caller now declares the exact headings its registered
# mapping revision names (`field_mapping_manifest.ExternalReference`), the
# default is none at all, and a heading nobody declared stays a retained
# unknown column.  #561 supplies real ones from a partner's own data
# dictionary.
EXTERNAL_REFERENCE_ROLES = ("external_system_id", "source_url")


class BaselineWorkbookUnsupported(ValueError):
    """Corridor operations cannot read this file as an adoptable baseline."""


@dataclass(frozen=True)
class WorksheetReading:
    """One worksheet of the offered workbook and what operations did with it."""

    name: str
    position: int
    state: str
    rows: int
    disposition: str
    reason: str


@dataclass(frozen=True)
class MappedColumn:
    """One printed heading a canonical field names."""

    column: str
    heading: str
    field: str
    material: bool


@dataclass(frozen=True)
class UnknownColumn:
    """One printed heading no canonical field names, with what it holds."""

    column: str
    heading: str
    populated_cells: int


@dataclass(frozen=True)
class ControlledVocabulary:
    """One column's controlled vocabulary and whether its values are inside it."""

    column: str
    heading: str
    allowed_values: tuple[str, ...]
    checked: bool
    out_of_vocabulary: tuple[str, ...]
    reference: str | None = None


@dataclass(frozen=True)
class CellValue:
    """One exact cell, at the locator its Source Segment will carry."""

    field: str
    sheet_name: str
    cell_range: str
    exact_text: str


@dataclass(frozen=True)
class UnsupportedValue:
    """One populated cell no released transformation can type."""

    field: str
    sheet_name: str
    cell_range: str
    exact_text: str
    reason: str
    material: bool


@dataclass(frozen=True)
class RetainedValue:
    """One populated cell under a column no canonical field names."""

    heading: str
    sheet_name: str
    cell_range: str
    exact_text: str


@dataclass(frozen=True)
class OperationsDiagnostic:
    """One importer diagnostic, addressed to Corridor operations only."""

    code: str
    detail: str
    locator: str | None = None
    blocking: bool = False


@dataclass(frozen=True)
class RoundTrip:
    """Whether the adopted values return exactly through the field mapping."""

    values_checked: int
    rows_checked: int
    mismatches: tuple[str, ...]

    @property
    def clean(self) -> bool:
        return not self.mismatches


@dataclass(frozen=True)
class BaselineRow:
    """One source row of the adopted worksheet, with nothing dropped."""

    sheet_name: str
    row_number: int
    source_row_key: str
    business_identity: str | None
    external_system_id: str | None
    source_url: str | None
    values: tuple[CellValue, ...]
    unsupported: tuple[UnsupportedValue, ...]
    retained: tuple[RetainedValue, ...]
    exclusion_reason: str | None

    @property
    def excluded(self) -> bool:
        return self.exclusion_reason is not None

    def value(self, field: str) -> str | None:
        for cell in self.values:
            if cell.field == field:
                return cell.exact_text
        return None


@dataclass(frozen=True)
class OperationsReading:
    """What Corridor operations proved about this file before anyone adopts it."""

    parser: str
    importer_identity: str
    importer_version: str
    worksheets: tuple[WorksheetReading, ...]
    adopted_sheet: str
    header_row_number: int
    column_mapping: tuple[MappedColumn, ...]
    unknown_columns: tuple[UnknownColumn, ...]
    controlled_vocabularies: tuple[ControlledVocabulary, ...]
    formula_cells: tuple[str, ...]
    hidden_content: tuple[str, ...]
    source_row_key_rule: str
    unsupported_values: tuple[UnsupportedValue, ...]
    diagnostics: tuple[OperationsDiagnostic, ...]
    round_trip: RoundTrip
    rows: tuple[BaselineRow, ...]

    @property
    def blocking_diagnostics(self) -> tuple[OperationsDiagnostic, ...]:
        return tuple(item for item in self.diagnostics if item.blocking)

    @property
    def resolved(self) -> bool:
        """Whether operations settled the mechanics and the file may be offered."""

        return not self.blocking_diagnostics and self.round_trip.clean


def read_baseline_workbook(
    path: Path | str,
    *,
    external_references: Mapping[str, str] | None = None,
) -> OperationsReading:
    """Read one offered workbook and report only what operations owns.

    ``external_references`` is the declared heading-to-role map of the mapping
    revision this file is being read through, casefolded and whitespace
    collapsed.  It defaults to none: an undeclared heading carries no meaning
    here whatever it is spelled (#597).

    Raises ``BaselineWorkbookUnsupported`` when the file is not a workbook this
    importer can read at all — there is nothing to report a diagnostic *about*,
    and reporting an empty adoption for an unreadable file is the failure
    ``sheets.NoConflictSheet`` exists to prevent.
    """

    path = Path(path)
    declared = dict(external_references or {})
    unknown = sorted(set(declared.values()) - set(EXTERNAL_REFERENCE_ROLES))
    if unknown:
        raise BaselineWorkbookUnsupported(
            f"{unknown} is not an external-reference role this reader carries"
        )
    try:
        sheets = read_workbook(path)
    except Exception as exc:  # openpyxl raises its own container errors
        raise BaselineWorkbookUnsupported(
            f"{path.name} cannot be read as a workbook: {exc}"
        ) from exc
    try:
        chosen = conflict_sheet(sheets)
    except NoConflictSheet as exc:
        raise BaselineWorkbookUnsupported(str(exc)) from exc

    segments = {
        (segment.sheet_name, segment.cell_range): segment.exact_text
        for segment in spreadsheet_segments(path)
    }
    formula_cells, hidden_content, sheet_states = _workbook_structure(path)
    diagnostics: list[OperationsDiagnostic] = []

    worksheets = tuple(
        WorksheetReading(
            name=sheet.name,
            position=position,
            state=sheet_states.get(sheet.name, "visible"),
            rows=len(sheet.rows),
            disposition=(
                "adopted" if sheet.name == chosen.sheet.name else "not_adopted"
            ),
            reason=(
                "the one sheet of this workbook that heads a utility conflict "
                "matrix"
                if sheet.name == chosen.sheet.name
                else "no canonical conflict columns, or a narrower rendering of "
                "the adopted sheet"
            ),
        )
        for position, sheet in enumerate(sheets, start=1)
    )

    headings = list(chosen.headings)
    for position, heading in enumerate(headings):
        if is_sequencing_header(heading):
            diagnostics.append(
                OperationsDiagnostic(
                    code="sequencing_column_unsupported",
                    detail=(
                        f"{heading!r} asserts work sequencing, which Corridor "
                        "has no model for; adopting the sheet would drop it"
                    ),
                    locator=f"{chosen.sheet.name}!{get_column_letter(position + 1)}",
                    blocking=True,
                )
            )

    column_mapping = tuple(
        MappedColumn(
            column=get_column_letter(position + 1),
            heading=" ".join(str(headings[position] or "").split()),
            field=field,
            material=field in MATERIAL_FIELDS,
        )
        for position, field in sorted(chosen.mapping.items())
    )
    reference_columns = _reference_columns(headings, chosen.mapping, declared)
    rows, unknown_counts, unsupported = _read_rows(
        chosen, segments, reference_columns, diagnostics
    )
    unknown_columns = tuple(
        UnknownColumn(
            column=get_column_letter(position + 1),
            heading=_reported_heading(position, headings[position]),
            populated_cells=unknown_counts.get(position, 0),
        )
        for position in range(len(headings))
        if position not in chosen.mapping and position not in reference_columns
    )
    if not rows:
        diagnostics.append(
            OperationsDiagnostic(
                code="no_source_rows",
                detail=(
                    f"{chosen.sheet.name!r} heads a conflict matrix but holds no "
                    "populated source rows"
                ),
                blocking=True,
            )
        )

    return OperationsReading(
        parser=PARSER,
        importer_identity=IMPORTER_IDENTITY,
        importer_version=IMPORTER_VERSION,
        worksheets=worksheets,
        adopted_sheet=chosen.sheet.name,
        header_row_number=chosen.header_index + 1,
        column_mapping=column_mapping,
        unknown_columns=unknown_columns,
        controlled_vocabularies=_controlled_vocabularies(path, chosen, rows),
        formula_cells=formula_cells,
        hidden_content=hidden_content,
        source_row_key_rule=SOURCE_ROW_KEY_RULE,
        unsupported_values=tuple(unsupported),
        diagnostics=tuple(diagnostics),
        round_trip=_round_trip(rows, segments),
        rows=tuple(rows),
    )


def _read_rows(
    chosen: ConflictSheet,
    segments: dict[tuple[str, str], str],
    reference_columns: dict[int, str],
    diagnostics: list[OperationsDiagnostic],
) -> tuple[list[BaselineRow], dict[int, int], list[UnsupportedValue]]:
    """Every populated source row, with nothing dropped and every reason stated."""

    sheet_name = chosen.sheet.name
    headings = list(chosen.headings)
    unknown_counts: dict[int, int] = {}
    unsupported_all: list[UnsupportedValue] = []
    rows: list[BaselineRow] = []

    for table_row, raw in enumerate(chosen.rows, start=1):
        row_number = chosen.header_index + 1 + table_row
        if not any(str(value or "").strip() for value in raw):
            continue
        values: list[CellValue] = []
        unsupported: list[UnsupportedValue] = []
        retained: list[RetainedValue] = []
        fields: dict[str, str] = {}
        external_system_id: str | None = None
        source_url: str | None = None

        for position, cell_text in enumerate(raw):
            text = str(cell_text or "").strip()
            if not text:
                continue
            cell_range = f"{get_column_letter(position + 1)}{row_number}"
            field = chosen.mapping.get(position)
            if field is None:
                role = reference_columns.get(position)
                if role == "external_system_id":
                    external_system_id = external_system_id or text
                    continue
                if role == "source_url":
                    source_url = source_url or text
                    continue
                unknown_counts[position] = unknown_counts.get(position, 0) + 1
                retained.append(
                    RetainedValue(
                        heading=_reported_heading(position, headings[position])
                        if position < len(headings)
                        else _reported_heading(position, ""),
                        sheet_name=sheet_name,
                        cell_range=cell_range,
                        exact_text=text,
                    )
                )
                continue
            fields[field] = text
            if segments.get((sheet_name, cell_range)) is None:
                diagnostics.append(
                    OperationsDiagnostic(
                        code="mapped_cell_has_no_segment",
                        detail=(
                            "the segment reader found no cell where the sheet "
                            "reader found a value"
                        ),
                        locator=f"{sheet_name}!{cell_range}",
                        blocking=True,
                    )
                )
                continue
            reason = _unsupported_reason(field, text)
            if reason is not None:
                item = UnsupportedValue(
                    field=field,
                    sheet_name=sheet_name,
                    cell_range=cell_range,
                    exact_text=text,
                    reason=reason,
                    material=field in MATERIAL_FIELDS,
                )
                unsupported.append(item)
                unsupported_all.append(item)
                continue
            values.append(
                CellValue(
                    field=field,
                    sheet_name=sheet_name,
                    cell_range=cell_range,
                    exact_text=text,
                )
            )

        rows.append(
            BaselineRow(
                sheet_name=sheet_name,
                row_number=row_number,
                source_row_key=f"{sheet_name}!{row_number}",
                business_identity=fields.get("utility_id"),
                external_system_id=external_system_id,
                source_url=source_url,
                values=tuple(values),
                unsupported=tuple(unsupported),
                retained=tuple(retained),
                exclusion_reason=_exclusion_reason(fields),
            )
        )
    return rows, unknown_counts, unsupported_all


def _exclusion_reason(fields: dict[str, str]) -> str | None:
    """Why this row would not become part of the accepted record, or None.

    Every reason is the form's own bookkeeping or an identity the record could
    not be built on. None of them removes the row: an excluded row is retained
    with this reason and reported to the coordinator before anyone adopts.
    """

    if is_retired_row(fields):
        return "retired_row"
    if len(fields) < MIN_ROW_FIELDS:
        return "insufficient_mapped_fields"
    if not all(fields.get(name) for name in REQUIRED):
        return "missing_required_fields"
    return None


def _unsupported_reason(field: str, text: str) -> str | None:
    """Why no released transformation can type this populated cell, or None."""

    if field not in BASELINE_FACT_FIELDS:
        return "no_released_baseline_transformation"
    contract = FACT_TYPE_CONTRACTS[field]
    try:
        validated_scalar_value(contract, text)
    except FactValidationError as exc:
        return str(exc)
    return None


def _reference_columns(
    headings: list[str], mapping: dict[int, str], declared: dict[str, str]
) -> dict[int, str]:
    """Which columns carry an external system id or a source URL, by exact name."""

    found: dict[int, str] = {}
    taken: set[str] = set()
    for position, heading in enumerate(headings):
        if position in mapping:
            continue
        role = declared.get(" ".join(str(heading or "").split()).casefold())
        if role is not None and role not in taken:
            found[position] = role
            taken.add(role)
    return found


def _round_trip(
    rows: list[BaselineRow], segments: dict[tuple[str, str], str]
) -> RoundTrip:
    """Re-read every adopted value from the segment reader and compare.

    Two independent readings of the same bytes have to agree before a value
    can become the accepted record: the sheet reader supplies the mapping and
    the segment reader supplies the exact text a Fact will materialize from.
    A disagreement is the one output check that has to happen before adoption
    rather than after, because after adoption the wrong value is the record.
    """

    mismatches: list[str] = []
    checked = 0
    for row in rows:
        for cell in row.values:
            checked += 1
            exact = segments.get((cell.sheet_name, cell.cell_range))
            if exact is None or exact.strip() != cell.exact_text:
                mismatches.append(
                    f"{cell.sheet_name}!{cell.cell_range} reads "
                    f"{cell.exact_text!r} through the mapping and "
                    f"{(exact or '').strip()!r} through its Source Segment"
                )
    return RoundTrip(
        values_checked=checked, rows_checked=len(rows), mismatches=tuple(mismatches)
    )


def _workbook_structure(
    path: Path,
) -> tuple[tuple[str, ...], tuple[str, ...], dict[str, str]]:
    """Formula cells, hidden content, and each worksheet's visibility state.

    Read from a second, formula-preserving load: the reading everything else
    uses is `data_only`, which is right — the value is what the document
    asserts, and `=B2*2` is not something a reviewer can check a quote against
    (ADR-0005) — but it also means a computed cell is indistinguishable from a
    typed one, and a person adopting a baseline is entitled to know which is
    which.
    """

    formulas: list[str] = []
    hidden: list[str] = []
    states: dict[str, str] = {}
    workbook = load_workbook(path, data_only=False)
    try:
        for sheet in workbook.worksheets:
            states[sheet.title] = sheet.sheet_state
            if sheet.sheet_state != "visible":
                hidden.append(f"worksheet {sheet.title!r} is {sheet.sheet_state}")
            for number, dimension in sheet.row_dimensions.items():
                if dimension.hidden:
                    hidden.append(f"{sheet.title}!row {number} is hidden")
            for letter, dimension in sheet.column_dimensions.items():
                if dimension.hidden:
                    hidden.append(f"{sheet.title}!column {letter} is hidden")
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas.append(f"{sheet.title}!{cell.coordinate}")
    finally:
        workbook.close()
    return tuple(formulas), tuple(hidden), states


def _controlled_vocabularies(
    path: Path, chosen: ConflictSheet, rows: list[BaselineRow]
) -> tuple[ControlledVocabulary, ...]:
    """Every list validation on the adopted sheet, and the values outside it.

    The published template ships a `Drop-Down Lists` sheet and binds its
    columns to it, so a controlled column is a fact about the file rather than
    a rule Corridor imposes.  A validation whose allowed values are a range
    reference is reported unchecked with the reference: resolving it would mean
    deciding what another sheet means, which is not a mechanical question.
    """

    headings = list(chosen.headings)
    readings: list[ControlledVocabulary] = []
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook[chosen.sheet.name]
        for validation in sheet.data_validations.dataValidation:
            if validation.type != "list" or not validation.formula1:
                continue
            formula = str(validation.formula1)
            inline = formula.startswith('"') and formula.endswith('"')
            allowed = (
                tuple(
                    part.strip()
                    for part in formula[1:-1].split(",")
                    if part.strip()
                )
                if inline
                else ()
            )
            for column in _validated_columns(validation):
                position = _column_index(column)
                if position is None or position >= len(headings):
                    continue
                values = tuple(
                    value.exact_text
                    for row in rows
                    for value in row.values
                    if value.cell_range.startswith(column)
                    and value.cell_range[len(column) :].isdigit()
                )
                readings.append(
                    ControlledVocabulary(
                        column=column,
                        heading=_reported_heading(position, headings[position]),
                        allowed_values=allowed,
                        checked=inline,
                        out_of_vocabulary=tuple(
                            sorted({value for value in values if value not in allowed})
                        )
                        if inline
                        else (),
                        reference=None if inline else formula,
                    )
                )
    finally:
        workbook.close()
    return tuple(readings)


def _validated_columns(validation) -> tuple[str, ...]:
    """The column letters one data validation covers, in sqref order."""

    columns: list[str] = []
    for reference in str(validation.sqref).split():
        for part in reference.split(":"):
            letters = "".join(char for char in part if char.isalpha())
            if letters and letters not in columns:
                columns.append(letters)
    return tuple(columns)


def _column_index(column: str) -> int | None:
    index = 0
    for char in column:
        if not char.isalpha():
            return None
        index = index * 26 + (ord(char.upper()) - ord("A") + 1)
    return index - 1 if index else None


def _reported_heading(position: int, value: str | None) -> str:
    cleaned = " ".join(str(value or "").split())
    return cleaned or f"Column {get_column_letter(position + 1)} (blank heading)"
