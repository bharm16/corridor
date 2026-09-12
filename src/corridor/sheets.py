"""Reading a spreadsheet source natively (ADR-0005).

TxDOT's Utility Conflict Matrix *is* a spreadsheet, and the PDFs everywhere
else in this corpus are printouts of it. Every defect ADR-0004 catalogues is
print damage to data that was structured before someone exported it —
separators lost between text spans, cells clipped at their boundaries, four
columns collapsed into `Start Station, Offset`, an owner pushed out of the
table into a page header.

So nothing here uses a model. The structure is explicit in the file: named
columns, a header row, one value per cell. Running a vision model over a
spreadsheet would be recovering from pixels what the file already states,
and it would reintroduce the transcription error class ADR-0006 removed by
construction.

Two things this deliberately does **not** do:

- **Assume a conflict id is unique.** The form's own data dictionary says
  `Utility Conflict ID` is "unique within the transportation project"; real
  documents disagree, and one Project A revision reuses 47 of them. The
  intent is worth recording precisely because the data does not honour it.
- **Assume a workbook has one table.** The published template has seven
  sheets, two of which head a column `Utility Conflict ID`.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from corridor.extraction_errors import NoMatrixFound
from corridor.vocabulary import (
    REQUIRED,
    STRUCTURED_RECORD_HEADINGS,
    TEMPLATE_FIELDS,
    UCM_CONFLICT_LIST_HEADINGS,
)

# How deep a header can sit. The published template puts a merged title
# band on row 1 and the real headings on row 2. I-35 NEX South's "UCM -
# Utility Conflict List" workbook is deeper: a title, then five rows of
# project identification (owner, CCSJ, description, route, developed/reviewed
# by), then the column header on row 8 (#365). Bounded so a "header" found
# far down is a data row that happens to read like one, not searched
# open-endedly.
HEADER_SEARCH_ROWS = 8

# ...and how many canonical fields a row must name to be the header rather
# than a title band. A band populates one cell of the row it spans, which
# is the shape being excluded.
MIN_HEADER_FIELDS = 2

# printed heading -> canonical field, for lookup by the name a form prints.
# The template half is inverted from `TEMPLATE_FIELDS` rather than restated,
# so a spreadsheet and a printed page cannot drift into two vocabularies —
# the whole argument of ADR-0005 is that they are one document in two forms.
# The second half is the exact headings of TxDOT's earlier "UCM - Utility
# Conflict List" form (I-35 NEX South, #365), added on the same terms: an
# exact column name read from the form's own data dictionary, never a synonym.
# Distinct forms name a field differently, and one sheet never carries two of
# those names — `column_mapping` maps each canonical field once regardless.
_BY_HEADING = {
    " ".join(column.split()).casefold(): field
    for field, column in TEMPLATE_FIELDS.items()
}
_BY_HEADING.update(
    {
        " ".join(heading.split()).casefold(): field
        for heading, field in UCM_CONFLICT_LIST_HEADINGS.items()
    }
)
_BY_HEADING.update(
    {
        " ".join(heading.split()).casefold(): field
        for heading, field in STRUCTURED_RECORD_HEADINGS.items()
    }
)

#: The released heading vocabulary as normalized-heading -> canonical field, so
#: a reader outside this module can carry it without touching a private name.
#: `record_capture_correction_result` needs the same lookup in the writing
#: transaction, and the migration derives its versioned, command-trusted copy
#: from this (#945); `tests/test_capture_correction_retirement.py` proves the two
#: never drift. It is `_BY_HEADING` itself rather than a copy on purpose: a
#: second dict is a second thing to keep in step.
HEADING_FIELD_VOCABULARY = _BY_HEADING


class NoConflictSheet(NoMatrixFound):
    """No sheet in this workbook is a utility conflict matrix.

    Raised rather than returning nothing, for the reason #59 story 2 gives
    about the page extractor: a workbook nobody could read must not report
    as a workbook with no conflicts.

    A `NoMatrixFound`, because it is the same answer about a different
    container — which is what lets `extract_project` report a workbook it
    could not read through the path it already has, rather than growing a
    second one that means the same thing.
    """


@dataclass(frozen=True)
class Sheet:
    """One worksheet, as rows of text.

    Text rather than cell values, because every reader downstream treats a
    cell as something a citation could quote. A `None` or a float leaking
    through would make each of them defend against it separately.
    """

    name: str
    rows: tuple[tuple[str, ...], ...]


def read_workbook(path: Path | str) -> list[Sheet]:
    """Every sheet of the workbook, in the order the file lists them.

    `data_only` so a formula cell yields its last computed value rather
    than its formula: the value is what the document asserts, and `=B2*2`
    is not something a reviewer can check a quote against.
    """
    workbook = load_workbook(Path(path), data_only=True, read_only=True)
    try:
        return [
            Sheet(
                name=name,
                rows=tuple(
                    tuple(_cell(value) for value in row)
                    for row in workbook[name].iter_rows(values_only=True)
                ),
            )
            for name in workbook.sheetnames
        ]
    finally:
        workbook.close()


def header_row(sheet: Sheet) -> int | None:
    """Which row carries the column headings, or None if none does.

    The first row naming at least `MIN_HEADER_FIELDS` canonical fields.
    Declining is the answer for a sheet whose headings this vocabulary does
    not carry — a project-information block, a dropdown list — because
    guessing at the first populated row would file its cells under whatever
    the row above happened to say.
    """
    for index, row in enumerate(sheet.rows[:HEADER_SEARCH_ROWS]):
        if len(column_mapping(row)) >= MIN_HEADER_FIELDS:
            return index
    return None


def column_mapping(headings) -> dict[int, str]:
    """Column index to canonical field, by printed name.

    Exact after collapsing whitespace, and no synonyms: this reads a
    published form whose columns are named in its own data dictionary, so a
    heading that does not match is a column the vocabulary has not ruled on
    rather than a spelling to guess at. The page extractor needs a model
    for this precisely because a printout does not carry the form's names.
    """
    mapping: dict[int, str] = {}
    for index, heading in enumerate(headings):
        field = _BY_HEADING.get(" ".join(str(heading or "").split()).casefold())
        if field and field not in mapping.values():
            mapping[index] = field
    return mapping


@dataclass(frozen=True)
class ConflictSheet:
    """The chosen sheet, with the reading that chose it.

    `header_index` and `mapping` are what the selection was made on, so
    they come back with it. The caller used to recompute both, and had to
    use `header_row`'s `int | None` unguarded — correct only because the
    selection had already proved it non-None, which is an invariant the
    return type could not state.

    `page_no` is the sheet's position in the workbook, which is the page
    number `ingest` gave the same sheet. Every spreadsheet citation rests
    on the two agreeing, so it is settled once, here.
    """

    sheet: Sheet
    header_index: int
    mapping: dict[int, str]
    page_no: int

    @property
    def headings(self) -> list[str]:
        return self.sheet.rows[self.header_index]

    @property
    def rows(self) -> list[list[str]]:
        """The conflict rows: everything under the header."""
        return self.sheet.rows[self.header_index + 1 :]


@dataclass(frozen=True)
class NativeEvidenceTable:
    """One known structured SUE table, mapped only by exact source headings."""

    sheet: Sheet
    header_index: int
    page_no: int
    kind: str
    mapping: dict[int, str]
    required_fields: tuple[str, ...]

    @property
    def headings(self) -> tuple[str, ...]:
        return self.sheet.rows[self.header_index]

    @property
    def rows(self) -> tuple[tuple[str, ...], ...]:
        return self.sheet.rows[self.header_index + 1 :]

    @property
    def unmapped_columns(self) -> tuple[tuple[int, str], ...]:
        return tuple(
            (index, _reported_heading(index, heading))
            for index, heading in enumerate(self.headings)
            if index not in self.mapping
        )

    @property
    def unmapped_headings(self) -> tuple[str, ...]:
        return tuple(heading for _, heading in self.unmapped_columns)


_SUE_TABLE_SCHEMAS = (
    (
        "sue_probe_depth",
        {
            "probe #": "probe_number",
            "utility name": "external_org",
            "diameter": "diameter",
            "northing": "northing",
            "easting": "easting",
            "natural ground elevation": "natural_ground_elevation",
            "top of utility elevation": "top_of_utility_elevation",
        },
        ("probe_number", "external_org"),
    ),
    (
        "sue_test_hole_index",
        {
            "csj": "csj",
            "northing": "northing",
            "easting": "easting",
            "station": "station",
            "offset": "offset",
            "test hole #": "test_hole_number",
            "utility owner": "external_org",
            "diameter": "diameter",
            "test hole depth (feet)": "test_hole_depth_feet",
            "elevation natural ground (ng)": "natural_ground_elevation",
            "elevation top of utility (tou)": "top_of_utility_elevation",
            "date": "observation_date",
        },
        ("test_hole_number", "external_org"),
    ),
)


def native_evidence_table(sheets: list[Sheet]) -> NativeEvidenceTable:
    """Find one exact known SUE schema; never infer a meaning for a heading."""

    found: list[NativeEvidenceTable] = []
    for page_no, sheet in enumerate(sheets, start=1):
        for header_index, headings in enumerate(sheet.rows[:HEADER_SEARCH_ROWS]):
            normalized = [_normalized_heading(value) for value in headings]
            for kind, schema, required in _SUE_TABLE_SCHEMAS:
                mapping = {
                    index: schema[heading]
                    for index, heading in enumerate(normalized)
                    if heading in schema
                }
                if len(mapping) >= 5 and all(
                    field in mapping.values() for field in required
                ):
                    found.append(
                        NativeEvidenceTable(
                            sheet=sheet,
                            header_index=header_index,
                            page_no=page_no,
                            kind=kind,
                            mapping=mapping,
                            required_fields=required,
                        )
                    )
    if len(found) != 1:
        raise NoConflictSheet(
            "workbook does not contain exactly one recognized conflict or SUE table"
        )
    return found[0]


def _normalized_heading(value: str) -> str:
    return " ".join((value or "").split()).casefold()


def _reported_heading(index: int, value: str) -> str:
    cleaned = " ".join((value or "").split())
    return cleaned or f"Column {get_column_letter(index + 1)} (blank heading)"


def conflict_sheet(sheets: list[Sheet]) -> ConflictSheet:
    """The one sheet that is a utility conflict matrix.

    Identified by the conflict id, which is ADR-0009's distinction
    mechanised: an inventory records what is out there and a conflict
    matrix records conflicts, and `Utility Feature ID` is not a conflict
    id. The published template ships both, and reading the inventory as a
    matrix is the error ADR-0009 was written about.

    Where more than one sheet qualifies — the template also carries
    `Utility Conflict List-Print`, a rendering of the same rows — the
    fuller sheet wins. That is ADR-0005 applied inside one workbook: a
    printed view does not outrank the structured original it was made
    from, and being narrower is what makes it the view.
    """
    best: ConflictSheet | None = None
    for page_no, sheet in enumerate(sheets, start=1):
        index = header_row(sheet)
        if index is None:
            continue
        mapping = column_mapping(sheet.rows[index])
        if not all(field in mapping.values() for field in REQUIRED):
            continue
        if best is None or len(mapping) > len(best.mapping):
            best = ConflictSheet(sheet, index, mapping, page_no)

    if best is None:
        raise NoConflictSheet(
            f"no sheet of this workbook heads a utility conflict matrix; "
            f"read {', '.join(s.name for s in sheets) or 'nothing'}. A "
            "layout variant is unhandled — do not treat this as a workbook "
            "with no conflicts."
        )
    return best


def row_text(row) -> str:
    """One row as the line it becomes in `sheet_text`.

    Shared with the extractor's citation quote deliberately. A quote
    against cells is checked at 1.0 (`verify.threshold_for`), so the two
    have to render a row identically — written twice, they could drift and
    every citation on every spreadsheet would fail at once.
    """
    return " ".join(cell for cell in row if cell)


def sheet_text(sheet: Sheet) -> str:
    """The sheet as the text a citation is checked against.

    A sheet has no page image, so this is the rendering that stands in for
    one — which is why it has to carry every cell a stored value could come
    from. Rows stay on their own lines: flattening them would let a quote
    match across two conflicts that never appeared on one row together.
    """
    return "\n".join(row_text(row) for row in sheet.rows).strip()


def _cell(value) -> str:
    """A cell as the text it shows.

    Integral floats lose the `.0` openpyxl gives them: a station reads
    `1149` on the form, and storing `1149.0` would put a value in the
    Ledger that appears nowhere in the document a reviewer opens.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
