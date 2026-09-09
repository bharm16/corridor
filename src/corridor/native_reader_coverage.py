"""Collect seven reader contracts from actual outputs and native authority (#458).

The former equivalence caller supplied its own coverage/provenance claims. This
collector accepts a database scope and cutoff, calls the public readers itself,
and reconciles their populations with independently queried native rows. Missing
classes, incomplete public lineage and changed authority remain explicit blockers.
It does not prove legacy equivalence, authorize cutover, create releases, or call
a model/provider by default. Rendering uses temporary private local files only.
"""
from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass, fields, is_dataclass
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
import json

from sqlalchemy import func, select, text
from sqlalchemy.exc import SQLAlchemyError

from corridor.accepted_field_reading import NativeReadingRefused
from corridor.config import settings
from corridor.models import Project, ProjectRecordRevision
from corridor.reader_coverage import CONTRACTS, SemanticRecord, SurfaceReading


# All names are implementation-owned. No caller chooses SQL, authority tables,
# semantic origins, observed-class flags or an already assembled reader output.
_TABLES = (
    "project_record_revisions", "facts", "fact_decisions", "fact_sources", "source_segments",
    "fact_applies_to", "fact_closure_results", "fact_closure_sources", "fact_statement_timings",
    "proposed_deltas", "delta_dispositions", "delta_record_decisions", "delta_deferrals", "delta_supersessions", "fact_dispositions",
    "delta_follow_up_plans", "delta_follow_up_plan_evidence", "support_assessments", "support_assessment_sources", "coordination_record_decisions",
    "coordination_record_reversals", "recorded_verbal_origins", "release_candidates",
    "release_packages", "external_report_artifacts", "external_report_releases",
    "project_baseline_formats", "project_baseline_source_rows",
)


def _plain(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Mapping):
        return {str(key): _plain(item) for key, item in value.items()}
    if isinstance(value, (tuple, list)):
        return [_plain(item) for item in value]
    if isinstance(value, (set, frozenset)):
        return sorted((_plain(item) for item in value), key=_encoded)
    if is_dataclass(value):
        return {item.name: _plain(getattr(value, item.name)) for item in fields(value)}
    raise TypeError(f"reader exposes unsupported value type {type(value).__name__}")


def _encoded(value):
    return json.dumps(_plain(value), sort_keys=True, separators=(",", ":"), allow_nan=False)


def _digest(value):
    return sha256(_encoded(value).encode()).hexdigest()


@dataclass(frozen=True)
class CollectedSurface:
    reading: SurfaceReading
    population_evidence: dict
    reader_outputs: dict
    blockers: tuple[str, ...]


@dataclass(frozen=True)
class NativeReaderCoverage:
    project_id: int
    revision_id: int
    as_of_revision_id: int
    evaluated_at: datetime
    authority_sha256: str
    surfaces: tuple[CollectedSurface, ...]

    @property
    def ready_for_comparison(self):
        """Read completeness only; an independent legacy comparison is still owed."""
        return len(self.surfaces) == len(CONTRACTS) and all(not item.blockers for item in self.surfaces)

    @property
    def readings(self):
        return tuple(item.reading for item in self.surfaces)

    @property
    def blockers(self):
        return tuple(f"{item.reading.surface}: {reason}" for item in self.surfaces for reason in item.blockers)

    def payload(self):
        return _plain(self)


class _Surface:
    def __init__(self, contract):
        self.contract = contract
        self.records = []
        self.observed = set()
        self.populations = {}
        self.outputs = {}
        self.blockers = []
        self.output_identity = None

    def population(self, kind, table, identities):
        identities = tuple(identities)
        self.populations[kind] = {"authority": table, "count": len(identities), "identities": identities}

    def add(self, kind, key, values, origins):
        if any(record.kind == kind and record.key == str(key) for record in self.records):
            self.blockers.append(f"{kind}: public reader duplicated identity {key}")
        self.records.append(SemanticRecord(kind, str(key), _plain(values), dict(origins)))
        missing = self.contract.fields - set(values)
        absent = self.contract.fields - set(origins)
        if missing or absent:
            self.blockers.append(f"{kind} {key}: missing semantic fields {sorted(missing)} or native lineage {sorted(absent)}")

    def finish_kind(self, kind, before):
        if len(self.blockers) == before:
            self.observed.add(kind)

    def result(self):
        for kind in sorted(self.contract.record_kinds - self.observed):
            if not any(reason.startswith(f"{kind}:") for reason in self.blockers):
                self.blockers.append(f"{kind}: public reading did not establish complete native coverage")
        if self.contract.unchanged_output and self.output_identity is None:
            self.blockers.append("intentionally unchanged output has no actual rendered/sealed identity")
        return CollectedSurface(SurfaceReading(self.contract.name, tuple(self.records),
            frozenset(self.observed), self.output_identity), self.populations, self.outputs,
            tuple(dict.fromkeys(self.blockers)))


def _attempt(surface, label, operation):
    """A refusal is evidence of a gap. Never convert it to an empty population."""
    try:
        return operation()
    except SQLAlchemyError as error:
        surface.blockers.append(f"{label}: database reader failed ({type(error).__name__}); no absence is established")
    except (ValueError, LookupError, OSError, TypeError, AttributeError, ImportError) as error:
        surface.blockers.append(f"{label}: {type(error).__name__}: {str(error)[:500]}")
    return None


def _inventory(session, project_id, maximum):
    result = {}
    for table in _TABLES:
        total = session.scalar(text(f"select count(*) from public.{table} where project_id=:project"), {"project": project_id})
        if total > maximum:
            raise NativeReadingRefused(f"{table} population {total} exceeds coverage limit {maximum}; no rows may be silently omitted")
        result[table] = tuple(session.scalars(text(
            f"select to_jsonb(r) from public.{table} r where project_id=:project order by id"
        ), {"project": project_id}))
    current = max((r["id"] for r in result["project_record_revisions"]), default=None)
    facts = {row["id"]: row for row in result["facts"]}
    result["native_current_values"] = tuple({"decision_id": decision["id"], "fact_id": decision["fact_id"],
        "subject": decision["subject_key"], "field": decision["fact_type"],
        "value": _fact_payload(result, facts[decision["fact_id"]])}
        for decision in _accepted_decisions_at(result, current)) if current is not None else ()
    return result


def _accepted_decisions_at(inventory, revision):
    """Derive the complete accepted decision set from raw custody rows."""
    decisions = {row["id"]: row for row in inventory["fact_decisions"]}
    active = tuple(row for row in decisions.values() if row["revision_id"] <= revision
        and (row["superseded_by"] is None
             or decisions[row["superseded_by"]]["revision_id"] > revision))
    suppressed = {row["subject_key"] for row in active
                  if row["fact_type"] == "statement_wording" and row["disposition"] == "do_not_add"}
    return tuple(sorted((row for row in active if row["disposition"] == "include"
                         and row["subject_key"] not in suppressed), key=lambda row: row["id"]))


def _fact_payload(inventory, fact):
    """Read typed values from retained Fact rows, independently of the reader."""
    if fact["fact_type"] == "statement_timing":
        return {"timings": [{"role": row["timing_role"], "text": row["text"],
            "precision": row["precision"], "start_date": row["start_date"], "end_date": row["end_date"]}
            for row in sorted(inventory["fact_statement_timings"], key=lambda row: row["timing_role"])
            if row["fact_id"] == fact["id"]]}
    if fact["fact_type"] == "applies_to":
        subjects = [row["record_subject_key"] for row in sorted(inventory["fact_applies_to"], key=lambda row: row["ordinal"])
                    if row["fact_id"] == fact["id"] and row["record_subject_key"] is not None]
        return {"mode": "selected" if subjects else "unknown", "subject_keys": subjects}
    if fact["fact_type"] == "closure_result":
        return {"closure_kind": next((row["closure_kind"] for row in inventory["fact_closure_results"]
                                       if row["fact_id"] == fact["id"]), None)}
    return fact["date_value"] or fact["text_value"]


def _ids(inventory, table):
    return {row["id"] for row in inventory[table]}


def _origin(table, identity):
    return f"source_segment:{identity}" if table == "source_segments" else f"native_decision:{table}:{identity}"


def _revision_origins(surface, revision_id):
    return {field: f"revision:{revision_id}" for field in surface.contract.fields}


def _native_fields(surface, population, inventory):
    """Verify real native field identities instead of relabelling legacy IDs."""
    decisions = {row["id"]: row for row in inventory["fact_decisions"]}
    segments = {row["id"]: row for row in inventory["source_segments"]}
    expected_values = {r["decision_id"]: r for r in inventory["native_current_values"]}
    edges = {(r["fact_id"], r["source_segment_id"]) for r in inventory["fact_sources"]}
    accepted = []
    for record in population.records:
        for name, field in record.fields.items():
            kind = getattr(field, "decision_kind", "fact_decision")
            held = decisions.get(field.decision_id) if kind == "fact_decision" else None
            if held is None or (held["fact_id"], held["revision_id"]) != (field.fact_id, field.revision_id):
                surface.blockers.append(f"accepted_values: {record.subject_key}/{name} lacks a verified {kind} identity {field.decision_id}")
            projected = expected_values.get(field.decision_id)
            if projected is None or projected["field"] != name or _plain(projected["value"]) != _plain(field.value):
                surface.blockers.append(f"accepted_values: {record.subject_key}/{name} differs from native authority")
            for source in field.sources:
                actual = segments.get(source.source_segment_id)
                if (actual is None or actual["document_id"] != source.document_id or actual["exact_text"] != source.quote
                    or (field.fact_id, source.source_segment_id) not in edges):
                    surface.blockers.append(f"source_support: segment {source.source_segment_id} does not corroborate its displayed passage")
            accepted.append({"subject": record.subject_key, "field": name, "value": _plain(field.value),
                "authority": {"kind": kind, "id": field.decision_id, "revision_id": field.revision_id,
                              "fact_id": field.fact_id}, "sources": _plain(field.sources)})
    return accepted


def _constraint_log(surface, reading, inventory):
    population = reading.native_population
    before = len(surface.blockers)
    fields_read = _native_fields(surface, population, inventory)
    expected = {record.subject_key for record in population.open_records}
    actual = [row.dependency.id for row in reading.rows]
    surface.population("constraint", "accepted native population at revision", sorted(expected))
    surface.outputs["ledger_rows"] = [{"identity": row.dependency.id, "fields": _plain(row.dependency.fields)} for row in reading.rows]
    if len(set(actual)) != len(actual) or set(actual) != expected:
        surface.blockers.append("constraint: public log population differs from the native frozen population")
    for row in reading.rows:
        record = row.dependency
        values = {"identity": record.subject_key,
            "accepted_values": [value for value in fields_read if value["subject"] == record.subject_key],
            "source_support": _plain(record.source_passages),
            "coordination": _plain([p for p in population.follow_up_plans if p.target_subject_identity == record.subject_key]),
            "check_results": _plain(reading.evaluation.for_dependency(record.id))}
        surface.add("constraint", record.subject_key, values, _revision_origins(surface, population.revision_id))
    surface.finish_kind("constraint", before)
    before = len(surface.blockers)
    checks = tuple(reading.evaluation.found)
    surface.population("check", "evaluate_native_population over verified native record IDs", [(c.dependency_id, c.rule) for c in checks])
    surface.outputs["evaluation"] = {"project_id": reading.evaluation.project_id, "ruleset_version": reading.evaluation.ruleset_version,
        "today": _plain(reading.evaluation.today), "results": _plain(checks)}
    if reading.evaluation.native_population is not population:
        surface.blockers.append("check: evaluation does not use the shared native population")
    for index, check in enumerate(checks):
        if check.dependency_id not in expected:
            surface.blockers.append(f"check: output names a subject outside the population: {check.dependency_id}")
        surface.add("check", f"{check.dependency_id}/{check.rule}/{index}", {
            "identity": {"subject": check.dependency_id, "rule": check.rule},
            "accepted_values": [v for v in fields_read if v["subject"] == check.dependency_id],
            "source_support": [v["sources"] for v in fields_read if v["subject"] == check.dependency_id],
            "coordination": _plain(population.follow_up_plans), "check_results": _plain(check)},
            _revision_origins(surface, population.revision_id))
    surface.finish_kind("check", before)


def _record_values(surface, session, project_id, revision, kind, inventory):
    from corridor.record_projection import read_native_record_values, record_value_payload
    before = len(surface.blockers)
    values = read_native_record_values(session, project_id, revision)
    # The reader is independent of Candidate/Dependency compatibility joins.
    surface.outputs[kind] = _plain(values)
    decisions = {row["id"]: row for row in _accepted_decisions_at(inventory, revision)}
    facts = {row["id"]: row for row in inventory["facts"]}
    surface.population(kind, f"raw fact_decisions at project={project_id},revision={revision}", sorted(decisions))
    actual_ids = [value.decision_id for value in values]
    if len(actual_ids) != len(set(actual_ids)) or set(actual_ids) != set(decisions):
        surface.blockers.append(f"{kind}: public accepted decision population differs from raw authority")
    for value in values:
        authority = decisions.get(value.decision_id)
        if authority is None or (authority["fact_id"], authority["revision_id"]) != (value.fact_id, value.revision_id):
            surface.blockers.append(f"{kind}: value {value.decision_id} does not bind an actual FactDecision")
            continue
        fact = facts[authority["fact_id"]]
        if (value.project_id != project_id or value.dependency_id is not None
            or value.subject_key != authority["subject_key"] or value.fact_type != authority["fact_type"]
            or value.subject_kind != fact["subject_kind"] or value.fact_subject_key != fact["subject_key"]
            or _plain(record_value_payload(value)) != _fact_payload(inventory, fact)
            or any(_plain(getattr(value, field)) != fact[field] for field in (
                "text_value", "date_value", "date_range_start", "date_range_end",
                "external_org_value_id", "document_value_id"))):
            surface.blockers.append(f"{kind}: public value or source identity differs for FactDecision {value.decision_id}")
        scope = sorted((row for row in inventory["fact_applies_to"] if row["fact_id"] == fact["id"]),
                       key=lambda row: row["ordinal"])
        closure = next((row for row in inventory["fact_closure_results"] if row["fact_id"] == fact["id"]), {})
        closure_sources = sorted((row for row in inventory["fact_closure_sources"] if row["fact_id"] == fact["id"]),
                                 key=lambda row: row["ordinal"])
        timing = sorted((row for row in inventory["fact_statement_timings"] if row["fact_id"] == fact["id"]),
                        key=lambda row: row["timing_role"])
        if (value.applies_to_dependency_ids != tuple(row["dependency_id"] for row in scope if row["dependency_id"] is not None)
            or value.applies_to_subject_keys != tuple(row["record_subject_key"] for row in scope if row["record_subject_key"] is not None)
            or value.closure_kind != closure.get("closure_kind")
            or value.closure_successor_dependency_id != closure.get("successor_dependency_id")
            or value.closure_governing_source_segment_ids != tuple(row["source_segment_id"] for row in closure_sources)
            or _plain(value.statement_timings) != [{key: row[key] for key in (
                "timing_role", "text", "precision", "start_date", "end_date")} for row in timing]):
            surface.blockers.append(f"{kind}: structured source metadata differs for FactDecision {value.decision_id}")
        sources = tuple(row["source_segment_id"] for row in sorted(inventory["fact_sources"], key=lambda r: (r["role"], r["ordinal"]))
                        if row["fact_id"] == value.fact_id)
        surface.add(kind, f"fact_decisions:{value.decision_id}", {"identity": {"subject": value.subject_key, "field": value.fact_type},
            "revision": revision, "accepted_values": record_value_payload(value),
            "authority": {"kind": "fact_decision", **{key: authority[key] for key in ("id", "fact_id", "revision_id", "disposition", "decided_at")}}, "source_support": list(sources)},
            {field: _origin("fact_decisions", value.decision_id) for field in surface.contract.fields})
    if inventory["coordination_record_decisions"]:
        surface.blockers.append(f"{kind}: coordination authority is a distinct class and is not enumerated by the source-value projection")
    if any(row["revision_id"] <= revision for row in inventory["delta_record_decisions"]):
        surface.blockers.append(f"{kind}: Resolve Delta authority and subject-lifecycle effects need their own projection inventory")
    surface.finish_kind(kind, before)


def _work_list(surface, session, project_id, revision, instant, inventory):
    from corridor.packet_review import read_review_items
    from corridor.work_list import build_work_list
    from corridor.native_follow_up_reading import read_adopted_follow_up_plans
    actual = read_review_items(session, project_id=project_id, as_of=instant)
    surface.outputs["review_items"] = _plain(actual)
    expected_plans = read_adopted_follow_up_plans(session, project_id, revision, current=True, as_of=instant)
    public_plans = tuple(getattr(actual, "follow_up_plans", ()))
    expected_by_id = {p.plan_id: p for p in expected_plans}
    public_by_id = {p.plan_id: p for p in public_plans}
    native_plans = {r["id"]: r for r in inventory["delta_follow_up_plans"]}
    native_deltas = {r["id"]: r for r in inventory["proposed_deltas"]}
    invalid_plans = set()
    seen_plans = set()
    for plan in public_plans:
        if plan.plan_id in seen_plans:
            invalid_plans.add(plan.plan_id)
        seen_plans.add(plan.plan_id)

    def recorded_time(value):
        return datetime.fromisoformat(value) if value is not None else None

    for identity, plan in public_by_id.items():
        native = native_plans.get(identity)
        delta = native_deltas.get(native["delta_id"]) if native is not None else None
        assessments = {e["support_assessment_id"] for e in inventory["delta_follow_up_plan_evidence"] if e["plan_id"] == identity}
        segments = {e["source_segment_id"] for e in inventory["support_assessment_sources"] if e["support_assessment_id"] in assessments}
        if (native is None or delta is None or plan != expected_by_id.get(identity) or plan.delta_id != native["delta_id"] or plan.revision_id != native["revision_id"]
            or plan.open_question != native["open_question"] or plan.recorded_by != native["recorded_by_principal"]
            or plan.responsible_principal != native["responsible_principal"]
            or plan.responsible_organization != native["responsible_organization"]
            or plan.return_date != recorded_time(native["return_date"])
            or plan.recorded_at != recorded_time(native["recorded_at"])
            or plan.target_subject_identity != delta["target_subject_identity"]
            or set(plan.support_assessment_ids) != assessments or set(plan.source_segment_ids) != segments):
            invalid_plans.add(identity)
    before = len(surface.blockers)
    for identity in sorted(invalid_plans):
        surface.blockers.append(f"proposed_delta: Follow-up Plan {identity} differs from its complete native act/target/evidence")
    expected = _ids(inventory, "proposed_deltas")
    surface.population("proposed_delta", "proposed_deltas scoped to project", sorted(expected))
    standings = {s.delta_id: s for s in actual.reading.standings}
    if set(standings) != expected or len(standings) != len(actual.reading.standings):
        surface.blockers.append(f"proposed_delta: standings omit {sorted(expected-set(standings))}, add {sorted(set(standings)-expected)}, or duplicate identities")
    if actual.accepted_revision_id != revision:
        surface.blockers.append("proposed_delta: Work List reads a different accepted revision")
    children = {c.delta_id: c for item in actual.items for c in item.children}
    for identity, standing in standings.items():
        child = children.get(identity)
        plans = [p for p in public_plans if p.delta_id == identity and p.plan_id not in invalid_plans]
        expected_ids = {p.plan_id for p in expected_plans if p.delta_id == identity}
        if {p.plan_id for p in plans} != expected_ids:
            surface.blockers.append(f"proposed_delta: {identity} has an undisplayed native Follow-up Plan")
        values = {"identity": identity, "band": standing.band,
            "scope": child.subject_identity if child else None,
            "assignment": [{"plan_id": p.plan_id, "person": p.responsible_principal, "organization": p.responsible_organization} for p in plans],
            "next_action": [{"plan_id": p.plan_id, "question": p.open_question, "return_date": _plain(p.return_date)} for p in plans],
            "deferral": {"returns_at": _plain(standing.returns_at), "wake_condition": standing.wake_condition},
            "disposition": standing.standing}
        # An empty plan census at this revision/workflow observation justifies
        # no recorded assignment/action. A present plan keeps its actual ID.
        origins = {name: f"revision:{revision}" for name in values}
        if plans:
            origins.update(assignment=_origin("delta_follow_up_plans", plans[0].plan_id),
                           next_action=_origin("delta_follow_up_plans", plans[0].plan_id))
        surface.add("proposed_delta", identity, values, origins)
    surface.finish_kind("proposed_delta", before)
    before = len(surface.blockers)
    for identity in sorted(invalid_plans):
        surface.blockers.append(f"constraint_work: Follow-up Plan {identity} differs from its complete native act/target/evidence")
    surface.population("constraint_work", "current native Follow-up Plans and coordination authority", sorted(expected_by_id))
    if set(public_by_id) != set(expected_by_id):
        surface.blockers.append("constraint_work: public Follow-up Plan population differs from its native current census")
    for identity, plan in public_by_id.items():
        if identity in invalid_plans:
            continue  # Never label unchecked public values as native authority.
        standing = standings.get(plan.delta_id)
        surface.add("constraint_work", f"delta_follow_up_plans:{identity}", {
            "identity": {"table": "delta_follow_up_plans", "id": identity, "delta_id": plan.delta_id},
            "band": standing.band if standing else None, "scope": plan.target_subject_identity,
            "assignment": {"person": plan.responsible_principal, "organization": plan.responsible_organization},
            "next_action": {"question": plan.open_question, "return_date": _plain(plan.return_date)},
            "deferral": {"returns_at": _plain(standing.returns_at), "wake_condition": standing.wake_condition} if standing else None,
            "disposition": standing.standing if standing else None},
            {name: _origin("delta_follow_up_plans", identity) for name in surface.contract.fields})
    if inventory["coordination_record_decisions"]:
        surface.blockers.append("constraint_work: accepted coordination is distinct from Follow-up Plans and needs explicit public subject/field lineage")
    surface.finish_kind("constraint_work", before)
    before = len(surface.blockers)
    legacy = build_work_list(session, project_id, today=instant.date())
    surface.outputs["statement_work_reader"] = _plain(legacy)
    native_statements = tuple(r for r in inventory["facts"] if "statement" in r["subject_kind"])
    surface.population("statement_work", "native statement facts and actual statement-work reader", [r["id"] for r in native_statements])
    if native_statements or actual.source_questions or legacy.immediate or legacy.backlog_total or legacy.candidate_backlog_total:
        surface.blockers.append("statement_work: statement work requires its own native item/assignment/Next Action population adapter")
    surface.finish_kind("statement_work", before)


def _check_citations(surface, report, inventory):
    """Bind displayed provenance to exact native acts and declared render inputs."""
    from corridor.accepted_field_reading import native_field_visible, visible_native_statements
    from corridor.exceptions import RULESET_VERSION
    from corridor.report import EMPTY_LEDGER
    from corridor.presentation import field_label, label

    segments = {row["id"]: row for row in inventory["source_segments"]}
    decisions = {row["id"]: row for row in inventory["fact_decisions"]}
    edges = {(row["fact_id"], row["source_segment_id"]) for row in inventory["fact_sources"]}
    revisions = {row["id"]: row for row in inventory["project_record_revisions"]}
    plans = {row["id"]: row for row in inventory["delta_follow_up_plans"]}
    evaluation = report.evaluation
    population = evaluation.native_population if evaluation is not None else None
    derived_inputs, record_inputs, plan_inputs = {}, {}, {}
    verified_refs = {f"revision:{identity}" for identity in revisions}
    verified_refs.update(f"native_decision:{identity}" for identity in decisions)
    # AcceptedField names its authority namespace explicitly; an ID only
    # verifies the fact_decision namespace, never an arbitrary decision kind.
    verified_refs.update(f"native_decision:fact_decision:{identity}" for identity in decisions)

    def derived(cell, record_ids, *, scope="", refs=()):
        derived_inputs[id(cell)] = (tuple(record_ids), scope, tuple(refs))

    def accepted(cell, record, field):
        if field is None:
            derived(cell, (record.id,), refs=(f"revision:{population.revision_id}",))
        elif not field.sources:
            derived(cell, (record.id,), refs=(f"revision:{field.revision_id}", field.origin))

    # The input manifest is bound by the containing rendered row/section, not
    # reconstructed from the citation's own claimed record IDs or free text.
    if population is not None:
        records = {record.ref_code: record for record in population.open_records}
        ids = tuple(population.record_ids)
        supported = tuple(record.id for record in population.open_records if record.checked_source_passages)
        for index, cell in enumerate(report.summary):
            if index < 4:
                derived(cell, (supported or ids) if index == 2 else ids, scope=EMPTY_LEDGER)
        statements = {statement.subject_key: statement for statement in visible_native_statements(population, document_only=report.document_only)}
        for section in report.sections:
            if section.title == "Appendix — constraint log":
                for row in section.rows:
                    record = records.get(row[0].value) if row else None
                    if record is None or len(row) != 9:
                        continue
                    derived(row[0], (record.id,), refs=(f"revision:{record.identity_revision_id}",))
                    derived(row[3], (record.id,), refs=tuple(record.fields[name].origin for name in ("utility_type", "external_org") if name in record.fields))
                    derived(row[7], (record.id,), refs=(f"revision:{population.revision_id}",))
                    for index, name in ((1, "utility_id"), (2, "external_org"), (4, "station_from"), (5, "station_to"), (6, "resolution_strategy"), (8, "need_date")):
                        accepted(row[index], record, record.fields.get(name))
            elif section.title == "Accepted Project Record values":
                for row in section.rows:
                    record = records.get(row[0].value) if row else None
                    if record is None or len(row) != 3:
                        continue
                    derived(row[0], (record.id,))
                    derived(row[1], (record.id,))
                    matching = [field for name, field in record.fields.items() if field_label(name) == row[1].value]
                    if len(matching) == 1:
                        accepted(row[2], record, matching[0])
            elif section.title == "Accepted statements":
                for row in section.rows:
                    statement = statements.get(row[0].value) if row else None
                    if statement is None or len(row) != 5:
                        continue
                    for index in (0, 1, 4):
                        record_inputs[id(row[index])] = statement.fields["statement_wording"].decision_id
                    for index, name in ((2, "statement_timing"), (3, "applies_to")):
                        field = statement.fields.get(name)
                        if field is not None and native_field_visible(field, document_only=report.document_only):
                            record_inputs[id(row[index])] = field.decision_id
                        else:
                            derived(row[index], (statement.subject_key,), refs=(f"revision:{population.revision_id}",))
            elif section.title == label("follow_up_plan"):
                for row in section.rows:
                    matching = [plan for plan in population.follow_up_plans if row and str(plan.delta_id) == row[0].value]
                    if len(matching) == 1:
                        for cell in row:
                            plan_inputs[id(cell)] = (matching[0].plan_id,)
            elif section.title == label("constraint_alerts"):
                facets = evaluation.facets()
                if len(section.rows) == len(facets):
                    for row, facet in zip(section.rows, facets):
                        if len(row) != 4:
                            continue
                        affected = tuple(item.dependency_id for item in facet.exceptions)
                        for cell in row[:2]:
                            derived(cell, affected)
                        for cell in row[2:]:
                            derived(cell, (facet.exceptions[0].dependency_id,) if facet.has_quantities else affected)
    for cell in report.cells:
        citation = cell.provenance
        kind = type(citation).__name__
        if kind == "Assertion":
            source = segments.get(citation.source_segment_id)
            decision = decisions.get(citation.decision_id)
            if (source is None or decision is None or citation.revision_id not in revisions
                or (decision["fact_id"], decision["revision_id"]) != (citation.fact_id, citation.revision_id)
                or (citation.fact_id, citation.source_segment_id) not in edges
                or (source["document_id"], source["exact_text"]) != (citation.document_id, citation.quote)):
                surface.blockers.append(f"citations: {cell.label!r} lacks its actual source segment/FactDecision/revision")
        elif kind == "RecordDecision":
            decision = decisions.get(citation.decision_id)
            authority = revisions.get(citation.revision_id, {})
            source_refs = tuple(f"source_segment:{row['source_segment_id']}" for row in sorted(
                inventory["fact_sources"], key=lambda row: (row["role"], row["ordinal"])) if row["fact_id"] == citation.fact_id)
            if (getattr(citation, "decision_kind", "fact_decision") != "fact_decision" or decision is None
                or (decision["fact_id"], decision["revision_id"]) != (citation.fact_id, citation.revision_id)
                or record_inputs.get(id(cell)) != citation.decision_id
                or citation.actor != (authority.get("human_principal") or authority.get("released_policy"))
                or citation.decided_at != datetime.fromisoformat(decision["decided_at"]).date()
                or tuple(citation.source_refs) != source_refs):
                surface.blockers.append(f"citations: {cell.label!r} RecordDecision attribution or source references differ from its exact native act")
        elif kind == "Derivation":
            declared = derived_inputs.get(id(cell))
            actual = (tuple(citation.record_ids), citation.scope, tuple(citation.input_refs))
            if (evaluation is None or population is None or evaluation.ruleset_version != RULESET_VERSION
                or citation.ruleset_version != evaluation.ruleset_version or declared is None or actual != declared
                or population.revision_id not in revisions or evaluation.project_id != population.project_id
                or any(reference not in verified_refs for reference in citation.input_refs)
                or not citation.resolves):
                surface.blockers.append(f"citations: {cell.label!r} derivation lacks its exact verified input identities, evaluation rule/version or scope")
        elif kind == "WorkDecision" and citation.decision_kind == "delta_follow_up_plan":
            named = tuple(citation.decision_ids)
            if (not named or len(set(named)) != len(named) or named != plan_inputs.get(id(cell))
                or any(identity not in plans or citation.recorded_by != plans[identity]["recorded_by_principal"]
                    or citation.recorded_at != datetime.fromisoformat(plans[identity]["recorded_at"]).date() for identity in named)):
                surface.blockers.append(f"citations: {cell.label!r} Follow-up Plan identities or attribution differ from their exact native acts")
        else:
            surface.blockers.append(f"citations: {cell.label!r} uses unsupported {kind}; no legacy ID is relabelled native")


def _report(surface, session, reading, inventory, fixture_client):
    from corridor.briefing import assemble_native_citables, brief_project
    from corridor.report import build_report
    population = reading.native_population
    before = len(surface.blockers)
    report = build_report(session, reading.project.id, today=reading.evaluation.today, frozen_reading=reading)
    _check_citations(surface, report, inventory)
    report_data = {"covered_records": _plain(report.covered_records), "summary": _plain(report.summary), "sections": _plain(report.sections)}
    surface.outputs["report"] = report_data
    surface.population("report", "actual covered_records against frozen native population", list(population.record_ids))
    if {r[0] for r in report.covered_records} != set(population.record_ids):
        surface.blockers.append("report: rendered population differs from the shared native reading")
    values = {"population": _plain(report.covered_records), "accepted_values": report_data,
        "coordination": _plain(population.follow_up_plans), "checks": _plain(reading.evaluation.found),
        "statements": _plain(getattr(population, "statements", ())), "citations": _plain([c.provenance for c in report.cells])}
    if any("statement" in row["subject_kind"] for row in inventory["facts"]) and not hasattr(population, "statements"):
        surface.blockers.append("report: accepted statement class is not represented by this native population")
    surface.add("report", f"project:{reading.project.id}/revision:{population.revision_id}", values, _revision_origins(surface, population.revision_id))
    surface.finish_kind("report", before)
    before = len(surface.blockers)
    citables, floor, committed = assemble_native_citables(population, reading.evaluation, reading.statement_publication, project_scope=True)
    output = {"citables": _plain(citables), "required_citation_floor": _plain(floor), "committed_dates": _plain(committed)}
    if fixture_client is not None:
        generated = brief_project(session, reading.project.id, client=fixture_client,
            today=reading.evaluation.today, frozen_reading=reading)
        output["fixture_briefing"] = _plain(generated)
        if generated.refused:
            surface.blockers.append(f"briefing: fixture reader refused: {generated.refusal_reason}")
    else:
        output["mode"] = "native citation inputs only; no model invoked"
    surface.outputs["briefing"] = output
    surface.population("briefing", "assemble_native_citables over the same native revision", list(population.record_ids))
    values = {"population": list(population.record_ids), "accepted_values": _plain(population.records),
        "coordination": _plain(population.follow_up_plans), "checks": _plain(reading.evaluation.found),
        "statements": _plain(getattr(population, "statements", ())), "citations": output}
    surface.add("briefing", f"project:{reading.project.id}/revision:{population.revision_id}", values, _revision_origins(surface, population.revision_id))
    if any("statement" in row["subject_kind"] for row in inventory["facts"]) and not hasattr(population, "statements"):
        surface.blockers.append("briefing: native statement citation population is not established")
    surface.finish_kind("briefing", before)


def _workbook_cells(content):
    from io import BytesIO
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(content), data_only=False)
    try:
        return {sheet.title: [[_plain(cell.value) for cell in row] for row in sheet.iter_rows()] for sheet in book}
    finally:
        book.close()


def _workbooks(surface, session, reading, inventory):
    from corridor.baseline_adoption import effective_baseline_formats
    from corridor.export import to_xlsx
    from corridor.models import BaselineFormatObject
    from corridor.object_storage import content_store
    from corridor.workbook_render import render_project_record_workbook, workbook_reader_input_manifest
    from corridor.accepted_field_reading import native_reader_input_manifest
    population = reading.native_population
    input_manifest = native_reader_input_manifest(population)
    surface.outputs["native_inputs"] = input_manifest
    output_ids = {}
    before = len(surface.blockers)
    with TemporaryDirectory(prefix="corridor-reader-coverage-") as temporary:
        path = to_xlsx(session, reading.project.id, Path(temporary)/"native.xlsx",
            evaluation=reading.evaluation, statement_publication=reading.statement_publication,
            frozen_reading=reading, internal_working_copy=True)
        cells = _workbook_cells(path.read_bytes())
    surface.outputs["internal_workbook"] = cells
    surface.population("workbook", "actual internal workbook over frozen native records", list(population.record_ids))
    values = {"population": list(population.record_ids), "accepted_values": cells,
        "coordination": _plain(population.follow_up_plans), "statements": _plain(getattr(population, "statements", ())),
        "source_support": [source for record in population.open_records for source in _plain(record.source_passages)]}
    # The artifact itself must carry the claimed field lineage, not merely be
    # accompanied by a second query returning the missing evidence.
    if "Accepted value sources" not in cells:
        surface.blockers.append("workbook: actual artifact omits its native accepted-value source inventory")
    if inventory["coordination_record_decisions"]:
        surface.blockers.append("workbook: coordination decisions need explicit artifact/native-subject reconciliation")
    surface.add("workbook", f"project:{reading.project.id}/revision:{population.revision_id}", values, _revision_origins(surface, population.revision_id))
    output_ids["workbook_cells"] = _digest(cells)
    surface.finish_kind("workbook", before)
    before = len(surface.blockers)
    formats = effective_baseline_formats(session, reading.project.id)
    template = formats.get("output_template")
    surface.population("customer_format", "project_baseline_formats output_template", () if template is None else (template.id,))
    if template is None:
        surface.blockers.append("customer_format: no approved output template; absence is not a successful export")
        return
    stored = session.get(BaselineFormatObject, template.id)
    if stored is None or stored.project_id != reading.project.id:
        surface.blockers.append("customer_format: approved template bytes are not retained")
        return
    content = content_store().get(stored.storage_key, sha256=template.content_sha256)
    rendered = render_project_record_workbook(session, project_id=reading.project.id,
        revision_id=population.revision_id, template_bytes=content)
    if sha256(rendered.content).hexdigest() != rendered.output_sha256:
        surface.blockers.append("customer_format: rendered workbook does not match its declared digest")
    artifact = {field.name: _plain(getattr(rendered, field.name)) for field in fields(rendered) if field.name != "content"}
    artifact["cells"] = _workbook_cells(rendered.content)
    mapped = workbook_reader_input_manifest(rendered)
    artifact["mapped_input_manifest"] = mapped
    surface.outputs["customer_format"] = artifact
    decisions = {r["id"]: r for r in inventory["fact_decisions"]}
    represented = set()
    support = []
    from openpyxl.utils.cell import range_boundaries
    for cell in mapped["mapped_cells"]:
        represented.add(cell["record_subject_key"])
        min_col, min_row, max_col, max_row = range_boundaries(cell["cell_range"])
        sheet = artifact["cells"].get(cell["sheet_name"], [])
        actual_value = sheet[min_row-1][min_col-1] if min_row <= len(sheet) and min_col <= len(sheet[min_row-1]) else None
        if min_col != max_col or min_row != max_row or str(actual_value or "") != str(cell["rendered_value"] or ""):
            surface.blockers.append(f"customer_format: mapped cell {cell['cell_range']} differs from actual output or needs a multi-cell adapter")
        for value in cell["accepted_fields"]:
            held = decisions.get(value["decision_id"])
            if (value.get("decision_kind") != "fact_decision" or held is None
                or (held["fact_id"], held["revision_id"], held["fact_type"]) != (value["fact_id"], value["revision_id"], value["fact_type"])):
                surface.blockers.append(f"customer_format: {cell['cell_range']} lacks its actual typed native field authority")
            support.append({"cell_range": cell["cell_range"], "decision_id": value["decision_id"], "fact_id": value["fact_id"],
                "source_segment_ids": [r["source_segment_id"] for r in inventory["fact_sources"] if r["fact_id"] == value["fact_id"]]})
    if set(population.record_ids) - represented:
        surface.blockers.append("customer_format: mapped input manifest omits native open subjects")
    if (mapped["project_id"], mapped["revision_id"], mapped["output_sha256"]) != (reading.project.id, population.revision_id, rendered.output_sha256):
        surface.blockers.append("customer_format: mapped input manifest belongs to another output/revision")
    values = {"population": sorted(represented), "accepted_values": artifact,
        "coordination": input_manifest["follow_up_plans"], "statements": input_manifest["statements"],
        "source_support": support}
    surface.add("customer_format", f"project:{reading.project.id}/revision:{population.revision_id}", values,
        _revision_origins(surface, population.revision_id))
    if input_manifest["follow_up_plans"] or input_manifest["statements"] or inventory["coordination_record_decisions"]:
        surface.blockers.append("customer_format: nonempty coordination/statement classes are not represented by mapped accepted cells")
    output_ids["customer_format_bytes"] = rendered.output_sha256
    surface.finish_kind("customer_format", before)
    surface.output_identity = _digest(output_ids)


def _retained_report_context(surface, session, project_id, context, inventory):
    from corridor.accepted_field_reading import native_reader_input_manifest, read_accepted_field_population
    from corridor.record_projection import read_native_record_values, record_value_payload
    manifest = context.get("native_reader_input_manifest")
    if not isinstance(manifest, dict) or manifest.get("project_id") != project_id:
        surface.blockers.append("release: retained report has no project-bound native input manifest")
        return None
    boundary = manifest.get("revision_id")
    if boundary not in _ids(inventory, "project_record_revisions"):
        surface.blockers.append("release: retained native manifest names an unknown project revision")
        return None
    if not context.get("report_cells"):
        surface.blockers.append("release: retained report cells/citations are unavailable")
    # Reconstruct the native input at the retained revision, never at today's
    # head. The retained object remains the published output, even on mismatch.
    # Complete serialization covers actor/time, typed value, exact sources and
    # parent record identity; matching decision IDs alone proves none of those.
    expected = _attempt(surface, "release: retained native revision", lambda: native_reader_input_manifest(
        read_accepted_field_population(session, project_id, revision_id=boundary),
        document_only=bool(context.get("document_only", False))))
    if expected is not None and _encoded(expected) != _encoded(manifest):
        surface.blockers.append("release: retained native input metadata differs from its exact as-of authority")
    expected_fields = {(kind, record["subject_key"], name): field
        for kind in ("records", "statements") for record in (expected or {}).get(kind, ())
        for name, field in record["fields"].items()}
    decisions = {r["id"]: r for r in inventory["fact_decisions"]}
    projected = {v.decision_id: v for v in read_native_record_values(session, project_id, boundary)}
    segments = {r["id"]: r for r in inventory["source_segments"]}
    edges = {(r["fact_id"], r["source_segment_id"], r["role"]) for r in inventory["fact_sources"]}
    represented = set()
    for kind in ("records", "statements"):
        for record in manifest.get(kind, ()):
            for name, field in record["fields"].items():
                identity = field["decision_id"]
                if identity in represented:
                    surface.blockers.append(f"release: retained decision {identity} is duplicated across displayed fields")
                represented.add(identity)
                held = decisions.get(identity)
                value = projected.get(identity)
                if (field.get("decision_kind") != "fact_decision" or held is None
                    or (held["fact_id"], held["revision_id"], held["fact_type"]) != (field["fact_id"], field["revision_id"], name)):
                    surface.blockers.append(f"release: retained field {name} has no matching typed native decision")
                payload = record_value_payload(value) if value is not None else None
                if kind == "statements" and name == "applies_to" and value is not None:
                    payload = {"mode": "selected" if value.applies_to_subject_keys or value.applies_to_dependency_ids else "unknown",
                        "subject_keys": value.applies_to_subject_keys, "legacy_dependency_ids": value.applies_to_dependency_ids}
                if (value is None or value.fact_type != name
                    or field.get("fact_subject_key") != (value.fact_subject_key or value.subject_key)
                    or _plain(field.get("value")) != _plain(payload)
                    or _encoded(field) != _encoded(expected_fields.get((kind, record["subject_key"], name)))):
                    surface.blockers.append(f"release: retained field {record['subject_key']}/{name} differs from its exact as-of value, metadata or parent identity")
                sources = field.get("sources", ())
                # Statement sources carry exact_text and a typed role, with
                # decision authority on their parent field. Constraint sources
                # carry quote plus their own fact/decision/revision reference.
                actual_edges = [(field["fact_id"], source["source_segment_id"],
                    source.get("role") if kind == "statements" else "value_source") for source in sources]
                expected_edges = {edge for edge in edges if edge[0] == field["fact_id"]
                    and (kind == "statements" or edge[2] == "value_source")}
                if len(actual_edges) != len(set(actual_edges)) or set(actual_edges) != expected_edges:
                    surface.blockers.append(f"release: retained field {name} source references differ from its native source edges")
                for source in sources:
                    segment = segments.get(source["source_segment_id"])
                    if kind == "statements":
                        wrong_reference = segment is not None and (
                            source.get("kind") != segment["kind"] or source.get("content_sha256") != segment["content_sha256"])
                    else:
                        wrong_reference = (source.get("fact_id") != field["fact_id"] or source.get("decision_id") != identity
                            or source.get("revision_id") != field["revision_id"])
                    if (segment is None or segment["document_id"] != source.get("document_id")
                        or segment["exact_text"] != source.get("exact_text" if kind == "statements" else "quote")
                        or wrong_reference):
                        surface.blockers.append(f"release: retained field {name} source reference differs from native source authority")
    if set(projected) != represented:
        surface.blockers.append("release: retained field population differs from its native revision; audience-filtered or other unrepresented classes need explicit accounting")
    if manifest.get("coverage_blockers"):
        surface.blockers.extend(f"release: {b}" for b in manifest["coverage_blockers"])
    return manifest


def _releases(surface, session, project_id, revision, instant, inventory):
    from corridor.models import ReleaseCandidate, ReleasePackage
    from corridor.object_storage import content_store
    from corridor.release_candidate import authorization_blockers
    from corridor.release_authorization import candidate_set, release_history, retrieve_released_artifact
    before = len(surface.blockers)
    packages = release_history(session, project_id, with_identifiers=True)
    surface.outputs["authorized_history"] = _plain(packages)
    surface.population("release", "release_candidates + release_packages + historical external-report artifacts/releases",
        [(table, row["id"]) for table in ("release_candidates", "release_packages", "external_report_artifacts", "external_report_releases") for row in inventory[table]])
    output_ids = []
    for candidate_row in inventory["release_candidates"]:
        candidate = session.get(ReleaseCandidate, candidate_row["id"])
        artifacts = candidate_set(session, candidate)
        retained = []
        for artifact in artifacts:
            body = content_store().get(artifact.storage_key, sha256=artifact.content_sha256)
            if len(body) != artifact.byte_count:
                surface.blockers.append(f"release: candidate {candidate.id} artifact size differs")
            retained.append(_plain(artifact.as_payload()))
        output = {"candidate_id": candidate.id, "accepted_revision_id": candidate.accepted_revision_id,
            "authorization_blockers": authorization_blockers(session, candidate, as_of=instant), "artifacts": retained}
        surface.outputs[f"candidate:{candidate.id}"] = _plain(output)
        surface.add("release", f"release_candidates:{candidate.id}", {
            "validation": output, "approval": None, "citations": retained},
            {name: f"revision:{candidate.accepted_revision_id}" for name in ("validation", "approval", "citations")})
    if {row.package_id for row in packages} != _ids(inventory, "release_packages"):
        surface.blockers.append("release: authorized package population differs from database authority")
    for row in packages:
        package = session.get(ReleasePackage, row.package_id)
        for artifact in row.artifacts:
            content = retrieve_released_artifact(session, package, artifact.artifact_type)
            output_ids.append((row.package_id, artifact.artifact_type, sha256(content).hexdigest()))
        surface.add("release", f"release_packages:{row.package_id}", {
            "validation": _plain(row), "approval": {"actor": row.authorized_by_principal, "at": _plain(row.authorized_at)},
            "citations": _plain(row.artifacts)},
            {name: f"revision:{row.accepted_revision_id}" for name in ("validation", "approval", "citations")})
    if inventory["release_candidates"] or inventory["release_packages"]:
        surface.blockers.append("release: generic ReleaseCandidate/ReleasePackage has no complete retained native input manifest; artifact hashes alone do not prove field coverage")
    if inventory["external_report_artifacts"] or inventory["external_report_releases"]:
        from corridor.report_release import (external_report_release_history, review_prepared_external_report,
            retrieve_prepared_external_report, retrieve_released_external_report)
        history = external_report_release_history(session, project_id)
        surface.outputs["external_report_history"] = _plain(history)
        if {r.release_id for r in history} != _ids(inventory, "external_report_releases"):
            surface.blockers.append("release: public external-report history omits native release identities")
        for table, rows in (("external_report_artifacts", inventory["external_report_artifacts"]),
                            ("external_report_releases", inventory["external_report_releases"])):
            for row in rows:
                if table == "external_report_artifacts":
                    actual = retrieve_prepared_external_report(session, project_id, row["id"])
                    review = _plain(review_prepared_external_report(session, project_id, row["id"]))
                    approval = None  # Prepared artifacts own no approval act.
                else:
                    actual = retrieve_released_external_report(session, project_id, row["id"])
                    review = {"digest_verified_by": "retrieve_released_external_report", "pdf_sha256": actual.pdf_sha256}
                    approval = {"actor": actual.released_by, "at": _plain(actual.released_at)}
                    if (actual.released_by != row["released_by"]
                        or actual.released_at != datetime.fromisoformat(row["released_at"])):
                        surface.blockers.append(f"release: external report {row['id']} approval actor/time differs from native release authority")
                context = actual.record_context_json
                manifest = _retained_report_context(surface, session, project_id, context, inventory)
                surface.outputs[f"{table}:{row['id']}"] = {"context": context, "validation": review}
                output_ids.append((table, row["id"], actual.pdf_sha256))
                if manifest is not None:
                    surface.add("release", f"{table}:{row['id']}", {"population": {
                        "records": [r["subject_key"] for r in manifest["records"]],
                        "statements": [r["subject_key"] for r in manifest.get("statements", ())]},
                        "accepted_values": manifest, "validation": review, "approval": approval,
                        "citations": context.get("report_cells", ())},
                        _revision_origins(surface, manifest["revision_id"]))
    surface.finish_kind("release", before)
    # An empty actual artifact population has an exact identity too. It proves
    # no release exists; it never pretends a release was approved or rendered.
    if output_ids or not surface.populations["release"]["identities"]:
        surface.output_identity = _digest({"released_artifact_population": output_ids})


def _source_occurrences(surface, session, project_id, inventory):
    """Read every retained source occurrence, including ones with no decision.

    A capture is not an accepted field. The record contains capture/origin
    metadata and verified source text only. Unknown PDF engines remain a
    stated locator-verification gap, not an implicit network or model request.
    """
    from collections import defaultdict
    from corridor.models import Document, SourceSegment
    from corridor.source_segments import dereference_source_segment, spreadsheet_replay, replay_recorded_verbal_statement
    from corridor.storage import stored_file
    before = len(surface.blockers)
    surface.population("source", "source_segments scoped to project", sorted(_ids(inventory, "source_segments")))
    by_document = defaultdict(list)
    for segment in session.scalars(select(SourceSegment).where(SourceSegment.project_id == project_id).order_by(SourceSegment.id)):
        by_document[segment.document_id].append(segment)
    verbal_origins = {r["id"]: r for r in inventory["recorded_verbal_origins"]}
    outputs = []
    def record(segment, recovered, validation):
        origin = verbal_origins.get(segment.recorded_verbal_origin_id)
        if sha256(segment.exact_text.encode()).hexdigest() != segment.content_sha256 or recovered != segment.exact_text:
            raise NativeReadingRefused(f"source segment {segment.id} text or digest differs")
        raw = next(r for r in inventory["source_segments"] if r["id"] == segment.id)
        output = {"source_segment_id": segment.id, "kind": segment.kind,
            "document_id": segment.document_id, "exact_text": recovered,
            "content_sha256": segment.content_sha256, "locator_validation": validation,
            "locator": {key: value for key, value in raw.items() if key not in {"exact_text", "content_sha256", "created_at", "id", "project_id"}},
            "recorded_verbal_origin": origin}
        outputs.append(output)
        surface.add("source", f"source_segments:{segment.id}", {
            "identity": {"table": "source_segments", "id": segment.id},
            "original_actor": origin["recorded_by"] if origin else None,
            "original_time": origin["recorded_at"] if origin else raw.get("created_at"),
            "decision_type": {"class": "source_capture", "segment_kind": segment.kind},
            "source_identity": output,
            "predecessor": origin.get("corrects_origin_id") if origin else None},
            {field: _origin("source_segments", segment.id) for field in surface.contract.fields})
    for document_id, segments in by_document.items():
        document = session.get(Document, document_id) if document_id is not None else None
        path = stored_file(document) if document is not None else None
        spreadsheet = [s for s in segments if s.kind == "spreadsheet_cell"]
        if spreadsheet:
            if path is None:
                surface.blockers.append(f"source: registered bytes for document {document_id} are unavailable")
                for segment in spreadsheet:
                    record(segment, segment.exact_text, "retained_text_only")
            else:
                with spreadsheet_replay(document, path) as replay:
                    for segment in spreadsheet:
                        record(segment, replay(segment), "replayed_registered_cells")
        for segment in segments:
            if segment.kind == "spreadsheet_cell":
                continue
            if segment.kind == "recorded_verbal_statement":
                if segment.recorded_verbal_origin_id not in verbal_origins:
                    surface.blockers.append(f"source: recorded verbal {segment.id} lacks its native origin")
                record(segment, replay_recorded_verbal_statement(segment), "native_attestation_digest")
            elif segment.kind == "email_span" and path is not None:
                record(segment, dereference_source_segment(document, segment, path), "replayed_registered_mime")
            else:
                record(segment, segment.exact_text, "retained_text_only")
                surface.blockers.append(f"source: {segment.kind} segment {segment.id} needs an explicitly available locator-check reader")
    surface.outputs["source_occurrence_readback"] = outputs
    surface.finish_kind("source", before)


def _history(surface, session, project_id, revision, selected_revision, inventory, maximum):
    from corridor.record_history import SearchTerms, read_record_history
    history = read_record_history(session, project_id=project_id,
        terms=SearchTerms(revision=selected_revision), value_limit=maximum)
    surface.outputs["history"] = _plain(history)
    if history.truncated or history.current_revision_id != revision:
        surface.blockers.append("history: public history is truncated or reads a different current revision")
    decisions = {row["id"]: row for row in inventory["fact_decisions"]}
    revisions = {row["id"]: row for row in inventory["project_record_revisions"]}
    read_decisions = {row.decision_id: row for row in history.source_decisions}
    predecessors = {row["superseded_by"]: row["id"] for row in decisions.values() if row["superseded_by"] is not None}
    for kind, expected in (("decision", set(decisions)), ("correction", set(predecessors)),
                           ("reversal", {row["id"] for row in decisions.values() if row["disposition"] == "restore"})):
        before = len(surface.blockers)
        surface.population(kind, "fact_decisions and actual supersession/disposition columns", sorted(expected))
        if expected - set(read_decisions):
            surface.blockers.append(f"{kind}: native decision IDs absent from public history: {sorted(expected-set(read_decisions))}")
        for identity in sorted(expected & set(read_decisions)):
            row = read_decisions[identity]
            native = decisions[identity]
            owner = revisions.get(row.revision_id, {})
            if (row.authority != (owner.get("human_principal") or owner.get("released_policy"))
                or row.command_type != owner.get("command_type") or row.revision_id != native["revision_id"]
                or row.disposition != native["disposition"]
                or row.decided_at != datetime.fromisoformat(native["decided_at"])):
                surface.blockers.append(f"{kind}: original actor or decision type differs for FactDecision {identity}")
            surface.add(kind, f"fact_decisions:{identity}", {"identity": {"table": "fact_decisions", "id": identity},
                "original_actor": row.authority, "original_time": _plain(row.decided_at),
                "decision_type": {"command": row.command_type, "disposition": row.disposition},
                "source_identity": {"fact_id": row.fact.fact_id if row.fact else None, "sources": _plain(row.fact.sources) if row.fact else []},
                "predecessor": predecessors.get(identity)},
                {field: _origin("fact_decisions", identity) for field in surface.contract.fields})
            if native["fact_id"] != (row.fact.fact_id if row.fact else None):
                surface.blockers.append(f"{kind}: source Fact differs for decision {identity}")
        if kind == "decision" and inventory["coordination_record_decisions"]:
            surface.blockers.append("decision: history exposes current/selected coordination values but not a complete native decision-chain inventory")
        if kind == "decision" and inventory["delta_record_decisions"]:
            surface.blockers.append("decision: DeltaRecordDecision IDs are not enumerated by the public source-decision history")
        if kind == "correction" and inventory["fact_dispositions"]:
            surface.blockers.append("correction: source-reading correction edges are not individually enumerated by public history")
        if kind == "correction" and any(row["corrects_origin_id"] is not None for row in inventory["recorded_verbal_origins"]):
            surface.blockers.append("correction: original verbal re-attestation chain is not enumerated by the public history")
        if kind == "reversal" and inventory["coordination_record_reversals"]:
            surface.blockers.append("reversal: native coordination reversals are not individually enumerated by the public history")
        surface.finish_kind(kind, before)
    _source_occurrences(surface, session, project_id, inventory)
    before = len(surface.blockers)
    expected = _ids(inventory, "support_assessments")
    observed = {a.assessment_id: a for row in history.source_decisions if row.fact for a in row.fact.assessments}
    surface.population("support", "support_assessments scoped to project", sorted(expected))
    if set(observed) != expected:
        surface.blockers.append(f"support: history omits {sorted(expected-set(observed))} or adds {sorted(set(observed)-expected)} assessments")
    assessments = {row["id"]: row for row in inventory["support_assessments"]}
    predecessors = {row["superseded_by"]: row["id"] for row in assessments.values() if row["superseded_by"] is not None}
    for identity, assessment in observed.items():
        native = assessments.get(identity)
        links = {row["source_segment_id"] for row in inventory["support_assessment_sources"] if row["support_assessment_id"] == identity}
        rendered_links = {getattr(source, "source_segment_id", None) for source in assessment.segments}
        if rendered_links != links:
            surface.blockers.append(f"support: assessment {identity} lacks exact public SourceSegment occurrence IDs")
        if native is None or assessment.assessment != native["assessment"]:
            surface.blockers.append(f"support: assessment {identity} differs from its stored judgment")
        surface.add("support", f"support_assessments:{identity}", {
            "identity": {"table": "support_assessments", "id": identity}, "original_actor": assessment.authority,
            "original_time": _plain(assessment.assessed_at), "decision_type": {"kind": "support_assessment", "assessment": assessment.assessment},
            "source_identity": _plain(assessment.segments), "predecessor": predecessors.get(identity)},
            {field: f"native_assessment:support_assessments:{identity}" for field in surface.contract.fields})
    surface.finish_kind("support", before)
    if history.publication_support_gaps:
        surface.blockers.extend(f"support: {message}" for message in history.publication_support_gaps)
        surface.observed.discard("support")
    if history.retained_history is not None and history.retained_history.gaps:
        from corridor.legacy_history import inventory_history
        from corridor.legacy_history_inventory import HISTORY_CLASSES
        legacy = inventory_history(session, project_id)
        # Audit is a mixed native/legacy event family. Its presence alone is
        # not an unknown predecessor population. Native source families already
        # have their own identity census above.
        compatibility = {item.table: legacy.counts[item.table] for item in HISTORY_CLASSES
                         if item.treatment != "native_lineage" and item.table != "audit_log"}
        surface.outputs["compatibility_population_census"] = {"counts": compatibility,
            "inventory_sha256": legacy.content_sha256}
        if any(compatibility.values()):
            surface.blockers.extend(f"history: {message}" for message in history.retained_history.gaps)


def collect_native_reader_coverage(session, project_id: int, *, as_of_revision_id: int,
                                   evaluated_at: datetime, briefing_fixture_client=None,
                                   maximum_population: int = 100_000) -> NativeReaderCoverage:
    """Observe all seven surfaces; callers cannot supply evidence or success flags.

    Use an otherwise clean Session. Each public reader runs behind a savepoint so
    a refusal cannot poison the following surface's read. Native revision/data
    census is checked again at the end; drift invalidates every coverage class.
    A provider-backed store is refused before any source or artifact is fetched.
    The returned object is evidence for comparison, never a cutover verdict.
    """
    if evaluated_at.tzinfo is None or evaluated_at.utcoffset() is None:
        raise ValueError("coverage needs an explicit timezone-aware evaluation instant")
    if type(maximum_population) is not int or maximum_population < 1:
        raise ValueError("coverage population limit must be positive")
    if session.new or session.dirty or session.deleted:
        raise ValueError("coverage cannot flush caller-owned pending writes")
    if settings.storage_backend != "filesystem":
        raise NativeReadingRefused("native coverage requires a local filesystem store; provider I/O is not enabled")
    project = session.get(Project, project_id)
    if project is None:
        raise ValueError("coverage project does not exist")
    revision = session.scalar(select(func.max(ProjectRecordRevision.id)).where(ProjectRecordRevision.project_id == project_id))
    selected = session.get(ProjectRecordRevision, as_of_revision_id)
    if revision is None or selected is None or selected.project_id != project_id:
        raise ValueError("both current and selected revisions must belong to the requested project")
    inventory = _inventory(session, project_id, maximum_population)
    authority_digest = _digest(inventory)
    surfaces = {c.name: _Surface(c) for c in CONTRACTS}
    from corridor.project_reading import freeze_project_reading
    # The freeze is made here, not accepted from a caller who can forge it.
    try:
        with session.begin_nested():
            reading = freeze_project_reading(session, project_id, today=evaluated_at.date())
            if reading.native_population is None or reading.native_population.revision_id != revision:
                raise NativeReadingRefused("the public freeze does not expose the pinned native population")
    except (ValueError, LookupError, SQLAlchemyError) as error:
        reading = None
        reason = f"shared native reading unavailable: {type(error).__name__}"
        if not isinstance(error, SQLAlchemyError):
            reason += f": {str(error)[:500]}"
        for surface in surfaces.values():
            surface.blockers.append(reason)
    def observe(name, operation):
        surface = surfaces[name]
        def isolated_read():
            with session.begin_nested() as savepoint:
                result = operation(surface)
                savepoint.rollback()
                return result
        _attempt(surface, name, isolated_read)
    if reading is not None:
        for blocker in (*reading.coverage_blockers, *getattr(reading.native_population, "coverage_blockers", ())):
            for surface in surfaces.values():
                surface.blockers.append(f"native population: {blocker}")
        observe("constraint_log", lambda s: _constraint_log(s, reading, inventory))
        observe("coordination_report", lambda s: _report(s, session, reading, inventory, briefing_fixture_client))
        observe("workbook_export", lambda s: _workbooks(s, session, reading, inventory))
    observe("work_list", lambda s: _work_list(s, session, project_id, revision, evaluated_at, inventory))
    for kind, boundary in (("current", revision), ("as_of", as_of_revision_id)):
        observe("current_and_as_of_record", lambda s, k=kind, r=boundary: _record_values(s, session, project_id, r, k, inventory))
    observe("report_release", lambda s: _releases(s, session, project_id, revision, evaluated_at, inventory))
    observe("source_and_decision_history", lambda s: _history(s, session, project_id, revision, as_of_revision_id, inventory, maximum_population))
    if _digest(_inventory(session, project_id, maximum_population)) != authority_digest:
        for surface in surfaces.values():
            surface.blockers.append("native authority/source population changed during collection; repeat the entire snapshot")
            surface.observed.clear()
    # A general population/snapshot blocker must not leave an observed-kind
    # declaration that another caller could mistake for complete coverage.
    for surface in surfaces.values():
        if any(reason.startswith(("shared native", "native population", "history:")) for reason in surface.blockers):
            surface.observed.clear()
    return NativeReaderCoverage(project_id, revision, as_of_revision_id, evaluated_at, authority_digest,
        tuple(surfaces[c.name].result() for c in CONTRACTS))
