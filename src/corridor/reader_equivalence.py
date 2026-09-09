"""Render diagnostics for the legacy Project Record reader transition.

The previous station-only overlay lived beside the plain current view and
reported identical outputs as successful cutover. Rendering belongs here,
above the projection and renderers; complete semantic coverage belongs to
reader_coverage. Output equality alone never authorizes a reader switch.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass, replace
from datetime import date
from io import BytesIO
from pathlib import Path
import re

from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy.orm import Session

from corridor.briefing import brief_project, render as render_briefing
from corridor.export import to_xlsx
from corridor.project_reading import FrozenProjectReading, freeze_project_reading
from corridor.record_projection import read_current_project_record

from corridor.report import build_report, render as render_report
from corridor.report_release import render_external_report_pdf
from corridor.reader_coverage import CONTRACTS, CoverageResult


@dataclass(frozen=True)
class ReaderEquivalence:
    coordination_report_identical: bool
    briefing_identical: bool
    workbook_cells_identical: bool
    release_pdf_text_identical: bool
    explanations: tuple[str, ...]
    coverage: tuple[CoverageResult, ...] = ()

    @property
    def passed(self) -> bool:
        return (len(self.coverage) == len(CONTRACTS)
                and {row.surface for row in self.coverage} == {item.name for item in CONTRACTS}
                and all(row.passed for row in self.coverage)
                and self.rendered_outputs_identical)

    @property
    def rendered_outputs_identical(self) -> bool:
        return all(
            (
                self.coordination_report_identical,
                self.briefing_identical,
                self.workbook_cells_identical,
                self.release_pdf_text_identical,
            )
        )


def freeze_project_reading_from_current_view(
    session: Session,
    project_id: int,
    *,
    today: date | None = None,
    legacy_reading: FrozenProjectReading | None = None,
) -> FrozenProjectReading:
    legacy = legacy_reading or freeze_project_reading(session, project_id, today=today)
    values = {
        (value.dependency_id, value.fact_type): value.text_value
        for value in read_current_project_record(session, project_id)
        if value.dependency_id is not None
    }
    rows = []
    for row in legacy.rows:
        dependency = copy(row.dependency)
        for fact_type in ("station_from", "station_to"):
            value = values.get((dependency.id, fact_type))
            if value is not None:
                setattr(dependency, fact_type, value)
        rows.append(replace(row, dependency=dependency))
    return FrozenProjectReading(
        project=legacy.project,
        rows=tuple(rows),
        evaluation=legacy.evaluation,
        statement_publication=legacy.statement_publication,
    )


def prove_reader_equivalence(
    session: Session,
    project_id: int,
    *,
    today: date,
    output_dir: Path,
    briefing_client_factory,
) -> ReaderEquivalence:
    """Diagnose old render equality; never certify the partial overlay as cutover.

    These consumers still receive legacy population, evaluation and statement
    publication. Coverage records that actual read path independently from
    their output equality. The seven-surface semantic engine lives in
    ``reader_coverage`` and must be fed independently native readings.
    """

    legacy = freeze_project_reading(session, project_id, today=today)
    viewed = freeze_project_reading_from_current_view(
        session, project_id, today=today, legacy_reading=legacy
    )
    legacy_report = build_report(
        session, project_id, today=today, frozen_reading=legacy
    )
    viewed_report = build_report(
        session, project_id, today=today, frozen_reading=viewed
    )
    viewed_report = replace(viewed_report, generated_at=legacy_report.generated_at)
    report_identical = render_report(legacy_report) == render_report(viewed_report)

    legacy_briefing = brief_project(
        session,
        project_id,
        client=briefing_client_factory(),
        today=today,
        frozen_reading=legacy,
    )
    viewed_briefing = brief_project(
        session,
        project_id,
        client=briefing_client_factory(),
        today=today,
        frozen_reading=viewed,
    )
    briefing_identical = render_briefing(legacy_briefing) == render_briefing(
        viewed_briefing
    )

    output_dir.mkdir(parents=True, exist_ok=True)
    legacy_xlsx = to_xlsx(
        session,
        project_id,
        output_dir / "legacy.xlsx",
        evaluation=legacy.evaluation,
        statement_publication=legacy.statement_publication,
        frozen_reading=legacy,
    )
    viewed_xlsx = to_xlsx(
        session,
        project_id,
        output_dir / "view.xlsx",
        evaluation=viewed.evaluation,
        statement_publication=viewed.statement_publication,
        frozen_reading=viewed,
    )
    workbook_identical = _workbook_cells(legacy_xlsx) == _workbook_cells(viewed_xlsx)

    legacy_release = render_external_report_pdf(
        session, project_id, today=today, frozen_reading=legacy
    )
    viewed_release = render_external_report_pdf(
        session, project_id, today=today, frozen_reading=viewed
    )
    release_identical = _pdf_text(legacy_release.pdf_bytes) == _pdf_text(
        viewed_release.pdf_bytes
    )
    return ReaderEquivalence(
        coordination_report_identical=report_identical,
        briefing_identical=briefing_identical,
        workbook_cells_identical=workbook_identical,
        release_pdf_text_identical=release_identical,
        coverage=tuple(CoverageResult(contract.name, False, 0,
            ("Legacy overlay supplies population, evaluation and statement publication; only station_from/station_to are projected.",))
            for contract in CONTRACTS),
        explanations=(
            "XLSX package timestamps are excluded; every workbook cell is compared.",
            "PDF container metadata is excluded; normalized rendered page text is compared.",
        ),
    )


def _workbook_cells(path: Path) -> tuple:
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        return tuple(
            (sheet.title, tuple(tuple(cell for cell in row) for row in sheet.iter_rows(values_only=True)))
            for sheet in workbook.worksheets
        )
    finally:
        workbook.close()


def _pdf_text(pdf_bytes: bytes) -> tuple[str, ...]:
    return tuple(
        re.sub(r"Generated .*? UTC", "Generated <normalized> UTC", page.extract_text())
        for page in PdfReader(BytesIO(pdf_bytes)).pages
    )
