"""Read and prove the plain current Project Record decision view.

The view is authority-preserving: it projects effective typed decisions and
never becomes a writer or cache.  This module supplies current/as-of reads, a
FrozenProjectReading overlay used only by the pre-cutover equivalence gate, and
measured query timing for the documented view-to-materialization escalation
decision (ADR-0071).
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass, replace
from datetime import date
from pathlib import Path
import re

import pymupdf
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from corridor.briefing import brief_project, render as render_briefing
from corridor.export import to_xlsx
from corridor.project_reading import FrozenProjectReading, freeze_project_reading
from corridor.report import build_report, render as render_report
from corridor.report_release import render_external_report_pdf


@dataclass(frozen=True)
class CurrentRecordValue:
    project_id: int
    dependency_id: int | None
    subject_key: str
    fact_type: str
    text_value: str | None
    date_value: date | None
    date_range_start: date | None
    date_range_end: date | None
    external_org_value_id: int | None
    document_value_id: int | None
    decision_id: int
    fact_id: int
    revision_id: int


@dataclass(frozen=True)
class ViewPerformance:
    project_id: int
    row_count: int
    planning_time_ms: float
    execution_time_ms: float
    materialized_view_needed: bool


@dataclass(frozen=True)
class ReaderEquivalence:
    coordination_report_identical: bool
    briefing_identical: bool
    workbook_cells_identical: bool
    release_pdf_text_identical: bool
    explanations: tuple[str, ...]

    @property
    def passed(self) -> bool:
        return all(
            (
                self.coordination_report_identical,
                self.briefing_identical,
                self.workbook_cells_identical,
                self.release_pdf_text_identical,
            )
        )


def read_current_project_record(
    session: Session, project_id: int
) -> tuple[CurrentRecordValue, ...]:
    rows = session.execute(
        text(
            "select project_id, dependency_id, subject_key, fact_type, text_value, "
            "date_value, date_range_start, date_range_end, external_org_value_id, "
            "document_value_id, decision_id, fact_id, revision_id "
            "from current_project_record "
            "where project_id = :project_id order by dependency_id, fact_type"
        ),
        {"project_id": project_id},
    ).all()
    return tuple(CurrentRecordValue(*row) for row in rows)


def read_project_record_as_of_revision(
    session: Session, project_id: int, revision_id: int
) -> tuple[CurrentRecordValue, ...]:
    rows = session.execute(
        text(
            "select decisions.project_id, candidates.merged_into, decisions.subject_key, "
            "decisions.fact_type, facts.text_value, facts.date_value, "
            "facts.date_range_start, facts.date_range_end, "
            "facts.external_org_value_id, facts.document_value_id, "
            "decisions.id, facts.id, decisions.revision_id "
            "from fact_decisions decisions "
            "join facts on facts.id = decisions.fact_id "
            "left join fact_decisions successor on successor.id = decisions.superseded_by "
            "left join extracted_proposals proposals "
            "on proposals.project_id = facts.project_id "
            "and proposals.document_id = facts.document_id "
            "and proposals.extraction_run_id = facts.extraction_run_id "
            "and proposals.subject_key = facts.subject_key "
            "left join candidates on candidates.id = proposals.candidate_id "
            "where decisions.project_id = :project_id "
            "and decisions.revision_id <= :revision_id "
            "and (successor.id is null or successor.revision_id > :revision_id) "
            "order by candidates.merged_into, decisions.fact_type"
        ),
        {"project_id": project_id, "revision_id": revision_id},
    ).all()
    return tuple(CurrentRecordValue(*row) for row in rows)


def freeze_project_reading_from_current_view(
    session: Session,
    project_id: int,
    *,
    today: date | None = None,
    legacy_reading: FrozenProjectReading | None = None,
) -> FrozenProjectReading:
    legacy = legacy_reading or freeze_project_reading(session, project_id, today=today)
    values = {
        (value.dependency_id, value.fact_type): value.text_value
        for value in read_current_project_record(session, project_id)
        if value.dependency_id is not None
    }
    rows = []
    for row in legacy.rows:
        dependency = copy(row.dependency)
        for fact_type in ("station_from", "station_to"):
            value = values.get((dependency.id, fact_type))
            if value is not None:
                setattr(dependency, fact_type, value)
        rows.append(replace(row, dependency=dependency))
    return FrozenProjectReading(
        project=legacy.project,
        rows=tuple(rows),
        evaluation=legacy.evaluation,
        statement_publication=legacy.statement_publication,
    )


def measure_current_record_view(
    session: Session,
    project_id: int,
    *,
    slow_threshold_ms: float = 50.0,
) -> ViewPerformance:
    plan = session.scalar(
        text(
            "explain (analyze, format json) "
            "select * from current_project_record where project_id = :project_id"
        ),
        {"project_id": project_id},
    )[0]
    row_count = len(read_current_project_record(session, project_id))
    execution = float(plan["Execution Time"])
    return ViewPerformance(
        project_id=project_id,
        row_count=row_count,
        planning_time_ms=float(plan["Planning Time"]),
        execution_time_ms=execution,
        materialized_view_needed=execution > slow_threshold_ms,
    )


def prove_reader_equivalence(
    session: Session,
    project_id: int,
    *,
    today: date,
    output_dir: Path,
    briefing_client_factory,
) -> ReaderEquivalence:
    """Render all four surfaces through legacy and view-fed frozen readings."""

    legacy = freeze_project_reading(session, project_id, today=today)
    viewed = freeze_project_reading_from_current_view(
        session, project_id, today=today, legacy_reading=legacy
    )
    legacy_report = build_report(
        session, project_id, today=today, frozen_reading=legacy
    )
    viewed_report = build_report(
        session, project_id, today=today, frozen_reading=viewed
    )
    viewed_report = replace(viewed_report, generated_at=legacy_report.generated_at)
    report_identical = render_report(legacy_report) == render_report(viewed_report)

    legacy_briefing = brief_project(
        session,
        project_id,
        client=briefing_client_factory(),
        today=today,
        frozen_reading=legacy,
    )
    viewed_briefing = brief_project(
        session,
        project_id,
        client=briefing_client_factory(),
        today=today,
        frozen_reading=viewed,
    )
    briefing_identical = render_briefing(legacy_briefing) == render_briefing(
        viewed_briefing
    )

    output_dir.mkdir(parents=True, exist_ok=True)
    legacy_xlsx = to_xlsx(
        session,
        project_id,
        output_dir / "legacy.xlsx",
        evaluation=legacy.evaluation,
        statement_publication=legacy.statement_publication,
        frozen_reading=legacy,
    )
    viewed_xlsx = to_xlsx(
        session,
        project_id,
        output_dir / "view.xlsx",
        evaluation=viewed.evaluation,
        statement_publication=viewed.statement_publication,
        frozen_reading=viewed,
    )
    workbook_identical = _workbook_cells(legacy_xlsx) == _workbook_cells(viewed_xlsx)

    legacy_release = render_external_report_pdf(
        session, project_id, today=today, frozen_reading=legacy
    )
    viewed_release = render_external_report_pdf(
        session, project_id, today=today, frozen_reading=viewed
    )
    release_identical = _pdf_text(legacy_release.pdf_bytes) == _pdf_text(
        viewed_release.pdf_bytes
    )
    return ReaderEquivalence(
        coordination_report_identical=report_identical,
        briefing_identical=briefing_identical,
        workbook_cells_identical=workbook_identical,
        release_pdf_text_identical=release_identical,
        explanations=(
            "XLSX package timestamps are excluded; every workbook cell is compared.",
            "PDF container metadata is excluded; normalized rendered page text is compared.",
        ),
    )


def _workbook_cells(path: Path) -> tuple:
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        return tuple(
            (sheet.title, tuple(tuple(cell for cell in row) for row in sheet.iter_rows(values_only=True)))
            for sheet in workbook.worksheets
        )
    finally:
        workbook.close()


def _pdf_text(pdf_bytes: bytes) -> tuple[str, ...]:
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as document:
        return tuple(
            re.sub(r"Generated .*? UTC", "Generated <normalized> UTC", page.get_text())
            for page in document
        )
