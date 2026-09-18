"""Deterministic conversion of retained binary spreadsheet values."""

from datetime import datetime, timezone
from io import BytesIO
import re
import types
import zipfile

from openpyxl import load_workbook
from openpyxl.writer import excel as excel_writer

from pathlib import Path

from corridor.spreadsheet_conversion import convert_xls_bytes


SOURCE = Path(__file__).with_name("fixtures") / "conversion-origins.xls"
# Synthetic BIFF8 workbook authored with xlwt 1.3.0 outside the tests. Each
# sheet holds the same three rows, at A1, A2, or C3, with a merged heading.


def _convert_with_clock_at(monkeypatch, instant):
    """Convert with the clock openpyxl reads during `save()` frozen."""

    with monkeypatch.context() as patched:
        patched.setattr(
            excel_writer,
            "datetime",
            types.SimpleNamespace(
                datetime=types.SimpleNamespace(now=lambda tz=None: instant),
                timezone=timezone,
            ),
        )
        return convert_xls_bytes(SOURCE.read_bytes())


def _archive_parts(archive: bytes) -> list[bytes]:
    """Every part of the archive, decompressed — the entries carry no plain text."""

    package = zipfile.ZipFile(BytesIO(archive))
    with package:
        return [package.read(name) for name in package.namelist()]


def _recorded_timestamps(archive: bytes) -> list[bytes]:
    """Every W3CDTF value the archive records, across all of its parts."""

    return [
        match.group(1)
        for part in _archive_parts(archive)
        for match in re.finditer(rb'xsi:type="dcterms:W3CDTF"[^<>]*>([^<]*)', part)
    ]


def test_binary_workbook_values_convert_to_stable_readable_xlsx(tmp_path, monkeypatch):
    # Two consecutive calls used to agree only because they landed in the same
    # second, which hid a wall clock in `docProps/core.xml` (#579). Separate the
    # conversions by years of mocked clock, and independently refuse any value
    # that tracks the real clock in case a later openpyxl reads a different one.
    first = _convert_with_clock_at(
        monkeypatch, datetime(2031, 5, 17, 12, 0, 0, tzinfo=timezone.utc)
    )
    second = _convert_with_clock_at(
        monkeypatch, datetime(1994, 8, 3, 23, 59, 59, tzinfo=timezone.utc)
    )
    unmocked = convert_xls_bytes(SOURCE.read_bytes())
    assert first == second == unmocked

    stamps = _recorded_timestamps(first)
    assert stamps, "the archive should still record its created and modified times"
    assert set(stamps) == {b"2000-01-01T00:00:00Z"}
    today = datetime.now(tz=timezone.utc).date().isoformat().encode()
    assert not [part for part in _archive_parts(unmocked) if today in part]

    path = tmp_path / "converted.xlsx"
    path.write_bytes(first)
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        sheet = workbook["THDS INDEX"]
        assert sheet["A1"].value == "Test Hole Index"
        assert str(sheet.merged_cells) == "A1:B1"
        assert sheet["A3"].value == "169-A"
        assert sheet["B3"].value == 4.54
        assert sheet["C3"].value == datetime(2024, 11, 13)
        assert sheet["C3"].number_format == "yyyy-mm-dd"
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:C3"
        assert sheet.column_dimensions["A"].width >= len("TEST HOLE #")
        assert sheet["B3"].number_format == "0.###"
    finally:
        workbook.close()


def test_conversion_preserves_values_and_merges_in_the_source_coordinate_frame():
    converted = convert_xls_bytes(SOURCE.read_bytes())
    workbook = load_workbook(BytesIO(converted), data_only=True)
    try:
        for name, title, identifier, depth, day, merged in (
            ("Leading row", "A2", "A4", "B4", "C4", "A2:B2"),
            ("Leading rows and columns", "C3", "C5", "D5", "E5", "C3:D3"),
        ):
            sheet = workbook[name]
            assert sheet[title].value == "Test Hole Index"
            assert sheet[identifier].value == "169-A"
            assert sheet[depth].value == 4.54
            assert sheet[day].value == datetime(2024, 11, 13)
            assert str(sheet.merged_cells) == merged
            assert sheet["A1"].value is None
    finally:
        workbook.close()
