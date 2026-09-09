"""Production UCM readers consume explicit native identities and accepted decisions."""
from dataclasses import FrozenInstanceError
from datetime import date
from hashlib import sha256
from pathlib import Path
import re
from uuid import uuid4

from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import event, select

from corridor.accepted_field_reading import NativeReadingRefused, read_accepted_field_population
from corridor.baseline_adoption import adopt_baseline, preview_baseline_adoption
from corridor.briefing import brief_project
from corridor.config import settings
from corridor.db import Session, engine
from corridor.exceptions import evaluate_project
from corridor.export import to_xlsx
from corridor.fact_decisions import record_human_fact_decision
from corridor.field_mapping_manifest import MappingDeclaration
from corridor.ledger import browse
from corridor.materializer import materialize_segment_value
from corridor.models import BaselineSourceRow, Fact, FactDecision, Project, SourceSegment
from corridor.principals import HumanPrincipal
from corridor.project_reading import freeze_project_reading
from corridor.reader_equivalence import native_constraint_log_surface
from corridor.report import Assertion, build_report, render
from corridor.source_append import append_fact
from corridor.source_intake import validate_and_stage


PRINCIPAL = HumanPrincipal("local:coordinator")
TODAY = date(2026, 9, 9)


@pytest.fixture
def session():
    connection = engine.connect()
    transaction = connection.begin()
    scoped = Session(bind=connection)
    yield scoped
    scoped.close()
    if transaction.is_active:
        transaction.rollback()
    connection.close()


@pytest.fixture
def adopted(session, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "corpus_store", str(tmp_path / "files"))
    monkeypatch.setattr(settings, "corpus_images", str(tmp_path / "images"))
    project = Project(slug=f"native-readers-{uuid4().hex[:8]}", name="Native Reader Test", is_synthetic=True)
    session.add(project)
    session.flush()
    path = tmp_path / "ucm.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Utility Conflicts"
    sheet.append(["Utility Conflict Management (UCM) — Utility Conflicts"])
    sheet.append(["Utility Conflict ID", "Utility Owner", "Utility Type", "Start Station", "End Station", "Promised For", "Action Due Date", "Comment"])
    sheet.append(["UC-1", "City Water", "Water", "100+00", "101+00", "2026-09-01", "2026-08-01", "first facility"])
    sheet.append(["UC-1", "Electric Company", "Electric", "200+00", "201+00", "2026-10-01", "2026-08-02", "different facility"])
    book.save(path)
    preview = preview_baseline_adoption(session, project=project,
        staged=validate_and_stage(path.read_bytes(), path.name), customer="Synthetic Customer",
        source_identity="UCM synthetic baseline", field_mapping=MappingDeclaration())
    result = adopt_baseline(session, preview=preview, principal=PRINCIPAL,
        idempotency_key="adopt-native-readers", images_dir=tmp_path / "images")
    return project, result


class CoveringClient:
    model = "synthetic-briefing"
    def complete(self, *, user, **kwargs):
        references = sorted(set(re.findall(r"\b(?:XB|X|E|A|V)\d+\b", user)))
        return {"sentences": [{"text": "Read the accepted values and their cited sources.", "cites": references}]}


def test_actual_readers_never_select_legacy_populations_or_values(session, adopted, tmp_path):
    project, result = adopted
    observed = []
    forbidden = re.compile(r"\b(?:dependencies|candidates|assertions|evidence_links|work_decisions|operative_support)\b", re.I)
    def inspect_sql(connection, cursor, statement, parameters, context, executemany):
        if statement.lstrip().lower().startswith("select"):
            observed.append(statement)
            assert not forbidden.search(statement), statement
    event.listen(session.get_bind(), "before_cursor_execute", inspect_sql)
    try:
        reading = freeze_project_reading(session, project.id, today=TODAY)
        assert reading.native_population.revision_id == result.revision_id
        assert len(reading.rows) == 2
        assert {row.dependency.source_ref for row in reading.rows} == {"UC-1"}
        assert len({row.dependency.id for row in reading.rows}) == 2
        assert all(isinstance(row.dependency.id, str) for row in reading.rows)
        assert {row.dependency.station_from for row in reading.rows} == {"100+00", "200+00"}
        assert all(row.dependency.action_due_date is None for row in reading.rows)
        assert not any(finding.rule.startswith("ACTION_") for finding in reading.evaluation.found)
        assert len(browse(session, project.id, evaluation=reading.evaluation, owner="unassigned")) == 2
        report = build_report(session, project.id, today=TODAY, frozen_reading=reading)
        body = render(report)
        assert "100+00" in body and "200+00" in body and "p.None" not in body
        assert "Utility Conflicts!" in body
        assert any(isinstance(cell.provenance, Assertion) and cell.provenance.decision_id for cell in report.cells)
        briefing = brief_project(session, project.id, client=CoveringClient(), today=TODAY, frozen_reading=reading)
        assert not briefing.refused and briefing.sentences
        assert any("decision " in item.text for item in briefing.citables)
        output = to_xlsx(session, project.id, tmp_path / "native.xlsx", evaluation=reading.evaluation,
            statement_publication=reading.statement_publication, frozen_reading=reading)
        book = load_workbook(output, data_only=True)
        try:
            assert book.active.max_row == 3
            assert book.active["F2"].value in {"100+00", "200+00"}
            assert "Accepted value sources" in book.sheetnames
        finally:
            book.close()
        surface = native_constraint_log_surface(reading)
        assert surface.observed_record_kinds == {"constraint", "check"}
        assert sum(row.kind == "constraint" for row in surface.records) == 2
    finally:
        event.remove(session.get_bind(), "before_cursor_execute", inspect_sql)
    assert observed


def test_frozen_native_rows_and_fields_cannot_be_mutated(session, adopted):
    project, _ = adopted
    reading = freeze_project_reading(session, project.id, today=TODAY)
    record = reading.rows[0].dependency
    with pytest.raises(FrozenInstanceError):
        record.subject_key = "changed"
    with pytest.raises(TypeError):
        record.fields["station_from"] = None
    assert reading.evaluation.native_population is reading.native_population


def _unselected_fact(session, project, *, subject_key):
    original = session.scalar(select(Fact).where(Fact.project_id == project.id, Fact.fact_type == "station_from").order_by(Fact.id))
    segment = session.scalar(select(SourceSegment).where(SourceSegment.document_id == original.document_id,
        SourceSegment.exact_text == original.text_value))
    return append_fact(session, project_id=project.id, document_id=original.document_id,
        extraction_run_id=original.extraction_run_id, subject_kind="source_row", subject_key=subject_key,
        recorded_by="extractor:synthetic", content_sha256=sha256(f"unselected:{subject_key}".encode()).hexdigest(),
        value=materialize_segment_value(session, "station_from", segment))


def test_source_check_does_not_make_an_undecided_fact_an_accepted_value(session, adopted):
    project, _ = adopted
    before = read_accepted_field_population(session, project.id)
    fact = _unselected_fact(session, project, subject_key="unmapped-new-source")
    after = read_accepted_field_population(session, project.id)
    assert before.fingerprint == after.fingerprint
    assert all(field.fact_id != fact.id for row in after.records for field in row.fields.values())
    record_human_fact_decision(session, fact, principal=PRINCIPAL, command_type="resolve_discrepancy", idempotency_key="decide-unmapped")
    with pytest.raises(NativeReadingRefused, match="no declared adopted-row mapping"):
        read_accepted_field_population(session, project.id)


def test_ambiguous_effective_source_and_record_alias_decisions_refuse_instead_of_last_wins(session, adopted):
    project, _ = adopted
    row = session.scalar(select(BaselineSourceRow).where(BaselineSourceRow.project_id == project.id).order_by(BaselineSourceRow.id))
    assert row.source_row_key != row.record_subject_key
    fact = _unselected_fact(session, project, subject_key=row.record_subject_key)
    record_human_fact_decision(session, fact, principal=PRINCIPAL, command_type="resolve_discrepancy", idempotency_key="ambiguous-alias")
    with pytest.raises(NativeReadingRefused, match="more than one accepted decision"):
        read_accepted_field_population(session, project.id)


def test_native_as_of_reading_preserves_the_exact_decision_boundary(session, adopted):
    project, result = adopted
    before = freeze_project_reading(session, project.id, today=TODAY)
    record = before.native_population.records[0]
    held = record.fields["station_from"]
    fact = session.get(Fact, held.fact_id)
    changed = record_human_fact_decision(session, fact, principal=PRINCIPAL, command_type="resolve_discrepancy",
        idempotency_key="withdraw-station", expected_predecessor=held.decision_id, disposition="do_not_add")
    now = freeze_project_reading(session, project.id, today=TODAY)
    original = freeze_project_reading(session, project.id, today=TODAY, revision_id=result.revision_id)
    assert now.native_population.revision_id == changed.revision.id
    assert next(row for row in now.native_population.records if row.subject_key == record.subject_key).station_from is None
    assert original.native_population.fingerprint == before.native_population.fingerprint
    assert original.native_population.records[0].station_from == record.station_from
    paired = freeze_project_reading(session, project.id, today=TODAY, evaluation=before.evaluation,
        statement_publication=before.statement_publication)
    assert paired.native_population is before.native_population


def test_native_revision_and_legacy_population_gaps_are_explicit(session, adopted):
    project, result = adopted
    other = Project(slug=f"unmapped-{uuid4().hex[:8]}", name="No native map", is_synthetic=True)
    session.add(other)
    session.flush()
    with pytest.raises(NativeReadingRefused, match="legacy accepted-field"):
        read_accepted_field_population(session, other.id)
    with pytest.raises(ValueError, match="revision"):
        from corridor.record_projection import read_native_record_values
        read_native_record_values(session, other.id, result.revision_id)
