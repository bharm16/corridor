
from corridor.extract_sheet import (
    PROMPT_VERSION as SHEET_PROMPT_VERSION,
    SCHEMA_VERSION as SHEET_SCHEMA_VERSION,
)
from corridor.pipeline import extraction_route

from test_native_matrix import _document, matrix_source, project, session
from test_native_pipeline import _client


def select_native_runtime(session, project, source, tmp_path):
    from corridor.native_matrix_runtime import NativeMatrixRuntime
    from corridor.native_pipeline import native_pipeline_configuration, register_pipeline_configuration
    from corridor.pipeline_contracts import AcceptedEvidence, MaintainerAcceptance
    from corridor.pipeline_qualification import record_acceptance, select_qualified_pipeline
    from test_native_pipeline import ACTOR, _plan, _scope

    scope = _scope(source)
    runtime = NativeMatrixRuntime(scope.deployment, _client(source), _plan(source=source), tmp_path)
    configuration = native_pipeline_configuration(runtime.client, runtime.plan)
    registered = register_pipeline_configuration(session, configuration)
    acceptance = record_acceptance(
        session, MaintainerAcceptance(
            decision="Synthetic route regression", configuration_sha256=registered.configuration_sha256,
            implementation_revision=configuration["code_revision"], scope_sha256=scope.identity,
            evidence=(AcceptedEvidence(name="Authored fixture", reference="tests/test_native_matrix.py",
                                       summary="The PDF fixture declares its cells and expected mapping."),),
            limits=("Synthetic regression only; no production authorization.",),
            words="Select this exact configuration for the synthetic route regression.",
            accepted_on="2026-09-08",
        ), project_id=project.id, scope=scope, actor=ACTOR,
    )
    selection = select_qualified_pipeline(session, acceptance_id=acceptance.id, actor=ACTOR,
        reason="Synthetic route regression", expected_selection_id=None)
    return runtime, selection


def test_selected_matrix_ingests_extracts_and_cites_one_sealed_run(
    session, project, matrix_source, tmp_path,
):
    from sqlalchemy import func, select
    from corridor.evidence_citations import cite_source_segments, evidence_quotation
    from corridor.extract_project import extract_project
    from corridor.ingest import ingest_document
    from corridor.models import ActiveExtractionRun, EvidenceLink, ExtractionRun, Fact, PipelineObservation
    from corridor.native_matrix import PROMPT_VERSION, SCHEMA_VERSION
    from corridor.object_storage import store_bytes
    from corridor.reader_segments import append_native_segments
    from corridor.retained_history import replay_retained_reading

    source = matrix_source
    store_bytes(source.path.read_bytes(), sha256=source.reading.rendition_sha256, suffix=".pdf")
    document = ingest_document(session, project_id=project.id, path=source.path,
        doc_type="matrix", images_dir=tmp_path / "ingest")
    # Retained before the production route changes: the immutable reading and
    # link must still resolve after the new extraction writes its own run.
    segments = append_native_segments(session, document, source.reading)
    retained = next(segment for segment in segments if segment.exact_text == "Exact Utilities")
    citation = EvidenceLink(document_id=document.id, page_no=1, quote="", verified=True)
    session.add(citation)
    session.flush()
    cite_source_segments(session, citation, (retained,))
    before = evidence_quotation(session, citation)
    runtime, selection = select_native_runtime(session, project, source, tmp_path / "run")
    [outcome] = extract_project(session, project, commit=False,
        select_route=lambda doc: extraction_route(doc, native_runtime=runtime))

    assert outcome.status == "extracted", outcome.detail
    run = session.get_one(ExtractionRun, outcome.extraction_run_id)
    assert (run.prompt_version, run.schema_version) == (PROMPT_VERSION, SCHEMA_VERSION)
    assert run.candidate_count == outcome.rows == 1
    assert session.scalar(select(func.count()).select_from(ExtractionRun)) == 1
    observation = session.scalar(select(PipelineObservation))
    assert observation.extraction_run_id == run.id
    assert observation.configuration_sha256 == selection.configuration_sha256
    assert session.scalar(select(func.count()).select_from(Fact)) > 0
    assert session.scalar(select(func.count()).select_from(ActiveExtractionRun)) == 0
    assert evidence_quotation(session, citation) == before
    assert replay_retained_reading(retained, document=document, path=source.path).exact_text == "Exact Utilities"

    [resumed] = extract_project(session, project, commit=False,
        select_route=lambda doc: extraction_route(doc, native_runtime=runtime))
    assert resumed.status == "skipped"
    assert runtime.client.consumed == 1

    from corridor.pipeline_qualification import select_qualified_pipeline
    from test_native_pipeline import ACTOR

    select_qualified_pipeline(session, acceptance_id=selection.acceptance_id, actor=ACTOR,
        reason="Disable the synthetic selection", expected_selection_id=selection.id, enabled=False)
    [disabled] = extract_project(session, project, commit=False,
        select_route=lambda doc: extraction_route(doc, native_runtime=runtime))
    assert disabled.status == "failed"
    assert "no enabled" in disabled.detail
    assert runtime.client.consumed == 1


def test_out_of_scope_matrix_fails_before_rendering_or_mapping(
    session, project, matrix_source, tmp_path,
):
    from corridor.extract_project import extract_project
    from corridor.models import Document

    runtime, _ = select_native_runtime(session, project, matrix_source, tmp_path / "never-created")
    document = Document(project_id=project.id, doc_type="matrix", filename="outside.pdf",
        sha256="d" * 64, pages=1, parse_status="parsed")
    session.add(document)
    session.flush()
    [outcome] = extract_project(session, project, commit=False,
        select_route=lambda doc: extraction_route(doc, native_runtime=runtime))
    assert outcome.status == "failed"
    assert "exact source/configuration/class" in outcome.detail
    assert runtime.client.consumed == 0
    assert not runtime.output_dir.exists()


def test_refused_native_attempt_preserves_its_observation_and_artifacts(
    session, project, matrix_source, tmp_path,
):
    from dataclasses import replace
    from pathlib import Path
    from sqlalchemy import func, select
    from corridor.extract_project import extract_project
    from corridor.models import Candidate, ExtractionRun, Fact, PipelineObservation, ProcessingArtifact
    from corridor.object_storage import store_bytes

    source = replace(matrix_source, answers=({**matrix_source.answers[0], "header_row": None},))
    store_bytes(source.path.read_bytes(), sha256=source.reading.rendition_sha256, suffix=".pdf")
    _document(session, project, source)
    runtime, _ = select_native_runtime(session, project, source, tmp_path / "refused")
    [outcome] = extract_project(session, project, commit=False,
        select_route=lambda doc: extraction_route(doc, native_runtime=runtime))
    assert outcome.status == "failed"
    assert session.scalar(select(func.count()).select_from(PipelineObservation)) == 1
    assert session.scalar(select(func.count()).select_from(ExtractionRun)) == 1
    assert session.scalar(select(func.count()).select_from(Candidate)) == 0
    assert session.scalar(select(func.count()).select_from(Fact)) == 0
    artifacts = session.scalars(select(ProcessingArtifact)).all()
    assert artifacts and all(Path(artifact.storage_path).is_file() for artifact in artifacts)


def test_runtime_rejects_a_transport_endpoint_outside_its_authorized_request(tmp_path, monkeypatch):
    import json
    import pytest
    from corridor.config import Settings
    from corridor.native_matrix_runtime import configured_native_matrix_runtime
    from corridor.native_provider_boundary import NativeProviderRefused
    from test_native_provider_boundary import SOURCE, _experiment, _request

    runtime_file = tmp_path / "runtime.json"
    runtime_file.write_text(json.dumps({
        "deployment": "development", "authorization": _experiment().as_dict(),
        "request": _request().as_dict(), "source_sha256s": [SOURCE], "campaign": "endpoint-test",
        "budget": {"max_calls": 1, "max_pages": 1, "max_total_tokens": 10000},
    }))
    def forbidden(*args):
        pytest.fail("endpoint mismatch reached transport construction")
    monkeypatch.setattr("corridor.native_matrix_runtime.live_transport", forbidden)
    settings = Settings(_env_file=None, openai_base_url="https://unapproved.invalid/v1",
        CORRIDOR_NATIVE_MATRIX_RUNTIME_FILE=str(runtime_file))
    with pytest.raises(NativeProviderRefused, match="endpoint"):
        configured_native_matrix_runtime(settings)


def test_rejected_native_result_rolls_back_the_successful_capture(
    session, project, matrix_source, tmp_path,
):
    from dataclasses import replace
    import pytest
    from sqlalchemy import func, select
    from corridor.extract_project import extract_project
    from corridor.models import Candidate, ExtractionRun, Fact, PipelineObservation
    from corridor.object_storage import store_bytes

    store_bytes(matrix_source.path.read_bytes(), sha256=matrix_source.reading.rendition_sha256, suffix=".pdf")
    document = _document(session, project, matrix_source)
    runtime, _ = select_native_runtime(session, project, matrix_source, tmp_path / "rejected")
    route = extraction_route(document, native_runtime=runtime)
    def incorrect_result(db, document):
        reading = route.extract(db, document)
        reading.rows[0].model = "wrong-model"
        return reading
    broken = replace(route, extract=incorrect_result)
    with pytest.raises(ValueError, match="Candidate model"):
        extract_project(session, project, commit=False, select_route=lambda doc: broken)
    for model in (Candidate, Fact, PipelineObservation):
        assert session.scalar(select(func.count()).select_from(model)) == 0
    [run] = session.scalars(select(ExtractionRun)).all()
    assert run.outcome == "failed"


def test_unselected_matrix_is_a_recorded_processing_failure(
    session, project, matrix_source,
):
    from corridor.extract_project import extract_project
    from corridor.models import ExtractionRun

    document = _document(session, project, matrix_source)
    client = _client(matrix_source)
    [outcome] = extract_project(
        session, project, commit=False,
        select_route=lambda doc: extraction_route(doc, client=client),
    )

    assert outcome.status == "failed"
    assert "selection" in outcome.detail
    assert client.consumed == 0
    run = session.get_one(ExtractionRun, outcome.extraction_run_id)
    assert run.document_id == document.id
    assert run.outcome == "failed"
    assert run.candidate_count == 0


def test_a_spreadsheet_uses_the_native_sheet_prompt_version(monkeypatch):
    monkeypatch.setattr("corridor.pipeline.stored_file", lambda document: "/tmp/a.xlsx")

    route = extraction_route(object())

    assert route.effective_prompt_version == SHEET_PROMPT_VERSION
    assert route.schema_version == SHEET_SCHEMA_VERSION
    assert route.extractor_config is not None
    assert route.extractor_config.prompt_version == SHEET_PROMPT_VERSION
    assert route.extractor_config.config_json["request_controls"] == {
        "provider": "native",
        "model_requests": 0,
    }


def test_an_unreadable_ingest_records_why_on_its_receipt(session, project, tmp_path):
    """The receipt carries the stage, the exception type and the message.

    `error_detail` used to read `ingest parse_status is 'failed'` for every
    whole-document failure, so a reader exception, a rendition with no pages and
    drifted source bytes were one indistinguishable line and diagnosing any of
    them meant running the reader again.
    """

    import json
    from sqlalchemy import select
    from corridor.models import ExtractionRun
    from corridor.pipeline import ingest_and_extract

    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 this is not a real pdf")
    document, candidates = ingest_and_extract(
        session, project_id=project.id, path=broken, images_dir=tmp_path / "images",
    )

    assert candidates == []
    run = session.scalars(select(ExtractionRun)).one()
    assert (run.document_id, run.outcome) == (document.id, "unreadable")
    detail = json.loads(run.error_detail)
    assert detail["parse_status"] == "failed"
    assert detail["stage"] == "read_document"
    assert detail["error_type"] and detail["error_message"]


def test_a_native_reading_with_no_rows_still_records_through_the_fact_command(
    session, project, tmp_path, monkeypatch
):
    """Zero rows is a reading, not the absence of one.

    `record_routed_run` used to pick the scoped Fact command by inspecting the
    rows it was handed. A native spreadsheet that reads as the published empty
    template produces no rows, so the sniff concluded there were no source Facts
    here and recorded the run through the legacy command instead: no append
    receipt, and lineage reconstructed from candidate existence, which is the one
    thing `extraction_runs` exists to prevent.
    """

    from openpyxl import Workbook
    from sqlalchemy import select
    from corridor.ingest import ingest_document
    from corridor.models import ExtractionRun, SourceFactAppendReceipt, SourceSegment
    from corridor.pipeline import SOURCE_FACTS, extract_any

    book = Workbook()
    book.remove(book.active)
    sheet = book.create_sheet("Utility Conflicts")
    sheet.append(["Utility Conflict Management (UCM) - Utility Conflicts"])
    sheet.append(["Utility Conflict ID", "Utility Owner", "Start Station"])
    path = tmp_path / "empty-template.xlsx"
    book.save(path)

    document = ingest_document(
        session, project_id=project.id, path=path,
        doc_type="matrix", images_dir=tmp_path / "images",
    )
    document._stored_path = str(path)
    monkeypatch.setattr(
        "corridor.extract_sheet.stored_file", lambda d: getattr(d, "_stored_path", None)
    )

    route = extraction_route(document)
    assert route.output == SOURCE_FACTS
    assert extract_any(session, document) == []

    run = session.scalars(select(ExtractionRun)).one()
    assert (run.outcome, run.candidate_count) == ("completed", 0)
    assert session.scalars(select(SourceFactAppendReceipt)).one().extraction_run_id == run.id
    assert session.scalars(
        select(SourceSegment).where(SourceSegment.document_id == document.id)
    ).all()


def test_a_route_declared_legacy_never_reaches_the_fact_command(session, project):
    """A declaration the rows contradict is refused, not silently reinterpreted.

    The route knows at selection time which command its reading belongs at.
    `rows_carry_source_facts` is now the consistency check on that declaration
    rather than the decision itself, so structured-cell rows arriving on a route
    that declared legacy Extracted Proposals fail here instead of quietly
    reaching one command or the other.
    """

    import pytest
    from sqlalchemy import select
    from corridor.candidates import propose
    from corridor.extractor_lineage import injected_extractor_config
    from corridor.models import Document, ExtractionRun
    from corridor.pipeline import (
        EXTRACTED_PROPOSALS,
        ExtractionRoute,
        record_routed_run,
    )

    document = Document(
        project_id=project.id, doc_type="matrix", filename="legacy.pdf",
        sha256="1" * 64, pages=1, parse_status="parsed",
    )
    session.add(document)
    session.flush()

    def extract(db, target):
        candidate = propose(
            target, kind="dependency", fields={"utility_id": "UC-1"}, page_no=1,
            quote="UC-1", quote_verified=True, whole_row=True, confidence=None,
            prompt_version="legacy_fixture_v1", dedupe="uc-1",
            text_source="native_pdf_segments", tier="native",
        )
        db.add(candidate)
        db.flush([candidate])
        return [candidate]

    route = ExtractionRoute(
        effective_prompt_version="legacy_fixture_v1",
        schema_version="legacy_fixture_v1",
        extract=extract,
        output=EXTRACTED_PROPOSALS,
        extractor_config=injected_extractor_config(
            extractor="legacy-fixture", prompt_version="legacy_fixture_v1",
            model=None, schema_version="legacy_fixture_v1",
            prompt_bytes=b"legacy route fixture", schema={"type": "object"},
            postprocessor_bytes=b"legacy route fixture rules",
            request_controls={"strict": True},
        ),
    )
    candidates = route.extract(session, document)

    with pytest.raises(ValueError, match="declared"):
        record_routed_run(
            session, document, route, None, candidate_count=len(candidates),
            page_errors=0, outcome="completed", candidates=candidates, model=None,
        )
    assert session.scalars(select(ExtractionRun)).all() == []
