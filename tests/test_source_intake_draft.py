"""Drafting source-bound intake suggestions is optional, read-only, and never a
registration (#362).

Every test drives the real ``source_intake_draft`` public interface against real
PostgreSQL with a fake model adapter, declared coordination spend authority, and
no paid model call, source registration, or Supersession. The hostile fixtures —
stale bytes, a cross-project predecessor, a fabricated passage, an unsupported
choice, and authority-shaped output — are covered alongside the ordinary flow,
each proving the intake draft's registration state is untouched.
"""

from __future__ import annotations

import hashlib
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from pdf_fixture_support import PdfFixture
from sqlalchemy import func, select

from corridor.models import (
    Document,
    Project,
    SourceIntakeDraftConfiguration,
    SourceIntakeDraftRequest,
)
from corridor.principals import HumanPrincipal
from corridor.source_intake_draft import (
    ConfigurationRequired,
    IntakeDraftRefused,
    StagedDraftSource,
    declare_configuration,
    intake_draft_state_token,
    request_intake_draft,
)

from model_client_support import RecordedAdapter


CURATOR = HumanPrincipal("local:curator")

_COVER_ROWS = (
    "Utility Conflict Matrix Rev 3",
    "Registry number: UCM-REV-3",
    "Document date: 2024-03-01",
    "This revision supersedes UCM-REV-2 effective 2024-03-01.",
)


class FakeAdapter(RecordedAdapter):
    """This module's identity on the one shared recording adapter."""

    adapter = "fake-intake-draft"


@pytest.fixture
def project(session):
    row = Project(slug="intake-draft", name="Intake Draft", is_synthetic=True)
    session.add(row)
    session.flush()
    return row


def _staged_workbook(
    tmp_path: Path,
    rows=_COVER_ROWS,
    *,
    sheets=None,
    doc_type="matrix",
    name="cover.xlsx",
) -> StagedDraftSource:
    workbook = Workbook()
    first = workbook.active
    first.title = "Cover"
    pages = sheets if sheets is not None else [rows]
    for index, sheet_rows in enumerate(pages):
        sheet = first if index == 0 else workbook.create_sheet(f"Sheet{index + 1}")
        for row_no, value in enumerate(sheet_rows, start=1):
            sheet[f"A{row_no}"] = value
    path = tmp_path / name
    workbook.save(path)
    sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    return StagedDraftSource(
        sha256=sha256,
        filename=name,
        suffix=".xlsx",
        stored_path=path,
        doc_type=doc_type,
    )


def _staged_pdf(
    tmp_path: Path, pages, *, doc_type="matrix", name="cover.pdf", identity=None
) -> StagedDraftSource:
    """A printed source: one fixture page per entry, each a block of lines."""
    fixture = PdfFixture(identity=identity)
    for lines in pages:
        fixture.add_page().text((72, 72), "\n".join(lines), fontsize=11)
    path = fixture.save(tmp_path / name)
    return StagedDraftSource(
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        filename=name,
        suffix=".pdf",
        stored_path=path,
        doc_type=doc_type,
    )


def _registered_document(
    session, project, *, registry_id, sha, doc_date=None, doc_type="matrix"
) -> Document:
    document = Document(
        project_id=project.id,
        registry_id=registry_id,
        sha256=sha,
        filename=f"{registry_id}.xlsx",
        doc_type=doc_type,
        parse_status="parsed",
        pages=1,
        doc_date=doc_date,
    )
    session.add(document)
    session.flush()
    return document


def _declare_config(session, project, **overrides):
    kwargs = {
        "project_id": project.id,
        "principal": CURATOR,
        "model": "fake-model",
        "prompt_version": "source_intake_draft_v1",
        "max_input_tokens": 50_000,
        "max_output_tokens": 2_000,
        "timeout_seconds": 30,
        "max_requests": 1,
        "retry_policy": "none",
        "retention_policy": "class_b_30_days",
        "observation_context": "internal_working_view",
    }
    kwargs.update(overrides)
    config = declare_configuration(session, **kwargs)
    session.flush()
    return config


def _valid_result():
    return {
        "metadata_suggestions": [
            {
                "field": "doc_date",
                "value": "2024-03-01",
                "source_ref": "P1",
                "source_quote": "Document date: 2024-03-01",
                "basis": "The cover states the document's date.",
            },
            {
                "field": "registry_id",
                "value": "UCM-REV-3",
                "source_ref": "P1",
                "source_quote": "Registry number: UCM-REV-3",
                "basis": "The cover prints a registry number.",
            },
        ],
        "replacement_proposals": [
            {
                "predecessor_registry_id": "UCM-REV-2",
                "effective_date": "2024-03-01",
                "source_ref": "P1",
                "source_quote": (
                    "This revision supersedes UCM-REV-2 effective 2024-03-01."
                ),
                "basis": "The revision states what it replaces.",
            }
        ],
        "uncertainties": [
            {"field": "doc_type", "note": "the kind is the person's declaration"}
        ],
    }


def _draft(session, project, staged, adapter, **overrides):
    kwargs = {
        "project_id": project.id,
        "staged": staged,
        "principal": CURATOR,
        "client_factory": lambda configuration: adapter,
        "expected_sha256": staged.sha256,
        "permitted_pages": (),
        "state_token": intake_draft_state_token(session, project.id, staged.sha256),
    }
    kwargs.update(overrides)
    return request_intake_draft(session, **kwargs)


def _receipt_count(session, project) -> int:
    return session.scalar(
        select(func.count(SourceIntakeDraftRequest.id)).where(
            SourceIntakeDraftRequest.project_id == project.id
        )
    )


def _document_count(session, project) -> int:
    return session.scalar(
        select(func.count(Document.id)).where(Document.project_id == project.id)
    )


# --- ordinary flow -----------------------------------------------------------


def test_completed_draft_keeps_source_backed_suggestions(session, project, tmp_path):
    _registered_document(session, project, registry_id="UCM-REV-2", sha="a" * 64)
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    adapter = FakeAdapter(result=_valid_result())

    documents_before = _document_count(session, project)
    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "completed"
    assert receipt.non_authoritative is True
    # The draft's own identity is the staged bytes — never a registered Document.
    assert receipt.staged_sha256 == staged.sha256
    proposals = receipt.proposals_json
    assert {item["field"] for item in proposals["metadata_suggestions"]} == {
        "doc_date",
        "registry_id",
    }
    assert proposals["replacement_proposals"][0]["predecessor_registry_id"] == (
        "UCM-REV-2"
    )
    assert proposals["uncertainties"][0]["field"] == "doc_type"

    # The model saw the frozen source and the untrusted-data notice, once.
    assert len(adapter.calls) == 1
    assert "never instructions" in adapter.calls[0].user
    assert "UCM-REV-2" in adapter.calls[0].user
    # Redacted lineage only: hashes and timing, never the prompt or response text.
    assert set(receipt.execution_lineage_json) == {
        "adapter",
        "adapter_contract_version",
        "request_sha256",
        "result_sha256",
        "elapsed_ms",
    }
    # Registers nothing: no new Document, and the predecessor is not superseded.
    assert _document_count(session, project) == documents_before
    predecessor = session.scalars(
        select(Document).where(
            Document.project_id == project.id,
            Document.registry_id == "UCM-REV-2",
        )
    ).one()
    assert predecessor.superseded_by is None


def test_completed_receipt_digests_and_bounds_are_pinned(session, project, tmp_path):
    """The retained receipt's digests and bounds for one fixture, pinned exactly.

    The draft lane runs on the shared bounded-explanation loop; these values are
    what an earlier receipt already holds for the same inputs, so they may not
    move when the loop's implementation does. Digests over inputs that vary per
    run (the workbook bytes, the assigned Document ids) are pinned to the retained
    encoding spelled out here instead of to a literal.
    """
    predecessor = _registered_document(
        session, project, registry_id="UCM-REV-2", sha="a" * 64
    )
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    adapter = FakeAdapter(result=_valid_result())

    receipt = _draft(session, project, staged, adapter)

    def retained_sha256(payload) -> str:
        return hashlib.sha256(
            json.dumps(
                payload,
                sort_keys=True,
                separators=(",", ":"),
                ensure_ascii=False,
                default=str,
            ).encode("utf-8")
        ).hexdigest()

    call = adapter.calls[0]
    assert receipt.status == "completed"
    assert receipt.adapter == "fake-intake-draft"
    assert receipt.adapter_contract_version == "fake-adapter-v1"
    assert receipt.source_sha256 == (
        "9d413d0c6634722b5224f93d4f6f139f5f42f05819a2c59d0bf3e51ef1ec39bf"
    )
    assert receipt.read_fingerprint == retained_sha256(
        {
            "project_id": project.id,
            "staged_sha256": staged.sha256,
            "doc_type": "matrix",
            "documents": [
                {
                    "id": predecessor.id,
                    "registry_id": "UCM-REV-2",
                    "sha256": "a" * 64,
                    "superseded_by": None,
                }
            ],
        }
    )
    assert receipt.state_token == retained_sha256(
        {
            "project_id": project.id,
            "staged_sha256": staged.sha256,
            "registered_registry_ids": ["UCM-REV-2"],
            "known_facts": {"already_registered": False},
        }
    )
    lineage = receipt.execution_lineage_json
    assert lineage["adapter"] == "fake-intake-draft"
    assert lineage["adapter_contract_version"] == "fake-adapter-v1"
    assert lineage["request_sha256"] == retained_sha256([call.system, call.user])
    assert lineage["result_sha256"] == (
        "8778c2e0303571ec038030f9d19cdfd824b59d4290d59bb175cc86cca76fc35d"
    )
    assert isinstance(lineage["elapsed_ms"], int)
    assert receipt.budget_json == {
        "max_input_tokens": 50_000,
        "max_output_tokens": 2_000,
        "timeout_seconds": 30,
        "max_requests": 1,
        "retry_policy": "none",
    }
    assert receipt.usage_json == {
        "estimated_input_tokens": (len(call.system) + len(call.user) + 3) // 4,
        "reported": {},
    }


def test_a_pdf_source_is_read_page_by_page_from_its_text_layer(session, project, tmp_path):
    """A printed source's pages are its text layer, and a quote is bound to its page.

    The registry number is printed on page 1 and the document date on page 2.
    Cited where they are printed, both are retained and the frozen source the
    receipt keeps shows each on its own page; the date cited on page 1, where
    it is not printed, is refused as not literal on that page. A reading that
    ran the pages together, or numbered them from zero, would accept the
    misplaced citation.
    """
    _declare_config(session, project)
    pages = (
        ("Utility Conflict Matrix Rev 3", "Registry number: UCM-REV-3"),
        ("Document date: 2024-03-01",),
    )

    def result(date_ref: str):
        return {
            "metadata_suggestions": [
                {
                    "field": "registry_id",
                    "value": "UCM-REV-3",
                    "source_ref": "P1",
                    "source_quote": "Registry number: UCM-REV-3",
                    "basis": "The cover prints a registry number.",
                },
                {
                    "field": "doc_date",
                    "value": "2024-03-01",
                    "source_ref": date_ref,
                    "source_quote": "Document date: 2024-03-01",
                    "basis": "The second page states the document's date.",
                },
            ],
            "replacement_proposals": [],
            "uncertainties": [],
        }

    staged = _staged_pdf(tmp_path, pages)
    adapter = FakeAdapter(result=result("P2"))
    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "completed"
    assert {
        (item["field"], item["source_ref"])
        for item in receipt.proposals_json["metadata_suggestions"]
    } == {("registry_id", "P1"), ("doc_date", "P2")}
    frozen_pages = receipt.source_json["pages"]
    assert [(page["page_ref"], page["page_no"], page["text_source"]) for page in frozen_pages] == [
        ("P1", 1, "text_layer"),
        ("P2", 2, "text_layer"),
    ]
    assert "Registry number: UCM-REV-3" in frozen_pages[0]["text"]
    assert "Document date" not in frozen_pages[0]["text"]
    assert "Document date: 2024-03-01" in frozen_pages[1]["text"]
    assert len(adapter.calls) == 1

    # A configuration permits one request, so the misplaced citation spends a
    # newly declared one on a source with the same pages and its own identity.
    _declare_config(session, project)
    misplaced = _staged_pdf(tmp_path, pages, name="misplaced.pdf", identity="misplaced")
    receipt = _draft(session, project, misplaced, FakeAdapter(result=result("P1")))

    assert receipt.status == "validation_refused"
    assert receipt.reason == "a cited passage is not literal text on its permitted page"


def test_repeated_request_reuses_receipt_and_spends_once(session, project, tmp_path):
    _registered_document(session, project, registry_id="UCM-REV-2", sha="a" * 64)
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    adapter = FakeAdapter(result=_valid_result())

    first = _draft(session, project, staged, adapter)
    second = _draft(session, project, staged, adapter)

    assert first.id == second.id
    assert _receipt_count(session, project) == 1
    assert len(adapter.calls) == 1


# --- spend / configuration boundary -----------------------------------------


def test_missing_configuration_refuses_before_any_model_call(
    session, project, tmp_path
):
    staged = _staged_workbook(tmp_path)
    adapter = FakeAdapter(result=_valid_result())

    with pytest.raises(ConfigurationRequired):
        _draft(session, project, staged, adapter)

    assert adapter.calls == []
    assert _receipt_count(session, project) == 0


def test_configuration_rejects_uninstalled_prompt_and_bad_bounds(session, project):
    from corridor.source_intake_draft import InvalidDraftConfiguration

    with pytest.raises(InvalidDraftConfiguration):
        _declare_config(session, project, prompt_version="not-installed")
    with pytest.raises(InvalidDraftConfiguration):
        _declare_config(session, project, max_input_tokens=0)


# --- hostile output ----------------------------------------------------------


def test_fabricated_passage_is_refused(session, project, tmp_path):
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    result = _valid_result()
    result["replacement_proposals"] = []
    result["metadata_suggestions"] = [
        {
            "field": "doc_date",
            "value": "1999-01-01",
            "source_ref": "P1",
            "source_quote": "Document date: 1999-01-01",  # never on the page
            "basis": "An invented reading.",
        }
    ]
    adapter = FakeAdapter(result=result)

    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "validation_refused"
    assert receipt.proposals_json is None
    assert "not literal text" in receipt.reason


def test_unsupported_metadata_field_is_refused(session, project, tmp_path):
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    result = _valid_result()
    result["replacement_proposals"] = []
    result["metadata_suggestions"] = [
        {
            "field": "doc_type",  # the person declares the kind, not the model
            "value": "matrix",
            "source_ref": "P1",
            "source_quote": "Utility Conflict Matrix Rev 3",
            "basis": "Classifying the source.",
        }
    ]
    adapter = FakeAdapter(result=result)

    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "validation_refused"
    assert "unsupported field" in receipt.reason


def test_metadata_value_not_in_its_quote_is_refused(session, project, tmp_path):
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    result = _valid_result()
    result["replacement_proposals"] = []
    result["metadata_suggestions"] = [
        {
            "field": "doc_date",
            "value": "2025-12-31",  # a real page quote, but not the value it states
            "source_ref": "P1",
            "source_quote": "Document date: 2024-03-01",
            "basis": "Mismatched value and passage.",
        }
    ]
    adapter = FakeAdapter(result=result)

    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "validation_refused"
    assert "not stated by its cited passage" in receipt.reason


def test_unregistered_replacement_predecessor_is_refused(session, project, tmp_path):
    # No UCM-REV-2 is registered in this project.
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    adapter = FakeAdapter(result=_valid_result())

    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "validation_refused"
    assert "not a document registered in this project" in receipt.reason


def test_cross_project_predecessor_is_refused(session, project, tmp_path):
    other = Project(slug="foreign", name="Foreign", is_synthetic=True)
    session.add(other)
    session.flush()
    # The predecessor registry id exists only in the foreign project.
    _registered_document(session, other, registry_id="UCM-REV-2", sha="e" * 64)
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    adapter = FakeAdapter(result=_valid_result())

    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "validation_refused"
    assert "not a document registered in this project" in receipt.reason


def test_authority_shaped_basis_is_refused(session, project, tmp_path):
    _registered_document(session, project, registry_id="UCM-REV-2", sha="a" * 64)
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    result = _valid_result()
    result["replacement_proposals"][0]["basis"] = (
        "UCM-REV-2 is superseded and this is now registered."
    )
    adapter = FakeAdapter(result=result)

    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "validation_refused"
    assert "asserts an effective act" in receipt.reason


def test_extra_structured_field_is_refused(session, project, tmp_path):
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    result = _valid_result()
    result["registered"] = True
    adapter = FakeAdapter(result=result)

    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "validation_refused"
    assert "authority-shaped field" in receipt.reason


def test_a_passage_on_an_unpermitted_page_is_refused(session, project, tmp_path):
    _declare_config(session, project)
    # Two sheets; only page 1 is permitted, but the model cites page 2.
    staged = _staged_workbook(
        tmp_path,
        sheets=[
            ["Cover only"],
            ["Registry number: UCM-REV-3"],
        ],
    )
    result = _valid_result()
    result["replacement_proposals"] = []
    result["metadata_suggestions"] = [
        {
            "field": "registry_id",
            "value": "UCM-REV-3",
            "source_ref": "P2",
            "source_quote": "Registry number: UCM-REV-3",
            "basis": "Reading a page it was not permitted.",
        }
    ]
    adapter = FakeAdapter(result=result)

    receipt = _draft(
        session, project, staged, adapter, permitted_pages=(1,)
    )

    assert receipt.status == "validation_refused"
    assert "page that was not permitted" in receipt.reason


def test_model_text_is_sanitized_before_storage(session, project, tmp_path):
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    result = _valid_result()
    result["replacement_proposals"] = []
    result["metadata_suggestions"] = [
        {
            "field": "doc_date",
            "value": "2024-03-01",
            "source_ref": "P1",
            "source_quote": "Document date: 2024-03-01",
            "basis": "Clean\x07 basis\x00 text.",
        }
    ]
    adapter = FakeAdapter(result=result)

    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "completed"
    stored = receipt.proposals_json["metadata_suggestions"][0]["basis"]
    assert "\x07" not in stored and "\x00" not in stored
    assert stored == "Clean basis text."


# --- binding: stale bytes / stale registry / overwrite ----------------------


def test_stale_bytes_refuse_without_a_model_call(session, project, tmp_path):
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    adapter = FakeAdapter(result=_valid_result())

    with pytest.raises(IntakeDraftRefused) as excinfo:
        _draft(session, project, staged, adapter, expected_sha256="0" * 64)

    assert excinfo.value.reason == "stale_input"
    assert adapter.calls == []
    assert _receipt_count(session, project) == 0


def test_missing_staged_bytes_refuse(session, project, tmp_path):
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    staged.stored_path.unlink()
    adapter = FakeAdapter(result=_valid_result())

    with pytest.raises(IntakeDraftRefused) as excinfo:
        _draft(session, project, staged, adapter)

    assert excinfo.value.reason == "bytes_missing"
    assert adapter.calls == []


def test_stale_registry_state_refuses(session, project, tmp_path):
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)
    adapter = FakeAdapter(result=_valid_result())

    with pytest.raises(IntakeDraftRefused) as excinfo:
        _draft(session, project, staged, adapter, state_token="f" * 64)

    assert excinfo.value.reason == "stale_input"
    assert adapter.calls == []
    assert _receipt_count(session, project) == 0


def test_registry_change_during_the_request_is_a_stale_receipt(
    session, project, tmp_path
):
    """A document registered while the model call is in flight is a stale read.

    Preparation passes with the registry as previewed; the adapter double then
    registers another document before it answers, so the read fingerprint no
    longer matches what the draft was read against. The receipt says so and
    keeps no proposal. A loop that skipped the after-call re-check would store
    this same answer as a completed draft.
    """
    _registered_document(session, project, registry_id="UCM-REV-2", sha="a" * 64)
    _declare_config(session, project)
    staged = _staged_workbook(tmp_path)

    def answer_after_registering_another_document(call):
        _registered_document(session, project, registry_id="UCM-REV-4", sha="b" * 64)
        return _valid_result()

    adapter = FakeAdapter(result=answer_after_registering_another_document)

    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "stale_input"
    assert receipt.reason == (
        "the staged bytes or registered documents changed during the request"
    )
    assert receipt.proposals_json is None
    assert len(adapter.calls) == 1
    assert receipt.execution_lineage_json is not None
    assert _receipt_count(session, project) == 1


def test_a_suggestion_would_not_overwrite_a_known_registered_fact(
    session, project, tmp_path
):
    staged = _staged_workbook(tmp_path)
    # These exact bytes are already registered with a known date.
    _registered_document(
        session,
        project,
        registry_id="UCM-REV-3",
        sha=staged.sha256,
        doc_date=date(2020, 1, 1),
    )
    _registered_document(session, project, registry_id="UCM-REV-2", sha="a" * 64)
    _declare_config(session, project)
    result = _valid_result()
    result["replacement_proposals"] = []
    result["metadata_suggestions"] = [
        {
            "field": "doc_date",
            "value": "2024-03-01",
            "source_ref": "P1",
            "source_quote": "Document date: 2024-03-01",
            "basis": "Would change a known registered date.",
        }
    ]
    adapter = FakeAdapter(result=result)

    receipt = _draft(session, project, staged, adapter)

    assert receipt.status == "validation_refused"
    assert "overwrite the known registered" in receipt.reason


# --- configuration is append-only, retained authority -----------------------


def test_configuration_is_recorded_with_its_actor(session, project):
    config = _declare_config(session, project)
    stored = session.scalars(
        select(SourceIntakeDraftConfiguration).where(
            SourceIntakeDraftConfiguration.project_id == project.id
        )
    ).one()
    assert stored.id == config.id
    assert stored.created_by == CURATOR.subject
    assert stored.prompt_version == "source_intake_draft_v1"
