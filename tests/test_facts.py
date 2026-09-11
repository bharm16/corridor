"""Typed structured-cell facts through completed extraction and replay seams."""

import os
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from hashlib import sha256

from openpyxl import Workbook
import pytest
from sqlalchemy import select
from sqlalchemy.exc import DBAPIError, IntegrityError

from corridor.candidates import propose
from corridor.current_record import (
    read_current_project_record,
    read_project_record_as_of_revision,
)
from corridor.adjudicate import ImmutableExtractedProposal, edit_candidate
from corridor.config import settings
from corridor.extract_sheet import (
    PROMPT_VERSION,
    SCHEMA_VERSION,
    extract_document,
)
from corridor.extraction_runs import (
    SourceFactAppendConflict,
    append_source_facts,
    record_extraction_run,
)
from corridor.fact_decisions import include_current_structured_cell_facts
from corridor.facts import (
    AppliesToFactValue,
    ClosureFactValue,
    FACT_TYPE_CONTRACTS,
    FactReplayMismatch,
    FactValidationError,
    correct_fact,
    proposal_input_snapshots,
    replay_fact,
    replay_spreadsheet_facts,
)
from corridor.ingest import ingest_document
from corridor.row_accounting import RowAccounting
from corridor.models import (
    ActiveExtractionRun,
    Fact,
    FactAppliesTo,
    FactClosureResult,
    FactClosureSource,
    FactSource,
    ExtractionRun,
    Document,
    Dependency,
    ExtractedProposal,
    ExtractedProposalFact,
    ExternalOrg,
    FactDisposition,
    Project,
    SourceSegment,
    SourceFactAppendReceipt,
)
from corridor.principals import HumanPrincipal
from corridor.sheets import conflict_sheet, read_workbook


HEADINGS = [
    "Utility Conflict ID",
    "Utility Owner",
    "Utility Type",
    "Start Station",
    "End Station",
    "Utility Conflict Description",
    "Promised For",
    "Action Due Date",
    "Required By",
    "Resolution Status",
    "Applies To",
]
REAL_WORKBOOK_SHA256 = (
    "3cd94fea058a3e61ac95ab1efd566e684e146f64ce6f25048d93cf9db55f83ba"
)
CORPUS_STORE = Path(
    os.environ.get("CORRIDOR_TEST_CORPUS_STORE", settings.corpus_store)
)
REAL_WORKBOOK = (
    CORPUS_STORE / REAL_WORKBOOK_SHA256[:2] / f"{REAL_WORKBOOK_SHA256}.xlsx"
)
needs_corpus = pytest.mark.skipif(
    not REAL_WORKBOOK.exists(),
    reason="run `make corpus` to fetch the I-35 NEX South workbook",
)


def _workbook(tmp_path, name, rows):
    path = tmp_path / name
    book = Workbook()
    sheet = book.active
    sheet.title = "Utility Conflicts"
    sheet.append(["Utility Conflict Management (UCM) - Utility Conflicts"])
    sheet.append(HEADINGS)
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def _ingest(session, project, path, tmp_path):
    document = ingest_document(
        session,
        project_id=project.id,
        path=path,
        doc_type="matrix",
        images_dir=tmp_path / "images",
    )
    document._stored_path = str(path)
    return document


def _complete_extraction(session, document, monkeypatch):
    import corridor.extract_sheet as extract_sheet

    monkeypatch.setattr(
        extract_sheet,
        "stored_file",
        lambda value: getattr(value, "_stored_path", None),
    )
    candidates = extract_document(session, document)
    result = _append_request(
        session,
        document,
        candidates,
        idempotency_key=f"test:document:{document.id}",
    )
    return result.run, candidates


def _append_request(session, document, candidates, **overrides):
    values = {
        "idempotency_key": "extract:stationing:test",
        "prompt_version": PROMPT_VERSION,
        "schema_version": SCHEMA_VERSION,
        "candidate_count": len(candidates),
        "page_errors": 0,
        "candidates": candidates,
        "model": None,
        "row_accounting_json": candidates.row_accounting,
        "allow_unsealed_legacy": True,
        "source_path": getattr(document, "_stored_path", None),
    }
    values.update(overrides)
    return append_source_facts(session, document, **values)


@pytest.mark.parametrize("stage", ["segments", "run", "facts", "receipt"])
def test_scoped_append_failure_after_each_stage_leaves_no_partial_spine_rows(
    session, project, tmp_path, monkeypatch, stage
):
    path = _workbook(
        tmp_path,
        f"failure-{stage}.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = Document(
        project_id=project.id,
        sha256=sha256(path.read_bytes()).hexdigest(),
        filename=path.name,
        doc_type="matrix",
        numbering_scheme="project-unique",
        pages=1,
        parse_status="parsed",
    )
    session.add(document)
    session.flush()
    document._stored_path = str(path)
    candidate = propose(
        document,
        kind="dependency",
        fields={"station_from": "1149+00", "station_to": "1150+00"},
        page_no=1,
        quote="UC-1 CenterPoint Electric 1149+00 1150+00 Pole",
        quote_verified=True,
        whole_row=True,
        confidence=None,
        prompt_version=PROMPT_VERSION,
        dedupe="UC-1",
        text_source="cells",
        tier="native",
    )
    candidate.payload_json["citations"][0].update(
        {"table_row": 1, "sheet_name": "Utility Conflicts"}
    )
    accounting = RowAccounting(
        reader_version=PROMPT_VERSION, reader_path="spreadsheet_cells"
    )
    accounting.detect("sheet:1:row:1", page=1, row_number=1)
    accounting.account(
        "sheet:1:row:1", disposition="extracted", reason="candidate_recorded"
    )
    candidates = accounting.finish([candidate])
    segment_count = len(
        session.scalars(
            select(SourceSegment).where(SourceSegment.document_id == document.id)
        ).all()
    )

    with pytest.raises(RuntimeError, match=f"failure after {stage}"):
        _append_request(
            session,
            document,
            candidates,
            idempotency_key=f"failure:{stage}",
            fail_after_stage=stage,
        )

    assert len(
        session.scalars(
            select(SourceSegment).where(SourceSegment.document_id == document.id)
        ).all()
    ) == 0
    assert session.scalars(select(Fact)).all() == []
    assert session.scalars(select(SourceFactAppendReceipt)).all() == []
    assert session.scalars(select(ExtractionRun)).all() == []


def test_scoped_append_replay_returns_original_rows_and_key_conflict_fails(
    session, project, tmp_path, monkeypatch
):
    path = _workbook(
        tmp_path,
        "idempotent.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)
    import corridor.extract_sheet as extract_sheet

    monkeypatch.setattr(extract_sheet, "stored_file", lambda _value: path)
    candidates = extract_document(session, document)

    first = _append_request(session, document, candidates)
    replayed = _append_request(session, document, candidates)
    replayed_with_other_key = _append_request(
        session, document, candidates, idempotency_key="extract:stationing:other"
    )

    assert first.created is True
    assert replayed.created is False
    assert replayed.run.id == first.run.id
    assert replayed_with_other_key.created is False
    assert replayed_with_other_key.run.id == first.run.id
    assert [fact.id for fact in replayed.facts] == [fact.id for fact in first.facts]
    assert len(session.scalars(select(ExtractionRun)).all()) == 1
    assert len(session.scalars(select(Fact)).all()) == 6
    assert len(session.scalars(select(SourceFactAppendReceipt)).all()) == 1

    with pytest.raises(SourceFactAppendConflict, match="different content"):
        _append_request(session, document, candidates, model="different-model")

    second_candidates = extract_document(session, document)
    second = _append_request(
        session,
        document,
        second_candidates,
        idempotency_key="extract:stationing:second-content",
        schema_version="sheet_candidate_shape_variant",
    )
    assert second.created is True
    with pytest.raises(SourceFactAppendConflict, match="different content"):
        _append_request(
            session,
            document,
            second_candidates,
            idempotency_key="extract:stationing:test",
            schema_version="sheet_candidate_shape_variant",
        )


def test_concurrent_identical_retries_return_one_original_result(
    runtime_database, tmp_path
):
    path = _workbook(
        tmp_path,
        "concurrent.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    factory = runtime_database.session_factory
    with factory.begin() as session:
        project = Project(slug="concurrent-facts", name="Concurrent Facts", is_synthetic=True)
        session.add(project)
        session.flush()
        document = ingest_document(
            session,
            project_id=project.id,
            path=path,
            doc_type="matrix",
            images_dir=tmp_path / "images",
        )
        document_id = document.id

    barrier = Barrier(2)

    def run_retry():
        with factory.begin() as session:
            document = session.get(Document, document_id)
            candidate = propose(
                document,
                kind="dependency",
                fields={"station_from": "1149+00", "station_to": "1150+00"},
                page_no=1,
                quote="UC-1 CenterPoint Electric 1149+00 1150+00 Pole",
                quote_verified=True,
                whole_row=True,
                confidence=None,
                prompt_version=PROMPT_VERSION,
                dedupe="UC-1",
                text_source="cells",
                tier="native",
            )
            candidate.payload_json["citations"][0].update(
                {"table_row": 1, "sheet_name": "Utility Conflicts"}
            )
            accounting = RowAccounting(
                reader_version=PROMPT_VERSION,
                reader_path="spreadsheet_cells",
            )
            accounting.detect("sheet:1:row:1", page=1, row_number=1)
            accounting.account(
                "sheet:1:row:1", disposition="extracted", reason="candidate_recorded"
            )
            accounted = accounting.finish([candidate])
            barrier.wait()
            result = append_source_facts(
                session,
                document,
                idempotency_key="concurrent:same",
                prompt_version=PROMPT_VERSION,
                schema_version=SCHEMA_VERSION,
                candidate_count=1,
                page_errors=0,
                candidates=accounted,
                model=None,
                row_accounting_json=accounted.row_accounting,
                allow_unsealed_legacy=True,
                source_path=path,
            )
            return result.run.id, result.created

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _index: run_retry(), range(2)))

    assert len({run_id for run_id, _created in results}) == 1
    assert sorted(created for _run_id, created in results) == [False, True]
    with factory() as session:
        assert len(session.scalars(select(ExtractionRun)).all()) == 1
        assert len(session.scalars(select(Fact)).all()) == 2
        assert len(session.scalars(select(SourceFactAppendReceipt)).all()) == 1


def test_concurrent_unkeyed_appends_of_one_reading_return_one_original_result(
    runtime_database, tmp_path
):
    """The production shape of the same overlap: no key, only the content (#925).

    `record_routed_run` passes `idempotency_key=None` for every `SOURCE_FACTS`
    route, so the key is `content:<digest>` over the document's sha256, the
    prompt, schema, model and sealed configuration, the token usage, the row
    accounting, and every Candidate payload. That is a different index from
    the test above and a different branch inside the command -- the lookup by
    `uq_source_fact_append_content` that runs when no receipt carries the key
    -- and the project-processing pass reaches only this one.

    Two overlapping workers matter here because an expired work lease is
    re-claimable under a worker that is still running (#918), and nothing in
    `extract_project` re-checks a claim. This says what `lock_project` plus the
    content digest are worth when that happens: one run, one receipt, one set
    of Facts, and the loser is told it created nothing.

    It is worth saying what this does *not* prove, because the digest is over
    the reading and not over the document: two workers that read the same bytes
    to different rows, or to the same rows with different token usage, produce
    different digests and both append. The reader here is the deterministic
    spreadsheet one, so identical content is the honest case to assert, and a
    model-backed route cannot rely on it.
    """

    path = _workbook(
        tmp_path,
        "unkeyed.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    factory = runtime_database.session_factory
    with factory.begin() as session:
        project = Project(slug="unkeyed-facts", name="Unkeyed Facts", is_synthetic=True)
        session.add(project)
        session.flush()
        document = ingest_document(
            session,
            project_id=project.id,
            path=path,
            doc_type="matrix",
            images_dir=tmp_path / "images",
        )
        document_id = document.id

    barrier = Barrier(2)

    def append_without_a_key():
        with factory.begin() as session:
            document = session.get(Document, document_id)
            document._stored_path = str(path)
            candidate = propose(
                document,
                kind="dependency",
                fields={"station_from": "1149+00", "station_to": "1150+00"},
                page_no=1,
                quote="UC-1 CenterPoint Electric 1149+00 1150+00 Pole",
                quote_verified=True,
                whole_row=True,
                confidence=None,
                prompt_version=PROMPT_VERSION,
                dedupe="UC-1",
                text_source="cells",
                tier="native",
            )
            candidate.payload_json["citations"][0].update(
                {"table_row": 1, "sheet_name": "Utility Conflicts"}
            )
            accounting = RowAccounting(
                reader_version=PROMPT_VERSION,
                reader_path="spreadsheet_cells",
            )
            accounting.detect("sheet:1:row:1", page=1, row_number=1)
            accounting.account(
                "sheet:1:row:1", disposition="extracted", reason="candidate_recorded"
            )
            accounted = accounting.finish([candidate])
            barrier.wait()
            result = append_source_facts(
                session,
                document,
                idempotency_key=None,
                prompt_version=PROMPT_VERSION,
                schema_version=SCHEMA_VERSION,
                candidate_count=1,
                page_errors=0,
                candidates=accounted,
                model=None,
                row_accounting_json=accounted.row_accounting,
                allow_unsealed_legacy=True,
                source_path=path,
            )
            return result.run.id, result.created

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _index: append_without_a_key(), range(2)))

    assert len({run_id for run_id, _created in results}) == 1
    assert sorted(created for _run_id, created in results) == [False, True]
    with factory() as session:
        receipts = session.scalars(select(SourceFactAppendReceipt)).all()
        assert len(receipts) == 1
        # The key was derived, not given: this is the branch under test.
        assert receipts[0].idempotency_key == f"content:{receipts[0].content_sha256}"
        assert len(session.scalars(select(ExtractionRun)).all()) == 1
        assert len(session.scalars(select(Fact)).all()) == 2


def test_completed_native_extraction_keeps_stationing_typed_inside_the_full_row(
    session, project, tmp_path, monkeypatch
):
    path = _workbook(
        tmp_path,
        "stationing.xlsx",
        [
            ["UC-1", "CenterPoint", "Electric", " 1149+00 ", "1150+00", "Pole"],
            ["UC-2", "AT&T", "Telecom", "1151+00", "1152+00", "Duct"],
        ],
    )
    document = _ingest(session, project, path, tmp_path)

    run, _candidates = _complete_extraction(session, document, monkeypatch)

    facts = session.scalars(
        select(Fact)
        .where(Fact.fact_type.in_(("station_from", "station_to")))
        .order_by(Fact.id)
    ).all()
    assert [
        (
            fact.fact_type,
            fact.subject_kind,
            fact.subject_key,
            fact.text_value,
            fact.transformation,
            fact.document_id,
            fact.extraction_run_id,
        )
        for fact in facts
    ] == [
        (
            "station_from",
            "source_row",
            "Utility Conflicts!3",
            "1149+00",
            "trim_cell_text_v1",
            document.id,
            run.id,
        ),
        (
            "station_to",
            "source_row",
            "Utility Conflicts!3",
            "1150+00",
            "trim_cell_text_v1",
            document.id,
            run.id,
        ),
        (
            "station_from",
            "source_row",
            "Utility Conflicts!4",
            "1151+00",
            "trim_cell_text_v1",
            document.id,
            run.id,
        ),
        (
            "station_to",
            "source_row",
            "Utility Conflicts!4",
            "1152+00",
            "trim_cell_text_v1",
            document.id,
            run.id,
        ),
    ]
    sources = session.scalars(select(FactSource).order_by(FactSource.fact_id)).all()
    assert [(row.role, row.ordinal) for row in sources] == [
        ("value_source", 1)
    ] * 12
    assert all(not hasattr(row, "exact_text") for row in sources)
    assert run.candidate_inputs_json is None
    assert {candidate.extraction_run_id for candidate in _candidates} == {run.id}
    proposals = session.scalars(select(ExtractedProposal).order_by(ExtractedProposal.id)).all()
    assert [(row.subject_key, row.kind) for row in proposals] == [
        ("Utility Conflicts!3", "dependency"),
        ("Utility Conflicts!4", "dependency"),
    ]
    assert len(session.scalars(select(ExtractedProposalFact)).all()) == 12
    sealed = [
        snapshot["payload_json"]["fields"]
        for snapshot in proposal_input_snapshots(session, run)
    ]
    assert sealed == [
        {
            "utility_id": "UC-1",
            "external_org": "CenterPoint",
            "utility_type": "Electric",
            "station_from": "1149+00",
            "station_to": "1150+00",
            "conflict_description": "Pole",
        },
        {
            "utility_id": "UC-2",
            "external_org": "AT&T",
            "utility_type": "Telecom",
            "station_from": "1151+00",
            "station_to": "1152+00",
            "conflict_description": "Duct",
        },
    ]
    _candidates[0].state = "rejected"
    _candidates[0].payload_json = {"fields": {"station_from": "drifted"}}
    session.flush()
    snapshots_after_live_change = proposal_input_snapshots(session, run)
    assert snapshots_after_live_change[0]["state"] == "pending"
    assert snapshots_after_live_change[0]["payload_json"]["fields"] == {
        "utility_id": "UC-1",
        "external_org": "CenterPoint",
        "utility_type": "Electric",
        "station_from": "1149+00",
        "station_to": "1150+00",
        "conflict_description": "Pole",
    }


def test_completed_native_extraction_appends_every_mapped_cell_as_a_typed_fact(
    session, project, tmp_path, monkeypatch
):
    path = _workbook(
        tmp_path,
        "structured-facts.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)

    run, _candidates = _complete_extraction(session, document, monkeypatch)

    facts = session.scalars(select(Fact).order_by(Fact.fact_type)).all()
    assert [fact.fact_type for fact in facts] == [
        "conflict_description",
        "external_org",
        "station_from",
        "station_to",
        "utility_id",
        "utility_type",
    ]
    assert {fact.extraction_run_id for fact in facts} == {run.id}
    assert {fact.subject_key for fact in facts} == {"Utility Conflicts!3"}
    assert {fact.transformation for fact in facts} == {"trim_cell_text_v1"}
    assert {fact.text_value for fact in facts} == {
        "UC-1",
        "CenterPoint",
        "Electric",
        "1149+00",
        "1150+00",
        "Pole",
    }


def test_every_controlled_structured_cell_type_has_a_complete_contract():
    expected = {
        "utility_id", "external_org", "external_org_contact", "utility_type",
        "utility_subtype", "utility_function", "operational_status", "size",
        "material", "oh_ug", "row_placement", "orientation", "baseline",
        "station_from", "station_to", "offset_from", "offset_to", "sue_level",
        "conflict_description", "resolution_strategy", "notes", "alignment",
        "location_start", "location_end", "offset_side", "potential_conflict",
        "data_source", "marked_resolution", "committed_date", "action_due_date",
        "need_date", "applies_to", "closure_result",
    }

    assert expected < set(FACT_TYPE_CONTRACTS)
    for fact_type in expected:
        contract = FACT_TYPE_CONTRACTS[fact_type]
        assert contract.value_class
        assert contract.transformation
        assert contract.required_roles == frozenset({"value_source"})
        assert contract.validation_rule
        assert contract.current_value_rule
        assert contract.inclusion_rule


def test_structured_date_and_marked_resolution_cells_use_typed_contracts(
    session, project, tmp_path, monkeypatch
):
    path = _workbook(
        tmp_path,
        "dates-and-resolution.xlsx",
        [[
            "UC-1",
            "CenterPoint",
            "Electric",
            "1149+00",
            "1150+00",
            "Pole",
            "2026-09-15",
            "2026-09-01",
            "2026-08-15",
            "Resolved",
        ]],
    )
    document = _ingest(session, project, path, tmp_path)

    _complete_extraction(session, document, monkeypatch)

    facts = {
        fact.fact_type: fact
        for fact in session.scalars(select(Fact).order_by(Fact.fact_type))
    }
    assert facts["committed_date"].date_value.isoformat() == "2026-09-15"
    assert facts["action_due_date"].date_value.isoformat() == "2026-09-01"
    assert facts["need_date"].date_value.isoformat() == "2026-08-15"
    assert facts["marked_resolution"].text_value == "Resolved"
    assert {
        FACT_TYPE_CONTRACTS[name].current_value_rule
        for name in ("committed_date", "action_due_date", "need_date")
    } == {"latest_effective_single_value"}


def test_external_org_wording_resolves_only_one_exact_registered_spelling(
    session, project, tmp_path, monkeypatch
):
    registered = ExternalOrg(
        name="CenterPoint Energy", org_type="utility", aliases=["CenterPoint"]
    )
    session.add(registered)
    session.flush()
    path = _workbook(
        tmp_path,
        "registered-party.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)

    _complete_extraction(session, document, monkeypatch)

    fact = session.scalar(select(Fact).where(Fact.fact_type == "external_org"))
    assert fact.text_value == "CenterPoint"
    assert fact.external_org_value_id == registered.id


def test_applies_to_and_closure_satellites_round_trip_with_typed_foreign_keys(
    session, project, tmp_path, monkeypatch
):
    from corridor.models import Dependency

    path = _workbook(
        tmp_path,
        "satellites.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)
    run, _candidates = _complete_extraction(session, document, monkeypatch)
    segment = session.scalar(
        select(SourceSegment).where(
            SourceSegment.document_id == document.id,
            SourceSegment.cell_range == "A3",
        )
    )
    first = Dependency(
        project_id=project.id,
        ref_code="DEP-00001",
        dep_type="utility_relocation",
        title="First constraint",
    )
    successor = Dependency(
        project_id=project.id,
        ref_code="DEP-00002",
        dep_type="utility_relocation",
        title="Successor constraint",
    )
    session.add_all((first, successor))
    session.flush()
    applies_to = Fact(
        project_id=project.id,
        document_id=document.id,
        extraction_run_id=run.id,
        fact_type="applies_to",
        subject_kind="source_row",
        subject_key="Utility Conflicts!3",
        text_value=None,
        date_value=None,
        date_range_start=None,
        date_range_end=None,
        external_org_value_id=None,
        document_value_id=None,
        transformation="structured_reference_set_v1",
        recorded_by="extractor:sheet_native_v2",
        content_sha256=sha256(b"satellite-applies-to").hexdigest(),
    )
    closure = Fact(
        project_id=project.id,
        document_id=document.id,
        extraction_run_id=run.id,
        fact_type="closure_result",
        subject_kind="source_row",
        subject_key="Utility Conflicts!3",
        text_value=None,
        date_value=None,
        date_range_start=None,
        date_range_end=None,
        external_org_value_id=None,
        document_value_id=None,
        transformation="typed_closure_result_v1",
        recorded_by="human:closure-reviewer",
        content_sha256=sha256(b"satellite-closure").hexdigest(),
    )
    session.add_all((applies_to, closure))
    session.flush()
    session.add_all(
        (
            FactSource(
                project_id=project.id,
                document_id=document.id,
                fact_id=applies_to.id,
                source_segment_id=segment.id,
                role="value_source",
                ordinal=1,
            ),
            FactSource(
                project_id=project.id,
                document_id=document.id,
                fact_id=closure.id,
                source_segment_id=segment.id,
                role="value_source",
                ordinal=1,
            ),
            FactAppliesTo(
                project_id=project.id,
                fact_id=applies_to.id,
                dependency_id=first.id,
                ordinal=1,
            ),
            FactClosureResult(
                project_id=project.id,
                fact_id=closure.id,
                closure_kind="source_marked_resolved",
                successor_dependency_id=successor.id,
            ),
            FactClosureSource(
                project_id=project.id,
                document_id=document.id,
                fact_id=closure.id,
                source_segment_id=segment.id,
                ordinal=1,
            ),
        )
    )
    session.flush()

    assert session.scalar(
        select(FactAppliesTo.dependency_id).where(FactAppliesTo.fact_id == applies_to.id)
    ) == first.id
    stored = session.scalar(
        select(FactClosureResult).where(FactClosureResult.fact_id == closure.id)
    )
    assert (stored.closure_kind, stored.successor_dependency_id) == (
        "source_marked_resolved",
        successor.id,
    )
    assert session.scalar(
        select(FactClosureSource.source_segment_id).where(
            FactClosureSource.fact_id == closure.id
        )
    ) == segment.id


def test_structured_satellites_extract_replay_include_and_project_current_and_as_of(
    session, project, tmp_path, monkeypatch
):
    targets = [
        Dependency(
            project_id=project.id,
            ref_code=f"DEP-{ordinal:05d}",
            dep_type="utility_relocation",
            title=f"Target {ordinal}",
        )
        for ordinal in (1, 2)
    ]
    subject = Dependency(
        project_id=project.id,
        ref_code="DEP-00003",
        dep_type="utility_relocation",
        title="Subject constraint",
    )
    session.add_all((*targets, subject))
    session.flush()
    path = _workbook(
        tmp_path,
        "structured-satellites.xlsx",
        [[
            "UC-1", "Unknown Utility", "Electric", "1149+00", "1150+00",
            "Pole", "", "", "", "Resolved", "DEP-00001, DEP-00002",
        ]],
    )
    document = _ingest(session, project, path, tmp_path)

    run, candidates = _complete_extraction(session, document, monkeypatch)

    applies_to = session.scalar(select(Fact).where(Fact.fact_type == "applies_to"))
    closure = session.scalar(select(Fact).where(Fact.fact_type == "closure_result"))
    assert replay_fact(session, document, applies_to, path) == AppliesToFactValue(
        dependency_ids=(targets[0].id, targets[1].id)
    )
    closure_value = replay_fact(session, document, closure, path)
    assert isinstance(closure_value, ClosureFactValue)
    assert closure_value.closure_kind == "source_marked_resolved"
    assert closure_value.successor_dependency_id is None
    assert closure_value.governing_source_segment_ids
    assert replay_spreadsheet_facts(session, document, (applies_to, closure), path) == (
        AppliesToFactValue(dependency_ids=(targets[0].id, targets[1].id)), closure_value,
    )

    [candidate] = candidates
    candidate.state = "accepted"
    candidate.merged_into = subject.id
    session.add(
        ActiveExtractionRun(document_id=document.id, extraction_run_id=run.id)
    )
    session.flush()
    decisions = include_current_structured_cell_facts(session, project.id)
    current = {
        value.fact_type: value
        for value in read_current_project_record(session, project.id)
    }
    as_of = {
        value.fact_type: value
        for value in read_project_record_as_of_revision(
            session, project.id, decisions[-1].revision.id
        )
    }

    assert current["applies_to"].applies_to_dependency_ids == (
        targets[0].id,
        targets[1].id,
    )
    assert as_of["applies_to"].applies_to_dependency_ids == (
        targets[0].id,
        targets[1].id,
    )
    assert current["closure_result"].closure_kind == "source_marked_resolved"
    assert as_of["closure_result"].closure_governing_source_segment_ids
    assert "external_org" not in current


def test_source_reading_correction_appends_successor_and_disposition(
    session, project, tmp_path, monkeypatch
):
    path = _workbook(
        tmp_path,
        "correction.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)
    _complete_extraction(session, document, monkeypatch)
    predecessor = session.scalar(select(Fact).where(Fact.fact_type == "station_from"))
    replacement_segment = session.scalar(
        select(SourceSegment).where(
            SourceSegment.document_id == document.id,
            SourceSegment.cell_range == "E3",
        )
    )

    disposition = correct_fact(
        session,
        predecessor,
        source_segment_id=replacement_segment.id,
        principal=HumanPrincipal("local:fact-editor"),
    )

    successor = session.get(Fact, disposition.successor_fact_id)
    assert predecessor.text_value == "1149+00"
    assert successor.text_value == "1150+00"
    assert disposition.predecessor_fact_id == predecessor.id
    assert disposition.kind == "source_reading_correction"
    assert session.scalars(select(FactDisposition)).one().id == disposition.id


def test_typed_date_correction_keeps_the_value_in_the_date_column(
    session, project, tmp_path, monkeypatch
):
    path = _workbook(
        tmp_path,
        "date-correction.xlsx",
        [
            [
                "UC-1", "CenterPoint", "Electric", "1149+00", "1150+00",
                "Pole", "2026-09-15", "", "", "",
            ],
            [
                "UC-2", "CenterPoint", "Electric", "1151+00", "1152+00",
                "Duct", "2026-10-01", "", "", "",
            ],
        ],
    )
    document = _ingest(session, project, path, tmp_path)
    _complete_extraction(session, document, monkeypatch)
    predecessor = session.scalar(
        select(Fact).where(
            Fact.fact_type == "committed_date",
            Fact.subject_key == "Utility Conflicts!3",
        )
    )
    replacement = session.scalar(
        select(SourceSegment).where(
            SourceSegment.document_id == document.id,
            SourceSegment.cell_range == "G4",
        )
    )

    disposition = correct_fact(
        session,
        predecessor,
        source_segment_id=replacement.id,
        principal=HumanPrincipal("local:fact-editor"),
    )

    successor = session.get(Fact, disposition.successor_fact_id)
    assert successor.text_value is None
    assert successor.date_value.isoformat() == "2026-10-01"


def test_new_proposals_and_correction_edges_have_no_update_path(
    session, project, tmp_path, monkeypatch
):
    path = _workbook(
        tmp_path,
        "proposal-immutable.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)
    _complete_extraction(session, document, monkeypatch)
    proposal = session.scalar(select(ExtractedProposal))
    proposal.subject_key = "rewritten"

    with pytest.raises(DBAPIError, match="Extracted Proposal spine is immutable"):
        session.flush()


def test_legacy_edit_path_refuses_new_spine_backed_proposal(
    session, project, tmp_path, monkeypatch
):
    path = _workbook(
        tmp_path,
        "no-edit.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)
    _run, candidates = _complete_extraction(session, document, monkeypatch)

    with pytest.raises(ImmutableExtractedProposal, match="append a Fact correction"):
        edit_candidate(
            session,
            candidates[0],
            {"station_from": "9999+99"},
            principal=HumanPrincipal("local:fact-editor"),
        )


def test_every_structured_cell_fact_replays_from_its_exact_cell(
    session, project, tmp_path, monkeypatch
):
    path = _workbook(
        tmp_path,
        "replay.xlsx",
        [["UC-1", "CenterPoint", "Electric", " 1149+00 ", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)
    _complete_extraction(session, document, monkeypatch)
    facts = session.scalars(select(Fact).order_by(Fact.id)).all()

    assert [replay_fact(session, document, fact, path) for fact in facts] == [
        "UC-1",
        "CenterPoint",
        "Electric",
        "1149+00",
        "1150+00",
        "Pole",
    ]


@needs_corpus
@pytest.mark.slow
def test_real_dev_corpus_matrix_yields_every_queryable_replayable_cell_fact(
    session, project, tmp_path, monkeypatch
):
    import corridor.source_segments as source_segments

    sheet = conflict_sheet(read_workbook(REAL_WORKBOOK))
    external_org_column = next(
        index for index, field in sheet.mapping.items() if field == "external_org"
    )
    names = sorted(
        {
            row[external_org_column]
            for row in sheet.rows
            if row[external_org_column].strip()
        }
    )
    session.add_all(
        ExternalOrg(name=name, org_type="utility", aliases=[]) for name in names
    )
    session.flush()
    document = _ingest(session, project, REAL_WORKBOOK, tmp_path)
    original_load = source_segments.load_workbook
    workbook_opens = []

    def counted_load(*args, **kwargs):
        workbook_opens.append(1)
        return original_load(*args, **kwargs)

    monkeypatch.setattr(source_segments, "load_workbook", counted_load)
    run, candidates = _complete_extraction(session, document, monkeypatch)
    assert len(workbook_opens) == 1
    facts = session.scalars(select(Fact).order_by(Fact.id)).all()

    assert len(candidates) == 68
    assert len(facts) == 1018
    assert {
        fact_type: sum(fact.fact_type == fact_type for fact in facts)
        for fact_type in {fact.fact_type for fact in facts}
    } == {
        "external_org": 68,
        "external_org_contact": 67,
        "utility_id": 68,
        "utility_type": 68,
        "conflict_description": 68,
        "orientation": 68,
        "row_placement": 68,
        "baseline": 67,
        "station_from": 68,
        "offset_from": 68,
        "station_to": 68,
        "offset_to": 68,
        "sue_level": 68,
        "resolution_strategy": 68,
        "notes": 68,
    }
    assert {fact.extraction_run_id for fact in facts} == {run.id}
    station_from = session.scalar(
        select(Fact)
        .where(Fact.fact_type == "station_from")
        .order_by(Fact.id)
        .limit(1)
    )
    station_to = session.scalar(
        select(Fact)
        .where(Fact.fact_type == "station_to")
        .order_by(Fact.id.desc())
        .limit(1)
    )
    assert (
        station_from.fact_type,
        station_from.subject_key,
        station_from.text_value,
        replay_fact(session, document, station_from, REAL_WORKBOOK),
    ) == ("station_from", "UCM-Conflict List!9", "329312.79", "329312.79")
    assert (
        station_to.fact_type,
        station_to.subject_key,
        station_to.text_value,
        replay_fact(session, document, station_to, REAL_WORKBOOK),
    ) == ("station_to", "UCM-Conflict List!76", "-", "-")
    workbook_opens.clear()
    replayed = replay_spreadsheet_facts(session, document, iter(facts), REAL_WORKBOOK)
    assert len(replayed) == 1018
    assert replayed == tuple(
        fact.date_value if fact.date_value is not None else fact.text_value
        for fact in facts
    )
    assert len(workbook_opens) == 1

    dependencies = []
    for ordinal, candidate in enumerate(candidates, 1):
        dependency = Dependency(
            project_id=project.id,
            ref_code=f"DEP-{ordinal:05d}",
            source_ref=candidate.payload_json["fields"].get("utility_id"),
            dep_type="utility_relocation",
            title=f"Dev corpus constraint {ordinal}",
        )
        session.add(dependency)
        dependencies.append(dependency)
    session.flush()
    for candidate, dependency in zip(candidates, dependencies, strict=True):
        candidate.state = "accepted"
        candidate.merged_into = dependency.id
    session.add(
        ActiveExtractionRun(document_id=document.id, extraction_run_id=run.id)
    )
    session.flush()

    decisions = include_current_structured_cell_facts(session, project.id)
    current = read_current_project_record(session, project.id)
    as_of = read_project_record_as_of_revision(
        session, project.id, decisions[-1].revision.id
    )

    assert len(decisions) == len(current) == len(as_of) == 1018
    expected_types = {fact.fact_type for fact in facts}
    assert {value.fact_type for value in current} == expected_types
    assert {value.fact_type for value in as_of} == expected_types


@pytest.mark.parametrize("batch", [False, True])
def test_fact_replay_fails_closed_when_materialized_value_does_not_reproduce(
    session, project, tmp_path, monkeypatch, batch,
):
    path = _workbook(
        tmp_path,
        "mismatch.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)
    _complete_extraction(session, document, monkeypatch)
    facts = session.scalars(select(Fact).order_by(Fact.id)).all()
    fact = facts[-1]
    fact.text_value = "9999+99"

    with pytest.raises(FactReplayMismatch, match="does not reproduce"):
        if batch:
            replay_spreadsheet_facts(session, document, iter(facts), path)
        else:
            replay_fact(session, document, fact, path)


def test_spreadsheet_fact_batch_refuses_source_changes_before_returning(
    session, project, tmp_path, monkeypatch,
):
    import corridor.facts as fact_replay
    from corridor.source_segments import SourceDocumentDigestMismatch

    path = _workbook(
        tmp_path, "batch.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)
    _complete_extraction(session, document, monkeypatch)
    facts = session.scalars(select(Fact).order_by(Fact.id)).all()
    original_transform = fact_replay.validated_scalar_value
    calls = []

    def change_source_after_first_value(contract, exact):
        value = original_transform(contract, exact)
        if not calls:
            path.write_bytes(path.read_bytes() + b"changed during batch")
        calls.append(exact)
        return value

    monkeypatch.setattr(fact_replay, "validated_scalar_value", change_source_after_first_value)
    with pytest.raises(SourceDocumentDigestMismatch, match="registered Document digest"):
        replay_spreadsheet_facts(session, document, iter(facts), path)
    assert len(calls) == len(facts)
    with pytest.raises(SourceDocumentDigestMismatch, match="registered Document digest"):
        replay_fact(session, document, facts[0], path)


def test_scoped_spreadsheet_append_refuses_changed_snapshot_before_appending_facts(
    session, project, tmp_path, monkeypatch,
):
    import corridor.extract_sheet as extract_sheet
    import corridor.source_segments as source_segments
    from corridor.source_segments import SourceDocumentDigestMismatch

    path = _workbook(
        tmp_path, "append-snapshot.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)
    monkeypatch.setattr(extract_sheet, "stored_file", lambda value: str(path))
    candidates = extract_document(session, document)
    original_load = source_segments.load_workbook

    def change_source_after_snapshot_decode(*args, **kwargs):
        book = original_load(*args, **kwargs)
        path.write_bytes(path.read_bytes() + b"changed during append validation")
        return book

    monkeypatch.setattr(source_segments, "load_workbook", change_source_after_snapshot_decode)
    with pytest.raises(SourceDocumentDigestMismatch, match="registered Document digest"):
        _append_request(session, document, candidates)
    assert session.scalars(select(ExtractionRun)).all() == []
    assert session.scalars(select(Fact)).all() == []
    assert session.scalars(select(SourceFactAppendReceipt)).all() == []


@pytest.mark.parametrize("document_id,project_id", [(2, 1), (1, 2)])
def test_spreadsheet_fact_batch_refuses_a_fact_from_another_document_or_project(
    tmp_path, document_id, project_id,
):
    document = Document(id=1, project_id=1, sha256="0" * 64)
    fact = Fact(document_id=document_id, project_id=project_id)
    with pytest.raises(FactValidationError, match="another Document"):
        replay_spreadsheet_facts(None, document, iter([fact]), tmp_path / "absent.xlsx")


def test_database_rejects_fact_source_from_another_rendition(
    session, project, tmp_path, monkeypatch
):
    first_path = _workbook(
        tmp_path,
        "first.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    first = _ingest(session, project, first_path, tmp_path)
    _complete_extraction(session, first, monkeypatch)
    fact = session.scalars(select(Fact).order_by(Fact.id)).first()

    second_path = _workbook(
        tmp_path,
        "second.xlsx",
        [["UC-2", "AT&T", "Telecom", "210+00", "220+00", "Duct"]],
    )
    second = _ingest(session, project, second_path, tmp_path)
    foreign = session.scalars(
        select(SourceSegment).where(SourceSegment.document_id == second.id)
    ).first()
    session.add(
        FactSource(
            project_id=project.id,
            document_id=first.id,
            fact_id=fact.id,
            source_segment_id=foreign.id,
            role="context",
            ordinal=2,
        )
    )

    with pytest.raises(IntegrityError):
        session.flush()


def test_stationing_fact_type_requires_only_its_typed_text_value(
    session, project, tmp_path
):
    path = _workbook(tmp_path, "typed.xlsx", [])
    document = _ingest(session, project, path, tmp_path)
    fact = Fact(
        project_id=project.id,
        document_id=document.id,
        extraction_run_id=None,
        fact_type="station_from",
        subject_kind="source_row",
        subject_key="Utility Conflicts!3",
        text_value=None,
        date_value=None,
        date_range_start=None,
        date_range_end=None,
        external_org_value_id=None,
        document_value_id=None,
        transformation="trim_cell_text_v1",
        recorded_by="extractor:sheet_native_v2",
        # The digest is supplied so the typed-value constraint is what
        # refuses, not the Fact identity `not null` added by #457.
        content_sha256=sha256(b"stationing-missing-text").hexdigest(),
    )
    session.add(fact)

    with pytest.raises(IntegrityError):
        session.flush()


@pytest.mark.parametrize("operation", ["update", "delete"])
def test_facts_have_no_update_or_delete_path(
    session, project, tmp_path, monkeypatch, operation
):
    path = _workbook(
        tmp_path,
        "immutable.xlsx",
        [["UC-1", "CenterPoint", "Electric", "1149+00", "1150+00", "Pole"]],
    )
    document = _ingest(session, project, path, tmp_path)
    _complete_extraction(session, document, monkeypatch)
    fact = session.scalars(select(Fact).order_by(Fact.id)).first()
    if operation == "update":
        fact.text_value = "rewritten"
    else:
        session.delete(fact)

    with pytest.raises(DBAPIError, match="facts are append-only"):
        session.flush()
