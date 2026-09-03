"""Rendering the accepted record into the approved native UCM workbook (#495).

The deliverable is the customer's own file with the accepted values written into
it, so nearly every assertion here is about what did **not** change: which
package parts stayed byte-identical, which cells were left alone, and which
workbook features refuse the render outright rather than being quietly dropped.

Every workbook is synthetic and exercises the renderer, never a claim about a
real customer form (ADR-0046). The one thing the fixtures do assume about a real
template is the shape #509 already assumes: a title band, a header row of the
published form's exact column names, and one data row per conflict.
"""

from __future__ import annotations

import hashlib
import re
import time
from datetime import date
import zipfile
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from corridor.baseline_adoption import (
    BaselineAdoptionRefused,
    FormatIdentity,
    adopt_baseline,
    effective_baseline_formats,
    effective_field_mapping_manifest,
    preview_baseline_adoption,
    register_baseline_format,
)
from corridor.config import settings
from corridor.db import Session, engine
from corridor.extraction_runs import record_extraction_run
from corridor.extractor_lineage import deployed_extractor_config, zero_token_usage
from corridor.fact_decisions import record_human_fact_decision
from corridor.field_mapping_manifest import (
    COMBINED_RANGE,
    DEMO_EXTERNAL_REFERENCES,
    MappingDeclaration,
    MappingManifestRefused,
    MaterialMapping,
    declared_field_mapping,
)
from corridor.materializer import (
    materialize_segment_value,
    materialize_typed_satellite,
)
from corridor.models import (
    BaselineFormat,
    BaselineFormatManifest,
    Document,
    FactDecision,
    Project,
    ProjectRecordRevision,
)
from corridor.principals import HumanPrincipal
from corridor.source_append import ClosureValues, SegmentValues, append_fact
from corridor.source_append import append_source_segments
from corridor.source_intake import validate_and_stage
from corridor.spreadsheet_conversion import canonical_package
from corridor.workbook_render import (
    PROVENANCE_SHEET_NAME,
    RENDERER_VERSION,
    RenderProfile,
    RetirementMapping,
    RowInsertionTemplate,
    UnapprovedFieldMapping,
    UnapprovedOutputTemplate,
    UndeclaredRenderBehaviour,
    UnsupportedWorkbookFeature,
    WorkbookRenderRefused,
    render_project_record_workbook,
)


PRINCIPAL = HumanPrincipal("local:coordinator")
SHEET = "Utility Conflicts"

# The synthetic profile these fixtures declare by name. The four headings under
# it stopped being production defaults with #597; a fixture that wants an
# external reference says so, exactly as a customer's own mapping revision does.
DEMO = MappingDeclaration(external_references=DEMO_EXTERNAL_REFERENCES)

# The successor form this ticket exists for: `Start Station` carries the whole
# range and `End Station` is blank. Same headings, same drop-downs, same
# canonical fields — nothing a digest over the template's bytes can see.
COMBINED_STATIONS = MaterialMapping(
    source_columns=("Start Station", "End Station"),
    target_fields=("station_from", "station_to"),
    composition=COMBINED_RANGE,
    delimiter=" - ",
    precision="source_stated_unit_v1",
    blank_behaviour="all_absent_is_unknown_v1",
    material=True,
)
COMBINED_DECLARATION = MappingDeclaration(
    version="v2",
    external_references=DEMO_EXTERNAL_REFERENCES,
    mappings=(COMBINED_STATIONS,),
)

HEADINGS = [
    "Utility Conflict ID",
    "Utility Owner",
    "Utility Type",
    "Size",
    "Material",
    "Station Origin",
    "Start Station",
    "End Station",
    "Resolution Strategy Selected (from Resolution Alternatives)",
    "Promised For",
    "Action Due Date",
    "Resolution Status",
    "Comment",
    "UCM Record ID",
    "Record URL",
    "Early TxDOT Utility Activity",
]

ROWS = [
    ["UC-1", "CenterPoint Energy", "Electric", "12 in", "Steel", "SR-BL",
     "1149+00", "1150+00", "Relocate", "2026-03-01", "2026-02-01", "",
     "pole at station", "UCM-1001", "https://ucm.example/records/1001", "Yes"],
    ["UC-2", "City of Austin", "Water", "8 in", "PVC", "SR-BL",
     "1160+00", "1161+00", "Adjust", "2026-04-01", "2026-03-01", "",
     "", "UCM-1002", "", ""],
    ["UC-3", "Oncor", "Electric", "4 in", "Copper", "SR-BL",
     "1180+00", "1181+00", "Relocate", "2026-05-01", "", "", "", "UCM-1003",
     "", ""],
]


@pytest.fixture
def session():
    connection = engine.connect()
    transaction = connection.begin()
    scoped = Session(bind=connection)
    yield scoped
    scoped.close()
    if transaction.is_active:
        transaction.rollback()
    connection.close()


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "corpus_store", str(tmp_path / "files"))
    monkeypatch.setattr(settings, "corpus_images", str(tmp_path / "images"))
    return tmp_path / "files"


@pytest.fixture
def project(session):
    row = Project(
        slug=f"render-{uuid4().hex[:8]}",
        name="Workbook Render Test",
        is_synthetic=True,
    )
    session.add(row)
    session.flush()
    return row


def _workbook_bytes(
    tmp_path,
    name="ucm.xlsx",
    *,
    headings=None,
    rows=None,
    formula_cell=None,
    table=False,
    features=False,
) -> bytes:
    path = tmp_path / name
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(["Utility Conflict Management (UCM) — Utility Conflicts"])
    sheet.append(list(HEADINGS if headings is None else headings))
    for row in ROWS if rows is None else rows:
        sheet.append(list(row))
    if formula_cell is not None:
        sheet[formula_cell] = "=A3"
    if table:
        from openpyxl.worksheet.table import Table

        sheet.add_table(Table(displayName="Conflicts", ref="A2:P5"))
    if features:
        from openpyxl.styles import Font
        from openpyxl.worksheet.datavalidation import DataValidation

        rule = DataValidation(type="list", formula1='"Electric,Water,Gas"')
        sheet.add_data_validation(rule)
        rule.add("C3:C99")
        sheet["B3"].font = Font(bold=True, italic=True)
        sheet.column_dimensions["N"].hidden = True
        sheet.auto_filter.ref = "A2:P5"
        sheet.freeze_panes = "A3"
    book.save(path)
    return _pinned_template(path.read_bytes())


# openpyxl stamps the current time into `docProps/core.xml` while it saves
# (#579), so the fixture template would carry a wall clock of its own and a test
# asserting no clock reaches the output could not tell the two apart. The
# template is pinned the same way `spreadsheet_conversion` pins its converted
# workbooks: every W3CDTF-typed core property is replaced, and the archive is
# rewritten through the shared canonical writer.
_PINNED_TIMESTAMP = b"2000-01-01T00:00:00Z"
_W3CDTF_VALUE = re.compile(rb'(<[^<>]*xsi:type="dcterms:W3CDTF"[^<>]*>)[^<]*')


def _pinned_template(data: bytes) -> bytes:
    parts = _parts(data)
    parts["docProps/core.xml"] = _W3CDTF_VALUE.sub(
        lambda match: match.group(1) + _PINNED_TIMESTAMP, parts["docProps/core.xml"]
    )
    return canonical_package(parts)


@pytest.fixture
def adopted(session, project, tmp_path, store):
    """One adopted baseline, and the exact bytes registered as its template."""

    body = _workbook_bytes(tmp_path)
    staged = validate_and_stage(body, "ucm.xlsx")
    preview = preview_baseline_adoption(
        session,
        project=project,
        staged=staged,
        customer="Lone Star Transit Authority",
        source_identity="UCM workbook revision C",
        field_mapping=DEMO,
    )
    result = adopt_baseline(
        session,
        preview=preview,
        principal=PRINCIPAL,
        idempotency_key="adopt-1",
        images_dir=tmp_path / "images",
    )
    return _Adopted(
        body=body,
        result=result,
        document_id=result.document_id,
        manifest=preview.field_mapping_manifest,
    )


class _Adopted:
    def __init__(self, body, result, document_id, manifest):
        self.body = body
        self.result = result
        self.document_id = document_id
        self.revision_id = result.revision_id
        self.manifest = manifest


def _render(session, project, adopted, **kwargs):
    return render_project_record_workbook(
        session,
        project_id=project.id,
        revision_id=kwargs.pop("revision_id", adopted.revision_id),
        template_bytes=kwargs.pop("template_bytes", adopted.body),
        field_mapping=kwargs.pop("field_mapping", adopted.manifest),
        **kwargs,
    )


def _parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _cells(data: bytes, sheet_name: str = SHEET) -> dict[str, object]:
    book = load_workbook(BytesIO(data), data_only=True)
    try:
        sheet = book[sheet_name]
        return {
            cell.coordinate: cell.value
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        }
    finally:
        book.close()


def _second_document(session, project, digest_seed: str):
    """A second registered source revision, so a later value has a segment.

    A structured-cell Fact stays document- and run-bound, so the reading that
    produced it is registered too rather than left null.
    """

    document = Document(
        project_id=project.id,
        sha256=hashlib.sha256(digest_seed.encode()).hexdigest(),
        filename=f"{digest_seed}.xlsx",
        doc_type="matrix",
    )
    session.add(document)
    session.flush()
    config = deployed_extractor_config("baseline", client=None)
    run = record_extraction_run(
        session,
        document,
        prompt_version=config.prompt_version,
        schema_version=config.schema_version,
        candidate_count=0,
        page_errors=0,
        outcome="completed",
        model=None,
        extractor_config=config,
        token_usage=zero_token_usage(document.id),
        row_accounting_json=None,
    )
    document.extraction_run_id = run.id
    return document


def _segment(session, project, document, *, text, cell_range, ordinal):
    [segment] = append_source_segments(
        session,
        project_id=project.id,
        document_id=document.id,
        recorded_verbal_origin_id=None,
        segments=[
            SegmentValues(
                kind="spreadsheet_cell",
                exact_text=text,
                content_sha256=hashlib.sha256(text.encode("utf-8")).hexdigest(),
                ordinal=ordinal,
                sheet_name=SHEET,
                cell_range=cell_range,
            )
        ],
    )
    return segment


def _effective_decision(session, project, subject_key, fact_type):
    return session.scalar(
        select(FactDecision).where(
            FactDecision.project_id == project.id,
            FactDecision.subject_key == subject_key,
            FactDecision.fact_type == fact_type,
            FactDecision.superseded_by.is_(None),
        )
    )


def _decide(session, project, document, *, subject_key, fact_type, text, cell_range,
            ordinal, key, predecessor=None):
    segment = _segment(
        session, project, document, text=text, cell_range=cell_range, ordinal=ordinal
    )
    fact = append_fact(
        session,
        project_id=project.id,
        document_id=document.id,
        extraction_run_id=document.extraction_run_id,
        subject_kind="source_row",
        subject_key=subject_key,
        recorded_by="human:coordinator",
        content_sha256=hashlib.sha256(
            f"{subject_key}{fact_type}{text}".encode("utf-8")
        ).hexdigest(),
        value=materialize_segment_value(session, fact_type, segment),
    )
    outcome = record_human_fact_decision(
        session,
        fact,
        principal=PRINCIPAL,
        command_type="resolve_discrepancy",
        idempotency_key=key,
        expected_predecessor=predecessor,
    )
    return outcome


# --- The record the template was adopted from -------------------------------


def test_rendering_the_adopted_revision_changes_no_cell_and_no_package_part(
    session, project, adopted
):
    rendered = _render(session, project, adopted)

    assert rendered.changed_cells == ()
    assert rendered.added_rows == ()
    assert rendered.retired_rows == ()
    assert rendered.changed_parts == ()
    template = _parts(adopted.body)
    output = _parts(rendered.content)
    assert set(output) == set(template)
    assert output == template
    assert set(rendered.unchanged_parts) == set(template)


def test_the_receipt_carries_every_bound_identity_without_reparsing_the_workbook(
    session, project, adopted
):
    formats = effective_baseline_formats(session, project.id)

    rendered = _render(session, project, adopted)

    assert rendered.revision_id == adopted.revision_id
    assert rendered.baseline_content_sha256 == hashlib.sha256(adopted.body).hexdigest()
    assert rendered.output_template_identity == "UCM workbook revision C"
    assert rendered.output_template_sha256 == formats["output_template"].content_sha256
    assert rendered.field_mapping_identity == "ucm-published-column-headings"
    assert rendered.field_mapping_sha256 == formats["field_mapping"].content_sha256
    assert rendered.field_mapping_version == "v1"
    # The receipt names the mapping revision the render was performed under,
    # and the manifest schema that revision was written against (#597).
    assert rendered.field_mapping_revision == "ucm-published-column-headings v1"
    assert rendered.field_mapping_schema_version == "field-mapping-manifest-v1"
    assert rendered.render_profile_sha256 == RenderProfile().content_sha256
    assert rendered.renderer_version == RENDERER_VERSION
    assert rendered.output_sha256 == hashlib.sha256(rendered.content).hexdigest()
    assert rendered.sheet_name == SHEET
    assert rendered.package_differences


def test_the_output_template_identity_is_not_the_data_baseline_identity(
    session, project, adopted
):
    """The two digests coincide today only because #509 registered the same bytes."""

    formats = effective_baseline_formats(session, project.id)
    rendered = _render(session, project, adopted)

    assert formats["output_template"].id != formats["field_mapping"].id
    assert rendered.output_template_sha256 == rendered.baseline_content_sha256
    assert rendered.field_mapping_sha256 != rendered.baseline_content_sha256


# --- Determinism ------------------------------------------------------------


def test_the_same_record_and_template_render_to_the_same_bytes_across_a_clock_change(
    session, project, adopted, monkeypatch
):
    """#579: openpyxl's `save()` stamped the wall clock and the archive drifted.

    Two renders are taken under clocks twenty-nine years apart. Identical bytes
    is the first half; the second half is that neither clock appears in any
    **decompressed** part, because every entry is DEFLATE-compressed and a
    substring check against the raw archive would pass against broken code.
    """

    early = time.struct_time((2001, 2, 3, 4, 5, 6, 5, 34, 0))
    late = time.struct_time((2030, 6, 15, 16, 17, 18, 5, 166, 0))

    def _render_at(moment):
        monkeypatch.setattr(time, "localtime", lambda *args: moment)
        monkeypatch.setattr(time, "gmtime", lambda *args: moment)
        return _render(session, project, adopted).content

    first = _render_at(early)
    second = _render_at(late)
    monkeypatch.undo()

    assert first == second
    with zipfile.ZipFile(BytesIO(first)) as archive:
        assert {info.date_time for info in archive.infolist()} == {
            (1980, 1, 1, 0, 0, 0)
        }
    parts = _parts(first)
    decompressed = b"".join(payload for _name, payload in sorted(parts.items()))
    moments = [
        f"{moment.tm_year:04d}-{moment.tm_mon:02d}-{moment.tm_mday:02d}"
        for moment in (early, late)
    ] + [
        f"{moment.tm_hour:02d}:{moment.tm_min:02d}:{moment.tm_sec:02d}"
        for moment in (early, late)
    ] + [date.today().isoformat()]
    for token in moments:
        assert token.encode() not in decompressed, token
    # The customer's own document properties are the template's, not restamped.
    assert parts["docProps/core.xml"] == _parts(adopted.body)["docProps/core.xml"]
    assert _PINNED_TIMESTAMP in parts["docProps/core.xml"]


def test_repeated_renders_of_one_revision_converge_on_one_digest(
    session, project, adopted
):
    digests = {
        _render(session, project, adopted).output_sha256 for _ in range(12)
    }

    assert len(digests) == 1


# --- The approved identities gate the render --------------------------------


def test_bytes_that_are_not_the_approved_output_template_are_refused(
    session, project, adopted, tmp_path
):
    other = _workbook_bytes(
        tmp_path,
        name="other.xlsx",
        rows=[[*ROWS[0][:1], "Somebody Else", *ROWS[0][2:]], *ROWS[1:]],
    )
    assert other != adopted.body

    with pytest.raises(UnapprovedOutputTemplate):
        _render(session, project, adopted, template_bytes=other)


def test_a_successor_template_that_changes_the_mapping_is_refused_until_approved(
    session, project, adopted, tmp_path
):
    """A material field added under a new heading is not a mapping anyone approved."""

    headings = list(HEADINGS)
    headings.insert(5, "Utility Subtype")
    rows = [[*row[:5], "Distribution", *row[5:]] for row in ROWS]
    successor = _workbook_bytes(
        tmp_path, name="successor.xlsx", headings=headings, rows=rows
    )
    register_baseline_format(
        session,
        project_id=project.id,
        identity=FormatIdentity(
            kind="output_template",
            identity="Customer standard UCM export",
            version="2026.1",
            content_sha256=hashlib.sha256(successor).hexdigest(),
        ),
        principal=PRINCIPAL,
        idempotency_key="register-successor-template",
    )

    with pytest.raises(UnapprovedFieldMapping):
        _render(session, project, adopted, template_bytes=successor)


def test_a_successor_template_that_changes_a_controlled_vocabulary_is_refused(
    session, project, adopted_rich, tmp_path
):
    """Same columns, different drop-down: still not the mapping anyone approved."""

    path = tmp_path / "narrowed.xlsx"
    path.write_bytes(adopted_rich.body)
    from openpyxl import load_workbook as _load

    book = _load(path)
    sheet = book[SHEET]
    sheet.data_validations.dataValidation[0].formula1 = '"Electric,Water"'
    book.save(path)
    successor = _pinned_template(path.read_bytes())
    # Every printed column and every cell is the same; only the drop-down moved.
    assert _cells(successor) == _cells(adopted_rich.body)
    _approve_template(session, project, successor)

    with pytest.raises(UnapprovedFieldMapping, match="controlled vocabulary"):
        _render(session, project, adopted_rich, template_bytes=successor)


def test_an_approved_successor_template_renders_without_re_adopting_the_baseline(
    session, project, adopted, tmp_path
):
    """A later template and mapping change no accepted value and open no adoption."""

    from corridor.baseline_adoption import adopted_baseline_source

    successor = _workbook_bytes(
        tmp_path,
        name="successor.xlsx",
        headings=[*HEADINGS[:-1], "Sheet No."],
        rows=[[*row[:-1], "12"] for row in ROWS],
    )
    _approve_template(session, project, successor)
    before = adopted_baseline_source(session, project.id)

    # The renamed column carries no canonical field either way, so the mapping
    # revision in force still describes this template exactly and nothing needs
    # to be registered for it.
    rendered = _render(session, project, adopted, template_bytes=successor)

    after = adopted_baseline_source(session, project.id)
    assert after.revision_id == before.revision_id
    assert after.content_sha256 == before.content_sha256
    assert rendered.output_template_identity == "Customer standard UCM export"
    assert rendered.output_template_version == "2026.1"
    assert rendered.field_mapping_revision == "ucm-published-column-headings v1"
    assert rendered.changed_cells == ()
    # The customer's own unmapped column travels untouched.
    assert _cells(rendered.content)["P3"] == "12"


def test_a_revision_this_project_never_opened_is_refused(session, project, adopted):
    with pytest.raises(WorkbookRenderRefused, match="not a Project Record revision"):
        _render(session, project, adopted, revision_id=adopted.revision_id + 1_000_000)


def test_a_project_that_never_adopted_a_baseline_is_refused(
    session, project, adopted, tmp_path
):
    other = Project(slug=f"other-{uuid4().hex[:8]}", name="Other", is_synthetic=True)
    session.add(other)
    session.flush()

    with pytest.raises(WorkbookRenderRefused, match="has not adopted a baseline"):
        render_project_record_workbook(
            session,
            project_id=other.id,
            revision_id=adopted.revision_id,
            template_bytes=adopted.body,
            field_mapping=adopted.manifest,
        )


# --- The mapping revision, not the template's bytes (#597) -------------------


def _combined_rows():
    """The same rows, with the station range carried as one combined value."""

    return [
        [*row[:6], f"{row[6]} - {row[7]}", "", *row[8:]] for row in ROWS
    ]


def _reading_of(tmp_path, body: bytes, name: str):
    from corridor.baseline_workbook import read_baseline_workbook

    path = tmp_path / name
    path.write_bytes(body)
    return read_baseline_workbook(
        path, external_references=DEMO.external_reference_headings
    )


def _combined_manifest(tmp_path, body: bytes):
    """The declared revision for a form that combines the station range."""

    return declared_field_mapping(
        _reading_of(tmp_path, body, "combined-read.xlsx"), COMBINED_DECLARATION
    )


def _coordinator(session, project):
    """One person holding the project-coordination designation on this project."""

    from corridor.access import COORDINATION, enroll_member

    principal = HumanPrincipal(f"local:coord-{uuid4().hex[:8]}")
    enroll_member(
        session,
        project_id=project.id,
        email=f"{principal.subject.split(':')[1]}@example.test",
        principal=principal,
        display_name="Dana Ruiz",
        designations=[COORDINATION],
        operator=PRINCIPAL,
    )
    return principal


def _register_mapping(session, project, manifest, principal, key):
    return register_baseline_format(
        session,
        project_id=project.id,
        identity=FormatIdentity(
            kind="field_mapping",
            identity=manifest.identity,
            version=manifest.version,
            content_sha256=manifest.content_sha256,
        ),
        principal=principal,
        idempotency_key=key,
        manifest=manifest,
    )


def test_a_template_that_combines_a_mapped_material_value_is_refused(
    session, project, adopted, tmp_path
):
    """Direction one: two values become one, and nothing printed changes.

    This is the hole #495 left. Every heading, every canonical field, every
    materiality flag and every drop-down is identical, so a digest over the
    template's own columns is identical too — the assertion below says so
    directly. Only the registered mapping revision can tell, because only it
    says how many values a column carries.
    """

    combined = _workbook_bytes(tmp_path, name="combined.xlsx", rows=_combined_rows())
    before = _reading_of(tmp_path, adopted.body, "before-read.xlsx")
    after = _reading_of(tmp_path, combined, "after-read.xlsx")
    assert [
        (column.column, column.heading, column.field, column.material)
        for column in after.column_mapping
    ] == [
        (column.column, column.heading, column.field, column.material)
        for column in before.column_mapping
    ]
    assert after.controlled_vocabularies == before.controlled_vocabularies
    _approve_template(session, project, combined)

    with pytest.raises(UnapprovedFieldMapping, match="combines two values") as raised:
        _render(session, project, adopted, template_bytes=combined)

    assert "one value per column" in str(raised.value)
    assert "project-coordination designation" in str(raised.value)


def test_a_template_that_splits_a_combined_mapped_material_value_is_refused(
    session, project, adopted, tmp_path
):
    """Direction two: one value becomes two, headings still unchanged.

    The project's approved revision now says the range is one combined value,
    so the split form it was adopted from is the successor — and it is refused
    on exactly the same terms.
    """

    combined = _workbook_bytes(tmp_path, name="combined.xlsx", rows=_combined_rows())
    manifest = _combined_manifest(tmp_path, combined)
    _register_mapping(
        session, project, manifest, _coordinator(session, project), "register-v2"
    )

    with pytest.raises(
        UnapprovedFieldMapping, match="declares the whole range"
    ) as raised:
        _render(session, project, adopted, field_mapping=manifest)

    assert "combined_range_v1, cardinality 1 → 2" in str(raised.value)
    assert "project-coordination designation" in str(raised.value)


def test_an_approved_mapping_revision_renders_the_combined_template(
    session, project, adopted, tmp_path
):
    """The attributable act is what unblocks it, and it moves no accepted value."""

    combined = _workbook_bytes(tmp_path, name="combined.xlsx", rows=_combined_rows())
    manifest = _combined_manifest(tmp_path, combined)
    _approve_template(session, project, combined)
    _register_mapping(
        session, project, manifest, _coordinator(session, project), "register-v2"
    )

    rendered = _render(
        session, project, adopted, template_bytes=combined, field_mapping=manifest
    )

    assert rendered.field_mapping_revision == "ucm-published-column-headings v2"
    # The accepted record still holds two station values; the customer's form
    # now prints them as one, and the render writes nothing to say so.
    assert rendered.changed_cells == ()
    cells = _cells(rendered.content)
    assert cells["G3"] == "1149+00 - 1150+00"
    assert "H3" not in cells


def test_a_mapping_revision_is_approved_by_the_project_coordination_designation(
    session, project, adopted, tmp_path
):
    """Operations constructs and validates it; a designated person approves it."""

    combined = _workbook_bytes(tmp_path, name="combined.xlsx", rows=_combined_rows())
    manifest = _combined_manifest(tmp_path, combined)

    with pytest.raises(
        BaselineAdoptionRefused, match="project-coordination designation"
    ):
        _register_mapping(session, project, manifest, PRINCIPAL, "register-v2")

    registered = _register_mapping(
        session, project, manifest, _coordinator(session, project), "register-v2-again"
    )

    assert registered.content_sha256 == manifest.content_sha256
    assert effective_baseline_formats(session, project.id)[
        "field_mapping"
    ].content_sha256 == manifest.content_sha256


def test_registering_a_mapping_revision_changes_no_accepted_value(
    session, project, adopted, tmp_path
):
    """A mapping revision says how the record is read out, never what it says."""

    combined = _workbook_bytes(tmp_path, name="combined.xlsx", rows=_combined_rows())
    manifest = _combined_manifest(tmp_path, combined)
    before_revision = session.scalar(
        select(func.max(ProjectRecordRevision.id)).where(
            ProjectRecordRevision.project_id == project.id
        )
    )
    before_decisions = _effective_decisions(session, project)

    _register_mapping(
        session, project, manifest, _coordinator(session, project), "register-v2"
    )

    assert session.scalar(
        select(func.max(ProjectRecordRevision.id)).where(
            ProjectRecordRevision.project_id == project.id
        )
    ) == before_revision
    assert _effective_decisions(session, project) == before_decisions


def _effective_decisions(session, project):
    return sorted(
        session.execute(
            select(
                FactDecision.subject_key, FactDecision.fact_type, FactDecision.fact_id
            ).where(
                FactDecision.project_id == project.id,
                FactDecision.superseded_by.is_(None),
            )
        ).all()
    )


def test_a_declaration_its_own_populated_example_contradicts_is_refused(
    session, project, adopted, tmp_path
):
    """The round trip is what proves a declaration, before anyone registers it."""

    with pytest.raises(MappingManifestRefused, match="declares the whole range"):
        declared_field_mapping(
            _reading_of(tmp_path, adopted.body, "split-read.xlsx"),
            COMBINED_DECLARATION,
        )


def test_a_manifest_that_is_not_the_registered_mapping_revision_is_refused(
    session, project, adopted, tmp_path
):
    """Offering a manifest is exactly as unprivileged as offering template bytes."""

    combined = _workbook_bytes(tmp_path, name="combined.xlsx", rows=_combined_rows())
    manifest = _combined_manifest(tmp_path, combined)

    with pytest.raises(
        UnapprovedFieldMapping, match="not the approved mapping revision"
    ):
        _render(session, project, adopted, field_mapping=manifest)


# --- Reproducing a render from stored state alone (#610) --------------------


def test_a_past_render_reproduces_without_the_caller_supplying_the_manifest(
    session, project, adopted, tmp_path
):
    """The registration now resolves its own mapping revision.

    #597 registered a revision by identity, version and digest, so reproducing
    a render depended on whoever declared it still holding the declaration.
    Omitting ``field_mapping`` here reproduces the same bytes from what the
    database stores, which is the whole claim.
    """

    supplied = _render(session, project, adopted)

    reproduced = render_project_record_workbook(
        session,
        project_id=project.id,
        revision_id=adopted.revision_id,
        template_bytes=adopted.body,
    )

    assert reproduced.output_sha256 == supplied.output_sha256
    assert reproduced.content == supplied.content
    assert reproduced.field_mapping_revision == supplied.field_mapping_revision
    assert reproduced.field_mapping_sha256 == supplied.field_mapping_sha256


def test_an_approved_successor_mapping_reproduces_the_render_it_was_approved_for(
    session, project, adopted, tmp_path
):
    """Stored state follows the registration, not the manifest last handed in."""

    combined = _workbook_bytes(tmp_path, name="combined.xlsx", rows=_combined_rows())
    manifest = _combined_manifest(tmp_path, combined)
    _register_mapping(
        session, project, manifest, _coordinator(session, project), "register-v2"
    )
    _approve_template(session, project, combined)

    reproduced = render_project_record_workbook(
        session,
        project_id=project.id,
        revision_id=adopted.revision_id,
        template_bytes=combined,
    )

    assert reproduced.field_mapping_sha256 == manifest.content_sha256
    assert reproduced.output_sha256 == _render(
        session,
        project,
        adopted,
        template_bytes=combined,
        field_mapping=manifest,
    ).output_sha256


def test_a_registration_that_stores_no_declaration_refuses_by_name(
    session, project, adopted, tmp_path
):
    """The shape every registration had before #610, and what it can prove.

    The command alone writes identity, version and digest — no declaration —
    so this is exactly a registration made before the declaration was stored
    beside it. The absence has to be stated, never read as an empty mapping.
    """

    combined = _workbook_bytes(tmp_path, name="combined.xlsx", rows=_combined_rows())
    manifest = _combined_manifest(tmp_path, combined)
    session.scalar(
        select(
            func.register_baseline_format(
                project.id,
                "field_mapping",
                manifest.identity,
                manifest.version,
                manifest.content_sha256,
                PRINCIPAL.subject,
                "register-without-a-declaration",
            )
        )
    )
    session.expire_all()

    assert effective_field_mapping_manifest(session, project.id) is None
    with pytest.raises(WorkbookRenderRefused, match="stores no declaration"):
        render_project_record_workbook(
            session,
            project_id=project.id,
            revision_id=adopted.revision_id,
            template_bytes=adopted.body,
        )


def test_an_identical_re_registration_stays_idempotent_and_stores_one_declaration(
    session, project, adopted, tmp_path
):
    """Registering the same declaration again is not a semantic change (#597).

    It needs no project-coordination designation, and it stores the same
    declaration rather than a second copy of it.
    """

    manifest = adopted.manifest
    registered = _register_mapping(
        session, project, manifest, PRINCIPAL, "register-the-same-again"
    )

    assert registered.content_sha256 == manifest.content_sha256
    assert effective_field_mapping_manifest(
        session, project.id
    ).declaration_json == manifest.declaration_json
    assert session.scalar(
        select(func.count())
        .select_from(BaselineFormatManifest)
        .where(BaselineFormatManifest.project_id == project.id)
    ) == session.scalar(
        select(func.count())
        .select_from(BaselineFormat)
        .where(
            BaselineFormat.project_id == project.id,
            BaselineFormat.format_kind == "field_mapping",
        )
    )


# --- Fail closed on what cannot be preserved --------------------------------


@pytest.mark.parametrize(
    "part,payload,expected",
    (
        ("xl/vbaProject.bin", b"\x00macro", "a macro project"),
        ("xl/externalLinks/externalLink1.xml", b"<x/>", "an external workbook link"),
        ("xl/pivotCache/pivotCacheDefinition1.xml", b"<x/>", "a pivot cache"),
        ("xl/ctrlProps/ctrlProp1.xml", b"<x/>", "a form control"),
        ("_xmlsignatures/sig1.xml", b"<x/>", "a digital signature"),
        ("xl/somethingNobodyEnumerated.xml", b"<x/>", "cannot read"),
    ),
)
def test_an_unsupported_package_part_fails_closed(
    session, project, adopted, part, payload, expected
):
    tampered = _with_part(adopted.body, part, payload)
    _approve_template(session, project, tampered)

    with pytest.raises(UnsupportedWorkbookFeature, match=expected):
        _render(session, project, adopted, template_bytes=tampered)


def test_worksheet_protection_fails_closed(session, project, adopted):
    parts = _parts(adopted.body)
    name = "xl/worksheets/sheet1.xml"
    parts[name] = parts[name].replace(
        b"</sheetData>", b"</sheetData><sheetProtection sheet=\"1\"/>"
    )
    protected = _repack(parts)
    _approve_template(session, project, protected)

    with pytest.raises(UnsupportedWorkbookFeature, match="worksheet protection"):
        _render(session, project, adopted, template_bytes=protected)


def test_a_formula_under_a_mapped_column_refuses_rather_than_being_overwritten(
    session, project, adopted
):
    """The customer's own calculation is never silently replaced by a value."""

    with_formula = _with_formula_cell(adopted.body, "J3", "2026-03-01")
    _approve_template(session, project, with_formula)
    outcome = _decide(
        session,
        project,
        _second_document(session, project, "later-promise"),
        subject_key=f"{SHEET}!3",
        fact_type="committed_date",
        text="2026-09-30",
        cell_range="J3",
        ordinal=1,
        key="change-promise",
        predecessor=_effective_decision(
            session, project, f"{SHEET}!3", "committed_date"
        ).id,
    )

    with pytest.raises(UnsupportedWorkbookFeature, match="formula cell"):
        _render(
            session,
            project,
            adopted,
            template_bytes=with_formula,
            revision_id=outcome.revision.id,
        )


# --- An existing subject changes only its approved mapped cells -------------


def test_an_existing_subject_writes_one_mapped_cell_and_nothing_else(
    session, project, adopted
):
    outcome = _decide(
        session,
        project,
        _second_document(session, project, "owner-correction"),
        subject_key=f"{SHEET}!3",
        fact_type="external_org",
        text="Oncor Electric Delivery",
        cell_range="B3",
        ordinal=1,
        key="change-owner",
        predecessor=_effective_decision(
            session, project, f"{SHEET}!3", "external_org"
        ).id,
    )

    rendered = _render(session, project, adopted, revision_id=outcome.revision.id)

    [changed] = rendered.changed_cells
    assert changed.field == "external_org"
    assert changed.cell_range == "B3"
    assert changed.before == "CenterPoint Energy"
    assert changed.after == "Oncor Electric Delivery"
    # Both identities travel, and they are not the same string.
    assert changed.source_row_key == f"{SHEET}!3"
    assert changed.record_subject_key == "UC-1"
    assert rendered.changed_parts == ("xl/worksheets/sheet1.xml",)

    template = _parts(adopted.body)
    output = _parts(rendered.content)
    assert {
        name: payload
        for name, payload in output.items()
        if name != "xl/worksheets/sheet1.xml"
    } == {
        name: payload
        for name, payload in template.items()
        if name != "xl/worksheets/sheet1.xml"
    }
    cells = _cells(rendered.content)
    assert cells["B3"] == "Oncor Electric Delivery"
    assert cells["B4"] == "City of Austin"
    assert cells["P3"] == "Yes"  # an unknown customer column, untouched
    assert cells["A1"] == "Utility Conflict Management (UCM) — Utility Conflicts"


def test_the_source_row_identity_keeps_its_external_references(
    session, project, adopted
):
    rendered = _render(session, project, adopted)

    assert rendered.external_references == (
        (f"{SHEET}!3", "UCM-1001", "https://ucm.example/records/1001"),
        (f"{SHEET}!4", "UCM-1002", ""),
        (f"{SHEET}!5", "UCM-1003", ""),
    )


def test_a_template_whose_rows_moved_refuses_rather_than_writing_another_line(
    session, project, adopted, tmp_path
):
    shuffled = _workbook_bytes(
        tmp_path, name="shuffled.xlsx", rows=[ROWS[1], ROWS[0], ROWS[2]]
    )
    _approve_template(session, project, shuffled)

    with pytest.raises(WorkbookRenderRefused, match="rows have moved"):
        _render(session, project, adopted, template_bytes=shuffled)


# --- A new subject, and an apparent removal ---------------------------------


def _add_new_subject(session, project) -> int:
    document = _second_document(session, project, "new-conflict")
    _decide(
        session,
        project,
        document,
        subject_key="proposed:UC-4",
        fact_type="utility_id",
        text="UC-4",
        cell_range="A9",
        ordinal=1,
        key="new-subject-id",
    )
    outcome = _decide(
        session,
        project,
        document,
        subject_key="proposed:UC-4",
        fact_type="external_org",
        text="Bluebonnet Electric Cooperative",
        cell_range="B9",
        ordinal=2,
        key="new-subject-owner",
    )
    return outcome.revision.id


def test_a_new_subject_without_an_approved_row_insertion_template_is_refused(
    session, project, adopted
):
    revision = _add_new_subject(session, project)

    with pytest.raises(UndeclaredRenderBehaviour, match="row-insertion template"):
        _render(session, project, adopted, revision_id=revision)


def test_a_declared_new_subject_adds_exactly_one_row(session, project, adopted):
    revision = _add_new_subject(session, project)
    profile = RenderProfile(row_insertion=RowInsertionTemplate(style_source_row=3))

    rendered = _render(
        session, project, adopted, revision_id=revision, profile=profile
    )

    [added] = rendered.added_rows
    assert added.record_subject_key == "proposed:UC-4"
    assert added.row_number == 6
    assert added.fields == ("external_org", "utility_id")
    assert rendered.changed_cells == ()
    cells = _cells(rendered.content)
    assert cells["A6"] == "UC-4"
    assert cells["B6"] == "Bluebonnet Electric Cooperative"
    assert cells["A5"] == "UC-3"
    assert "C6" not in cells


def _close_first_conflict(session, project) -> int:
    document = _second_document(session, project, "closure")
    segment = _segment(
        session,
        project,
        document,
        text="Conflict resolved in the field",
        cell_range="L3",
        ordinal=1,
    )
    fact = append_fact(
        session,
        project_id=project.id,
        document_id=document.id,
        extraction_run_id=document.extraction_run_id,
        subject_kind="source_row",
        subject_key=f"{SHEET}!3",
        recorded_by="human:coordinator",
        content_sha256=hashlib.sha256(b"closure-uc-1").hexdigest(),
        value=materialize_typed_satellite("closure_result", segment),
        closure=ClosureValues(
            closure_kind="constraint_closed",
            successor_dependency_id=None,
            governing_source_segment_ids=(segment.id,),
        ),
    )
    outcome = record_human_fact_decision(
        session,
        fact,
        principal=PRINCIPAL,
        command_type="resolve_discrepancy",
        idempotency_key="close-uc-1",
    )
    return outcome.revision.id


def test_an_apparent_removal_without_a_declared_retirement_mapping_is_refused(
    session, project, adopted
):
    revision = _close_first_conflict(session, project)

    with pytest.raises(UndeclaredRenderBehaviour, match="retirement"):
        _render(session, project, adopted, revision_id=revision)


def test_a_retired_row_keeps_its_line_and_takes_the_customers_own_wording(
    session, project, adopted
):
    revision = _close_first_conflict(session, project)
    profile = RenderProfile(
        retirement=RetirementMapping(field="marked_resolution", wording="Not Used")
    )

    rendered = _render(
        session, project, adopted, revision_id=revision, profile=profile
    )

    [retired] = rendered.retired_rows
    assert retired.source_row_key == f"{SHEET}!3"
    assert retired.record_subject_key == "UC-1"
    assert retired.cell_range == "L3"
    assert retired.wording == "Not Used"
    cells = _cells(rendered.content)
    # The row is still there, with everything it said.
    assert cells["A3"] == "UC-1"
    assert cells["B3"] == "CenterPoint Energy"
    assert cells["L3"] == "Not Used"
    assert len(_cells(adopted.body)) + 1 == len(cells)


# --- The one configured workbook variant ------------------------------------


def test_the_provenance_worksheet_is_configured_rather_than_chosen_each_issue(
    session, project, adopted
):
    plain = RenderProfile()
    configured = RenderProfile(provenance_worksheet=True)
    assert plain.content_sha256 != configured.content_sha256

    rendered = _render(session, project, adopted, profile=configured)

    assert rendered.render_profile_sha256 == configured.content_sha256
    book = load_workbook(BytesIO(rendered.content))
    try:
        assert PROVENANCE_SHEET_NAME in book.sheetnames
        recorded = {
            row[0]: row[1] for row in book[PROVENANCE_SHEET_NAME].iter_rows(
                values_only=True
            )
        }
    finally:
        book.close()
    assert recorded["Project Record revision"] == str(adopted.revision_id)
    assert recorded["Accepted baseline digest"] == rendered.baseline_content_sha256
    assert recorded["Renderer"] == RENDERER_VERSION
    assert rendered.changed_parts == (
        "[Content_Types].xml",
        "xl/_rels/workbook.xml.rels",
        "xl/workbook.xml",
        "xl/worksheets/sheet2.xml",
    )
    # The customer's own worksheet is untouched by the variant.
    assert _parts(rendered.content)["xl/worksheets/sheet1.xml"] == _parts(
        adopted.body
    )["xl/worksheets/sheet1.xml"]


# --- Fixture plumbing -------------------------------------------------------


def _repack(parts: dict[str, bytes]) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


def _with_part(data: bytes, name: str, payload: bytes) -> bytes:
    parts = _parts(data)
    parts[name] = payload
    return _repack(parts)


def _with_formula_cell(data: bytes, reference: str, cached: str) -> bytes:
    import re

    parts = _parts(data)
    name = "xl/worksheets/sheet1.xml"
    pattern = re.compile(
        rb'<c r="' + reference.encode() + rb'"[^>]*?(?:/>|>.*?</c>)', re.S
    )
    replacement = (
        f'<c r="{reference}" t="str"><f>A1</f><v>{cached}</v></c>'
    ).encode()
    updated, count = pattern.subn(replacement, parts[name], count=1)
    assert count == 1, f"{reference} is not in the fixture worksheet"
    parts[name] = updated
    return _repack(parts)


def _approve_template(session, project, template: bytes) -> None:
    register_baseline_format(
        session,
        project_id=project.id,
        identity=FormatIdentity(
            kind="output_template",
            identity="Customer standard UCM export",
            version="2026.1",
            content_sha256=hashlib.sha256(template).hexdigest(),
        ),
        principal=PRINCIPAL,
        idempotency_key=f"approve-{hashlib.sha256(template).hexdigest()[:16]}",
    )


# --- Round trip over the complete package -----------------------------------


@pytest.fixture
def adopted_rich(session, project, tmp_path, store):
    """An adopted workbook carrying the features a real UCM actually uses."""

    body = _workbook_bytes(tmp_path, name="rich.xlsx", table=True, features=True)
    staged = validate_and_stage(body, "rich.xlsx")
    preview = preview_baseline_adoption(
        session,
        project=project,
        staged=staged,
        customer="Lone Star Transit Authority",
        source_identity="UCM workbook revision C",
        field_mapping=DEMO,
    )
    result = adopt_baseline(
        session,
        preview=preview,
        principal=PRINCIPAL,
        idempotency_key="adopt-rich",
        images_dir=tmp_path / "images",
    )
    return _Adopted(
        body=body,
        result=result,
        document_id=result.document_id,
        manifest=preview.field_mapping_manifest,
    )


def test_formatting_validation_tables_filters_and_hidden_content_survive_a_write(
    session, project, adopted_rich
):
    from corridor.baseline_workbook import read_baseline_workbook

    outcome = _decide(
        session,
        project,
        _second_document(session, project, "rich-correction"),
        subject_key=f"{SHEET}!3",
        fact_type="external_org",
        text="Oncor Electric Delivery",
        cell_range="B3",
        ordinal=1,
        key="rich-change-owner",
        predecessor=_effective_decision(
            session, project, f"{SHEET}!3", "external_org"
        ).id,
    )

    rendered = _render(
        session, project, adopted_rich, revision_id=outcome.revision.id
    )

    template = _parts(adopted_rich.body)
    output = _parts(rendered.content)
    assert rendered.changed_parts == ("xl/worksheets/sheet1.xml",)
    for name in sorted(template):
        if name != "xl/worksheets/sheet1.xml":
            assert output[name] == template[name], name
    assert any(name.startswith("xl/tables/") for name in template)

    sheet = output["xl/worksheets/sheet1.xml"]
    original = template["xl/worksheets/sheet1.xml"]
    # Everything the customer's sheet declares outside <sheetData> is untouched.
    for marker in (b"<dataValidation", b"<autoFilter", b"<pane", b"<tableParts"):
        assert marker in original
        assert sheet.split(b"<sheetData>")[0] == original.split(b"<sheetData>")[0]
        assert sheet.split(b"</sheetData>")[1] == original.split(b"</sheetData>")[1]
    # The written cell keeps the style index the customer's formatting gave it.
    style = re.search(rb'<c r="B3"[^>]*\bs="(\d+)"', original).group(1)
    assert re.search(rb'<c r="B3"[^>]*\bs="(\d+)"', sheet).group(1) == style

    # And the workbook still reads as the same matrix, with the new value in it.
    import tempfile

    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "rendered.xlsx"
        path.write_bytes(rendered.content)
        again = read_baseline_workbook(
            path, external_references=DEMO.external_reference_headings
        )
    assert again.adopted_sheet == SHEET
    assert again.round_trip.clean
    assert [row.value("external_org") for row in again.rows] == [
        "Oncor Electric Delivery",
        "City of Austin",
        "Oncor",
    ]
    assert again.hidden_content == ("Utility Conflicts!column N is hidden",)
    assert [column.heading for column in again.unknown_columns] == [
        "Early TxDOT Utility Activity"
    ]


def test_a_declared_new_row_carries_only_the_approved_table_metadata(
    session, project, adopted_rich
):
    """One row, plus exactly the associated metadata the template declares."""

    revision = _add_new_subject(session, project)
    template = _parts(adopted_rich.body)
    [table_part] = [name for name in template if name.startswith("xl/tables/")]

    without = _render(
        session,
        project,
        adopted_rich,
        revision_id=revision,
        profile=RenderProfile(row_insertion=RowInsertionTemplate()),
    )
    with_metadata = _render(
        session,
        project,
        adopted_rich,
        revision_id=revision,
        profile=RenderProfile(
            row_insertion=RowInsertionTemplate(
                style_source_row=3, extend_table_ref=True, extend_auto_filter=True
            )
        ),
    )

    # Undeclared metadata is left exactly as the customer wrote it.
    assert _parts(without.content)[table_part] == template[table_part]
    assert (
        b'<autoFilter ref="A2:P5"'
        in _parts(without.content)["xl/worksheets/sheet1.xml"]
    )
    # Declared metadata follows the row, and nothing else does.
    assert b'ref="A2:P6"' in _parts(with_metadata.content)[table_part]
    assert (
        b'<autoFilter ref="A2:P6"'
        in _parts(with_metadata.content)["xl/worksheets/sheet1.xml"]
    )
    assert with_metadata.changed_parts == (table_part, "xl/worksheets/sheet1.xml")
    # The new row takes the formatting the declared source row carries.
    sheet = _parts(with_metadata.content)["xl/worksheets/sheet1.xml"]
    style = re.search(rb'<c r="B3"[^>]*\bs="(\d+)"', sheet).group(1)
    assert re.search(rb'<c r="B6"[^>]*\bs="(\d+)"', sheet).group(1) == style
    assert _cells(with_metadata.content)["B6"] == "Bluebonnet Electric Cooperative"


# --- What a real Excel-authored template looks like --------------------------


def _shared_string_template(data: bytes) -> bytes:
    """Rewrite an openpyxl package the way Excel writes one: shared strings.

    openpyxl stores every string inline, and Excel stores strings in
    `xl/sharedStrings.xml` and points cells at it with `t="s"`. Only the second
    shape reaches Corridor from a customer, so the renderer is exercised against
    it here rather than only against the shape the fixture writer happens to use.
    """

    parts = _parts(data)
    name = "xl/worksheets/sheet1.xml"
    sheet = parts[name]
    table: list[bytes] = []

    def _share(match):
        text = match.group(2)
        if text not in table:
            table.append(text)
        return match.group(1) + b't="s"><v>' + str(
            table.index(text)
        ).encode() + b"</v></c>"

    sheet = re.sub(
        rb'(<c r="[A-Z]+\d+"(?: s="\d+")? )t="inlineStr"><is><t[^>]*>(.*?)</t></is></c>',
        _share,
        sheet,
    )
    parts[name] = sheet
    parts["xl/sharedStrings.xml"] = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        b'count="%d" uniqueCount="%d">' % (len(table), len(table))
        + b"".join(b"<si><t>" + value + b"</t></si>" for value in table)
        + b"</sst>"
    )
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/'
        b'vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
        b"</Types>",
        1,
    )
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="rIdSharedStrings" Type="http://schemas.openxmlformats'
        b'.org/officeDocument/2006/relationships/sharedStrings" '
        b'Target="sharedStrings.xml"/></Relationships>',
        1,
    )
    return canonical_package(parts)


def test_a_shared_string_template_renders_the_way_excel_writes_one(
    session, project, tmp_path, store
):
    body = _shared_string_template(_workbook_bytes(tmp_path, name="shared.xlsx"))
    assert b't="s"' in _parts(body)["xl/worksheets/sheet1.xml"]
    assert _cells(body)["B3"] == "CenterPoint Energy"

    staged = validate_and_stage(body, "shared.xlsx")
    preview = preview_baseline_adoption(
        session,
        project=project,
        staged=staged,
        customer="Lone Star Transit Authority",
        source_identity="UCM workbook revision C",
        field_mapping=DEMO,
    )
    result = adopt_baseline(
        session,
        preview=preview,
        principal=PRINCIPAL,
        idempotency_key="adopt-shared",
        images_dir=tmp_path / "images",
    )
    adopted = _Adopted(
        body=body,
        result=result,
        document_id=result.document_id,
        manifest=preview.field_mapping_manifest,
    )
    outcome = _decide(
        session,
        project,
        _second_document(session, project, "shared-correction"),
        subject_key=f"{SHEET}!3",
        fact_type="external_org",
        text="Oncor Electric Delivery",
        cell_range="B3",
        ordinal=1,
        key="shared-change-owner",
        predecessor=_effective_decision(
            session, project, f"{SHEET}!3", "external_org"
        ).id,
    )

    rendered = _render(session, project, adopted, revision_id=outcome.revision.id)

    assert rendered.changed_parts == ("xl/worksheets/sheet1.xml",)
    # The customer's own string table is left exactly as it was.
    assert _parts(rendered.content)["xl/sharedStrings.xml"] == _parts(body)[
        "xl/sharedStrings.xml"
    ]
    cells = _cells(rendered.content)
    assert cells["B3"] == "Oncor Electric Delivery"
    assert cells["B4"] == "City of Austin"
    assert cells["A3"] == "UC-1"
