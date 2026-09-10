"""Actual reader outputs cannot certify omitted classes or fabricated native IDs."""
from dataclasses import replace
from datetime import datetime, timezone
from uuid import uuid4

import pytest
from sqlalchemy import func, select

from corridor.config import settings
from corridor.db import Session, engine
from corridor.fact_decisions import record_human_fact_decision
from corridor.models import Fact, FactDecision, Project, ProjectRecordRevision, SourceSegment
from corridor.native_reader_coverage import collect_native_reader_coverage
from corridor.reader_coverage import CONTRACTS, compare_all_surfaces
from test_native_accepted_readers import _adopt_native_workbook, PRINCIPAL, TODAY, covering_client


@pytest.fixture
def session():
    with engine.connect() as connection:
        transaction = connection.begin()
        with Session(bind=connection) as session:
            yield session
        if transaction.is_active:
            transaction.rollback()


@pytest.fixture
def adopted(session, tmp_path, monkeypatch):
    return _adopt_native_workbook(session, tmp_path, monkeypatch)


def collect(session, project, adoption, **options):
    return collect_native_reader_coverage(session, project.id,
        as_of_revision_id=adoption.revision_id,
        evaluated_at=datetime.combine(TODAY, datetime.min.time(), timezone.utc), **options)


def surfaces(result):
    return {s.reading.surface: s for s in result.surfaces}


def test_collector_observes_all_seven_public_surfaces_and_native_populations(session, adopted):
    project, adoption = adopted
    before = {model.__name__: session.scalar(select(func.count(model.id)).where(model.project_id == project.id))
              for model in (Fact, FactDecision, ProjectRecordRevision, SourceSegment)}
    result = collect(session, project, adoption)
    views = surfaces(result)
    assert tuple(views) == tuple(c.name for c in CONTRACTS)
    source_ids = set(session.scalars(select(SourceSegment.id).where(SourceSegment.project_id == project.id)))
    assert set(views['source_and_decision_history'].population_evidence['source']['identities']) == source_ids
    assert views['report_release'].population_evidence['release']['count'] == 0
    assert views['report_release'].reading.records == ()
    assert views['coordination_report'].reader_outputs['briefing']['mode'].startswith('native citation inputs')
    assert result.revision_id == adoption.revision_id
    assert 'release' in views['report_release'].reading.observed_record_kinds
    assert views['report_release'].reading.output_identity
    assert views['source_and_decision_history'].reader_outputs['source_occurrence_readback']
    if result.blockers:
        assert not all(r.passed for r in compare_all_surfaces(result.readings, result.readings))
    after = {model.__name__: session.scalar(select(func.count(model.id)).where(model.project_id == project.id)) for model in (Fact, FactDecision, ProjectRecordRevision, SourceSegment)}
    assert after == before


def test_current_and_selected_as_of_preserve_exact_fact_decision_identity(session, adopted):
    project, adoption = adopted
    original = session.scalar(select(FactDecision).where(FactDecision.project_id == project.id,
        FactDecision.fact_type == 'station_from').order_by(FactDecision.id))
    changed = record_human_fact_decision(session, session.get(Fact, original.fact_id), principal=PRINCIPAL,
        command_type='resolve_discrepancy', idempotency_key='coverage-withdraw',
        expected_predecessor=original.id, disposition='do_not_add')
    result = collect(session, project, adoption)
    observed = surfaces(result)['current_and_as_of_record'].reading
    current = {r.key for r in observed.records if r.kind == 'current'}
    historical = {r.key for r in observed.records if r.kind == 'as_of'}
    assert f'fact_decisions:{original.id}' not in current
    assert f'fact_decisions:{original.id}' in historical
    assert result.revision_id == changed.revision.id
    row = next(r for r in observed.records if r.kind == 'as_of' and r.key == f'fact_decisions:{original.id}')
    assert row.fields['authority']['fact_id'] == original.fact_id
    assert set(row.field_origins.values()) == {f'native_decision:fact_decisions:{original.id}'}


def test_reader_omission_is_a_gap_not_an_empty_class(session, adopted, monkeypatch):
    import corridor.record_history as history
    project, adoption = adopted
    reader = history.read_record_history
    monkeypatch.setattr(history, 'read_record_history', lambda *a, **kw: replace(reader(*a, **kw), source_decisions=()))
    result = collect(session, project, adoption)
    actual = surfaces(result)['source_and_decision_history']
    assert actual.population_evidence['decision']['count'] > 0
    assert 'decision' not in actual.reading.observed_record_kinds
    assert any('native decision IDs absent' in b for b in actual.blockers)
    from corridor.reader_equivalence import compare_native_reader_surfaces
    # Even identical supplied semantic rows cannot erase a collector's actual
    # omission findings. The wrapper must collect again and carry those findings.
    compared = compare_native_reader_surfaces(session, project.id, legacy=result.readings,
        as_of_revision_id=adoption.revision_id,
        evaluated_at=datetime.combine(TODAY, datetime.min.time(), timezone.utc))
    history_result = next(row for row in compared if row.surface == 'source_and_decision_history')
    assert not history_result.passed
    assert any('native decision IDs absent' in b for b in history_result.blockers)


def test_foreign_project_revision_is_refused_before_any_surface(session, adopted):
    project, adoption = adopted
    other = Project(slug=f'foreign-coverage-{uuid4().hex}', name='Other', is_synthetic=True)
    session.add(other)
    session.flush()
    with pytest.raises(ValueError, match='both current and selected revisions'):
        collect_native_reader_coverage(session, other.id, as_of_revision_id=adoption.revision_id,
            evaluated_at=datetime.now(timezone.utc))


def test_provider_store_and_pending_writes_are_refused_without_fetches(session, adopted, monkeypatch):
    from corridor.accepted_field_reading import NativeReadingRefused
    project, adoption = adopted
    monkeypatch.setattr(settings, 'storage_backend', 's3')
    with pytest.raises(NativeReadingRefused, match='provider I/O is not enabled'):
        collect(session, project, adoption)
    monkeypatch.setattr(settings, 'storage_backend', 'filesystem')
    project.name = 'unflushed caller change'
    with pytest.raises(ValueError, match='caller-owned pending writes'):
        collect(session, project, adoption)


def test_population_limit_never_silently_truncates_evidence(session, adopted):
    from corridor.accepted_field_reading import NativeReadingRefused
    with pytest.raises(NativeReadingRefused, match='exceeds coverage limit'):
        collect(session, *adopted, maximum_population=1)


def test_public_reader_failure_does_not_prevent_other_surface_observations(session, adopted, monkeypatch):
    import corridor.report as report
    monkeypatch.setattr(report, 'build_report', lambda *a, **kw: (_ for _ in ()).throw(ValueError('unsupported field class')))
    result = collect(session, *adopted)
    views = surfaces(result)
    assert any('unsupported field class' in b for b in views['coordination_report'].blockers)
    assert views['current_and_as_of_record'].reading.records
    assert 'history' in views['source_and_decision_history'].reader_outputs


def test_source_citation_cannot_borrow_another_existing_fact_decision(session, adopted, monkeypatch):
    import corridor.report as report
    project, adoption = adopted
    reader = report.build_report
    def forged(*args, **kwargs):
        result = reader(*args, **kwargs)
        citations = [c for c in result.cells if isinstance(c.provenance, report.Assertion)]
        left = citations[0]
        right = next(c for c in citations if c.provenance.fact_id != left.provenance.fact_id)
        left.provenance = replace(left.provenance, decision_id=right.provenance.decision_id)
        return result
    monkeypatch.setattr(report, 'build_report', forged)
    result = collect(session, project, adoption)
    actual = surfaces(result)['coordination_report']
    assert any('actual source segment/FactDecision/revision' in b for b in actual.blockers)
    assert 'report' not in actual.reading.observed_record_kinds


def test_briefing_runs_only_the_explicit_fixture_client(session, adopted):
    client = covering_client()
    result = collect(session, *adopted, briefing_fixture_client=client)
    assert len(client.calls) == 1
    assert 'fixture_briefing' in surfaces(result)['coordination_report'].reader_outputs['briefing']


def test_changed_authority_invalidates_all_observed_classes(session, adopted, monkeypatch):
    import corridor.native_reader_coverage as coverage
    original = coverage._inventory
    calls = 0
    def changed(*args):
        nonlocal calls
        calls += 1
        actual = original(*args)
        if calls == 2:
            # Simulate a concurrent append detected by the final native census;
            # no synthetic domain row is written or passed to the collector.
            actual = {**actual, 'concurrent_observation': ({'id': 'changed'},)}
        return actual
    monkeypatch.setattr(coverage, '_inventory', changed)
    result = collect(session, *adopted)
    assert all(not s.reading.observed_record_kinds for s in result.surfaces)
    assert all(any('changed during collection' in b for b in s.blockers) for s in result.surfaces)


@pytest.mark.parametrize('omission', ['record', 'field'])
def test_product_freeze_must_exhaust_the_independent_native_population(session, adopted, monkeypatch, omission):
    import corridor.project_reading as project_reading
    original = project_reading.freeze_project_reading
    def incomplete(*args, **kwargs):
        reading = original(*args, **kwargs)
        population = reading.native_population
        if omission == 'record':
            records = population.records[1:]
        else:
            first = population.records[0]
            values = dict(first.fields)
            values.pop(next(iter(values)))
            records = (replace(first, fields=values), *population.records[1:])
        changed = replace(population, records=tuple(records))
        identities = {r.id: r for r in records}
        rows = tuple(replace(row, dependency=identities[row.dependency.id]) for row in reading.rows if row.dependency.id in identities)
        # Keep the product's own metadata internally consistent, reproducing
        # the original hole: its population cannot serve as its own census.
        return replace(reading, native_population=changed, rows=rows,
            evaluation=replace(reading.evaluation, native_population=changed))
    monkeypatch.setattr(project_reading, 'freeze_project_reading', incomplete)
    result = collect(session, *adopted)
    for name in ('constraint_log', 'coordination_report', 'workbook_export'):
        observed = surfaces(result)[name]
        assert not observed.reading.observed_record_kinds
        assert any('product freeze omits' in b for b in observed.blockers)


@pytest.mark.parametrize('omission', ['accepted_row', 'accepted_section'])
def test_report_requires_the_actual_complete_accepted_field_rows(session, adopted, monkeypatch, omission):
    import corridor.report as report
    original = report.build_report
    def incomplete(*args, **kwargs):
        output = original(*args, **kwargs)
        section = next(s for s in output.sections if s.columns == ['Ref', 'Field', 'Accepted value'])
        if omission == 'accepted_row':
            section.rows.pop()
        else:
            output.sections.remove(section)
        # covered_records and all remaining citations remain valid.
        return output
    monkeypatch.setattr(report, 'build_report', incomplete)
    result = collect(session, *adopted)
    observed = surfaces(result)['coordination_report']
    assert not observed.reading.observed_record_kinds
    assert any('actual accepted-field rows' in b for b in observed.blockers)


@pytest.mark.parametrize('damage', ['source_rows', 'main_rows', 'source_value', 'main_value'])
def test_workbook_requires_actual_population_and_full_native_source_rows(session, adopted, monkeypatch, damage):
    import corridor.export as export
    from openpyxl import load_workbook
    original = export.to_xlsx
    def incomplete(*args, **kwargs):
        path = original(*args, **kwargs)
        workbook = load_workbook(path)
        source = workbook['Accepted value sources']
        if damage == 'source_rows':
            source.delete_rows(2, source.max_row)
        elif damage == 'main_rows':
            workbook.active.delete_rows(2, workbook.active.max_row)
        elif damage == 'source_value':
            source['D2'] = 'changed accepted value'
        else:
            workbook.active['F2'] = 'changed station'
        workbook.save(path)
        workbook.close()
        return path
    monkeypatch.setattr(export, 'to_xlsx', incomplete)
    result = collect(session, *adopted)
    observed = surfaces(result)['workbook_export']
    assert not observed.reading.observed_record_kinds
    assert any('actual accepted-value source rows' in b or 'actual main-sheet population' in b or 'main-sheet field column' in b for b in observed.blockers)


def test_a_late_surface_failure_cannot_publish_observed_classes(session, adopted, monkeypatch):
    import corridor.legacy_history as history
    def failed(*args, **kwargs):
        raise ValueError('late compatibility census is unavailable')
    monkeypatch.setattr(history, 'inventory_history', failed)
    result = collect(session, *adopted)
    observed = surfaces(result)['source_and_decision_history']
    assert any('late compatibility census is unavailable' in b for b in observed.blockers)
    assert observed.reading.records  # Preserve evidence, never the coverage claim.
    assert not observed.reading.observed_record_kinds
    assert not next(r for r in compare_all_surfaces(result.readings, result.readings)
                    if r.surface == 'source_and_decision_history').passed


@pytest.mark.parametrize('changed', ['actor', 'time', 'evidence_role', 'source_text', 'source_document', 'source_locator', 'source_filename'])
def test_support_history_corroborates_every_attributable_public_field(session, adopted, monkeypatch, changed):
    import corridor.record_history as history
    from datetime import timedelta
    original = history.read_record_history
    def incorrect(*args, **kwargs):
        output = original(*args, **kwargs)
        decisions = []
        for row in output.source_decisions:
            if row.fact is None or not row.fact.assessments:
                decisions.append(row)
                continue
            assessment = row.fact.assessments[0]
            if changed == 'actor':
                assessment = replace(assessment, authority='local:wrong-reviewer')
            elif changed == 'time':
                assessment = replace(assessment, assessed_at=assessment.assessed_at + timedelta(days=1))
            elif changed == 'evidence_role':
                assessment = replace(assessment, evidence_role='wrong-evidence-role')
            else:
                source = assessment.segments[0]
                change = {'source_text': {'exact_text': 'changed supporting wording'},
                    'source_document': {'document_id': (source.document_id or 0)+1},
                    'source_locator': {'locator': 'sheet other, cell A1'},
                    'source_filename': {'filename': 'another.xlsx'}}[changed]
                assessment = replace(assessment, segments=(replace(source, **change), *assessment.segments[1:]))
            fact = replace(row.fact, assessments=(assessment, *row.fact.assessments[1:]))
            decisions.append(replace(row, fact=fact))
        return replace(output, source_decisions=tuple(decisions))
    monkeypatch.setattr(history, 'read_record_history', incorrect)
    result = collect(session, *adopted)
    observed = surfaces(result)['source_and_decision_history']
    assert not observed.reading.observed_record_kinds
    expected = 'actor, time, role or judgment' if changed in {'actor', 'time', 'evidence_role'} else 'source text/document/locator'
    assert any(expected in b for b in observed.blockers)
