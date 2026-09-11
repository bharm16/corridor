"""The Source Passage Check: one mechanical status, presented as itself.

ADR-0082 renames the quote-on-page flag ``locator_validation_status`` and
strips "verified" from customer displays unless the word names the object.
These tests hold three separate promises: the status is mechanical and
replayable from the original bytes, ``EvidenceLink.verified`` survives only as
the projection ``status == valid``, and no customer surface prints a bare
"verified" beside a value.
"""

from __future__ import annotations

import ast
from email.message import EmailMessage
from hashlib import sha256
from pathlib import Path
import re

from openpyxl import Workbook
import pytest
from sqlalchemy import select

from corridor.exceptions import RULES, format_exception_name
from corridor.export import COLUMNS
from corridor.ingest import ingest_document
from corridor.locator_validation import (
    INVALID,
    LOCATOR_VALIDATION_STATUSES,
    LocatorValidation,
    NOT_CHECKED,
    NOT_RE_READABLE,
    VALID,
    cited_passage_locator_validation_status,
    evidence_link_locator_validation_status,
    evidence_link_verified,
    recorded_verbal_statement_locator_validation,
    recorded_verbal_statement_locator_validation_status,
    source_passage_checks,
    source_segment_locator_validation,
    source_segment_locator_validation_status,
)
from corridor.email_segments import append_email_segments
from corridor.models import Document, SourceSegment
from corridor.source_append import SegmentValues, append_source_segments
from corridor.presentation import (
    documentation_review_label,
    field_label,
    label,
    provenance_label,
    resolution_strategy_label,
    source_passage_check_label,
    statement_type_label,
)
import corridor.source_segments as source_segments
from corridor.source_segments import recorded_verbal_statement_segment

from pdf_fixture_support import PdfFixture

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "src" / "corridor"
TEMPLATES = SOURCE / "web" / "templates"


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "utility-conflicts.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Utility Conflicts"
    sheet.append(["Conflict ID", "Owner"])
    sheet.append(["UC-1", "CenterPoint"])
    book.save(path)
    return path


def _ingest(session, project, path, tmp_path, doc_type="matrix"):
    return ingest_document(
        session,
        project_id=project.id,
        path=path,
        doc_type=doc_type,
        images_dir=tmp_path / "images",
    )


def _segments(session, document):
    return list(
        session.scalars(
            select(SourceSegment)
            .where(SourceSegment.document_id == document.id)
            .order_by(SourceSegment.ordinal)
        ).all()
    )


class _StoredCheck:
    """The one attribute a legacy evidence link contributes to the check."""

    def __init__(self, verified: bool) -> None:
        self.verified = verified


# --- The status is mechanical and replayable ------------------------------------


def test_a_workbook_cell_locator_that_recovers_its_text_is_valid(
    session, project, workbook, tmp_path
):
    document = _ingest(session, project, workbook, tmp_path)

    statuses = [
        source_segment_locator_validation_status(document, segment, workbook)
        for segment in _segments(session, document)
    ]

    assert statuses and set(statuses) == {VALID}


def test_the_same_bytes_and_locator_always_give_the_same_answer(
    session, project, workbook, tmp_path
):
    """Replayable: the status is recomputed, never remembered."""

    document = _ingest(session, project, workbook, tmp_path)
    segment = _segments(session, document)[0]

    first = source_segment_locator_validation_status(document, segment, workbook)
    second = source_segment_locator_validation_status(document, segment, workbook)

    assert (first, second) == (VALID, VALID)


def test_a_locator_that_no_longer_recovers_its_stored_text_is_invalid(
    session, project, workbook, tmp_path
):
    document = _ingest(session, project, workbook, tmp_path)
    segment = _segments(session, document)[0]
    segment.cell_range = "Z99"

    assert (
        source_segment_locator_validation_status(document, segment, workbook)
        == INVALID
    )


def test_a_tampered_digest_is_invalid_rather_than_raising(
    session, project, workbook, tmp_path
):
    """A caller asking for a status gets one; the replay still fails closed."""

    document = _ingest(session, project, workbook, tmp_path)
    segment = _segments(session, document)[0]
    segment.content_sha256 = "0" * 64

    assert (
        source_segment_locator_validation_status(document, segment, workbook)
        == INVALID
    )


def test_bytes_that_are_not_the_registered_document_are_invalid(
    session, project, workbook, tmp_path
):
    document = _ingest(session, project, workbook, tmp_path)
    segment = _segments(session, document)[0]
    other = tmp_path / "other.xlsx"
    other.write_bytes(workbook.read_bytes() + b"tampered")

    assert (
        source_segment_locator_validation_status(document, segment, other) == INVALID
    )


def test_a_pdf_span_locator_validates_against_the_registered_pdf(
    session, project, tmp_path
):
    path = tmp_path / "coordination-minutes.pdf"
    fixture = PdfFixture()
    fixture.add_page().text(
        (72, 72),
        "Meeting notes and attendance.\n"
        "Action Items:\n"
        "1. Equistar will submit the signed exhibit by March 2025.\n"
        "Meeting Notes",
    )
    fixture.save(path)
    document = _ingest(session, project, path, tmp_path, doc_type="minutes")

    segments = _segments(session, document)

    assert segments
    assert {
        source_segment_locator_validation_status(document, segment, path)
        for segment in segments
    } == {VALID}


def _retained_prose_citation(session, project, tmp_path):
    """A registered Document and one ``prose_span`` written by the retired reader."""

    path = tmp_path / "retained-minutes.pdf"
    fixture = PdfFixture()
    fixture.add_page().text((72, 72), "Equistar will submit the signed exhibit.")
    fixture.save(path)
    document = Document(
        project_id=project.id,
        sha256=sha256(path.read_bytes()).hexdigest(),
        filename=path.name,
        doc_type="minutes",
    )
    session.add(document)
    session.flush()
    words = "Equistar will submit the signed exhibit."
    (segment,) = append_source_segments(
        session,
        project_id=project.id,
        document_id=document.id,
        recorded_verbal_origin_id=None,
        segments=[
            SegmentValues(
                kind="prose_span",
                exact_text=words,
                content_sha256=sha256(words.encode("utf-8")).hexdigest(),
                ordinal=1,
                page_no=1,
                start_offset=0,
                end_offset=len(words),
            )
        ],
    )
    return document, segment, path


def test_a_retained_prose_citation_is_not_re_readable_and_never_invalid(
    session, project, tmp_path
):
    """The reader that wrote this locator left the product (#741, ADR-0094).

    A ``prose_span`` is a pair of offsets into the page string the retired
    reader produced, so no reader here can say whether that location still
    yields these words. What the check must not do is call the passage
    ``invalid``: that word means the locator was replayed and did not reach the
    passage, which is a statement about the source, and nothing replayed it.
    """

    document, segment, path = _retained_prose_citation(session, project, tmp_path)

    assert source_segment_locator_validation_status(document, segment, path) == (
        NOT_RE_READABLE
    )
    assert source_segment_locator_validation_status(document, segment, path) != INVALID

    # Everything the check can still establish without a reader, it still
    # establishes. A retained citation whose stored words no longer match the
    # digest recorded with them is an integrity failure that no page needed to
    # be opened to find, and it stays ``invalid``.
    segment.exact_text = "Equistar will submit something else."
    assert source_segment_locator_validation_status(document, segment, path) == INVALID


def test_without_registered_bytes_the_check_still_answers_what_needs_none(
    session, project, workbook, tmp_path
):
    """A reader that cannot stage the file passes ``None`` and gets the truth.

    The accepted statement reader used to keep its own ladder for this case.
    Now the one home answers: a cell it cannot open is ``not_checked`` with no
    reason of its own (the caller knows why it had no bytes), a stored text
    that disagrees with its digest is ``invalid`` without opening anything,
    and a retired ``prose_span`` is ``not_re_readable`` whether or not the
    file is here.
    """

    document = _ingest(session, project, workbook, tmp_path)
    segment = _segments(session, document)[0]
    retained_document, prose, _ = _retained_prose_citation(session, project, tmp_path)

    assert source_segment_locator_validation(document, segment, None) == (
        LocatorValidation(NOT_CHECKED, None)
    )

    retained = source_segment_locator_validation(retained_document, prose, None)
    assert retained.status == NOT_RE_READABLE
    assert "prose_span locator can only be re-read" in retained.reason

    # In memory only: the rows themselves are append-only.
    segment.exact_text = "different words"
    tampered = source_segment_locator_validation(document, segment, None)
    assert (tampered.status, tampered.reason) == (
        INVALID, "stored segment digest does not match its text"
    )


def test_every_answer_carries_the_replay_own_reason(
    session, project, workbook, tmp_path
):
    document = _ingest(session, project, workbook, tmp_path)
    segment = _segments(session, document)[0]

    assert source_segment_locator_validation(document, segment, workbook).reason is None

    segment.cell_range = "Z99"
    moved = source_segment_locator_validation(document, segment, workbook)
    assert moved.status == INVALID
    assert moved.reason == "spreadsheet locator does not exist: Utility Conflicts!Z99"

    verbal = recorded_verbal_statement_segment(
        project_id=1, recorded_verbal_origin_id=2, exact_text="They will pull the pole."
    )
    assert recorded_verbal_statement_locator_validation(verbal).reason is None
    verbal.content_sha256 = sha256(b"other words").hexdigest()
    assert recorded_verbal_statement_locator_validation(verbal).reason == (
        "stored segment digest does not match its recorded words"
    )


def _registered_email(session, project, tmp_path):
    """A registered ``.eml`` Document and the ``email_span`` rows read from it."""

    message = EmailMessage()
    message["From"] = "Utility Person <utility@example.test>"
    message["To"] = "project@example.test"
    message["Message-ID"] = "<one@example.test>"
    message.set_content("We will finish the relocation in October.\n")
    raw = message.as_bytes()
    path = tmp_path / "thread.eml"
    path.write_bytes(raw)
    document = Document(
        project_id=project.id,
        sha256=sha256(raw).hexdigest(),
        filename=path.name,
        doc_type="email",
    )
    session.add(document)
    session.flush()
    return document, append_email_segments(session, document, raw), path


def _second_workbook(tmp_path):
    path = tmp_path / "key-dates.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Key Dates"
    sheet.append(["Milestone", "Date"])
    sheet.append(["Letting", "2026-03-01"])
    sheet.append(["Clearance", "2025-12-15"])
    book.save(path)
    return path


def _counted_decodes(monkeypatch):
    """Count every workbook decode ``source_segments`` performs, either way in."""

    original = source_segments.load_workbook
    decodes = []

    def counted(source, **options):
        decodes.append(options)
        return original(source, **options)

    monkeypatch.setattr(source_segments, "load_workbook", counted)
    return decodes


def test_the_batch_answers_exactly_as_the_one_shot_on_every_branch(
    session, project, workbook, tmp_path
):
    """One ladder, two entrances: ``check`` may never disagree with the function.

    Each row below is one segment kind or one failure branch the one-shot
    function decides: a cell that replays, a cell whose stored digest was
    tampered, a cell moved to an empty location, a cell that belongs to
    another Document, bytes that are not the registered Document, a retired
    ``prose_span``, an email span with and without its bytes, and a cell with
    no bytes at all.  The batch must return the same status *and* the same
    reason for every one, since a reader prints the reason as the limitation.
    """

    document = _ingest(session, project, workbook, tmp_path)
    cells = _segments(session, document)
    other_bytes = tmp_path / "other.xlsx"
    other_bytes.write_bytes(workbook.read_bytes() + b"tampered")
    retained_document, prose, retained_path = _retained_prose_citation(
        session, project, tmp_path
    )
    email_document, email_spans, email_path = _registered_email(session, project, tmp_path)

    # In memory only, after every row is written: the rows are append-only.
    tampered, moved, foreign = cells[1], cells[2], cells[3]
    tampered.content_sha256 = "0" * 64
    moved.cell_range = "Z99"
    foreign.document_id = document.id + 1

    cases = [
        (document, cells[0], workbook),
        (document, tampered, workbook),
        (document, moved, workbook),
        (document, foreign, workbook),
        (document, cells[0], other_bytes),
        (document, cells[0], None),
        (document, tampered, None),
        (retained_document, prose, retained_path),
        (retained_document, prose, None),
        (email_document, email_spans[0], email_path),
        (email_document, email_spans[0], None),
    ]

    one_shot = [source_segment_locator_validation(*case) for case in cases]
    with source_passage_checks() as checks:
        batched = [checks.check(*case) for case in cases]

    assert batched == one_shot
    # The rows cover every status the check can return, so the equality above
    # is not vacuous.
    assert {answer.status for answer in one_shot} == {
        VALID, INVALID, NOT_CHECKED, NOT_RE_READABLE
    }
    assert [answer.reason for answer in one_shot[1:5]] == [
        "stored segment digest does not match its text",
        "spreadsheet locator does not exist: Utility Conflicts!Z99",
        "source segment does not belong to the supplied Document",
        "source bytes do not match the registered Document digest",
    ]

    # What the one-shot lets through, the batch lets through unchanged.
    missing = tmp_path / "not-staged.xlsx"
    with pytest.raises(FileNotFoundError):
        source_segment_locator_validation(document, cells[0], missing)
    with source_passage_checks() as checks:
        with pytest.raises(FileNotFoundError):
            checks.check(document, cells[0], missing)


def test_the_batch_decodes_each_spreadsheet_document_once(
    session, project, workbook, tmp_path, monkeypatch
):
    """N cells of one Document cost one decode; two Documents cost two.

    The accepted field reader runs the check over every accepted field of a
    project, and after it moved onto the one-shot function a Constraint Log
    read decoded the same workbook once per cell.  ``spreadsheet_replay``
    already owned the one-decode batch; this is the batch offered at the
    check's own interface.
    """

    first = _ingest(session, project, workbook, tmp_path)
    second_path = _second_workbook(tmp_path)
    second = _ingest(session, project, second_path, tmp_path)
    first_cells = _segments(session, first)
    second_cells = _segments(session, second)
    assert len(first_cells) >= 4 and len(second_cells) >= 4
    decodes = _counted_decodes(monkeypatch)

    one_shot = [
        source_segment_locator_validation(first, cell, workbook) for cell in first_cells
    ]
    assert len(decodes) == len(first_cells)
    assert set(one_shot) == {LocatorValidation(VALID)}

    decodes.clear()
    with source_passage_checks() as checks:
        batched = [checks.check(first, cell, workbook) for cell in first_cells]
    assert batched == one_shot
    assert len(decodes) == 1

    decodes.clear()
    with source_passage_checks() as checks:
        interleaved = [
            checks.check(document, cell, path)
            for cell_pair in zip(first_cells, second_cells)
            for document, cell, path in (
                (first, cell_pair[0], workbook), (second, cell_pair[1], second_path),
            )
        ]
    assert set(interleaved) == {LocatorValidation(VALID)}
    assert len(decodes) == 2

    # A closed batch answers nothing rather than an answer from an expired
    # replay.
    with pytest.raises(RuntimeError, match="closed"):
        checks.check(first, first_cells[0], workbook)


def test_the_not_re_readable_state_is_labelled_without_claiming_a_check(
    session, project, tmp_path
):
    assert source_passage_check_label(NOT_RE_READABLE) == (
        "Cited location cannot be re-read"
    )
    assert NOT_RE_READABLE in LOCATOR_VALIDATION_STATUSES
    # The compatibility projection stays exactly ``status == valid``: a
    # passage nobody could re-read is not a verified one.
    assert evidence_link_verified(NOT_RE_READABLE) is False


def test_a_recorded_verbal_statement_validates_on_its_own_digest():
    """ADR-0033: a verbal has no document, so the digest is the whole check."""

    segment = recorded_verbal_statement_segment(
        project_id=1,
        recorded_verbal_origin_id=2,
        exact_text="They will pull the pole in March.",
    )

    assert recorded_verbal_statement_locator_validation_status(segment) == VALID

    segment.content_sha256 = sha256(b"different words").hexdigest()

    assert recorded_verbal_statement_locator_validation_status(segment) == INVALID


def test_the_legacy_page_locator_replays_through_the_same_predicate():
    page = "The utility owner will relocate the 12-inch main by March 2025."

    assert (
        cited_passage_locator_validation_status(
            "relocate the 12-inch main by March 2025", page
        )
        == VALID
    )
    assert (
        cited_passage_locator_validation_status("abandon the main in place", page)
        == INVALID
    )


def test_no_locator_to_dereference_is_not_run_rather_than_failed():
    """"Never checked" and "checked and absent" are different facts."""

    assert cited_passage_locator_validation_status(None, "any page text") == (
        NOT_CHECKED
    )
    assert cited_passage_locator_validation_status("a quote", None) == NOT_CHECKED
    assert evidence_link_locator_validation_status(None) == NOT_CHECKED


# --- The retired flag is only a projection ---------------------------------------


def test_verified_is_exactly_the_projection_of_a_valid_status():
    """ADR-0082 keeps the column as ``locator_validation_status == valid``."""

    assert evidence_link_verified(VALID) is True
    assert evidence_link_verified(INVALID) is False
    assert evidence_link_verified(NOT_CHECKED) is False
    assert evidence_link_verified(NOT_RE_READABLE) is False
    assert set(LOCATOR_VALIDATION_STATUSES) == {
        VALID,
        INVALID,
        NOT_CHECKED,
        NOT_RE_READABLE,
    }

    with pytest.raises(ValueError, match="unknown locator validation status"):
        evidence_link_verified("passed")


def test_a_stored_link_reads_back_as_a_status_not_a_boolean():
    assert evidence_link_locator_validation_status(_StoredCheck(True)) == VALID
    assert evidence_link_locator_validation_status(_StoredCheck(False)) == INVALID


def test_the_legacy_stored_check_can_never_report_a_retired_locator_scheme():
    """A boolean has no way to say ``not_re_readable``, so this reader cannot.

    The Constraint log's ``AssertionView`` carries this reader's answer, and
    its comment used to describe the field as the Source Passage Check itself.
    It is the stored column projected into that vocabulary: whether the check
    ran, and what it recorded when it did. Discovering that a locator scheme
    has been retired takes the replay, which this projection never performs.

    The status family still holds all four -- the readers that do replay need
    the fourth -- so what is pinned here is the reachable set of this one
    function, over its whole input domain: no link, and both booleans.
    """
    reachable = {
        evidence_link_locator_validation_status(link)
        for link in (None, _StoredCheck(True), _StoredCheck(False))
    }

    assert reachable == {NOT_CHECKED, VALID, INVALID}
    assert NOT_RE_READABLE in LOCATOR_VALIDATION_STATUSES


def test_no_screen_reads_the_retired_column():
    """"No new code reads it" is a property of the tree, not a convention.

    ADR-0081 froze the legacy writers and readers where they stand, and #458
    removes the column with them, so this does not claim the whole codebase
    has moved. It claims the display path has: every screen and every label
    adapter asks ``locator_validation`` for a status, so what a customer sees
    can no longer be a raw boolean.
    """

    attribute = re.compile(
        r"(?<![\w.])(?:EvidenceLink|link|_link|a|assertion)\.verified\b"
    )
    display = [
        *sorted(TEMPLATES.glob("*.html")),
        SOURCE / "presentation.py",
        SOURCE / "export.py",
        SOURCE / "briefing.py",
        SOURCE / "web" / "app.py",
        SOURCE / "web" / "queue.py",
    ]
    offenders = {
        path.name for path in display if attribute.search(path.read_text())
    }

    assert offenders == set(), f"these screens read the retired flag: {offenders}"


def test_locator_validation_never_reads_the_support_relation():
    """ADR-0082: the check is mechanical and may not stand in for support."""

    source = (SOURCE / "locator_validation.py").read_text()
    tree = ast.parse(source)
    docstring_end = tree.body[0].end_lineno if ast.get_docstring(tree) else 0
    body = "\n".join(
        line
        for line in source.splitlines()[docstring_end:]
        if not line.strip().startswith("#")
    )

    for forbidden in (
        "support_assessment",
        "SupportAssessment",
        "operative_support",
        "satisfies_requirement",
    ):
        assert forbidden not in body, f"the check reads {forbidden}"


# --- Nothing shows the flag as a bare "verified" ---------------------------------


_WORD = re.compile(r"\b(?:un)?verif(?:y|ied|ication|ies)\b", re.IGNORECASE)
_MARKUP = re.compile(r"\{%.*?%\}|\{\{.*?\}\}|\{#.*?#\}|<[^>]*>", re.DOTALL)

# ADR-0082 permits "verified" where it names the object. The revision-change
# explanation names one completed document-revision comparison, not the
# passage check, so the word is the object's name and not a status beside a
# value. Every other customer-facing occurrence is a defect.
NAMES_ITS_OBJECT = ("the verified comparison",)


def _customer_labels() -> dict[str, str]:
    """Every customer word these adapters can return, by where it comes from."""

    labels: dict[str, str] = {}
    for key in (
        "constraint",
        "constraint_log",
        "cited_passage",
        "documentation_review",
        "evidence",
        "provenance",
        "source_passage_check",
        "supporting_documents",
        "required_documents",
        "constraint_alerts",
        "record_conclusion",
        "coordination_decision",
        "assertion",
        "derivation",
        "verbal",
    ):
        labels[f"label({key})"] = label(key)
    for rule in RULES:
        labels[f"exception_name({rule})"] = format_exception_name(rule)
    for status in LOCATOR_VALIDATION_STATUSES:
        labels[f"source_passage_check_label({status})"] = source_passage_check_label(
            status
        )
    for field in (
        "external_org",
        "evidence_required",
        "committed_date",
        "resolution_strategy",
        "need_date",
    ):
        labels[f"field_label({field})"] = field_label(field)
    for kind in ("assertion", "derivation", "verbal", "evidence", "cited"):
        labels[f"provenance_label({kind})"] = provenance_label(kind)
    for event in ("commitment", "committed_date_change", "closure"):
        labels[f"statement_type_label({event})"] = statement_type_label(event)
    for strategy in ("relocate", "remove", "abandon_in_place"):
        labels[f"resolution_strategy_label({strategy})"] = resolution_strategy_label(
            strategy
        )
    for sufficient in (True, False):
        labels[f"documentation_review_label({sufficient})"] = (
            documentation_review_label(sufficient)
        )
    for index, column in enumerate(COLUMNS):
        labels[f"export column {index}"] = column
    return labels


def _receipt_and_explanation_text() -> dict[str, str]:
    """The literal customer sentences receipts and explanations carry.

    A Constraint Alert's detail is the sentence a Coordination Report, an
    export, a release receipt, and a change explanation all quote, so reading
    the engine's own string literals covers every one of them at the source.
    """

    text: dict[str, str] = {}
    for name in ("exceptions.py", "presentation.py", "export.py"):
        tree = ast.parse((SOURCE / name).read_text())
        # A docstring explains the code to the next engineer and is never
        # printed, so only the literals a reader can end up holding count.
        docstrings = {
            id(node.body[0].value)
            for node in ast.walk(tree)
            if isinstance(
                node,
                (ast.Module, ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef),
            )
            and node.body
            and isinstance(node.body[0], ast.Expr)
            and isinstance(node.body[0].value, ast.Constant)
            and isinstance(node.body[0].value.value, str)
        }
        for node in ast.walk(tree):
            if (
                isinstance(node, ast.Constant)
                and isinstance(node.value, str)
                and id(node) not in docstrings
            ):
                text[f"{name}:{node.lineno}"] = node.value
    return text


def test_no_customer_surface_shows_the_flag_as_verified():
    """The gate ADR-0082 asks for, over templates, labels, exports, receipts.

    Templates are read as the prose a reader sees: Jinja tags and HTML tags
    are stripped first, so an attribute or a status identifier in the code is
    not mistaken for a word on the page.
    """

    offenders: list[str] = []

    for path in sorted(TEMPLATES.glob("*.html")):
        prose = _MARKUP.sub(" ", path.read_text())
        for match in _WORD.finditer(prose):
            window = prose[max(0, match.start() - 60) : match.end() + 60]
            if any(phrase in " ".join(window.split()) for phrase in NAMES_ITS_OBJECT):
                continue
            offenders.append(f"{path.name}: {' '.join(window.split())}")

    for where, value in {
        **_customer_labels(),
        **_receipt_and_explanation_text(),
    }.items():
        if _WORD.search(value) and not any(
            phrase in " ".join(value.split()) for phrase in NAMES_ITS_OBJECT
        ):
            offenders.append(f"{where}: {' '.join(value.split())}")

    assert offenders == [], "\n".join(offenders)


def test_the_flag_is_presented_as_the_source_passage_check_with_its_state():
    """The state names where the passage was looked for, not a verdict (#600).

    Passed, Failed, and Not run each read as a judgment on the value or on
    inspected work, and "Not run" also implied a scheduled check that was
    skipped rather than a locator there was never anything to replay.
    """

    assert label("source_passage_check") == "Source passage check"
    assert [
        source_passage_check_label(s)
        for s in (VALID, INVALID, NOT_CHECKED, NOT_RE_READABLE)
    ] == [
        "Found at cited location",
        "Not found at cited location",
        "No cited location recorded",
        "Cited location cannot be re-read",
    ]


def test_the_three_concepts_stay_visibly_distinct():
    """The check, Supporting Documentation in Use, and Record Inclusion.

    A passed check is not a designation that a passage backs this value, and
    neither is a record decision. The Constraint page keeps them in separate
    columns, and the words themselves do not overlap.
    """

    distinct = {
        label("source_passage_check"),
        label("documentation_review"),
        label("supporting_documents"),
        label("record_conclusion"),
    }
    assert len(distinct) == 4

    dependency_page = (TEMPLATES / "dependency.html").read_text()
    header = dependency_page[
        dependency_page.index("<th>Source document</th>") :
    ].split("</tr>")[0]
    assert "<th>Source passage check</th>" in header
    assert "<th>Documentation review</th>" in header
