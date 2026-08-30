"""Internal working Coordination Report: view, facets, and workbook.

These exercise the public HTTP surface on real PostgreSQL. The internal view
reuses the coherent project reading, Evaluation, report and export behavior of
the external-release path (ADR-0040) but seals nothing and advances no
comparison baseline (ADR-0053): viewing, refreshing, filtering, or downloading
here reads only.
"""

import hashlib
from datetime import date, datetime, timedelta, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from corridor import audit
from corridor.adjudicate import accept_candidate
from corridor.db import Session, engine
from corridor.exceptions import evaluate_project
from corridor.extraction_runs import declare_active_run, record_extraction_run
from corridor.external_statements import (
    CitedStatementEvidence,
    StatementScope,
    StatementTiming,
    record_external_party_statement,
)
from corridor.ledger import mark_satisfies
from corridor.models import (
    AuditLog,
    Candidate,
    Dependency,
    DocPage,
    Document,
    EvidenceLink,
    ExternalOrg,
    ExternalReportArtifact,
    Project,
    ReportRun,
)
from corridor.principals import HumanPrincipal
from corridor.verbal import record_verbal
from access_support import seed_membership

TEST_PRINCIPAL = HumanPrincipal("local:test-reviewer")


@pytest.fixture
def session():
    connection = engine.connect()
    trans = connection.begin()
    s = Session(bind=connection)
    yield s
    s.close()
    trans.rollback()
    connection.close()


@pytest.fixture
def client(session):
    """The app shares the test's transaction, so nothing is committed."""
    from corridor.web.app import app, get_human_principal, get_session

    app.dependency_overrides[get_session] = lambda: session
    app.dependency_overrides[get_human_principal] = lambda: TEST_PRINCIPAL
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


@pytest.fixture
def client_without_session(session):
    """No signed-in session: exercise the real fail-closed identity gate (#331)."""
    from corridor.web.app import app, get_session

    app.dependency_overrides.clear()
    app.dependency_overrides[get_session] = lambda: session
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


@pytest.fixture
def project(session):
    project = Project(slug="internal-report-test", name="Internal Report Test", is_synthetic=True)
    session.add(project)
    session.flush()
    seed_membership(session, project, TEST_PRINCIPAL)
    return project


def _document(session, project, *, filename, text, page_no=1):
    document = Document(
        project_id=project.id,
        sha256=hashlib.sha256(f"{project.id}:{filename}".encode()).hexdigest(),
        filename=filename,
        doc_type="matrix",
        parse_status="parsed",
        pages=28,
    )
    session.add(document)
    session.flush()
    session.add(
        DocPage(
            document_id=document.id,
            page_no=page_no,
            text=text,
            image_path="/tmp/corridor-missing-page.png",
        )
    )
    session.flush()
    return document


def _accept_dependency(session, project, document, *, ref, org, page_no=1):
    """One accepted constraint with a verified citation."""
    candidate = Candidate(
        project_id=project.id,
        kind="dependency",
        payload_json={
            "kind": "dependency",
            "fields": {
                "utility_id": ref,
                "external_org": org,
                "utility_type": "Telecom",
                "station_from": "1149+00",
                "station_to": "1153+17",
            },
            "citations": [
                {
                    "document_id": document.id,
                    "page": page_no,
                    "quote": f"{ref} {org} Telecom",
                    "verified": True,
                }
            ],
            "confidence": 1.0,
            "dedupe_hint": ref,
        },
        source_document_id=document.id,
        source_pages=[page_no],
        confidence=1.0,
        prompt_version="txdot_ucm_v1",
        citations_verified=True,
    )
    session.add(candidate)
    run = record_extraction_run(
        session,
        document,
        prompt_version=candidate.prompt_version,
        candidate_count=1,
        page_errors=0,
        candidates=(candidate,),
        model=candidate.model,
        allow_unsealed_legacy=True,
    )
    session.flush()
    declare_active_run(session, document.id, run.id, principal=TEST_PRINCIPAL)
    session.flush()
    session.add(ExternalOrg(name=org, aliases=[]))
    session.flush()
    accept_candidate(session, candidate, principal=TEST_PRINCIPAL)
    return session.scalars(
        select(Dependency).where(
            Dependency.project_id == project.id, Dependency.source_ref == ref
        )
    ).one()


def _seed_one_overdue_constraint(session, project):
    """A constraint with an authoritative, exact, past-due committed date."""
    document = _document(
        session,
        project,
        filename="overdue-inventory.pdf",
        text=(
            "FOC1-1 Overdue Utility Telecom\n"
            "Overdue Utility will finish relocation on 2026-06-01."
        ),
    )
    dependency = _accept_dependency(
        session, project, document, ref="FOC1-1", org="Overdue Utility"
    )
    party = session.get(ExternalOrg, dependency.external_org_id)
    record_external_party_statement(
        session,
        project_id=project.id,
        affected_external_org_id=party.id,
        stated_party=party.name,
        stated_external_org_id=party.id,
        source_kind="cited",
        event_date=date(2026, 5, 1),
        description="Overdue Utility will finish relocation on 2026-06-01.",
        new_timing=StatementTiming.day("2026-06-01", date(2026, 6, 1)),
        scope=StatementScope.selected((dependency.id,)),
        created_by="corridor:event-admission",
        evidence=CitedStatementEvidence(
            document.id, 1, "Overdue Utility will finish relocation on 2026-06-01."
        ),
    )
    session.flush()
    return dependency


def _bare_dependencies(session, project, count, *, prefix):
    """Cheap constraints with no evidence; each raises MISSING_EVIDENCE."""
    made = []
    for i in range(count):
        dependency = Dependency(
            project_id=project.id,
            ref_code=f"DEP-{prefix}-{i:03d}",
            source_ref=f"{prefix}{i:03d}",
            dep_type="utility_relocation",
            title=f"{prefix} constraint {i}",
        )
        session.add(dependency)
        made.append(dependency)
    session.flush()
    return made


def _report_run_count(session, project_id):
    return session.scalar(
        select(func.count()).select_from(ReportRun).where(
            ReportRun.project_id == project_id
        )
    )


def _artifact_count(session, project_id):
    return session.scalar(
        select(func.count()).select_from(ExternalReportArtifact).where(
            ExternalReportArtifact.project_id == project_id
        )
    )


def _frontend_receipt_count(session, project_id):
    return session.scalar(
        select(func.count()).select_from(AuditLog).where(
            AuditLog.action == audit.PRODUCT_PROVING_FRONTEND_REQUEST,
            AuditLog.entity_type == audit.PROJECT,
            AuditLog.entity_id == project_id,
        )
    )


# --- AC1: the ordinary entry shows current facts, no approval step ----------


def test_internal_entry_shows_current_facts_and_is_marked_internal(
    client, session, project
):
    _seed_one_overdue_constraint(session, project)
    response = client.get(f"/internal-report/{project.slug}")

    assert response.status_code == 200
    body = response.text
    assert "Internal working report" in body
    # Clearly not an external release, and it explains it advances no baseline.
    assert "not an approved external release" in body.lower()
    assert "ADR-0040" in body and "ADR-0053" in body
    # The Evaluation date and ruleset are shown. build_report evaluates at the
    # UTC date (report.py), so the expected date is UTC, not the local day.
    assert datetime.now(timezone.utc).date().isoformat() in body
    assert "ruleset v0.4" in body
    # No approval control on the internal entry.
    assert "Prepare fixed PDF" not in body
    assert 'action="/reports/' not in body


# --- AC2 / AC9: viewing and downloading move no retained record -------------


def test_viewing_and_downloading_append_no_run_or_baseline(client, session, project):
    _seed_one_overdue_constraint(session, project)
    runs_before = _report_run_count(session, project.id)
    artifacts_before = _artifact_count(session, project.id)
    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    committed_before = dependency.committed_date

    # View, refresh with a provenance filter, drill a facet, and download.
    assert client.get(f"/internal-report/{project.slug}").status_code == 200
    assert client.get(f"/internal-report/{project.slug}").status_code == 200
    assert (
        client.get(f"/internal-report/{project.slug}?document_only=true").status_code
        == 200
    )
    assert client.get(f"/internal-report/{project.slug}/full").status_code == 200
    assert (
        client.get(f"/internal-report/{project.slug}/alerts/OVERDUE").status_code == 200
    )
    workbook = client.get(f"/internal-report/{project.slug}/workbook.xlsx")
    assert workbook.status_code == 200

    session.expire_all()
    assert _report_run_count(session, project.id) == runs_before == 0
    assert _artifact_count(session, project.id) == artifacts_before == 0
    # It cannot alter project facts.
    dependency = session.get(Dependency, dependency.id)
    assert dependency.committed_date == committed_before


def test_the_internal_view_records_no_external_release_history(client, session, project):
    _seed_one_overdue_constraint(session, project)
    for _ in range(2):
        client.get(f"/internal-report/{project.slug}")
        client.get(f"/internal-report/{project.slug}/full")
        client.get(f"/internal-report/{project.slug}/workbook.xlsx")
    from corridor.report_release import external_report_release_history

    assert external_report_release_history(session, project.id) == ()


# --- AC3: one coherent bound reading across every surface -------------------


def test_every_surface_reads_one_coherent_bound_reading(client, session, project):
    _seed_one_overdue_constraint(session, project)
    _bare_dependencies(session, project, 2, prefix="COH")

    landing = client.get(f"/internal-report/{project.slug}").text
    alerts = client.get(f"/internal-report/{project.slug}/alerts/OVERDUE").text
    workbook = client.get(f"/internal-report/{project.slug}/workbook.xlsx").content

    sheet = load_workbook(BytesIO(workbook))
    meta = {row[0]: row[1] for row in sheet["Source traceability"].values}
    covered = 3  # one overdue + two bare

    # Same evaluation date, ruleset, population across the view and the export.
    # build_report evaluates at the UTC date (report.py), so expect UTC here.
    evaluated_on = datetime.now(timezone.utc).date().isoformat()
    assert meta["Evaluated on"] == evaluated_on
    assert meta["Ruleset version"] == "v0.4"
    assert meta["Records"] == covered
    assert evaluated_on in landing
    assert evaluated_on in alerts
    assert "Covered constraints" in landing
    assert f"<dd>{covered}</dd>" in landing


# --- AC4: facets reach the complete matching population, not one page -------


def test_facets_reach_the_complete_population_with_bounded_navigation(
    client, session, project
):
    # 55 constraints all raise MISSING_EVIDENCE — more than one alert page.
    made = _bare_dependencies(session, project, 55, prefix="EVID")
    complete = {
        e.dependency_id
        for e in evaluate_project(session, project.id).found
        if e.rule == "MISSING_EVIDENCE"
    }
    assert len(complete) == 55

    landing = client.get(f"/internal-report/{project.slug}").text
    # The facet count is the complete population, not a page of the Work List.
    assert ">55<" in landing

    page1 = client.get(
        f"/internal-report/{project.slug}/alerts/MISSING_EVIDENCE"
    ).text
    assert "Matching constraints</dt><dd>55</dd>" in page1
    assert "Page 1 of 2" in page1
    refs_page1 = {d.ref_code for d in made if d.ref_code in page1}
    assert len(refs_page1) == 50

    page2 = client.get(
        f"/internal-report/{project.slug}/alerts/MISSING_EVIDENCE?page=2"
    ).text
    assert "Page 2 of 2" in page2
    refs_page2 = {d.ref_code for d in made if d.ref_code in page2}
    assert len(refs_page2) == 5
    # Every matching constraint is reachable across the bounded pages.
    assert refs_page1 | refs_page2 == {d.ref_code for d in made}


# --- AC5: navigation preserves scope, drills to source, empty explains -------


def test_alert_rows_drill_through_to_the_constraint_within_the_project(
    client, session, project
):
    _bare_dependencies(session, project, 1, prefix="DRILL")
    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    page = client.get(
        f"/internal-report/{project.slug}/alerts/MISSING_EVIDENCE"
    ).text
    assert f'href="/ledger/{project.slug}/{dependency.id}"' in page


def test_an_empty_alert_population_explains_itself(client, session, project):
    _bare_dependencies(session, project, 1, prefix="EMPTY")
    # Nothing is overdue here, so the OVERDUE population is empty.
    page = client.get(f"/internal-report/{project.slug}/alerts/OVERDUE")
    assert page.status_code == 200
    body = page.text
    assert "Matching constraints</dt><dd>0</dd>" in body
    assert "No constraint currently raises this alert" in body
    # It must not invent a favorable conclusion.
    assert "all clear" not in body.lower()
    assert "complete" in body.lower()  # "...not that the underlying work is complete."


def test_an_unknown_alert_rule_is_refused(client, session, project):
    _bare_dependencies(session, project, 1, prefix="RULE")
    assert (
        client.get(f"/internal-report/{project.slug}/alerts/NOT_A_RULE").status_code
        == 404
    )


# --- AC6: internal XLSX uses the export contract, marked internal ------------


def test_internal_workbook_uses_export_contract_and_is_marked_internal(
    client, session, project
):
    _seed_one_overdue_constraint(session, project)
    response = client.get(f"/internal-report/{project.slug}/workbook.xlsx")

    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    assert "internal-working-constraint-log" in disposition

    workbook = load_workbook(BytesIO(response.content))
    assert "Constraint log" in workbook.sheetnames
    meta = {row[0]: row[1] for row in workbook["Source traceability"].values}
    assert meta["Working view"] == (
        "Internal working copy — not an approved external release."
    )
    assert meta["Ruleset version"] == "v0.4"


# --- AC7: honest visibility according to the selected provenance mode --------


def test_provenance_mode_selects_verbal_visibility(client, session, project):
    document = _document(
        session,
        project,
        filename="verbal-inventory.pdf",
        text="FOC1-1 Verbal Utility Telecom",
    )
    dependency = _accept_dependency(
        session, project, document, ref="FOC1-1", org="Verbal Utility"
    )
    # A verbal-only past-due commitment: visible under all sources, excluded
    # under documents-only.
    record_verbal(
        session,
        dependency,
        stated_party="Verbal Utility",
        description="Verbal Utility said relocation finished long ago.",
        conversation_date=date.today() - timedelta(days=40),
        committed_date=date.today() - timedelta(days=30),
        principal=TEST_PRINCIPAL,
    )
    session.flush()

    all_sources = client.get(
        f"/internal-report/{project.slug}/alerts/OVERDUE"
    ).text
    assert dependency.ref_code in all_sources

    documents_only = client.get(
        f"/internal-report/{project.slug}/alerts/OVERDUE?document_only=true"
    ).text
    assert "Matching constraints</dt><dd>0</dd>" in documents_only
    assert dependency.ref_code not in documents_only


def test_approximate_timing_is_not_converted_to_an_invented_overdue_date(
    client, session, project
):
    document = _document(
        session,
        project,
        filename="month-inventory.pdf",
        text=(
            "FOC1-1 Month Utility Telecom\n"
            "Month Utility now expects completion in February 2026."
        ),
    )
    dependency = _accept_dependency(
        session, project, document, ref="FOC1-1", org="Month Utility"
    )
    party = session.get(ExternalOrg, dependency.external_org_id)
    # A month-precision commitment in the past must not become an exact overdue.
    record_external_party_statement(
        session,
        project_id=project.id,
        affected_external_org_id=party.id,
        stated_party=party.name,
        stated_external_org_id=party.id,
        source_kind="cited",
        event_date=date(2026, 1, 5),
        description="Month Utility now expects completion in February 2026.",
        new_timing=StatementTiming.month("February 2026", 2026, 2),
        scope=StatementScope.selected((dependency.id,)),
        created_by="corridor:event-admission",
        evidence=CitedStatementEvidence(
            document.id, 1, "Month Utility now expects completion in February 2026."
        ),
    )
    session.flush()

    overdue = client.get(f"/internal-report/{project.slug}/alerts/OVERDUE").text
    assert "Matching constraints</dt><dd>0</dd>" in overdue
    assert dependency.ref_code not in overdue


def test_full_report_shows_unknown_scope_party_commitments(client, session, project):
    party = ExternalOrg(name="Unknown Scope Utility")
    session.add(party)
    session.flush()
    quote = "Unknown Scope Utility to provide chain of title (Due date of 01/2026)."
    document = _document(
        session, project, filename="unknown-scope-minutes.pdf", text=quote
    )
    record_external_party_statement(
        session,
        project_id=project.id,
        affected_external_org_id=party.id,
        stated_party=party.name,
        stated_external_org_id=party.id,
        source_kind="cited",
        event_date=date(2026, 1, 16),
        description=quote,
        new_timing=StatementTiming.month("01/2026", 2026, 1),
        scope=StatementScope.unknown(),
        created_by="local:coordinator",
        evidence=CitedStatementEvidence(document.id, 1, quote),
    )
    session.flush()

    full = client.get(f"/internal-report/{project.slug}/full").text
    assert "Unknown Scope Utility" in full
    # Unknown Applies To is shown honestly, not resolved to a constraint.
    assert "not yet known" in full.lower()
    # The full report is also marked internal.
    assert "not an approved external release" in full.lower()


# --- AC8: current-source rules apply; access contract participates ----------


def test_an_unverified_cited_date_never_reappears_through_the_workbook(
    client, session, project
):
    document = _document(
        session,
        project,
        filename="unverified-inventory.pdf",
        text="FOC1-1 Unverified Utility Telecom",
    )
    dependency = _accept_dependency(
        session, project, document, ref="FOC1-1", org="Unverified Utility"
    )
    from corridor.models import DependencyEvent, DependencyEventScope, DependencyEventTiming

    event = DependencyEvent(
        project_id=project.id,
        affected_external_org_id=dependency.external_org_id,
        stated_external_org_id=dependency.external_org_id,
        scope_mode="selected",
        event_type="commitment",
        source_kind="cited",
        stated_party="Unverified Utility",
        event_date=date(2026, 5, 8),
        description="Unverified Utility will finish on a date with no evidence.",
        created_by="corridor:event-admission",
    )
    session.add(event)
    session.flush()
    session.add_all(
        (
            DependencyEventScope(event_id=event.id, dependency_id=dependency.id),
            DependencyEventTiming(
                event_id=event.id,
                kind="new",
                text="past",
                precision="day",
                start_date=date(2026, 6, 1),
                end_date=date(2026, 6, 1),
            ),
        )
    )
    session.flush()

    workbook = client.get(f"/internal-report/{project.slug}/workbook.xlsx").content
    sheet = load_workbook(BytesIO(workbook))["Constraint log"]
    headers = [c.value for c in sheet[1]]
    row = {h: c.value for h, c in zip(headers, sheet[2])}
    assert row["Promised for"] is None
    assert "OVERDUE" not in (row["Constraint alerts"] or "")


def test_new_routes_write_a_project_access_receipt(client, session, project):
    _bare_dependencies(session, project, 1, prefix="RCPT")
    before = _frontend_receipt_count(session, project.id)
    client.get(f"/internal-report/{project.slug}")
    client.get(f"/internal-report/{project.slug}/full")
    client.get(f"/internal-report/{project.slug}/alerts/MISSING_EVIDENCE")
    client.get(f"/internal-report/{project.slug}/workbook.xlsx")
    session.expire_all()
    assert _frontend_receipt_count(session, project.id) == before + 4


@pytest.mark.parametrize(
    "path",
    [
        "/internal-report/{slug}",
        "/internal-report/{slug}/full",
        "/internal-report/{slug}/alerts/MISSING_EVIDENCE",
        "/internal-report/{slug}/workbook.xlsx",
    ],
)
def test_new_routes_fail_closed_without_a_seeded_identity(
    client_without_session, project, path
):
    response = client_without_session.get(path.format(slug=project.slug))
    assert response.status_code == 401


def test_unknown_project_is_not_found(client, session, project):
    assert client.get("/internal-report/no-such-project").status_code == 404
    assert (
        client.get("/internal-report/no-such-project/workbook.xlsx").status_code == 404
    )


# --- cross-project output is not combined -----------------------------------


def test_one_projects_alert_population_excludes_another_projects_constraints(
    client, session, project
):
    _bare_dependencies(session, project, 2, prefix="MINE")
    other = Project(slug="internal-report-other", name="Other", is_synthetic=True)
    session.add(other)
    session.flush()
    foreign = _bare_dependencies(session, other, 2, prefix="THEIRS")

    page = client.get(
        f"/internal-report/{project.slug}/alerts/MISSING_EVIDENCE"
    ).text
    assert "Matching constraints</dt><dd>2</dd>" in page
    for dependency in foreign:
        assert dependency.ref_code not in page
        assert f"/ledger/{project.slug}/{dependency.id}" not in page
