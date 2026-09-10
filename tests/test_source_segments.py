"""Spreadsheet and Minutes Source Segments through ingest and replay seams.

The workbooks are written at test time so CI exercises real XLSX bytes without
depending on the separately fetched corpus.  The release rehearsal also reads a
fetched development-corpus workbook; these tests pin the same byte-to-segment
contract deterministically.
"""

from hashlib import sha256
import os
from pathlib import Path

from openpyxl import Workbook
import pytest
from sqlalchemy import select
from sqlalchemy.exc import DBAPIError, IntegrityError

from corridor.config import settings
from corridor.ingest import ingest_document
from corridor.models import Document, Project, SourceSegment
from corridor.prose_spans import is_prose_segment
from corridor.source_segments import (
    SourceDocumentDigestMismatch,
    SourceSegmentDigestMismatch,
    SourceSegmentLocatorMismatch,
    dereference_source_segment,
    spreadsheet_replay,
)

from pdf_fixture_support import PdfFixture

REAL_WORKBOOK_SHA256 = (
    "3cd94fea058a3e61ac95ab1efd566e684e146f64ce6f25048d93cf9db55f83ba"
)
CORPUS_STORE = Path(
    os.environ.get("CORRIDOR_TEST_CORPUS_STORE", settings.corpus_store)
)
REAL_WORKBOOK = (
    CORPUS_STORE
    / REAL_WORKBOOK_SHA256[:2]
    / f"{REAL_WORKBOOK_SHA256}.xlsx"
)
REAL_MINUTES_SHA256 = (
    "ada13da24574950264d974f8a352fd07a3dba042ac2b37e825f8409730182eaf"
)
REAL_MINUTES = (
    CORPUS_STORE / REAL_MINUTES_SHA256[:2] / f"{REAL_MINUTES_SHA256}.pdf"
)
# The exact wording the reader in the product recovers for this statement.
# The retired reader put a space before the line break inside "(Due \ndate" and
# cut this document into 135 page-prose spans; the paired-rendition reader's
# page-text projection does not, and cuts it into 82 (#741). Both are readings
# of the same registered bytes, so these numbers describe the reader rather
# than the document, and a citation retained under the old reading keeps the
# words it was written with.
REAL_MINUTES_STATEMENT = (
    "Equistar to provide a chain of title on the ROW agreement that is in DOW’s name "
    "(Due\ndate of 01/2025)."
)
needs_corpus = pytest.mark.skipif(
    not REAL_WORKBOOK.exists(),
    reason="run `make corpus` to fetch the I-35 NEX South workbook",
)
needs_minutes_corpus = pytest.mark.skipif(
    not REAL_MINUTES.exists(),
    reason="run `make corpus` to fetch the SH99 Equistar meeting notes",
)

MINUTES_STATEMENT = "Equistar will submit the signed exhibit by March 2025."


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "utility-conflicts.xlsx"
    book = Workbook()
    first = book.active
    first.title = "Summary"
    first.append(["Project", "IH 45", "  exact spacing  "])
    first.append(["Blank stays absent", None])
    conflicts = book.create_sheet("Utility Conflicts")
    conflicts.append(["Conflict ID", "Owner", "Station"])
    conflicts.append(["UC-1", "CenterPoint", 1149])
    book.save(path)
    return path


def _ingest(session, project, workbook, tmp_path):
    return ingest_document(
        session,
        project_id=project.id,
        path=workbook,
        doc_type="matrix",
        images_dir=tmp_path / "images",
    )


def _minutes_pdf(tmp_path):
    path = tmp_path / "coordination-minutes.pdf"
    fixture = PdfFixture()
    fixture.add_page().text(
        (72, 72),
        "Meeting notes and attendance.\n"
        "Action Items:\n"
        f"1. {MINUTES_STATEMENT}\n"
        "Meeting Notes",
    )
    return fixture.save(path)


def test_workbook_ingest_appends_one_exact_segment_per_populated_cell(
    session, project, workbook, tmp_path
):
    document = _ingest(session, project, workbook, tmp_path)

    segments = session.scalars(
        select(SourceSegment)
        .where(SourceSegment.document_id == document.id)
        .order_by(SourceSegment.ordinal)
    ).all()

    assert [
        (segment.kind, segment.sheet_name, segment.cell_range, segment.exact_text)
        for segment in segments
    ] == [
        ("spreadsheet_cell", "Summary", "A1", "Project"),
        ("spreadsheet_cell", "Summary", "B1", "IH 45"),
        ("spreadsheet_cell", "Summary", "C1", "  exact spacing  "),
        ("spreadsheet_cell", "Summary", "A2", "Blank stays absent"),
        ("spreadsheet_cell", "Utility Conflicts", "A1", "Conflict ID"),
        ("spreadsheet_cell", "Utility Conflicts", "B1", "Owner"),
        ("spreadsheet_cell", "Utility Conflicts", "C1", "Station"),
        ("spreadsheet_cell", "Utility Conflicts", "A2", "UC-1"),
        ("spreadsheet_cell", "Utility Conflicts", "B2", "CenterPoint"),
        ("spreadsheet_cell", "Utility Conflicts", "C2", "1149"),
    ]
    assert [segment.ordinal for segment in segments] == list(
        range(1, len(segments) + 1)
    )
    assert len({(segment.kind, segment.sheet_name, segment.cell_range) for segment in segments}) == len(segments)
    assert all(
        segment.content_sha256 == sha256(segment.exact_text.encode("utf-8")).hexdigest()
        for segment in segments
    )


@needs_corpus
def test_real_dev_corpus_workbook_registers_and_replays_from_its_addressed_bytes(
    session, project, tmp_path
):
    assert REAL_WORKBOOK.parent.name == REAL_WORKBOOK_SHA256[:2]
    assert REAL_WORKBOOK.stem == REAL_WORKBOOK_SHA256
    document = ingest_document(
        session,
        project_id=project.id,
        path=REAL_WORKBOOK,
        doc_type="matrix",
        images_dir=tmp_path / "images",
        filename="I-35 NEX SOUTH Potential Utility Conflicts.xlsx",
        expected_sha256=REAL_WORKBOOK_SHA256,
    )
    segments = session.scalars(
        select(SourceSegment)
        .where(SourceSegment.document_id == document.id)
        .order_by(SourceSegment.ordinal)
    ).all()

    assert document.parse_status == "parsed"
    assert len(segments) == 1668
    assert len({(row.kind, row.sheet_name, row.cell_range) for row in segments}) == 1668
    assert (
        segments[0].sheet_name,
        segments[0].cell_range,
        dereference_source_segment(document, segments[0], REAL_WORKBOOK),
    ) == (
        "UCM-Conflict List",
        "A1",
        "TxDOT Utility Conflict Management (UCM) - Potential Utility Conflict List",
    )
    assert (
        segments[-1].sheet_name,
        segments[-1].cell_range,
        dereference_source_segment(document, segments[-1], REAL_WORKBOOK),
    ) == ("Drop-Down Lists", "A57", "Rejected")


def test_every_segment_replays_from_the_original_workbook_bytes(
    session, project, workbook, tmp_path
):
    document = _ingest(session, project, workbook, tmp_path)
    segments = session.scalars(
        select(SourceSegment)
        .where(SourceSegment.document_id == document.id)
        .order_by(SourceSegment.ordinal)
    ).all()

    assert [
        dereference_source_segment(document, segment, workbook) for segment in segments
    ] == [segment.exact_text for segment in segments]


def test_minutes_pdf_ingest_appends_non_overlapping_replayable_prose_spans(
    session, project, tmp_path
):
    """The prose of a Minutes page, under the reader that now reads it.

    The locator scheme changed with the reader: a page's exact prose is a
    ``pdf_span`` on the page stream, cut by the same boundaries the retired
    ``prose_span`` used (#736, #741). What the segments have to be is
    unchanged -- non-overlapping, in page order, and replayable to their own
    stored words.
    """

    path = _minutes_pdf(tmp_path)

    document = ingest_document(
        session,
        project_id=project.id,
        path=path,
        doc_type="minutes",
        images_dir=tmp_path / "images",
    )
    segments = session.scalars(
        select(SourceSegment)
        .where(
            SourceSegment.document_id == document.id,
            SourceSegment.kind == "pdf_span",
            SourceSegment.span_stream == "page",
        )
        .order_by(SourceSegment.ordinal)
    ).all()

    assert MINUTES_STATEMENT in [segment.exact_text for segment in segments]
    assert all(is_prose_segment(segment) for segment in segments)
    assert all(
        segment.sheet_name is None
        and segment.cell_range is None
        and segment.page_no == 1
        and segment.start_offset is not None
        and segment.end_offset is not None
        and segment.start_offset < segment.end_offset
        for segment in segments
    )
    ordered_ranges = [
        (segment.start_offset, segment.end_offset) for segment in segments
    ]
    assert all(
        previous_end <= next_start
        for (_previous_start, previous_end), (next_start, _next_end) in zip(
            ordered_ranges, ordered_ranges[1:]
        )
    )
    assert [
        dereference_source_segment(document, segment, path) for segment in segments
    ] == [segment.exact_text for segment in segments]


@needs_minutes_corpus
def test_real_dev_corpus_minutes_replay_exact_statement_spans(
    session, project, tmp_path
):
    document = ingest_document(
        session,
        project_id=project.id,
        path=REAL_MINUTES,
        filename="Meeting Notes/Equistar/2025.02.12 GPB1 Equistar notes final.pdf",
        doc_type="minutes",
        images_dir=tmp_path / "images",
        expected_sha256=REAL_MINUTES_SHA256,
    )
    segments = session.scalars(
        select(SourceSegment)
        .where(
            SourceSegment.document_id == document.id,
            SourceSegment.kind == "pdf_span",
            SourceSegment.span_stream == "page",
        )
        .order_by(SourceSegment.ordinal)
    ).all()
    statement = next(
        segment for segment in segments if segment.exact_text == REAL_MINUTES_STATEMENT
    )

    assert len(segments) == 82
    assert statement.page_no == 2
    assert dereference_source_segment(document, statement, REAL_MINUTES) == (
        REAL_MINUTES_STATEMENT
    )


def test_replay_fails_closed_when_a_segment_digest_is_tampered(
    session, project, workbook, tmp_path
):
    document = _ingest(session, project, workbook, tmp_path)
    segment = session.scalars(
        select(SourceSegment).where(SourceSegment.document_id == document.id)
    ).first()
    segment.content_sha256 = "0" * 64

    with pytest.raises(SourceSegmentDigestMismatch, match="stored segment digest"):
        dereference_source_segment(document, segment, workbook)


def test_replay_refuses_bytes_other_than_the_registered_workbook(
    session, project, workbook, tmp_path
):
    document = _ingest(session, project, workbook, tmp_path)
    segment = session.scalars(
        select(SourceSegment).where(SourceSegment.document_id == document.id)
    ).first()
    other = tmp_path / "other.xlsx"
    other.write_bytes(workbook.read_bytes() + b"tampered")

    with pytest.raises(ValueError, match="registered Document digest"):
        dereference_source_segment(document, segment, other)


@pytest.fixture
def spreadsheet_replay_source(workbook):
    document = Document(id=1, project_id=1, sha256=sha256(workbook.read_bytes()).hexdigest())
    segment = SourceSegment(
        document_id=1, project_id=1, kind="spreadsheet_cell",
        sheet_name="Summary", cell_range="B1", exact_text="IH 45",
        content_sha256=sha256(b"IH 45").hexdigest(), ordinal=1,
    )
    return document, segment


def test_spreadsheet_replay_decodes_once_and_expires(
    workbook, spreadsheet_replay_source, monkeypatch,
):
    import corridor.source_segments as source_segments

    document, segment = spreadsheet_replay_source
    original_load = source_segments.load_workbook
    opens = []

    def counted_load(source, **options):
        opens.append(options)
        assert source.getvalue() == workbook.read_bytes()
        return original_load(source, **options)

    monkeypatch.setattr(source_segments, "load_workbook", counted_load)
    with spreadsheet_replay(document, workbook) as replay:
        assert [replay(segment) for _ in range(25)] == ["IH 45"] * 25
    assert opens == [{"data_only": True, "read_only": False}]
    with pytest.raises(SourceSegmentLocatorMismatch, match="closed"):
        replay(segment)


def test_spreadsheet_replay_refuses_changed_bytes_before_decoding(
    workbook, spreadsheet_replay_source, monkeypatch,
):
    import corridor.source_segments as source_segments

    document, _ = spreadsheet_replay_source
    workbook.write_bytes(workbook.read_bytes() + b"changed")

    def forbidden_decode(*args, **kwargs):
        pytest.fail("changed source bytes must be refused before workbook decoding")

    monkeypatch.setattr(source_segments, "load_workbook", forbidden_decode)
    with pytest.raises(SourceDocumentDigestMismatch, match="registered Document digest"):
        with spreadsheet_replay(document, workbook):
            pytest.fail("changed source bytes must never enter replay")


def test_spreadsheet_replay_refuses_changes_during_and_between_operations(
    workbook, spreadsheet_replay_source,
):
    document, segment = spreadsheet_replay_source
    with pytest.raises(SourceDocumentDigestMismatch, match="registered Document digest"):
        with spreadsheet_replay(document, workbook) as replay:
            assert replay(segment) == "IH 45"
            workbook.write_bytes(workbook.read_bytes() + b"changed")
            # The operation reads its already-verified snapshot, then refuses
            # the complete batch on exit because its source changed underneath it.
            assert replay(segment) == "IH 45"
    with pytest.raises(SourceDocumentDigestMismatch, match="registered Document digest"):
        with spreadsheet_replay(document, workbook):
            pytest.fail("a new operation must revalidate the source at the same path")


def test_spreadsheet_replay_refuses_changed_document_identity_on_exit(
    workbook, spreadsheet_replay_source,
):
    document, segment = spreadsheet_replay_source
    with pytest.raises(SourceSegmentLocatorMismatch, match="Document identity changed"):
        with spreadsheet_replay(document, workbook) as replay:
            assert replay(segment) == "IH 45"
            document.sha256 = "0" * 64


@pytest.mark.parametrize("changes,error,reason", [
    ({"document_id": 2}, SourceSegmentLocatorMismatch, "supplied Document"),
    ({"project_id": 2}, SourceSegmentLocatorMismatch, "supplied Document"),
    ({"kind": "prose_span"}, SourceSegmentLocatorMismatch, "cell segment"),
    ({"content_sha256": "0" * 64}, SourceSegmentDigestMismatch, "stored segment digest"),
    ({"sheet_name": None}, SourceSegmentLocatorMismatch, "incomplete"),
    ({"cell_range": "A0"}, SourceSegmentLocatorMismatch, "invalid spreadsheet cell"),
    ({"cell_range": "A1"}, SourceSegmentLocatorMismatch, "does not recover"),
    ({"sheet_name": "Absent"}, SourceSegmentLocatorMismatch, "does not exist"),
    ({"cell_range": "B2"}, SourceSegmentLocatorMismatch, "does not exist"),
])
def test_spreadsheet_replay_keeps_each_segment_integrity_check(
    workbook, spreadsheet_replay_source, changes, error, reason,
):
    document, segment = spreadsheet_replay_source
    for name, value in changes.items():
        setattr(segment, name, value)
    with spreadsheet_replay(document, workbook) as replay:
        with pytest.raises(error, match=reason):
            replay(segment)


def test_database_refuses_overlapping_spreadsheet_cell_locators(
    session, project, workbook, tmp_path
):
    document = _ingest(session, project, workbook, tmp_path)
    first = session.scalars(
        select(SourceSegment).where(SourceSegment.document_id == document.id)
    ).first()
    session.add(
        SourceSegment(
            project_id=project.id,
            document_id=document.id,
            kind=first.kind,
            exact_text=first.exact_text,
            content_sha256=first.content_sha256,
            ordinal=100,
            sheet_name=first.sheet_name,
            cell_range=first.cell_range,
        )
    )

    with pytest.raises(IntegrityError):
        session.flush()


def test_database_refuses_overlapping_prose_span_locators(
    session, project, tmp_path
):
    path = _minutes_pdf(tmp_path)
    document = ingest_document(
        session,
        project_id=project.id,
        path=path,
        doc_type="minutes",
        images_dir=tmp_path / "images",
    )
    first = session.scalars(
        select(SourceSegment)
        .where(
            SourceSegment.document_id == document.id,
            SourceSegment.kind == "pdf_span",
            SourceSegment.span_stream == "page",
        )
        .order_by(SourceSegment.ordinal)
    ).first()
    assert first.start_offset is not None
    assert first.end_offset is not None
    overlapping_text = first.exact_text[1:]
    session.add(
        SourceSegment(
            project_id=project.id,
            document_id=document.id,
            kind="pdf_span",
            exact_text=overlapping_text,
            content_sha256=sha256(overlapping_text.encode()).hexdigest(),
            ordinal=100,
            sheet_name=None,
            cell_range=None,
            page_no=first.page_no,
            span_stream=first.span_stream,
            start_offset=first.start_offset + 1,
            end_offset=first.end_offset,
            rendition_sha256=first.rendition_sha256,
            reading_sha256=first.reading_sha256,
            reader_identity=first.reader_identity,
            location_json=first.location_json,
        )
    )

    with pytest.raises(DBAPIError, match="prose source segments cannot overlap"):
        session.flush()


def test_database_refuses_a_segment_scoped_to_another_project(
    session, project, workbook, tmp_path
):
    document = _ingest(session, project, workbook, tmp_path)
    other = Project(slug="other-segment-project", name="Other", is_synthetic=True)
    session.add(other)
    session.flush()
    session.add(
        SourceSegment(
            project_id=other.id,
            document_id=document.id,
            kind="spreadsheet_cell",
            exact_text="foreign",
            content_sha256=sha256(b"foreign").hexdigest(),
            ordinal=100,
            sheet_name="Summary",
            cell_range="C1",
        )
    )

    with pytest.raises(IntegrityError):
        session.flush()


@pytest.mark.parametrize("operation", ["update", "delete"])
def test_source_segments_have_no_update_or_delete_path(
    session, project, workbook, tmp_path, operation
):
    document = _ingest(session, project, workbook, tmp_path)
    segment = session.scalars(
        select(SourceSegment).where(SourceSegment.document_id == document.id)
    ).first()
    if operation == "update":
        segment.exact_text = "rewritten"
    else:
        session.delete(segment)

    with pytest.raises(DBAPIError, match="source segments are append-only"):
        session.flush()
