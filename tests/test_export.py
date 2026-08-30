from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from corridor.adjudicate import accept_candidate
from corridor.db import Session, engine
from corridor.dependency_events import published_dependency_statements
from corridor.exceptions import Thresholds, evaluate_project, format_exception_label
from corridor.extraction_runs import declare_active_run, record_extraction_run
from corridor.export import COLUMNS, to_pdf, to_xlsx
from corridor.ledger import mark_satisfies
from corridor.models import (
    Candidate,
    Dependency,
    DependencyEvent,
    DependencyEventScope,
    DependencyEventTiming,
    DocPage,
    Document,
    EvidenceLink,
    ExternalOrg,
    Project,
)
from corridor.operative_support import designate_publication_support
from corridor.principals import HumanPrincipal
from corridor.report import build_report, render
from corridor.verbal import record_verbal

TEST_PRINCIPAL = HumanPrincipal("local:tester")


@pytest.fixture
def session():
    connection = engine.connect()
    trans = connection.begin()
    s = Session(bind=connection)
    yield s
    s.close()
    trans.rollback()
    connection.close()


@pytest.fixture
def project(session):
    project = Project(slug="exp-test", name="Export Test", is_synthetic=True)
    session.add(project)
    session.flush()
    session.add(ExternalOrg(name="Export Test Utility", aliases=[]))
    session.flush()
    doc = Document(
        project_id=project.id,
        sha256="x9" * 32,
        filename="nhhip-seg3c2-utilities-inventory-2-13-2026.pdf",
        doc_type="matrix",
        parse_status="parsed",
        pages=28,
    )
    session.add(doc)
    session.flush()
    session.add(
        DocPage(
            document_id=doc.id,
            page_no=4,
            text=(
                "FOC1-1 Export Test Utility Telecom\n"
                "Export Test Utility will finish relocation on August 15.\n"
                "Export Test Utility now expects completion in August 2026."
            ),
            image_path="/tmp/corridor-missing-page.png",
        )
    )
    session.flush()

    candidate = Candidate(
        project_id=project.id,
        kind="dependency",
        payload_json={
            "kind": "dependency",
            "fields": {
                "utility_id": "FOC1-1",
                "external_org": "Export Test Utility",
                "utility_type": "Telecom",
                "station_from": "1149+00",
                "station_to": "1153+17",
            },
            "citations": [
                {
                    "document_id": doc.id,
                    "page": 4,
                    "quote": "FOC1-1 Export Test Utility Telecom",
                    "verified": True,
                }
            ],
            "confidence": 1.0,
            "dedupe_hint": "x",
        },
        source_document_id=doc.id,
        source_pages=[4],
        confidence=1.0,
        prompt_version="txdot_ucm_v1",
        citations_verified=True,
    )
    session.add(candidate)
    run = record_extraction_run(
        session,
        doc,
        prompt_version=candidate.prompt_version,
        candidate_count=1,
        page_errors=0,
        candidates=(candidate,),
        model=candidate.model,
        allow_unsealed_legacy=True,
    )
    session.flush()
    declare_active_run(session, doc.id, run.id, principal=TEST_PRINCIPAL)
    session.flush()
    accept_candidate(session, candidate, principal=TEST_PRINCIPAL)
    return project


def _statement_publication(session, project_id):
    dependency_ids = session.scalars(
        select(Dependency.id).where(
            Dependency.project_id == project_id,
            Dependency.dismissed_at.is_(None),
        )
    ).all()
    return published_dependency_statements(
        session, dependency_ids, project_id=project_id
    )


@pytest.mark.parametrize(
    ("strategy", "strategy_label"),
    [
        ("protect_in_place", "Protect in place"),
        ("policy_exception", "Exception to policy"),
    ],
)
def test_the_xlsx_keeps_the_citation_columns(
    session, project, tmp_path, strategy, strategy_label
):
    """A spreadsheet that drops the provenance is just the matrix they had."""
    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    link = session.scalars(
        select(EvidenceLink).where(EvidenceLink.dependency_id == dependency.id)
    ).one()
    source_text = "Ready Milestone Road: Evidence, Assertion and Verbal"
    document = session.get(Document, link.document_id)
    document.filename = "Ready_Milestone_Evidence.pdf"
    dependency.title = source_text
    dependency.evidence_required = source_text
    dependency.resolution_strategy = strategy
    link.quote = source_text
    page = session.scalar(
        select(DocPage).where(
            DocPage.document_id == link.document_id,
            DocPage.page_no == link.page_no,
        )
    )
    page.text = source_text
    path = to_xlsx(
        session,
        project.id,
        tmp_path / "ledger.xlsx",
        evaluation=evaluate_project(session, project.id),
        statement_publication=_statement_publication(session, project.id),
    )
    sheet = load_workbook(path)["Constraint log"]

    headers = [c.value for c in sheet[1]]
    assert headers == COLUMNS
    assert "Supporting document" in headers
    assert "Source page" in headers
    assert "Cited passage" in headers
    assert "Required by" in headers
    assert "Promised for" in headers
    assert "Ready" not in headers

    row = {h: c.value for h, c in zip(headers, sheet[2])}
    assert row["Ref"].startswith("DEP-")
    assert row["Source ID"] == "FOC1-1"
    assert row["Organization"] == "Export Test Utility"
    assert row["Title"] == source_text
    assert row["Source page"] == 4
    assert row["Supporting document"] == "Ready_Milestone_Evidence.pdf"
    assert row["Cited passage"] == source_text
    assert row["Documents required for this condition"] == source_text
    assert row["Documentation review"] == "Not confirmed"
    assert row["Resolution method"] == strategy_label
    assert dependency.resolution_strategy == strategy


def test_the_xlsx_records_what_produced_it(session, project, tmp_path):
    """A snapshot with no ruleset version cannot be reproduced or dated."""
    evaluation = evaluate_project(
        session,
        project.id,
        today=date(2026, 8, 1),
        thresholds=Thresholds(stale_days=21, due_soon_days=9),
    )
    path = to_xlsx(
        session,
        project.id,
        tmp_path / "ledger.xlsx",
        evaluation=evaluation,
        statement_publication=_statement_publication(session, project.id),
    )
    workbook = load_workbook(path)
    assert "Source traceability" in workbook.sheetnames

    meta = {row[0]: row[1] for row in workbook["Source traceability"].values}
    assert meta["Project"] == "Export Test"
    assert meta["Ruleset version"]
    assert meta["Evaluated on"] == "2026-08-01"
    assert meta["STALE threshold (days)"] == 21
    assert meta["DUE_SOON threshold (days)"] == 9
    assert meta["Records"] == 1


def test_the_xlsx_refuses_an_evaluation_from_another_statement_reading(
    session, project, tmp_path
):
    """Date cells and date-derived Exceptions must describe one snapshot."""
    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    publication = _statement_publication(session, project.id)
    mismatched = evaluate_project(
        session,
        project.id,
        today=date(2026, 8, 4),
        committed_dates={dependency.id: date(2026, 8, 3)},
    )

    with pytest.raises(ValueError, match="different Committed Date readings"):
        to_xlsx(
            session,
            project.id,
            tmp_path / "mismatched-reading.xlsx",
            evaluation=mismatched,
            statement_publication=publication,
        )


def test_the_xlsx_refuses_a_same_date_from_a_different_statement(
    session, project, tmp_path
):
    """A paired evaluation cannot be reused with a later statement event."""
    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    record_verbal(
        session,
        dependency,
        stated_party="Export Test Utility",
        description="Export Test Utility said relocation will finish on August 15.",
        conversation_date=date(2026, 8, 1),
        committed_date=date(2026, 8, 15),
        principal=TEST_PRINCIPAL,
    )
    first_publication = _statement_publication(session, project.id)
    evaluation = evaluate_project(
        session,
        project.id,
        today=date(2026, 8, 4),
        committed_dates=first_publication.committed_dates,
    )
    record_verbal(
        session,
        dependency,
        stated_party="Export Test Utility",
        description="Export Test Utility repeated the August 15 commitment.",
        conversation_date=date(2026, 8, 2),
        committed_date=date(2026, 8, 15),
        principal=TEST_PRINCIPAL,
    )

    with pytest.raises(ValueError, match="different statement provenance"):
        to_xlsx(
            session,
            project.id,
            tmp_path / "different-statement.xlsx",
            evaluation=evaluation,
            statement_publication=_statement_publication(session, project.id),
        )


def test_the_xlsx_refuses_a_ledger_population_change_after_the_paired_reading(
    session, project, tmp_path
):
    """A new record cannot be mixed into an export that did not evaluate it."""
    existing = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    publication = _statement_publication(session, project.id)
    evaluation = evaluate_project(
        session, project.id, committed_dates=publication.committed_dates
    )
    session.add(
        Dependency(
            project_id=project.id,
            ref_code="DEP-export-population-change",
            source_ref="population-change",
            dep_type=existing.dep_type,
            title="A record admitted after the paired reading",
            external_org_id=existing.external_org_id,
        )
    )
    session.flush()
    path = tmp_path / "population-change.xlsx"

    with pytest.raises(ValueError, match="Ledger population changed"):
        to_xlsx(
            session,
            project.id,
            path,
            evaluation=evaluation,
            statement_publication=publication,
        )

    assert path.exists() is False


def test_an_evaluation_statement_reading_cannot_be_mutated_after_computation(
    session, project
):
    """Frozen evaluation facts keep a later exporter from changing their input."""
    publication = _statement_publication(session, project.id)
    evaluation = evaluate_project(
        session, project.id, committed_dates=publication.committed_dates
    )
    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()

    with pytest.raises(TypeError):
        evaluation.committed_dates[dependency.id] = date(2026, 8, 15)


def test_the_xlsx_refuses_a_statement_publication_from_another_project(
    session, tmp_path
):
    target = Project(slug="export-target-empty", name="Target", is_synthetic=True)
    foreign = Project(slug="export-foreign-empty", name="Foreign", is_synthetic=True)
    session.add_all((target, foreign))
    session.flush()
    evaluation = evaluate_project(session, target.id)
    publication = published_dependency_statements(
        session, (), project_id=foreign.id
    )

    with pytest.raises(ValueError, match="publication belongs to another project"):
        to_xlsx(
            session,
            target.id,
            tmp_path / "foreign-publication.xlsx",
            evaluation=evaluation,
            statement_publication=publication,
        )


def test_the_xlsx_carries_computed_exceptions(session, project, tmp_path):
    dep = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    record_verbal(
        session,
        dep,
        stated_party="Export Test Utility",
        description="Export Test Utility said relocation will finish on July 28.",
        conversation_date=date(2026, 7, 27),
        committed_date=date(2026, 7, 28),
        principal=TEST_PRINCIPAL,
    )
    dep.need_date = date(2026, 8, 8)
    session.flush()
    evaluation = evaluate_project(session, project.id, today=date(2026, 8, 5))

    path = to_xlsx(
        session,
        project.id,
        tmp_path / "ledger.xlsx",
        evaluation=evaluation,
        statement_publication=_statement_publication(session, project.id),
    )
    sheet = load_workbook(path)["Constraint log"]
    headers = [c.value for c in sheet[1]]
    row = {h: c.value for h, c in zip(headers, sheet[2])}
    by_rule = {
        e.rule: e
        for e in evaluation.for_dependency(dep.id)
    }
    assert format_exception_label(by_rule["OVERDUE"]) in (row["Constraint alerts"] or "")
    assert format_exception_label(by_rule["DUE_SOON"]) in (row["Constraint alerts"] or "")


def test_the_xlsx_attributes_a_verbal_backed_committed_date(
    session, project, tmp_path
):
    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    record_verbal(
        session,
        dependency,
        stated_party="Export Test Utility",
        description="Export Test Utility said relocation will finish in August.",
        conversation_date=date(2026, 5, 8),
        committed_date=date(2026, 8, 15),
        principal=TEST_PRINCIPAL,
    )

    path = to_xlsx(
        session,
        project.id,
        tmp_path / "ledger.xlsx",
        evaluation=evaluate_project(session, project.id),
        statement_publication=_statement_publication(session, project.id),
    )
    sheet = load_workbook(path)["Constraint log"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: cell.value for header, cell in zip(headers, sheet[2])}

    assert "Promised timing source" in headers
    assert row["Promised timing source"] == (
        "Recorded verbal statement — Export Test Utility told local:tester on 2026-05-08"
    )


def test_the_xlsx_uses_an_exact_day_statement_over_a_stale_scalar(
    session, project, tmp_path
):
    """The workbook is a reader, so its date and statement provenance agree."""
    from corridor.external_statements import (
        CitedStatementEvidence,
        StatementScope,
        StatementTiming,
        record_external_party_statement,
    )

    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    document = session.scalars(
        select(Document).where(Document.project_id == project.id)
    ).one()
    party = session.scalars(
        select(ExternalOrg).where(ExternalOrg.name == "Export Test Utility")
    ).one()
    record_external_party_statement(
        session,
        project_id=project.id,
        affected_external_org_id=party.id,
        stated_party=party.name,
        stated_external_org_id=party.id,
        source_kind="cited",
        event_date=date(2026, 5, 8),
        description="Export Test Utility will finish relocation on August 15.",
        new_timing=StatementTiming.day("August 15", date(2026, 8, 15)),
        scope=StatementScope.selected((dependency.id,)),
        created_by="corridor:event-admission",
        evidence=CitedStatementEvidence(
            document.id, 4, "Export Test Utility will finish relocation on August 15."
        ),
    )
    dependency.committed_date = date(2026, 6, 3)
    session.flush()

    path = to_xlsx(
        session,
        project.id,
        tmp_path / "statement-projection.xlsx",
        evaluation=evaluate_project(session, project.id),
        statement_publication=_statement_publication(session, project.id),
    )
    sheet = load_workbook(path)["Constraint log"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: cell.value for header, cell in zip(headers, sheet[2])}

    assert row["Promised for"].date() == date(2026, 8, 15)
    assert row["Promised timing source"] == (
        "Cited statement — nhhip-seg3c2-utilities-inventory-2-13-2026.pdf "
        "p.4: “Export Test Utility will finish relocation on August 15.”"
    )


def test_the_xlsx_suppresses_a_stale_scalar_after_a_month_statement(
    session, project, tmp_path
):
    """A source month cannot become a day in an external-ready workbook."""
    from corridor.external_statements import (
        CitedStatementEvidence,
        StatementScope,
        StatementTiming,
        record_external_party_statement,
    )

    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    document = session.scalars(
        select(Document).where(Document.project_id == project.id)
    ).one()
    party = session.scalars(
        select(ExternalOrg).where(ExternalOrg.name == "Export Test Utility")
    ).one()
    record_external_party_statement(
        session,
        project_id=project.id,
        affected_external_org_id=party.id,
        stated_party=party.name,
        stated_external_org_id=party.id,
        source_kind="cited",
        event_date=date(2026, 5, 8),
        description="Export Test Utility now expects completion in August 2026.",
        new_timing=StatementTiming.month("August 2026", 2026, 8),
        scope=StatementScope.selected((dependency.id,)),
        created_by="corridor:event-admission",
        evidence=CitedStatementEvidence(
            document.id,
            4,
            "Export Test Utility now expects completion in August 2026.",
        ),
    )
    dependency.committed_date = date(2026, 6, 3)
    session.flush()

    path = to_xlsx(
        session,
        project.id,
        tmp_path / "month-statement-projection.xlsx",
        evaluation=evaluate_project(session, project.id),
        statement_publication=_statement_publication(session, project.id),
    )
    sheet = load_workbook(path)["Constraint log"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: cell.value for header, cell in zip(headers, sheet[2])}

    assert row["Promised for"] is None
    assert row["Promised timing source"] is None
    assert "No exact promised date for this check" in row["Constraint alerts"]
    assert "No promised timing" not in row["Constraint alerts"]


def test_the_xlsx_withholds_an_unverified_cited_statement_date(
    session, project, tmp_path
):
    """A standalone workbook cannot publish a cited date without Evidence."""
    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    event = DependencyEvent(
        project_id=project.id,
        affected_external_org_id=dependency.external_org_id,
        stated_external_org_id=dependency.external_org_id,
        scope_mode="selected",
        event_type="commitment",
        source_kind="cited",
        stated_party="Export Test Utility",
        event_date=date(2026, 5, 8),
        description="Export Test Utility will finish relocation on August 15.",
        created_by="corridor:event-admission",
    )
    session.add(event)
    session.flush()
    session.add_all(
        (
            DependencyEventScope(event_id=event.id, dependency_id=dependency.id),
            DependencyEventTiming(
                event_id=event.id,
                kind="new",
                text="August 3",
                precision="day",
                start_date=date(2026, 8, 3),
                end_date=date(2026, 8, 3),
            ),
        )
    )
    session.flush()
    publication = _statement_publication(session, project.id)

    path = to_xlsx(
        session,
        project.id,
        tmp_path / "unverified-statement.xlsx",
        evaluation=evaluate_project(session, project.id, today=date(2026, 8, 4)),
        statement_publication=publication,
    )
    sheet = load_workbook(path)["Constraint log"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: cell.value for header, cell in zip(headers, sheet[2])}

    assert row["Promised for"] is None
    assert row["Promised timing source"] is None
    assert "DUE_SOON" not in (row["Constraint alerts"] or "")
    assert "OVERDUE" not in (row["Constraint alerts"] or "")


def test_the_xlsx_does_not_treat_a_scalar_only_date_as_statement_authority(
    session, project, tmp_path
):
    """A stale materialized scalar cannot publish a commitment by itself."""
    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    dependency.committed_date = date(2026, 8, 15)
    session.flush()

    path = to_xlsx(
        session,
        project.id,
        tmp_path / "legacy-scalar-projection.xlsx",
        evaluation=evaluate_project(
            session,
            project.id,
            statement_publication=_statement_publication(session, project.id),
        ),
        statement_publication=_statement_publication(session, project.id),
    )
    sheet = load_workbook(path)["Constraint log"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: cell.value for header, cell in zip(headers, sheet[2])}

    assert row["Promised for"] is None
    assert row["Promised timing source"] is None


def test_the_pdf_renders(session, project, tmp_path):
    html = render(build_report(session, project.id))
    path = to_pdf(html, tmp_path / "report.pdf")
    assert path.exists()
    assert path.read_bytes()[:5] == b"%PDF-"


def test_an_empty_ledger_still_exports(session, tmp_path):
    """A project with nothing adjudicated yet must not crash the export."""
    empty = Project(slug="exp-empty", name="Empty", is_synthetic=True)
    session.add(empty)
    session.flush()
    path = to_xlsx(
        session,
        empty.id,
        tmp_path / "empty.xlsx",
        evaluation=evaluate_project(session, empty.id),
        statement_publication=_statement_publication(session, empty.id),
    )
    sheet = load_workbook(path)["Constraint log"]
    assert [c.value for c in sheet[1]] == COLUMNS
    assert sheet.max_row == 1


def test_xlsx_uses_publication_support_while_readiness_stays_independent(
    session, project, tmp_path
):
    dependency = session.scalars(
        select(Dependency).where(Dependency.project_id == project.id)
    ).one()
    original = session.scalars(
        select(EvidenceLink).where(EvidenceLink.dependency_id == dependency.id)
    ).one()
    completion = EvidenceLink(
        dependency_id=dependency.id,
        document_id=original.document_id,
        page_no=4,
        quote="completion evidence only",
        verified=True,
    )
    publication = EvidenceLink(
        dependency_id=dependency.id,
        document_id=original.document_id,
        page_no=4,
        quote="explicit publication evidence",
        verified=True,
    )
    session.add_all([completion, publication])
    session.flush()
    mark_satisfies(
        session,
        dependency.id,
        completion.id,
        principal=TEST_PRINCIPAL,
    )
    designate_publication_support(
        session,
        dependency.id,
        publication.id,
        principal=TEST_PRINCIPAL,
    )

    path = to_xlsx(
        session,
        project.id,
        tmp_path / "role-scoped.xlsx",
        evaluation=evaluate_project(session, project.id),
        statement_publication=_statement_publication(session, project.id),
    )
    sheet = load_workbook(path)["Constraint log"]
    headers = [cell.value for cell in sheet[1]]
    row = {header: cell.value for header, cell in zip(headers, sheet[2])}

    assert row["Documentation review"] == "Documents marked sufficient"
    assert row["Cited passage"] == "explicit publication evidence"
    assert row["Cited passage"] != "completion evidence only"
