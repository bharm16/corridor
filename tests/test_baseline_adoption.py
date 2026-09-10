"""Adopt Baseline: two bounded readings, and one atomic attributable Save (#509).

The operations reading proves the workbook mechanics; the coordinator reading
carries only material project questions; the Save writes one Project Record
revision with separately identified baseline decisions and moves the project
into adopted-baseline mode, or writes nothing at all.

Every workbook here is synthetic. It exercises the importer, never a claim
about a real customer form (ADR-0046).
"""

from __future__ import annotations

import hashlib
import json
from uuid import uuid4

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select, text
from sqlalchemy.exc import DBAPIError

from corridor import audit
from corridor.baseline_adoption import (
    COORDINATOR_QUESTION_KINDS,
    BaselineAdoptionRefused,
    BaselineOperationsUnresolved,
    FormatIdentity,
    StaleBaselinePreview,
    adopt_baseline,
    adopted_baseline_source,
    adopted_source_rows,
    effective_baseline_formats,
    effective_field_mapping_manifest,
    preview_baseline_adoption,
    register_baseline_format,
    stored_mapping_revision,
)
from corridor.baseline_workbook import (
    IMPORTER_IDENTITY,
    IMPORTER_VERSION,
    BaselineWorkbookUnsupported,
    read_baseline_workbook,
)
from corridor.config import settings
from corridor.field_mapping_manifest import (
    DEMO_EXTERNAL_REFERENCES,
    MappingDeclaration,
    MappingManifestRefused,
    declared_field_mapping,
    manifest_from_declaration,
)
from corridor.models import (
    AuditLog,
    BaselineFormat,
    BaselineFormatManifest,
    BaselineSource,
    BaselineSourceRow,
    Dependency,
    Document,
    Fact,
    FactDecision,
    Project,
    ProjectRecordRevision,
    SourceSegment,
    SupportAssessment,
)
from corridor.operating_mode import (
    ADOPTED_BASELINE,
    LEGACY,
    baseline_adoption,
    project_operating_mode,
)
from corridor.principals import HumanPrincipal
from corridor.source_intake import validate_and_stage
from corridor.support_assessments import FactProposition, current_support_assessments


PRINCIPAL = HumanPrincipal("local:coordinator")

# The four headings #509 guessed at are no longer production defaults (#597).
# A fixture that wants them declares this profile by its name, which is exactly
# what a customer does with headings read from their own data dictionary.
DEMO = MappingDeclaration(external_references=DEMO_EXTERNAL_REFERENCES)
DEMO_HEADINGS = DEMO.external_reference_headings

HEADINGS = [
    "Utility Conflict ID",
    "Utility Owner",
    "Utility Type",
    "Size",
    "Material",
    "Station Origin",
    "Start Station",
    "End Station",
    "Resolution Strategy Selected (from Resolution Alternatives)",
    "Promised For",
    "Action Due Date",
    "Comment",
    "UCM Record ID",
    "Record URL",
    "Early TxDOT Utility Activity",
]

# UC-1 and UC-2 are ordinary rows; UC-2 repeats under a different owner and
# location, UC-4 is the form's own retired numbering, UC-5 names no owner,
# UC-6 promises a date nothing can read, and UC-7's *non*-material due date is
# equally unreadable — the two must land on opposite sides of the reading split.
ROWS = [
    ["UC-1", "CenterPoint Energy", "Electric", "12 in", "Steel", "SR-BL",
     "1149+00", "1150+00", "Relocate", "2026-03-01", "2026-02-01",
     "pole at station", "UCM-1001", "https://ucm.example/records/1001", "Yes"],
    ["UC-2", "City of Austin", "Water", "8 in", "PVC", "SR-BL",
     "1160+00", "1161+00", "Adjust", "2026-04-01", "2026-03-01",
     "", "UCM-1002", "", ""],
    ["UC-2", "Oncor", "Electric", "4 in", "Copper", "SR-BL",
     "1180+00", "1181+00", "Relocate", "2026-05-01", "", "", "UCM-1003", "", ""],
    ["UC-4", "Not Used", "", "", "", "", "", "", "", "", "", "", "", "", ""],
    ["UC-5", "", "Gas", "6 in", "Steel", "SR-BL", "1200+00", "1201+00",
     "", "", "", "", "", "", ""],
    ["UC-6", "Atmos Energy", "Gas", "6 in", "Steel", "SR-BL",
     "1210+00", "1211+00", "Relocate", "TBD", "2026-06-01", "", "", "", ""],
    ["UC-7", "Google Fiber", "Telecom", "2 in", "HDPE", "SR-BL",
     "1220+00", "1221+00", "Adjust", "2026-07-01", "TBD", "", "", "", ""],
]


@pytest.fixture
def store(tmp_path, monkeypatch):
    """Point the content-addressed store and page renders at a temp directory."""

    monkeypatch.setattr(settings, "corpus_store", str(tmp_path / "files"))
    monkeypatch.setattr(settings, "corpus_images", str(tmp_path / "images"))
    return tmp_path / "files"


@pytest.fixture
def project(session):
    row = Project(
        slug=f"adopt-baseline-{uuid4().hex[:8]}",
        name="Adopt Baseline Test",
        is_synthetic=True,
    )
    session.add(row)
    session.flush()
    return row


def _workbook_bytes(
    tmp_path,
    name="ucm.xlsx",
    *,
    headings=None,
    rows=None,
    hide_row=None,
    formula_cell=None,
    validation=None,
    extra_sheet=None,
) -> bytes:
    path = tmp_path / name
    book = Workbook()
    sheet = book.active
    sheet.title = "Utility Conflicts"
    sheet.append(["Utility Conflict Management (UCM) — Utility Conflicts"])
    sheet.append(list(HEADINGS if headings is None else headings))
    for row in ROWS if rows is None else rows:
        sheet.append(list(row))
    if hide_row is not None:
        sheet.row_dimensions[hide_row].hidden = True
    if formula_cell is not None:
        sheet[formula_cell] = "=A3"
    if validation is not None:
        column, formula = validation
        rule = DataValidation(type="list", formula1=formula)
        sheet.add_data_validation(rule)
        rule.add(f"{column}3:{column}99")
    if extra_sheet is not None:
        other = book.create_sheet(extra_sheet)
        other.append(["Drop-Down Lists"])
    book.save(path)
    return path.read_bytes()


def _stage(body: bytes, filename: str = "ucm.xlsx"):
    return validate_and_stage(body, filename)


def _preview(session, project, staged, **overrides):
    values = {
        "customer": "Lone Star Transit Authority",
        "source_identity": "UCM workbook revision C",
    }
    values.setdefault("field_mapping", DEMO)
    values.update(overrides)
    return preview_baseline_adoption(
        session, project=project, staged=staged, **values
    )


def _adopt(session, preview, tmp_path, key="adopt-1"):
    return adopt_baseline(
        session,
        preview=preview,
        principal=PRINCIPAL,
        idempotency_key=key,
        images_dir=tmp_path / "images",
    )


# --- The Corridor operations reading ---------------------------------------


def test_the_operations_reading_reports_the_workbook_mechanics(tmp_path, store):
    body = _workbook_bytes(tmp_path, extra_sheet="Drop-Down Lists")
    staged = _stage(body)

    reading = read_baseline_workbook(
        staged.stored_path, external_references=DEMO_HEADINGS
    )

    assert reading.parser == "openpyxl:data_only"
    assert reading.importer_identity == IMPORTER_IDENTITY
    assert reading.importer_version == IMPORTER_VERSION
    assert reading.adopted_sheet == "Utility Conflicts"
    assert reading.header_row_number == 2
    assert reading.source_row_key_rule == "sheet_name!worksheet_row_number"
    assert [sheet.name for sheet in reading.worksheets] == [
        "Utility Conflicts",
        "Drop-Down Lists",
    ]
    assert [sheet.disposition for sheet in reading.worksheets] == [
        "adopted",
        "not_adopted",
    ]
    mapped = {column.field: column.column for column in reading.column_mapping}
    assert mapped["utility_id"] == "A"
    assert mapped["external_org"] == "B"
    assert mapped["committed_date"] == "J"
    assert [column.heading for column in reading.unknown_columns] == [
        "Early TxDOT Utility Activity"
    ]
    assert reading.unknown_columns[0].populated_cells == 1
    assert reading.formula_cells == ()
    assert reading.hidden_content == ()
    assert reading.round_trip.clean
    assert reading.round_trip.rows_checked == len(ROWS)
    assert reading.round_trip.values_checked > 0
    assert reading.resolved


def test_no_heading_carries_an_external_reference_by_its_spelling_alone(
    tmp_path, store
):
    """#597: the four guessed headings are gone from the production default.

    `UCM Record ID`, `Document Control No.`, `Record URL` and `Document Link`
    were provisional names with no customer form behind them. Read with nothing
    declared, they are retained unknown columns like any other heading nobody
    named — and read through a declared profile they carry their roles again.
    """

    staged = _stage(_workbook_bytes(tmp_path))

    undeclared = read_baseline_workbook(staged.stored_path)

    assert [column.heading for column in undeclared.unknown_columns] == [
        "UCM Record ID",
        "Record URL",
        "Early TxDOT Utility Activity",
    ]
    assert all(row.external_system_id is None for row in undeclared.rows)
    assert all(row.source_url is None for row in undeclared.rows)
    # Nothing is lost by declining to guess: the cells are retained.
    retained = {
        value.heading
        for row in undeclared.rows
        for value in row.retained
    }
    assert {"UCM Record ID", "Record URL"} <= retained

    declared = read_baseline_workbook(
        staged.stored_path, external_references=DEMO_HEADINGS
    )

    assert declared.rows[0].external_system_id == "UCM-1001"
    assert declared.rows[0].source_url == "https://ucm.example/records/1001"


def test_a_declared_external_reference_role_must_be_one_this_reader_carries(
    tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))

    with pytest.raises(BaselineWorkbookUnsupported, match="external-reference role"):
        read_baseline_workbook(
            staged.stored_path,
            external_references={"ucm record id": "whatever_the_caller_liked"},
        )


def test_the_operations_reading_reports_formulas_and_hidden_content(
    tmp_path, store
):
    body = _workbook_bytes(tmp_path, formula_cell="P3", hide_row=4)
    staged = _stage(body)

    reading = read_baseline_workbook(
        staged.stored_path, external_references=DEMO_HEADINGS
    )

    assert reading.formula_cells == ("Utility Conflicts!P3",)
    assert reading.hidden_content == ("Utility Conflicts!row 4 is hidden",)
    # A populated column whose heading is blank is reported, never dropped.
    assert "Column P (blank heading)" in [
        column.heading for column in reading.unknown_columns
    ]
    assert reading.resolved


def test_the_operations_reading_reports_a_controlled_vocabulary_and_its_outliers(
    tmp_path, store
):
    body = _workbook_bytes(tmp_path, validation=("C", '"Electric,Water,Gas"'))
    staged = _stage(body)

    reading = read_baseline_workbook(
        staged.stored_path, external_references=DEMO_HEADINGS
    )

    controlled = {item.column: item for item in reading.controlled_vocabularies}
    assert controlled["C"].heading == "Utility Type"
    assert controlled["C"].checked is True
    assert controlled["C"].allowed_values == ("Electric", "Water", "Gas")
    assert controlled["C"].out_of_vocabulary == ("Telecom",)


def test_a_controlled_vocabulary_held_in_another_sheet_is_reported_unresolved(
    tmp_path, store
):
    body = _workbook_bytes(
        tmp_path,
        validation=("C", "'Drop-Down Lists'!$A$1:$A$9"),
        extra_sheet="Drop-Down Lists",
    )
    staged = _stage(body)

    reading = read_baseline_workbook(
        staged.stored_path, external_references=DEMO_HEADINGS
    )

    controlled = {item.column: item for item in reading.controlled_vocabularies}
    assert controlled["C"].checked is False
    assert controlled["C"].reference == "'Drop-Down Lists'!$A$1:$A$9"
    assert controlled["C"].out_of_vocabulary == ()


def test_a_file_that_heads_no_conflict_matrix_is_refused_by_operations(
    tmp_path, store
):
    body = _workbook_bytes(
        tmp_path,
        headings=["Notes", "More Notes"],
        rows=[["a", "b"]],
    )
    staged = _stage(body)

    with pytest.raises(BaselineWorkbookUnsupported):
        read_baseline_workbook(
        staged.stored_path, external_references=DEMO_HEADINGS
    )


def test_a_sequencing_column_blocks_adoption_before_the_coordinator_sees_it(
    session, project, tmp_path, store
):
    headings = list(HEADINGS) + ["Dependent Activity"]
    rows = [list(row) + ["UC-2"] for row in ROWS]
    staged = _stage(_workbook_bytes(tmp_path, headings=headings, rows=rows))

    reading = read_baseline_workbook(
        staged.stored_path, external_references=DEMO_HEADINGS
    )
    assert not reading.resolved
    assert [item.code for item in reading.blocking_diagnostics] == [
        "sequencing_column_unsupported"
    ]

    with pytest.raises(BaselineOperationsUnresolved):
        _preview(session, project, staged)


# --- The coordinator reading ------------------------------------------------


def test_the_coordinator_reading_reports_only_material_project_questions(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))

    preview = _preview(session, project, staged)

    assert {item.kind for item in preview.questions} <= set(
        COORDINATOR_QUESTION_KINDS
    )
    # No importer diagnostic, unknown column, formula, or hidden row reaches
    # the coordinator: those are operations' to resolve and they did.
    mechanics = (
        [item.code for item in preview.operations.diagnostics]
        + [column.heading for column in preview.operations.unknown_columns]
        + list(preview.operations.formula_cells)
        + list(preview.operations.hidden_content)
        + ["column", "worksheet", "formula", "parser", "round trip"]
    )
    rendered = " ".join(
        f"{item.kind} {item.subject} {item.detail}" for item in preview.questions
    ).casefold()
    for term in mechanics:
        assert term.casefold() not in rendered, term

    scope = preview.questions_of("adopted_scope")
    assert len(scope) == 1
    assert scope[0].subject == "Lone Star Transit Authority: UCM workbook revision C"


def test_a_repeated_identifier_with_different_facts_asks_about_distinct_facilities(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))

    preview = _preview(session, project, staged)

    distinct = preview.questions_of("likely_distinct_facilities")
    assert [item.subject for item in distinct] == ["UC-2"]
    assert distinct[0].source_rows == ("Utility Conflicts!4", "Utility Conflicts!5")
    subjects = [
        item.record_subject_key for item in preview.adopted_rows
        if item.row.business_identity == "UC-2"
    ]
    assert subjects == ["UC-2", "UC-2#2"]


def test_a_repeated_identifier_describing_one_facility_asks_about_a_duplicate(
    session, project, tmp_path, store
):
    rows = [ROWS[0], list(ROWS[0])]
    staged = _stage(_workbook_bytes(tmp_path, rows=rows))

    preview = _preview(session, project, staged)

    duplicate = preview.questions_of("duplicate_business_identity")
    assert [item.subject for item in duplicate] == ["UC-1"]
    assert preview.questions_of("likely_distinct_facilities") == ()
    assert [item.record_subject_key for item in preview.adopted_rows] == [
        "UC-1",
        "UC-1#2",
    ]


def test_a_populated_material_value_that_cannot_map_is_a_coordinator_question(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))

    preview = _preview(session, project, staged)

    unmappable = preview.questions_of("unmappable_material_value")
    assert [item.source_rows for item in unmappable] == [("Utility Conflicts!8",)]
    assert "committed_date" in unmappable[0].subject
    assert "TBD" in unmappable[0].subject


def test_the_same_failure_on_a_non_material_column_stays_with_operations(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))

    preview = _preview(session, project, staged)

    reported = {
        (item.cell_range, item.field, item.material)
        for item in preview.operations.unsupported_values
    }
    assert ("J8", "committed_date", True) in reported
    assert ("K9", "action_due_date", False) in reported
    asked = {item.source_rows for item in preview.questions_of("unmappable_material_value")}
    assert ("Utility Conflicts!9",) not in asked


def test_every_excluded_row_is_a_proposed_exclusion_with_a_plain_reason(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))

    preview = _preview(session, project, staged)

    exclusions = {item.subject: item.detail for item in preview.questions_of("proposed_exclusion")}
    assert set(exclusions) == {"Utility Conflicts!6", "Utility Conflicts!7"}
    assert "retires this number" in exclusions["Utility Conflicts!6"]
    assert "utility owner" in exclusions["Utility Conflicts!7"]


def test_more_than_one_station_origin_is_a_conflicting_plan_basis_question(
    session, project, tmp_path, store
):
    rows = [list(ROWS[0]), list(ROWS[1])]
    rows[1][5] = "US-290-BL"
    staged = _stage(_workbook_bytes(tmp_path, rows=rows))

    preview = _preview(session, project, staged)

    basis = preview.questions_of("conflicting_plan_basis")
    assert len(basis) == 1
    assert basis[0].subject == "SR-BL, US-290-BL"


def test_a_coordinator_is_never_asked_to_accept_one_row_at_a_time(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))

    preview = _preview(session, project, staged)

    per_row = [
        item for item in preview.questions if item.kind not in
        ("proposed_exclusion", "unmappable_material_value")
    ]
    assert len(per_row) < len(preview.rows)
    assert len(preview.questions) < preview.adopted_value_count


# --- One atomic Save --------------------------------------------------------


def test_a_named_person_adopts_the_whole_preview_in_one_save(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)

    result = _adopt(session, preview, tmp_path)

    revision = session.get_one(ProjectRecordRevision, result.revision_id)
    assert revision.command_type == "adopt_baseline"
    assert revision.human_principal == "local:coordinator"
    assert revision.released_policy is None
    decisions = session.scalars(
        select(FactDecision).where(FactDecision.revision_id == revision.id)
    ).all()
    assert len(decisions) == preview.adopted_value_count == len(result.fact_ids)
    assert {decision.disposition for decision in decisions} == {"include"}
    # Separately identified: one decision per adopted value, each naming its
    # own Source Fact and subject, inside one revision.
    assert len({decision.fact_id for decision in decisions}) == len(decisions)
    assert session.scalar(
        select(func.count()).select_from(ProjectRecordRevision).where(
            ProjectRecordRevision.project_id == project.id
        )
    ) == 1

    source = adopted_baseline_source(session, project.id)
    assert source.content_sha256 == staged.sha256
    assert source.byte_size == staged.size_bytes
    assert source.customer == "Lone Star Transit Authority"
    assert source.source_identity == "UCM workbook revision C"
    assert source.source_kind == "ucm_workbook"
    assert source.importer_identity == IMPORTER_IDENTITY
    assert source.importer_version == IMPORTER_VERSION
    assert source.preview_fingerprint == preview.binding_fingerprint
    assert source.revision_id == revision.id
    assert source.worksheet_scope["adopted_sheet"] == "Utility Conflicts"

    entry = session.scalars(
        select(AuditLog).where(
            AuditLog.action == audit.ADOPT_BASELINE,
            AuditLog.entity_id == project.id,
        )
    ).one()
    assert entry.after_json["preview_fingerprint"] == preview.binding_fingerprint
    assert entry.human_principal == "local:coordinator"


def test_adoption_activates_the_one_way_operating_mode_in_the_same_transaction(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    assert project_operating_mode(session, project.id) == LEGACY
    preview = _preview(session, project, staged)

    result = _adopt(session, preview, tmp_path)

    assert project_operating_mode(session, project.id) == ADOPTED_BASELINE
    receipt = baseline_adoption(session, project.id)
    assert receipt.id == result.adoption_id
    assert receipt.revision_id == result.revision_id
    assert receipt.baseline_source_sha256 == staged.sha256
    assert receipt.importer_identity == IMPORTER_IDENTITY


def test_every_source_backed_baseline_value_carries_an_effective_support_assessment(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)

    result = _adopt(session, preview, tmp_path)

    assert session.scalar(
        select(func.count()).select_from(SupportAssessment).where(
            SupportAssessment.project_id == project.id
        )
    ) == len(result.fact_ids)
    for fact_id in result.fact_ids:
        assessments = current_support_assessments(
            session, project.id, FactProposition(fact_id=fact_id)
        )
        assert len(assessments) == 1
        assert assessments[0].evidence_role == "value_support"
        assert assessments[0].assessment == "supported"
        assert assessments[0].human_principal == "local:coordinator"
        assert assessments[0].released_policy is None


def test_source_row_identity_and_record_subject_identity_stay_distinct(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)

    _adopt(session, preview, tmp_path)

    rows = {row.source_row_key: row for row in adopted_source_rows(session, project.id)}
    assert set(rows) == {item.row.source_row_key for item in preview.rows}
    assert rows["Utility Conflicts!4"].business_identity == "UC-2"
    assert rows["Utility Conflicts!5"].business_identity == "UC-2"
    assert rows["Utility Conflicts!4"].record_subject_key == "UC-2"
    assert rows["Utility Conflicts!5"].record_subject_key == "UC-2#2"
    # And the source rows are reachable from the accepted values: a Fact's
    # subject key is its source row, never the record subject.
    subject_keys = set(
        session.scalars(
            select(Fact.subject_key).where(Fact.project_id == project.id)
        ).all()
    )
    assert subject_keys <= set(rows)
    assert "UC-2#2" not in subject_keys


def test_external_identifiers_and_source_urls_are_preserved_on_source_row_identity(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)

    _adopt(session, preview, tmp_path)

    rows = {row.source_row_key: row for row in adopted_source_rows(session, project.id)}
    assert rows["Utility Conflicts!3"].external_system_id == "UCM-1001"
    assert rows["Utility Conflicts!3"].source_url == "https://ucm.example/records/1001"
    assert rows["Utility Conflicts!4"].external_system_id == "UCM-1002"
    assert rows["Utility Conflicts!4"].source_url is None


def test_unknown_columns_and_excluded_rows_are_retained_rather_than_discarded(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)

    _adopt(session, preview, tmp_path)

    source = adopted_baseline_source(session, project.id)
    assert [column["heading"] for column in source.unknown_columns] == [
        "Early TxDOT Utility Activity"
    ]
    assert source.unknown_columns[0]["populated_cells"] == 1
    excluded = [
        row for row in adopted_source_rows(session, project.id) if row.excluded
    ]
    assert {row.source_row_key: row.exclusion_reason for row in excluded} == {
        "Utility Conflicts!6": "retired_row",
        "Utility Conflicts!7": "missing_required_fields",
    }
    assert all(row.record_subject_key is None for row in excluded)
    # The unknown column's cell is still exactly quotable: every populated cell
    # of the adopted file is a Source Segment.
    assert session.scalar(
        select(func.count()).select_from(SourceSegment).where(
            SourceSegment.project_id == project.id,
            SourceSegment.exact_text == "Yes",
        )
    ) == 1
    # ...and so is the material value nothing could type. It is retained on the
    # adoption, kept out of the accepted value, and still exactly citable.
    unsupported = {
        (item["cell_range"], item["field"], item["material"])
        for item in source.operations_summary["unsupported_values"]
    }
    assert ("J8", "committed_date", True) in unsupported
    assert ("K9", "action_due_date", False) in unsupported
    assert session.scalar(
        select(func.count()).select_from(SourceSegment).where(
            SourceSegment.project_id == project.id,
            SourceSegment.cell_range == "J8",
            SourceSegment.exact_text == "TBD",
        )
    ) == 1
    assert session.scalars(
        select(Fact.id).where(
            Fact.project_id == project.id,
            Fact.subject_key == "Utility Conflicts!8",
            Fact.fact_type == "committed_date",
        )
    ).all() == []


def test_the_preview_creates_no_accepted_authority(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    before = _spine_counts(session, project.id)

    preview = _preview(session, project, staged)

    assert preview.questions
    assert _spine_counts(session, project.id) == before
    assert project_operating_mode(session, project.id) == LEGACY


def test_after_adoption_the_legacy_path_cannot_replace_an_accepted_value(
    session, project, tmp_path, store
):
    """The composition #520 exists for: the shutdown is atomic with the Save."""

    staged = _stage(_workbook_bytes(tmp_path))
    _adopt(session, _preview(session, project, staged), tmp_path)

    with pytest.raises(DBAPIError, match="may not write an accepted value"):
        with session.begin_nested():
            session.add(
                Dependency(
                    project_id=project.id,
                    ref_code="UC-99",
                    dep_type="utility_relocation",
                    title="A legacy admission after adoption",
                )
            )
            session.flush()


def test_replaying_the_same_adoption_writes_nothing_further(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)
    first = _adopt(session, preview, tmp_path)
    before = _spine_counts(session, project.id)

    again = _adopt(session, preview, tmp_path)

    assert again.revision_id == first.revision_id
    assert again.baseline_source_id == first.baseline_source_id
    assert again.adoption_id == first.adoption_id
    assert again.created is False
    assert _spine_counts(session, project.id) == before
    # A replay records no second attributable act, of either kind.
    assert session.scalar(
        select(func.count()).select_from(AuditLog).where(
            AuditLog.entity_id == project.id,
            AuditLog.action == audit.ADOPT_BASELINE,
        )
    ) == 1
    assert session.scalar(
        select(func.count()).select_from(AuditLog).where(
            AuditLog.action == audit.CONFIRM_SOURCE_INTAKE,
            AuditLog.entity_id == first.document_id,
        )
    ) == 1


def test_a_stale_preview_is_refused(session, project, tmp_path, store):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)
    tampered = type(preview)(
        **{
            **{
                name: getattr(preview, name)
                for name in type(preview).__dataclass_fields__
            },
            "binding_fingerprint": "0" * 64,
        }
    )

    with pytest.raises(StaleBaselinePreview):
        _adopt(session, tampered, tmp_path)

    assert _spine_counts(session, project.id)["project_record_revisions"] == 0


def test_changed_bytes_are_refused(session, project, tmp_path, store):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)
    staged.stored_path.write_bytes(b"PK\x03\x04 not the workbook you previewed")

    with pytest.raises(StaleBaselinePreview):
        _adopt(session, preview, tmp_path)

    assert _spine_counts(session, project.id)["project_record_revisions"] == 0


def test_a_preview_of_another_project_cannot_adopt_this_ones_baseline(
    session, project, tmp_path, store
):
    other = Project(
        slug=f"adopt-baseline-other-{uuid4().hex[:8]}",
        name="Other",
        is_synthetic=True,
    )
    session.add(other)
    session.flush()
    staged = _stage(_workbook_bytes(tmp_path))
    foreign = _preview(session, other, staged)
    crossed = type(foreign)(
        **{
            **{
                name: getattr(foreign, name)
                for name in type(foreign).__dataclass_fields__
            },
            "project_id": project.id,
        }
    )

    with pytest.raises(StaleBaselinePreview):
        _adopt(session, crossed, tmp_path)

    assert _spine_counts(session, project.id)["project_record_revisions"] == 0


def test_a_second_different_baseline_is_refused(session, project, tmp_path, store):
    staged = _stage(_workbook_bytes(tmp_path))
    _adopt(session, _preview(session, project, staged), tmp_path)
    replacement = _stage(
        _workbook_bytes(tmp_path, "later.xlsx", rows=[ROWS[0]]), "later.xlsx"
    )

    later = _preview(
        session, project, replacement, source_identity="UCM workbook revision D"
    )
    assert later.already_adopted is True
    with pytest.raises(BaselineAdoptionRefused, match="already adopted"):
        _adopt(session, later, tmp_path, key="adopt-2")


def test_a_nonempty_project_record_is_not_silently_adopted_over(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    session.add(
        Dependency(
            project_id=project.id,
            ref_code="UC-1",
            dep_type="utility_relocation",
            title="An accepted Constraint the legacy path already wrote",
        )
    )
    session.flush()

    # A project that already holds an accepted record is refused before any
    # preview reaches a person...
    with pytest.raises(BaselineAdoptionRefused, match="accepted record decisions"):
        _preview(session, project, staged)

    # ...and an adoption that reaches the Save anyway is refused there too.
    fresh = Project(
        slug=f"adopt-baseline-fresh-{uuid4().hex[:8]}",
        name="Fresh",
        is_synthetic=True,
    )
    session.add(fresh)
    session.flush()
    preview = _preview(session, fresh, staged)
    assert preview.already_adopted is False
    session.add(
        Dependency(
            project_id=fresh.id,
            ref_code="UC-9",
            dep_type="utility_relocation",
            title="Written between the preview and the Save",
        )
    )
    session.flush()
    with pytest.raises(BaselineAdoptionRefused, match="accepted record decisions"):
        with session.begin_nested():
            _adopt(session, preview, tmp_path, key="adopt-race")
    assert _spine_counts(session, fresh.id)["project_record_revisions"] == 0


def test_the_database_itself_refuses_adoption_over_an_accepted_record(
    session, project, tmp_path, store
):
    """The refusal is not only the Python guard in front of the command."""

    document = Document(
        project_id=project.id,
        sha256=hashlib.sha256(b"registered").hexdigest(),
        filename="ucm.xlsx",
        doc_type="matrix",
        parse_status="parsed",
        pages=1,
    )
    session.add_all(
        [
            document,
            Dependency(
                project_id=project.id,
                ref_code="UC-1",
                dep_type="utility_relocation",
                title="An accepted Constraint the legacy path already wrote",
            ),
        ]
    )
    session.flush()

    with pytest.raises(DBAPIError, match="already holds an accepted Project Record"):
        with session.begin_nested():
            session.execute(
                text(
                    "select adopt_project_record_baseline("
                    ":project_id, 'local:coordinator', 'forced', "
                    "jsonb_build_object('document_id', :document_id), "
                    "'[]'::jsonb, '[]'::jsonb, '{}'::bigint[])"
                ),
                {"project_id": project.id, "document_id": document.id},
            )


def test_a_refused_adoption_leaves_no_partial_spine_rows(
    session, project, tmp_path, store, monkeypatch
):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)

    import corridor.baseline_adoption as module

    def explode(*args, **kwargs):
        raise RuntimeError("failure after capture")

    monkeypatch.setattr(module.func, "adopt_project_record_baseline", explode)
    with pytest.raises(RuntimeError, match="failure after capture"):
        with session.begin_nested():
            _adopt(session, preview, tmp_path)

    counts = _spine_counts(session, project.id)
    assert counts["project_record_revisions"] == 0
    assert counts["facts"] == 0
    assert counts["fact_decisions"] == 0
    assert counts["project_baseline_sources"] == 0
    assert project_operating_mode(session, project.id) == LEGACY


# --- Output template and field mapping identities ---------------------------


def test_the_three_identities_are_recorded_separately(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)

    _adopt(session, preview, tmp_path)

    formats = effective_baseline_formats(session, project.id)
    assert set(formats) == {"output_template", "field_mapping"}
    assert formats["output_template"].format_identity == "UCM workbook revision C"
    assert formats["output_template"].content_sha256 == staged.sha256
    assert formats["field_mapping"].format_identity == "ucm-published-column-headings"
    assert formats["field_mapping"].content_sha256 != staged.sha256
    source = adopted_baseline_source(session, project.id)
    assert source.content_sha256 == staged.sha256
    assert "format_identity" not in source.__table__.c


def test_a_replacement_output_template_changes_no_accepted_value(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)
    result = _adopt(session, preview, tmp_path)
    before = _spine_counts(session, project.id)
    previous = effective_baseline_formats(session, project.id)["output_template"]

    registered = register_baseline_format(
        session,
        project_id=project.id,
        identity=FormatIdentity(
            kind="output_template",
            identity="Customer standard UCM export",
            version="2026.1",
            content_sha256=hashlib.sha256(b"template").hexdigest(),
        ),
        principal=PRINCIPAL,
        idempotency_key="register-template-1",
        template_bytes=b"template",
    )

    after = _spine_counts(session, project.id)
    assert after["project_record_revisions"] == before["project_record_revisions"]
    assert after["fact_decisions"] == before["fact_decisions"]
    assert after["facts"] == before["facts"]
    formats = effective_baseline_formats(session, project.id)
    assert formats["output_template"].id == registered.id
    assert formats["output_template"].format_version == "2026.1"
    assert session.get_one(BaselineFormat, previous.id).superseded_by == registered.id
    # The data baseline is untouched, and no second adoption happened.
    assert adopted_baseline_source(session, project.id).revision_id == result.revision_id
    assert session.scalar(
        select(func.count()).select_from(BaselineSource).where(
            BaselineSource.project_id == project.id
        )
    ) == 1


def test_a_format_cannot_be_registered_before_the_data_baseline(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    manifest = declared_field_mapping(
        read_baseline_workbook(
            staged.stored_path, external_references=DEMO_HEADINGS
        ),
        DEMO,
    )

    with pytest.raises(DBAPIError, match="adopts its data baseline"):
        with session.begin_nested():
            register_baseline_format(
                session,
                project_id=project.id,
                identity=FormatIdentity(
                    kind="field_mapping",
                    identity=manifest.identity,
                    version=manifest.version,
                    content_sha256=manifest.content_sha256,
                ),
                principal=PRINCIPAL,
                idempotency_key="register-early",
                manifest=manifest,
            )


# --- The stored mapping-revision declaration (#610) -------------------------


def test_adopting_stores_what_the_registered_mapping_revision_declares(
    session, project, tmp_path, store
):
    """The registration records the digest; this records what it digests.

    Before #610 a registration named a mapping revision by identity, version
    and digest, so Corridor could prove *that* a render used a revision and
    not *what* that revision declared. The assertions below read the stored
    bytes rather than a rebuilt object, because storing the declaration is
    exactly the thing that would otherwise be assumed.
    """

    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)

    _adopt(session, preview, tmp_path)

    registration = effective_baseline_formats(session, project.id)["field_mapping"]
    stored = session.get_one(BaselineFormatManifest, registration.id)
    assert stored.content_sha256 == registration.content_sha256
    assert stored.manifest_schema_version == "field-mapping-manifest-v1"
    # The declaration itself, not a digest of it: the exact source column, the
    # canonical field it carries, and the composition rule that says how many
    # values it carries are all readable out of the stored bytes.
    assert '"Start Station"' in stored.declaration
    assert '"station_from"' in stored.declaration
    assert '"one_value_per_column_v1"' in stored.declaration
    assert stored.declaration == preview.field_mapping_manifest.declaration_json


def test_the_stored_declaration_reads_back_as_the_mapping_revision_registered(
    session, project, tmp_path, store
):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)
    _adopt(session, preview, tmp_path)

    read_back = effective_field_mapping_manifest(session, project.id)

    # Two manifests are the same mapping revision when they declare the same
    # thing, which is what the canonical declaration and its digest say; the
    # rebuilt one holds its mappings in that canonical order rather than the
    # order the importer happened to construct them in.
    assert read_back.declaration_json == (
        preview.field_mapping_manifest.declaration_json
    )
    assert read_back.content_sha256 == (
        preview.field_mapping_manifest.content_sha256
    )
    assert read_back.mapping_for("station_from").source_columns == ("Start Station",)
    assert read_back.mapping_for("station_from").delimiter == " - "
    assert read_back.external_reference_headings == DEMO_HEADINGS


def test_a_stored_declaration_is_retrievable_by_its_identity_and_digest(
    session, project, tmp_path, store
):
    """And a digest nothing stored resolves to nothing, explicitly."""

    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)
    _adopt(session, preview, tmp_path)
    manifest = preview.field_mapping_manifest

    resolved = stored_mapping_revision(
        session,
        identity=manifest.identity,
        version=manifest.version,
        content_sha256=manifest.content_sha256,
    )
    assert resolved.declaration_json == manifest.declaration_json
    assert stored_mapping_revision(
        session,
        identity=manifest.identity,
        version=manifest.version,
        content_sha256=hashlib.sha256(b"another revision").hexdigest(),
    ) is None


def test_a_declaration_that_is_not_the_registration_s_digest_is_refused(
    session, project, tmp_path, store
):
    """A stored declaration disagreeing with its registration is unrepresentable."""

    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)
    _adopt(session, preview, tmp_path)
    registration = effective_baseline_formats(session, project.id)["field_mapping"]
    tampered = preview.field_mapping_manifest.declaration_json.replace(
        "Start Station", "Beginning Station"
    )
    assert tampered != preview.field_mapping_manifest.declaration_json

    with pytest.raises(DBAPIError, match="digests to"):
        with session.begin_nested():
            session.execute(
                select(
                    func.attach_baseline_format_manifest(
                        project.id, registration.id, tampered
                    )
                )
            )

    assert (
        session.get_one(BaselineFormatManifest, registration.id).declaration
        == preview.field_mapping_manifest.declaration_json
    )


def test_stored_bytes_that_do_not_reconstruct_the_declaration_are_refused(
    session, project, tmp_path, store
):
    """Stored bytes are read back and re-digested, never trusted for being stored."""

    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)
    canonical = preview.field_mapping_manifest.declaration_json
    # The same declaration, written non-canonically: it parses to the same
    # object and digests to something else, so it is not the revision the
    # registration names.
    padded = canonical.replace('","', '" , "', 1)
    assert padded != canonical
    assert json.loads(padded) == json.loads(canonical)

    with pytest.raises(MappingManifestRefused, match="does not reconstruct"):
        manifest_from_declaration(padded)


@pytest.mark.parametrize(
    "statement",
    (
        # Every one of these would satisfy the table's own constraints: the
        # inserted row names a real registration and carries a declaration
        # that digests to it, and neither the update nor the delete touches a
        # constrained column. What refuses them is the trigger, and nothing
        # else — a case the constraints could refuse anyway would prove
        # nothing about who may write here.
        "insert into project_baseline_format_manifests (format_id, project_id, "
        "format_identity, format_version, content_sha256, "
        "manifest_schema_version, declaration) values (:format_id, "
        ":project_id, :identity, 'v2', :digest, 'field-mapping-manifest-v1', "
        ":declaration)",
        "update project_baseline_format_manifests "
        "set manifest_schema_version = 'forged' where project_id = :project_id",
        "delete from project_baseline_format_manifests "
        "where project_id = :project_id",
    ),
)
def test_a_stored_declaration_cannot_be_written_around_the_command(
    session, project, tmp_path, store, statement
):
    staged = _stage(_workbook_bytes(tmp_path))
    preview = _preview(session, project, staged)
    _adopt(session, preview, tmp_path)
    manifest = preview.field_mapping_manifest
    registration = effective_baseline_formats(session, project.id)["field_mapping"]
    # A second registration of the same declaration under another version, so
    # the insert below has a registration to name and a free primary key.
    second = session.scalar(
        select(
            func.register_baseline_format(
                project.id,
                "field_mapping",
                manifest.identity,
                "v2",
                manifest.content_sha256,
                PRINCIPAL.subject,
                "register-a-second-version",
            )
        )
    )

    with pytest.raises(DBAPIError, match="registration command|immutable"):
        with session.begin_nested():
            session.execute(
                text(statement),
                {
                    "project_id": project.id,
                    "format_id": int(second["format_id"]),
                    "identity": manifest.identity,
                    "digest": manifest.content_sha256,
                    "declaration": manifest.declaration_json,
                },
            )

    session.expire_all()
    assert session.scalar(
        select(func.count())
        .select_from(BaselineFormatManifest)
        .where(BaselineFormatManifest.project_id == project.id)
    ) == 1
    stored = session.get_one(BaselineFormatManifest, registration.id)
    assert stored.declaration == manifest.declaration_json
    assert stored.manifest_schema_version == "field-mapping-manifest-v1"


# --- The database holds the boundary ---------------------------------------


@pytest.mark.parametrize(
    "statement",
    (
        "insert into project_baseline_sources (project_id, revision_id, "
        "document_id, content_sha256, byte_size, filename, source_identity, "
        "customer, source_kind, worksheet_scope, unknown_columns, "
        "coordinator_questions, operations_summary, importer_identity, "
        "importer_version, preview_fingerprint, adopted_by_principal, "
        "idempotency_key) values (:project_id, 1, 1, :digest, 1, 'f', 'i', 'c', "
        "'ucm_workbook', '{}', '[]', '[]', '{}', 'raw', 'v0', :digest, "
        "'local:forged', 'forged')",
        "update project_baseline_sources set customer = 'Someone Else' "
        "where project_id = :project_id",
        "delete from project_baseline_sources where project_id = :project_id",
    ),
)
def test_the_adopted_baseline_cannot_be_written_around_the_command(
    session, project, tmp_path, store, statement
):
    staged = _stage(_workbook_bytes(tmp_path))
    _adopt(session, _preview(session, project, staged), tmp_path)

    with pytest.raises(DBAPIError):
        with session.begin_nested():
            session.execute(
                text(statement),
                {
                    "project_id": project.id,
                    "digest": hashlib.sha256(b"forged").hexdigest(),
                },
            )

    assert adopted_baseline_source(session, project.id).customer == (
        "Lone Star Transit Authority"
    )


def _spine_counts(session, project_id: int) -> dict[str, int]:
    return {
        table.name: session.scalar(
            select(func.count()).select_from(table).where(
                table.c.project_id == project_id
            )
        )
        for table in (
            ProjectRecordRevision.__table__,
            Fact.__table__,
            FactDecision.__table__,
            BaselineSource.__table__,
            BaselineSourceRow.__table__,
            BaselineFormat.__table__,
            SupportAssessment.__table__,
        )
    }
