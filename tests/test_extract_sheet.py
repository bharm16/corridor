"""Extracting conflicts from a spreadsheet source (ADR-0005, #60).

No model, no client, no stub for one — which is the point. The structure is
explicit in the file, so the mapping is a name lookup rather than a
judgement, and the values are cells rather than a reading of a layout.
"""

from pathlib import Path
from datetime import datetime

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from corridor.db import Session, engine
from corridor.extract_sheet import PROMPT_VERSION, TIER_NATIVE, extract_document
from corridor.geometry import NoMatrixFound
from corridor.ingest import ingest_document
from corridor.models import Candidate, Document, ExtractionRun, Project
from corridor.row_accounting import RowAccounting, RowAccountingFailure

HEADINGS = [
    "Utility Conflict ID",
    "Utility Owner",
    "Utility Type",
    "Start Station",
    "End Station",
    "Utility Conflict Description",
    "Parcel U-Number",
]


@pytest.fixture
def session():
    connection = engine.connect()
    trans = connection.begin()
    s = Session(bind=connection)
    yield s
    s.close()
    trans.rollback()
    connection.close()


@pytest.fixture
def project(session):
    p = Project(slug="sheet-test", name="Sheet Extraction Test", is_synthetic=True)
    session.add(p)
    session.flush()
    return p


def make_workbook(tmp_path, rows, *, name="Utility Conflicts", extra_sheets=None):
    book = Workbook()
    book.remove(book.active)
    sheet = book.create_sheet(name)
    sheet.append(["Utility Conflict Management (UCM) - Utility Conflicts"])
    for row in rows:
        sheet.append(row)
    for other, other_rows in (extra_sheets or {}).items():
        made = book.create_sheet(other)
        for row in other_rows:
            made.append(row)
    path = tmp_path / "ucm.xlsx"
    book.save(path)
    return path


def ingest(session, project, tmp_path, rows, **kwargs):
    path = make_workbook(tmp_path, rows, **kwargs)
    document = ingest_document(
        session,
        project_id=project.id,
        path=path,
        doc_type="matrix",
        images_dir=tmp_path / "images",
    )
    document._stored_path = str(path)
    return document


def ingest_native_table(session, project, tmp_path, headings, row, *, sheet_name):
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    sheet.append(headings)
    sheet.append(row)
    path = tmp_path / f"{sheet_name}.xlsx"
    book.save(path)
    document = ingest_document(
        session,
        project_id=project.id,
        path=path,
        doc_type="plan",
        images_dir=tmp_path / "images",
    )
    document._stored_path = str(path)
    return document


@pytest.fixture(autouse=True)
def stored_file_points_at_the_fixture(monkeypatch):
    """The content-addressed store is not populated in tests."""
    import corridor.extract_sheet as module

    monkeypatch.setattr(
        module, "stored_file", lambda d: getattr(d, "_stored_path", None)
    )


ROWS = [
    HEADINGS,
    ["UC-1", "CenterPoint Energy", "Electric", "1149+00", "1150+00", "Pole in ROW", "U-4"],
    ["UC-2", "AT&T Texas", "Communications", "1151+00", "1152+00", "Duct bank", "U-9"],
]


def test_sue_probe_rows_become_cited_evidence_proposals_without_heading_guesses(
    session, project, tmp_path
):
    document = ingest_native_table(
        session,
        project,
        tmp_path,
        [
            "PROBE #",
            "UTILITY NAME",
            "DIAMETER",
            "NORTHING",
            "EASTING",
            "NATURAL GROUND ELEVATION",
            "PD",
            "DOC",
            "TOP OF UTILITY ELEVATION",
            "MYSTERY HEADING",
        ],
        ["46-A", "ENERGY TRANSFER", '18"', 13735576.02, 3195793.33, 21.02, 3.92, 4.1, 17.1, "retained"],
        sheet_name="PROBES (WITH COORD)",
    )

    proposals = extract_document(session, document)

    assert len(proposals) == 1
    proposal = proposals[0]
    assert proposal.kind == "evidence"
    assert proposal.payload_json["fields"] == {
        "probe_number": "46-A",
        "external_org": "ENERGY TRANSFER",
        "diameter": '18"',
        "northing": "13735576.02",
        "easting": "3195793.33",
        "natural_ground_elevation": "21.02",
        "top_of_utility_elevation": "17.1",
        "unmapped:PD": "3.92",
        "unmapped:DOC": "4.1",
        "unmapped:MYSTERY HEADING": "retained",
    }
    assert proposal.payload_json["unmapped_columns"] == [
        "PD",
        "DOC",
        "MYSTERY HEADING",
    ]
    assert proposal.payload_json["citations"] == [
        {
            "document_id": document.id,
            "page": 1,
            "quote": (
                '46-A ENERGY TRANSFER 18" 13735576.02 3195793.33 '
                "21.02 3.92 4.1 17.1 retained"
            ),
            "verified": True,
            "whole_row": True,
            "table_row": 1,
        }
    ]
    assert proposal.citations_verified is True
    assert proposals.row_accounting["detected_row_count"] == 1


def test_converted_test_hole_rows_preserve_blank_heading_values_as_unmapped(
    session, project, tmp_path
):
    document = ingest_native_table(
        session,
        project,
        tmp_path,
        [
            "CSJ",
            "NORTHING",
            "EASTING",
            "STATION",
            "OFFSET",
            "",
            "TEST  \nHOLE #",
            "UTILITY OWNER",
            "DIAMETER",
            "TEST HOLE\nDEPTH\n(FEET)",
            "ELEVATION NATURAL GROUND\n(NG)",
            "ELEVATION TOP OF UTILITY (TOU)",
            "DATE",
        ],
        [
            "3510-01-003",
            13739979.81,
            3212072.81,
            "147+64.72",
            162.34,
            "RT",
            "169-A",
            "VERIZON FIBER OPTIC",
            '2" DIA',
            4.54,
            18.16,
            13.78,
            datetime(2024, 11, 13),
        ],
        sheet_name="THDS INDEX (WITH COORD)",
    )

    [proposal] = extract_document(session, document)

    assert proposal.kind == "evidence"
    assert proposal.payload_json["tier"] == "native:sue_test_hole_index"
    assert proposal.payload_json["fields"]["test_hole_number"] == "169-A"
    assert proposal.payload_json["fields"]["station"] == "147+64.72"
    assert proposal.payload_json["fields"]["observation_date"] == (
        "2024-11-13 00:00:00"
    )
    assert proposal.payload_json["unmapped_columns"] == [
        "Column F (blank heading)"
    ]
    assert proposal.payload_json["fields"][
        "unmapped:Column F (blank heading)"
    ] == "RT"


def test_sue_table_pipeline_records_exact_zero_model_usage(
    session, project, tmp_path, monkeypatch
):
    from corridor import pipeline

    document = ingest_native_table(
        session,
        project,
        tmp_path,
        ["PROBE #", "UTILITY NAME", "DIAMETER", "NORTHING", "EASTING"],
        ["1", "INEOS", '8"', 13730631.63, 3168635.19],
        sheet_name="PROBES (WITH COORD)",
    )
    monkeypatch.setattr(
        pipeline,
        "stored_file",
        lambda value: getattr(value, "_stored_path", None),
    )

    class ModelMustNotRun:
        def complete(self, **_):
            raise AssertionError("native SUE extraction must not call a model")

    [proposal] = pipeline.extract_any(
        session,
        document,
        client=ModelMustNotRun(),
    )

    assert proposal.kind == "evidence"
    run = session.scalar(
        select(ExtractionRun).where(ExtractionRun.document_id == document.id)
    )
    assert run.model is None
    assert run.candidate_count == 1
    assert run.token_usage_json == {
        "scope": "run",
        "document_ids": [document.id],
        "measurement": "exact",
        "prompt_tokens": 0,
        "completion_tokens": 0,
        "reasoning_tokens": 0,
        "cached_tokens": 0,
    }


# ------------------------------------------------------- values from cells


def test_every_row_becomes_a_candidate(session, project, tmp_path):
    document = ingest(session, project, tmp_path, ROWS)

    candidates = extract_document(session, document)

    assert len(candidates) == 2
    assert {c.payload_json["fields"]["utility_id"] for c in candidates} == {
        "UC-1",
        "UC-2",
    }


def test_values_are_the_cells_themselves(session, project, tmp_path):
    """The claim ADR-0005 rests on: no layout was interpreted, so nothing
    was recovered — stationing stays in its own column and the owner is
    never pushed into a page header."""
    document = ingest(session, project, tmp_path, ROWS)

    fields = extract_document(session, document)[0].payload_json["fields"]

    assert fields == {
        "utility_id": "UC-1",
        "external_org": "CenterPoint Energy",
        "utility_type": "Electric",
        "station_from": "1149+00",
        "station_to": "1150+00",
        "conflict_description": "Pole in ROW",
    }


def test_a_column_with_no_canonical_field_is_reported_not_stored(
    session, project, tmp_path
):
    """`Parcel U-Number` is a real template column with no canonical home.
    The trigger for a deliberate vocabulary extension, exactly as on the
    page path — never something the reader decides for itself."""
    document = ingest(session, project, tmp_path, ROWS)

    payload = extract_document(session, document)[0].payload_json

    assert payload["unmapped_columns"] == ["Parcel U-Number"]
    assert "U-4" not in payload["fields"].values()


def test_the_reading_is_marked_native(session, project, tmp_path):
    """A run has to be able to say how each row was read. This one used no
    model at all, which is a different claim from either model tier."""
    document = ingest(session, project, tmp_path, ROWS)

    candidate = extract_document(session, document)[0]

    assert candidate.payload_json["tier"] == TIER_NATIVE
    assert candidate.payload_json["text_source"] == "cells"
    assert candidate.prompt_version == PROMPT_VERSION
    assert candidate.model is None


# ------------------------------------------------------------- citations


def test_a_citation_quotes_the_row_and_verifies(session, project, tmp_path):
    document = ingest(session, project, tmp_path, ROWS)

    citation = extract_document(session, document)[0].payload_json["citations"][0]

    assert citation["quote"].startswith("UC-1 CenterPoint Energy")
    assert citation["verified"] is True
    assert citation["page"] == 1
    assert citation["table_row"] == 1


def test_every_stored_value_is_on_the_page(session, project, tmp_path):
    """Nothing can be unverified here by construction — the text was
    generated from the same cells the values came from. Asserted anyway,
    because that is the property the exact threshold rests on."""
    document = ingest(session, project, tmp_path, ROWS)

    for candidate in extract_document(session, document):
        assert candidate.payload_json["unverified_fields"] == []
        assert candidate.citations_verified is True


# ------------------------------------------------------------- edge cases


def test_a_repeated_conflict_id_is_two_candidates(session, project, tmp_path):
    """The trap named on the ticket and in ADR-0005.

    The form's data dictionary says `Utility Conflict ID` is "unique within
    the transportation project". Real documents disagree — one Project A
    revision reuses 47 ids, and two distinct Comcast conflicts share
    `FOC14-69` — so the intent is worth recording precisely because the
    data does not honour it. Collapsing them would silently merge two
    conflicts into one.
    """
    rows = [
        HEADINGS,
        ["UC-1", "CenterPoint Energy", "Electric", "1149+00", "1150+00", "Pole", "U-4"],
        ["UC-1", "Comcast", "Communications", "1400+00", "1401+00", "Vault", "U-7"],
    ]
    document = ingest(session, project, tmp_path, rows)

    candidates = extract_document(session, document)

    assert len(candidates) == 2
    assert {c.payload_json["fields"]["external_org"] for c in candidates} == {
        "CenterPoint Energy",
        "Comcast",
    }


def test_a_blank_form_extracts_no_rows_and_does_not_raise(session, project, tmp_path):
    """A form with a schema and no conflicts is a real answer.

    The published template ships exactly this way, and the distinction is
    #59 story 2's: a document with no conflicts must not read the same as
    one nobody could parse.
    """
    document = ingest(session, project, tmp_path, [HEADINGS])

    assert extract_document(session, document) == []


def test_a_row_missing_an_owner_is_not_a_dependency(session, project, tmp_path):
    """Same bar as the page path: a row needs an identifier to be tracked
    by and an External Party to be owed by."""
    rows = [
        HEADINGS,
        ["UC-1", "", "Electric", "1149+00", "1150+00", "Pole", "U-4"],
        ["UC-2", "AT&T Texas", "Communications", "1151+00", "1152+00", "Duct", "U-9"],
    ]
    document = ingest(session, project, tmp_path, rows)

    candidates = extract_document(session, document)

    assert [c.payload_json["fields"]["utility_id"] for c in candidates] == ["UC-2"]


def test_blank_and_skipped_spreadsheet_rows_are_visible_in_accounting(
    session, project, tmp_path
):
    rows = [
        HEADINGS,
        ["UC-1", "CenterPoint", "Electric", "1+00", "2+00", "Pole", ""],
        ["", "", "", "", "", "", ""],
        ["UC-2", "", "Telecom", "3+00", "4+00", "Duct", ""],
    ]
    document = ingest(session, project, tmp_path, rows)

    candidates = extract_document(session, document)

    assert len(candidates) == 1
    assert candidates.row_accounting["detected_row_count"] == 3
    assert candidates.row_accounting["accounted_row_count"] == 3
    assert candidates.row_accounting["extracted_row_count"] == 1
    assert candidates.row_accounting["blank_row_count"] == 1
    assert candidates.row_accounting["skipped_row_count"] == 1
    assert [item["reason"] for item in candidates.row_accounting["rows"]] == [
        "candidate_recorded",
        "blank_source_row",
        "missing_required_fields",
    ]


def test_spreadsheet_reader_fails_when_a_detected_row_is_unaccounted(
    session, project, tmp_path, monkeypatch
):
    document = ingest(session, project, tmp_path, ROWS)
    real_account = RowAccounting.account
    dropped = False

    def drop_one(self, *args, **kwargs):
        nonlocal dropped
        if not dropped:
            dropped = True
            return None
        return real_account(self, *args, **kwargs)

    monkeypatch.setattr(RowAccounting, "account", drop_one)

    with pytest.raises(RowAccountingFailure, match="unaccounted"):
        extract_document(session, document)


def test_spreadsheet_pipeline_persists_dropped_row_failure_without_candidates(
    session, project, tmp_path, monkeypatch
):
    from corridor import pipeline

    document = ingest(session, project, tmp_path, ROWS)
    monkeypatch.setattr(
        pipeline,
        "stored_file",
        lambda value: getattr(value, "_stored_path", None),
    )
    real_account = RowAccounting.account
    dropped = False

    def drop_one(self, *args, **kwargs):
        nonlocal dropped
        if not dropped:
            dropped = True
            return None
        return real_account(self, *args, **kwargs)

    monkeypatch.setattr(RowAccounting, "account", drop_one)

    with pytest.raises(RowAccountingFailure, match="unaccounted"):
        pipeline.extract_any(session, document, client=object())

    run = session.scalar(
        select(ExtractionRun).where(ExtractionRun.document_id == document.id)
    )
    assert run.outcome == "failed"
    assert run.candidate_count == 0
    assert run.row_accounting_json["detected_row_count"] == 2
    assert run.row_accounting_json["accounted_row_count"] == 1
    assert run.row_accounting_json["unaccounted_rows"] == ["sheet:1:row:1"]
    assert session.scalar(
        select(Candidate).where(Candidate.source_document_id == document.id)
    ) is None


def test_spreadsheet_pipeline_persists_blank_and_skip_reasons(
    session, project, tmp_path, monkeypatch
):
    from corridor import pipeline

    rows = [
        HEADINGS,
        ["UC-1", "CenterPoint", "Electric", "1+00", "2+00", "Pole", ""],
        ["", "", "", "", "", "", ""],
        ["UC-2", "", "Telecom", "3+00", "4+00", "Duct", ""],
    ]
    document = ingest(session, project, tmp_path, rows)
    monkeypatch.setattr(
        pipeline,
        "stored_file",
        lambda value: getattr(value, "_stored_path", None),
    )

    pipeline.extract_any(session, document, client=object())

    run = session.scalar(
        select(ExtractionRun).where(ExtractionRun.document_id == document.id)
    )
    assert run.outcome == "completed"
    assert [item["reason"] for item in run.row_accounting_json["rows"]] == [
        "candidate_recorded",
        "blank_source_row",
        "missing_required_fields",
    ]


def test_a_workbook_with_no_conflict_sheet_is_unreadable_not_empty(
    session, project, tmp_path
):
    """`NoMatrixFound`, so `extract_project` reports it as unreadable
    through the path it already has — a layout variant nobody handled, not
    a project with no conflicts."""
    document = ingest(
        session, project, tmp_path, [["Category", "Item", "Information"]],
        name="Project Information",
    )

    with pytest.raises(NoMatrixFound):
        extract_document(session, document)


def test_the_conflict_sheet_is_chosen_over_an_inventory(session, project, tmp_path):
    """A workbook holds several tables and only one is the conflict matrix.

    ADR-0009's whole subject: an inventory records what is out there and
    never how it resolves, and reading one as a matrix is the error that
    ADR superseded ADR-0007 to correct.
    """
    document = ingest(
        session, project, tmp_path, ROWS,
        extra_sheets={"Utility Inventory": [
            ["Utility Feature ID", "Utility Owner", "Utility Type", "Size"],
            ["UF-1", "CenterPoint Energy", "Electric", '12"'],
        ]},
    )

    candidates = extract_document(session, document)

    assert len(candidates) == 2
    assert all(
        c.payload_json["fields"]["utility_id"].startswith("UC-")
        for c in candidates
    )


def test_candidates_are_persisted(session, project, tmp_path):
    document = ingest(session, project, tmp_path, ROWS)

    extract_document(session, document)

    stored = session.scalars(
        select(Candidate).where(Candidate.source_document_id == document.id)
    ).all()
    assert len(stored) == 2


# ------------------------------- one command reads both forms (#60)


def test_the_router_sends_a_workbook_to_the_native_reader(
    session, project, tmp_path, monkeypatch
):
    """`make extract` reads a project, not a file format.

    ADR-0005 makes both forms first-class, so which reader runs is a
    property of the document rather than something a caller has to know.
    Handing a workbook to the page extractor fails in a way that reads as
    "no page image" — a document nobody could collect — rather than as the
    wrong reader.
    """
    from corridor import pipeline

    document = ingest(session, project, tmp_path, ROWS)
    monkeypatch.setattr(pipeline, "stored_file", lambda d: getattr(d, "_stored_path", None))

    candidates = pipeline.extract_any(session, document, client=object())

    assert len(candidates) == 2
    assert all(c.payload_json["tier"] == TIER_NATIVE for c in candidates)


def test_the_router_sends_a_pdf_to_the_page_extractor(session, project, monkeypatch):
    """And the page path is untouched — it still gets its client."""
    from corridor import pipeline

    document = Document(
        project_id=project.id, sha256="c" * 64, filename="ucm.pdf",
        doc_type="matrix", parse_status="parsed", pages=1,
    )
    session.add(document)
    session.flush()

    seen = {}
    monkeypatch.setattr(pipeline, "stored_file", lambda d: Path("x.pdf"))
    def extract_pdf(session, target, client=None, **_runtime):
        seen["client"] = client
        accounting = RowAccounting(
            reader_version="matrix_tiered_v4",
            reader_path="page_geometry_and_transcription",
        )
        return accounting.finish([])

    monkeypatch.setattr("corridor.extract_matrix.extract_document", extract_pdf)

    class Client:
        model = "gpt-sheet-route-test"
        effort = "none"
        flex = False
        base_url = "https://provider.example/v1"

    client = Client()
    pipeline.extract_any(session, document, client=client)

    assert seen["client"] is client


def test_a_retired_row_is_excluded_by_rule_not_luck(session, project, tmp_path):
    """ADR-0012 on the spreadsheet path: a row whose only content is an
    identifier plus a retirement phrase is the form's bookkeeping, and a
    populated row carrying the phrase is a conflict."""
    rows = [
        HEADINGS[:2] + ["Utility Type", "Start Station", "End Station",
                        "Utility Conflict Description", "Comment"],
        ["UC-1", "CenterPoint Energy", "Electric", "1149+00", "1150+00", "Pole", ""],
        ["UC-2", "", "", "", "", "", "Not Used"],
        ["UC-3", "PSE", "Electric", "1151+00", "1152+00", "Vault", "Not used"],
    ]
    document = ingest(session, project, tmp_path, rows)

    ids = {
        c.payload_json["fields"]["utility_id"]
        for c in extract_document(session, document)
    }

    assert ids == {"UC-1", "UC-3"}


# --------- a second published TxDOT form: "UCM - Utility Conflict List" (#365)

# I-35 NEX South publishes its Utility Conflict Matrix as this earlier TxDOT
# workbook (ADR-0005). Its header sits on row 8 beneath a project-information
# block, and its columns are named by the workbook's own data dictionary.
UCM_LIST_HEADINGS = [
    "Utility Company",
    "Utility Company Contact",
    "Utility Conflict ID",
    "Drawing or Sheet No.",
    "Line Style",
    "Utility Type",
    "Size and/or Material",
    "Base or Ultimate",
    "Utility Conflict Description",
    "Longitudinal or Crossing",
    "Utility Placement in Relation to Existing TxDOT Right of Way",
    "Highway\nAlignment",
    "Station Origin",
    "Start Station",
    "Start Offset",
    "End Station",
    "End Offset",
    "Level of Utility Investigation  Needed",
    "Test Hole No.",
    "Test Hole Depth",
    "Recommended Action or Resolution",
    "Estimated Resolution Date",
    "Resolution Status",
    "Comments",
]

UCM_LIST_ROW = [
    "CPS Electric", "John Offer", "41", "N/A", "N/A", "Electric",
    "Pullbox (2B)", "I-35 NEX South",
    "In conflict with proposed sidewalk improvements", "Longitudinal",
    "Inside", "IH-35", "RT", "327869.22", "131.97", "-", "-", "QLC", "N/A",
    "N/A", "Accommodate - Relocation", "", "", "IH 35 E ROW",
]


def ingest_ucm_list(session, project, tmp_path, data_rows=(UCM_LIST_ROW,)):
    book = Workbook()
    book.remove(book.active)
    sheet = book.create_sheet("UCM-Conflict List")
    for row in (
        ["TxDOT Utility Conflict Management (UCM) - Utility Conflict List"],
        [""],
        ["Project Owner:", "TxDOT"],
        ["CCSJ/RCSJ.:", "0016-05-111"],
        ["Project Description:", "I-35 NEX South"],
        ["Highway or Route:", "I-35 From FM 1103 to AT&T Center Drive"],
        [""],
        UCM_LIST_HEADINGS,
        *data_rows,
    ):
        sheet.append(row)
    book.create_sheet("Field_Column Descriptions").append(["Field", "Description"])
    path = tmp_path / "i35nex-ucm.xlsx"
    book.save(path)
    document = ingest_document(
        session,
        project_id=project.id,
        path=path,
        doc_type="matrix",
        images_dir=tmp_path / "images",
    )
    document._stored_path = str(path)
    return document


def test_the_ucm_conflict_list_form_becomes_a_cited_dependency_proposal(
    session, project, tmp_path
):
    """The structured original read on its own exact terms (ADR-0005, #365):
    the form's own column names become canonical fields, the values are the
    cells, and the citation quotes the whole row and verifies exactly."""
    document = ingest_ucm_list(session, project, tmp_path)

    [proposal] = extract_document(session, document)

    assert proposal.kind == "dependency"
    assert proposal.payload_json["fields"] == {
        "external_org": "CPS Electric",
        "external_org_contact": "John Offer",
        "utility_id": "41",
        "utility_type": "Electric",
        "conflict_description": "In conflict with proposed sidewalk improvements",
        "orientation": "Longitudinal",
        "row_placement": "Inside",
        "baseline": "RT",
        "station_from": "327869.22",
        "offset_from": "131.97",
        "station_to": "-",
        "offset_to": "-",
        "sue_level": "QLC",
        "resolution_strategy": "Accommodate - Relocation",
        "notes": "IH 35 E ROW",
    }
    assert proposal.payload_json["tier"] == TIER_NATIVE
    assert proposal.payload_json["text_source"] == "cells"
    assert proposal.model is None
    citation = proposal.payload_json["citations"][0]
    assert citation["verified"] is True
    assert citation["page"] == 1
    assert citation["table_row"] == 1
    assert citation["quote"].startswith("CPS Electric John Offer 41")
    assert proposal.citations_verified is True


def test_the_ucm_conflict_list_form_reports_its_extra_columns_unmapped(
    session, project, tmp_path
):
    """Nine of the form's columns have no canonical field — including the
    two the vocabulary declines, `Resolution Status` (workflow state,
    ADR-0002) and `Estimated Resolution Date` (the project's own estimate).
    Every one is reported, never guessed into a field."""
    document = ingest_ucm_list(session, project, tmp_path)

    payload = extract_document(session, document)[0].payload_json

    assert payload["unmapped_columns"] == [
        "Drawing or Sheet No.",
        "Line Style",
        "Size and/or Material",
        "Base or Ultimate",
        "Highway\nAlignment",
        "Test Hole No.",
        "Test Hole Depth",
        "Estimated Resolution Date",
        "Resolution Status",
    ]
    # The unmapped cell values are not smuggled into a canonical field.
    assert "N/A" not in payload["fields"].values()
    assert "Accommodate - Relocation" in payload["fields"].values()


def test_the_ucm_conflict_list_form_reads_with_no_model(
    session, project, tmp_path, monkeypatch
):
    """The whole reason it is read natively (ADR-0005): a spreadsheet states
    its own structure, so no model runs and the run records an exact zero."""
    from corridor import pipeline

    document = ingest_ucm_list(session, project, tmp_path)
    monkeypatch.setattr(
        pipeline,
        "stored_file",
        lambda value: getattr(value, "_stored_path", None),
    )

    class ModelMustNotRun:
        def complete(self, **_):
            raise AssertionError("native UCM extraction must not call a model")

    proposals = pipeline.extract_any(session, document, client=ModelMustNotRun())

    assert [p.kind for p in proposals] == ["dependency"]
    run = session.scalar(
        select(ExtractionRun).where(ExtractionRun.document_id == document.id)
    )
    assert run.model is None
    assert run.token_usage_json["measurement"] == "exact"
    assert run.token_usage_json["prompt_tokens"] == 0
    assert run.token_usage_json["completion_tokens"] == 0
